"""Deterministic synthetic enterprise workload and offline benchmark (never real evidence).

Packages are sharded to preserve the existing 256-source input bound. Measurements
are aggregate across distinct runs, not a fictitious combined 143-account GUI run.
"""

import argparse
import json
import time
from collections import Counter
from pathlib import Path

from offline_analyzer.contracts.r0 import Package, Scope, Source
from offline_analyzer.ingestion.safety import sha256_file
from tools.fixtures.r1_demo import database, group, instance, interface, permission, rid
from tools.fixtures.r12_demo import allow, deny, doc, gaad_role, managed, trust

TIME = "2026-10-10T07:00:00+00:00"
REGIONS = ("us-east-1", "eu-west-1", "ap-south-1")


def generate(
    output: Path,
    accounts: int,
    resources: int,
    regions: int,
    shard_size: int,
    scanner_copies: int = 1,
) -> dict:
    if not 1 <= accounts <= 1000 or not 1 <= resources <= 2000 or not 1 <= regions <= 3:
        raise ValueError("Workload dimensions out of bounds")
    # Per account: 6 per Region, IAM, CloudTrail failure. Plus one Organization source/shard.
    if not 1 <= scanner_copies <= 50:
        raise ValueError("Scanner copies out of bounds")
    if not 1 <= shard_size or shard_size * (regions * 6 + 2) + 1 > 256:
        raise ValueError("Shard exceeds source bound")
    output.mkdir(parents=True, exist_ok=False)
    expected = Counter()
    packages = []
    for start in range(0, accounts, shard_size):
        directory = output / f"package-{start // shard_size:03d}"
        directory.mkdir()
        sources, targets = [], []

        def add(
            name,
            operation,
            account,
            region,
            body=None,
            coverage="COMPLETE",
            errors=(),
            directory=directory,
            sources=sources,
        ):
            path = directory / name if body is not None else None
            if path:
                path.write_text(
                    json.dumps(body, sort_keys=True, separators=(",", ":")), encoding="utf-8"
                )
            sources.append(
                Source(
                    path=name if path else None,
                    sha256=sha256_file(path) if path else None,
                    format="prowler-ocsf-5.10" if operation == "prowler:Scan" else "aws-json-v1",
                    operation=operation,
                    account_id=account,
                    region=region,
                    collected_at=TIME,
                    coverage=coverage,
                    errors=errors,
                    collector_version="synthetic-scale-v1",
                    request_parameters={} if operation != "prowler:Scan" else None,
                )
            )

        for ai in range(start, min(accounts, start + shard_size)):
            account = f"{800000000000 + ai:012d}"
            previous_index = (ai - start - 1) % min(shard_size, accounts - start) + start
            previous = f"{800000000000 + previous_index:012d}"
            admin_policy = f"arn:aws:iam::{account}:policy/admin"
            roles = [
                gaad_role(
                    account,
                    "worker",
                    trust({"Service": "ec2.amazonaws.com"}),
                    inline={
                        "operations": doc(
                            allow(["iam:PassRole", "ec2:RunInstances", "sts:AssumeRole"])
                        )
                    },
                    profiles=["worker"],
                ),
                gaad_role(
                    account,
                    "admin",
                    trust({"AWS": f"arn:aws:iam::{previous}:root", "Service": "ec2.amazonaws.com"}),
                    attached=[admin_policy],
                    profiles=["admin"],
                ),
                gaad_role(
                    account,
                    "blocked",
                    trust({"Service": "lambda.amazonaws.com"}),
                    inline={"deny": doc(deny("iam:PassRole"), allow("s3:Get*"))},
                    boundary=f"arn:aws:iam::{account}:policy/not-exported" if ai % 3 == 0 else None,
                ),
            ]
            add(
                f"{ai}-iam.json",
                "iam:GetAccountAuthorizationDetails",
                account,
                "global",
                {
                    "RoleDetailList": roles,
                    "Policies": [managed(account, "admin", doc(allow("*")))],
                    "UserDetailList": [],
                    "GroupDetailList": [],
                    "IsTruncated": False,
                },
            )
            expected["raw_records"] += 4
            add(
                f"{ai}-trail-error.json",
                "cloudtrail:DescribeTrails",
                account,
                REGIONS[0],
                coverage="ERROR",
                errors=("AccessDenied",),
            )
            expected["collection_errors"] += 1
            for ri, region in enumerate(REGIONS[:regions]):
                targets.append(Scope(account_id=account, region=region))
                instances, groups, interfaces, databases, events = [], [], [], [], []
                for n in range(resources):
                    number = (ai * 3 + ri) * 10000 + n + 1
                    vm = instance(number, "optional" if n % 3 == 0 else "required", 2)
                    vm.update(
                        IamInstanceProfile={
                            "Arn": f"arn:aws:iam::{account}:instance-profile/worker"
                        },
                        SubnetId=rid("subnet", number),
                        VpcId=rid("vpc", ai * 3 + ri + 1),
                    )
                    if n % 11 == 0:
                        vm.pop("MetadataOptions")  # Missing evidence never yields PASS.
                    instances.append(vm)
                    sg = group(
                        account,
                        number,
                        f"sg-{number}",
                        [permission(22, "0.0.0.0/0" if n % 4 == 0 else "192.0.2.0/24")],
                    )
                    sg["VpcId"] = vm["VpcId"]
                    groups.append(sg)
                    eni = interface(
                        account,
                        number,
                        [number],
                        public="198.51.100.10" if n % 4 == 0 else None,
                        instance=number,
                    )
                    eni.update(SubnetId=vm["SubnetId"], VpcId=vm["VpcId"])
                    interfaces.append(eni)
                    if n % 5 == 0:
                        db = database(account, f"db-{number}", "postgres", n % 10 == 0, False)
                        db["DBInstanceArn"] = f"arn:aws:rds:{region}:{account}:db:db-{number}"
                        databases.append(db)
                        # Contradictory PASS scanner evidence remains independent of native FAIL.
                        events.append(
                            {
                                "class_uid": 2004,
                                "time_dt": TIME,
                                "status_code": "PASS",
                                "status_id": 1,
                                "status": "New",
                                "severity": "Low",
                                "metadata": {
                                    "event_code": "rds_instance_storage_encrypted",
                                    "product": {"name": "Prowler", "version": "5.10.0"},
                                },
                                "finding_info": {"uid": f"synthetic-{number}"},
                                "cloud": {
                                    "provider": "aws",
                                    "account": {"uid": account},
                                    "region": region,
                                },
                                "resources": [
                                    {
                                        "uid": db["DBInstanceArn"],
                                        "type": "AwsRdsDbInstance",
                                        "region": region,
                                        "cloud_partition": "aws",
                                    }
                                ],
                            }
                        )
                # Repeated occurrences have separate locators and one canonical observation.
                instances += instances[: max(1, resources // 10)]
                expanded = []
                for event in events:
                    for copy in range(scanner_copies):
                        expanded.append(
                            {
                                **event,
                                "finding_info": {
                                    "uid": event["finding_info"]["uid"] + f"-copy-{copy}"
                                },
                            }
                        )
                events = expanded + expanded[:1]
                prefix = f"{ai}-{ri}"
                add(
                    prefix + "-instances.json",
                    "ec2:DescribeInstances",
                    account,
                    region,
                    {"Reservations": [{"OwnerId": account, "Instances": instances}]},
                )
                add(
                    prefix + "-groups.json",
                    "ec2:DescribeSecurityGroups",
                    account,
                    region,
                    {"SecurityGroups": groups},
                )
                add(
                    prefix + "-interfaces.json",
                    "ec2:DescribeNetworkInterfaces",
                    account,
                    region,
                    {"NetworkInterfaces": interfaces},
                )
                add(
                    prefix + "-routes.json",
                    "ec2:DescribeRouteTables",
                    account,
                    region,
                    {
                        "RouteTables": [
                            {
                                "RouteTableId": rid("rtb", ai * 3 + ri + 1),
                                "VpcId": rid("vpc", ai * 3 + ri + 1),
                                "OwnerId": account,
                                "Associations": [{"Main": True}],
                                "Routes": [
                                    {
                                        "DestinationCidrBlock": "0.0.0.0/0",
                                        "GatewayId": "igw-synthetic",
                                        "State": "active",
                                    }
                                ],
                            }
                        ],
                        **({"NextToken": "synthetic-continuation"} if ai % 7 == 0 else {}),
                    },
                )
                add(
                    prefix + "-rds.json",
                    "rds:DescribeDBInstances",
                    account,
                    region,
                    {"DBInstances": databases},
                )
                add(
                    prefix + "-prowler.json",
                    "prowler:Scan",
                    account,
                    region,
                    events,
                    coverage="PARTIAL",
                )
                expected["raw_records"] += (
                    len(instances)
                    + len(groups)
                    + len(interfaces)
                    + 1
                    + len(databases)
                    + len(events)
                )
                expected["duplicate_occurrences"] += max(1, resources // 10) + 1
        owner = f"{800000000000 + start:012d}"
        add(
            "organization.json",
            "organizations:DescribeOrganization",
            owner,
            "global",
            {
                "Organization": {
                    "Id": "o-synthetic001",
                    "FeatureSet": "ALL",
                    "ManagementAccountId": "899999999999",
                }
            },
            coverage="PARTIAL",
        )
        expected["raw_records"] += 1
        package = Package(targets=tuple(targets), sources=tuple(sources))
        (directory / "manifest.json").write_text(
            package.model_dump_json(indent=2), encoding="utf-8"
        )
        packages.append(directory.name)
    info = {
        "origin": "SYNTHETIC",
        "generator": "scale-v1",
        "accounts": accounts,
        "regions": list(REGIONS[:regions]),
        "instances_per_account_region": resources,
        "scanner_copies": scanner_copies,
        "packages": packages,
        "expected": dict(expected),
    }
    (output / "workload.json").write_text(json.dumps(info, indent=2), encoding="utf-8")
    return info


def benchmark(inputs: Path, output: Path) -> dict:
    """Poll process RSS throughout real service calls; no estimates labelled measurements."""
    import threading

    import psutil
    from openpyxl import load_workbook

    from offline_analyzer.api.services import Views
    from offline_analyzer.findings.generate import record_decision
    from offline_analyzer.ingestion.importer import import_package
    from offline_analyzer.reporting.excel import export_excel
    from offline_analyzer.reporting.metrics import assessment_metrics
    from offline_analyzer.rules.engine import assess
    from offline_analyzer.storage.store import Store

    info = json.loads((inputs / "workload.json").read_text())
    if info["origin"] != "SYNTHETIC":
        raise ValueError("Benchmark accepts its synthetic generator only")
    output.mkdir(parents=True, exist_ok=False)
    peak = [0]
    stop = threading.Event()
    process = psutil.Process()

    def monitor():
        while not stop.wait(0.02):
            peak[0] = max(peak[0], process.memory_info().rss)

    thread = threading.Thread(target=monitor, daemon=True)
    thread.start()
    result = {
        "origin": "SYNTHETIC",
        "platform": __import__("platform").platform(),
        "python": __import__("platform").python_version(),
        "workload": info,
        "files": len(list(inputs.rglob("*.json"))),
        "input_bytes": sum(p.stat().st_size for p in inputs.rglob("*.json")),
        "runs": [],
    }
    workspace = output / "workspace"
    views = Views(workspace)
    try:
        with Store(workspace) as store:
            for name in info["packages"]:
                print(f"Processing synthetic shard {name}", flush=True)
                times = {}

                def timed(stage, fn, times=times):
                    begin = time.perf_counter()
                    value = fn()
                    times[stage] = round(time.perf_counter() - begin, 4)
                    return value

                run = timed(
                    "import_seconds", lambda name=name: import_package(store, inputs / name)
                )
                timed("analysis_seconds", lambda run=run: assess(store, run))
                metrics = timed("metrics_seconds", lambda run=run: assessment_metrics(store, run))
                candidate = next(store.candidates(run), None)
                if candidate:
                    record_decision(
                        store,
                        run,
                        candidate.finding_id,
                        "NEEDS_VALIDATION",
                        "synthetic-benchmark",
                        "Persistence check",
                    )
                snapshot = [(f.finding_id, f.model_dump_json()) for f in store.candidates(run)]
                counts = metrics["normalized_observations"], metrics["raw_observation_occurrences"]
                assert (
                    timed(
                        "reimport_seconds", lambda name=name: import_package(store, inputs / name)
                    )
                    == run
                )
                timed("reanalysis_seconds", lambda run=run: assess(store, run))
                assert [
                    (f.finding_id, f.model_dump_json()) for f in store.candidates(run)
                ] == snapshot
                assert (
                    store.db.execute(
                        "SELECT count(*) FROM observations WHERE run_id=?", [run]
                    ).fetchone()[0],
                    store.db.execute(
                        "SELECT count(*) FROM evidence_refs WHERE run_id=?", [run]
                    ).fetchone()[0],
                ) == counts
                if candidate:
                    assert store.latest_decision(candidate.finding_id).status == "NEEDS_VALIDATION"
                for kind in ("accounts", "findings", "evidence", "coverage", "network"):
                    data = timed(
                        "query_" + kind + "_seconds",
                        lambda k=kind, run=run: views.records(store, run, k, "", "", "", 0, 50),
                    )
                    assert len(data["items"]) <= 50
                evidence = views.records(store, run, "evidence", "", "", "", 0, 1)["items"]
                if evidence:
                    timed(
                        "query_evidence_detail_seconds",
                        lambda run=run, evidence=evidence: views.evidence_detail(
                            store, run, evidence[0]["observation_id"]
                        ),
                    )
                timed("query_dashboard_seconds", lambda run=run: views.dashboard(store, run))
                report = timed(
                    "excel_seconds",
                    lambda run=run, name=name: export_excel(store, run, output / f"{name}.xlsx"),
                )
                wb = load_workbook(report, read_only=True)
                # Write-only reports can omit worksheet dimensions; iterate without materializing.
                for sheet, key in (
                    ("Findings Register", "candidate_findings"),
                    ("Control Assessment", "control_evaluation_records"),
                    ("Evidence Register", "raw_observation_occurrences"),
                ):
                    rows = sum(1 for _ in wb[sheet].iter_rows()) - 3
                    expected_rows = metrics[key] + (
                        metrics["evidence_files"] if sheet == "Evidence Register" else 0
                    )
                    assert rows == expected_rows, (sheet, rows, expected_rows)
                wb.close()
                print(json.dumps({"package": name, "timings": times}), flush=True)
                result["runs"].append(
                    {"run_id": run, "package": name, "metrics": metrics, "timings": times}
                )
                (output / "progress.json").write_text(
                    json.dumps(result, indent=2), encoding="utf-8"
                )
        # Prove durable decisions after closing and reopening the DB.
        with Store(workspace) as reopened:
            decisions = reopened.db.execute("SELECT count(*) FROM decisions").fetchone()[0]
            assert decisions == sum(r["metrics"]["candidate_findings"] > 0 for r in result["runs"])
            result["persisted_decisions_after_restart"] = decisions
        total = Counter()
        for item in result["runs"]:
            for k in (
                "normalized_observations",
                "raw_observation_occurrences",
                "unique_identified_resources",
                "control_evaluation_records",
                "candidate_findings",
                "source_collections_with_errors",
            ):
                total[k] += item["metrics"][k]
        assert total["raw_observation_occurrences"] == info["expected"]["raw_records"]
        assert (
            total["raw_observation_occurrences"] - total["normalized_observations"]
            == info["expected"]["duplicate_occurrences"]
        )
        result["totals"] = dict(total)
        result["duplicate_occurrences"] = info["expected"]["duplicate_occurrences"]
        result["rejected_records"] = 0  # Whole-import failure raises; no ignored rejects.
        result["reimport_stable"] = True
        result["accounts_represented"] = sum(
            r["metrics"]["unique_accounts"] for r in result["runs"]
        )
        result["database_bytes"] = (workspace / "assessment.duckdb").stat().st_size
    finally:
        stop.set()
        thread.join()
    result["peak_rss_bytes"] = peak[0]
    (output / "results.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    gen = sub.add_parser("generate")
    gen.add_argument("--output", type=Path, required=True)
    gen.add_argument("--accounts", type=int, default=143)
    gen.add_argument("--resources", type=int, default=25)
    gen.add_argument("--regions", type=int, default=2)
    gen.add_argument("--shard-size", type=int, default=16)
    gen.add_argument("--scanner-copies", type=int, default=10)
    bench = sub.add_parser("benchmark")
    bench.add_argument("--input", type=Path, required=True)
    bench.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    if args.command == "generate":
        print(
            json.dumps(
                generate(
                    args.output,
                    args.accounts,
                    args.resources,
                    args.regions,
                    args.shard_size,
                    args.scanner_copies,
                ),
                indent=2,
            )
        )
    else:
        result = benchmark(args.input, args.output)
        print(
            json.dumps(
                {k: result[k] for k in ("totals", "peak_rss_bytes", "database_bytes")}, indent=2
            )
        )


if __name__ == "__main__":
    main()
