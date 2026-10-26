"""Synthetic assessment exports only. Never accepts or embeds client identifiers."""

import hashlib
import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from clogcp_report_finalizer.formatting import add_table, chunks
from clogcp_report_finalizer.schema import COLUMNS, REGISTER

INDEX = (
    "Evidence mapping ID",
    "Record ID",
    "GCP project",
    "Resource identifier",
    "Source operation",
    "Dataset",
    "Original evidence reference",
    "JSON path",
    "Collection time",
    "Artifact ID",
    "SHA-256",
    "Collection status",
    "Evidence availability",
    "Configuration excerpt",
)
DETAILS = (
    "Evidence mapping ID",
    "Record ID",
    "GCP project",
    "Resource identifier",
    "Original evidence reference",
    "JSON path",
    "Excerpt part",
    "Redacted original configuration",
    "Evidence limitation",
    "Back to evidence index",
)


def build(output: Path) -> dict:
    output.mkdir(parents=True, exist_ok=False)
    project = "synthetic-project"
    bindings = [
        {"role": "roles/viewer", "members": [f"user:unrelated-{i}@example.invalid"]}
        for i in range(400)
    ]
    bindings += [
        {
            "role": "roles/iam.serviceAccountUser",
            "members": ["user:deployer@example.invalid"],
            "condition": {"expression": "resource.name.startsWith('projects/')"},
        },
        {
            "role": "roles/iam.serviceAccountTokenCreator",
            "members": ["serviceAccount:caller@synthetic-project.iam.gserviceaccount.com"],
        },
    ]
    cases = [
        (
            "IAM",
            {
                "bindings": bindings,
                "version": 3,
                "etag": "synthetic-etag",
                "denyRules": [
                    {
                        "deniedPermissions": ["iam.serviceAccounts.getAccessToken"],
                        "denialCondition": {"expression": "request.time != null"},
                    }
                ],
            },
            ["/bindings/400", "/bindings/401"],
            "gcp:projects.getIamPolicy",
            "GCP-IAM-IMPERSONATION",
        ),
        (
            "Network",
            {
                "allow_rule": {
                    "sourceRanges": ["0.0.0.0/0"],
                    "allowed": [{"IPProtocol": "tcp", "ports": ["22"]}],
                    "priority": 1000,
                },
                "deny_rule": None,
            },
            ["/allow_rule", "/deny_rule"],
            "gcp:compute.firewalls.list",
            "GCP-NET-FW-ADMIN-INTERNET",
        ),
        (
            "Network blocker",
            {
                "allow_rule": {
                    "sourceRanges": ["0.0.0.0/0"],
                    "allowed": [{"IPProtocol": "tcp", "ports": ["22"]}],
                    "priority": 1000,
                },
                "deny_rule": {
                    "denied": [{"IPProtocol": "tcp", "ports": ["22"]}],
                    "priority": 900,
                    "targetTags": ["review-required"],
                },
            },
            ["/allow_rule", "/deny_rule"],
            "gcp:compute.firewalls.list",
            "GCP-NET-FW-ADMIN-INTERNET",
        ),
        (
            "GKE Standard",
            {
                "autopilot": {"enabled": False},
                "privateClusterConfig": {"enablePrivateEndpoint": False},
                "workloadIdentityConfig": {"workloadPool": "synthetic-project.svc.id.goog"},
            },
            ["/autopilot", "/privateClusterConfig", "/workloadIdentityConfig"],
            "gcp:expanded.gke_clusters",
            "",
        ),
        (
            "GKE Autopilot",
            {
                "autopilot": {"enabled": True},
                "privateClusterConfig": {"enablePrivateEndpoint": True},
            },
            ["/autopilot", "/privateClusterConfig"],
            "gcp:expanded.gke_clusters",
            "",
        ),
        (
            "Serverless",
            {
                "ingress": "INGRESS_TRAFFIC_INTERNAL_ONLY",
                "iam": {"bindings": [{"role": "roles/run.invoker", "members": ["allUsers"]}]},
                "template": {"serviceAccount": "runtime@synthetic-project.iam.gserviceaccount.com"},
            },
            ["/ingress", "/iam/bindings/0", "/template"],
            "gcp:run.services.list",
            "",
        ),
        (
            "Storage enforced",
            {
                "bindings": [{"role": "roles/storage.objectViewer", "members": ["allUsers"]}],
                "iamConfiguration": {
                    "publicAccessPrevention": "enforced",
                    "uniformBucketLevelAccess": {"enabled": True},
                },
            },
            ["/bindings/0", "/iamConfiguration"],
            "gcp:storage.buckets.getIamPolicy",
            "GCP-STORAGE-PUBLIC",
        ),
        (
            "Storage inherited",
            {
                "bindings": [
                    {"role": "roles/storage.objectViewer", "members": ["allAuthenticatedUsers"]}
                ],
                "iamConfiguration": {"publicAccessPrevention": "inherited"},
            },
            ["/bindings/0", "/iamConfiguration"],
            "gcp:storage.buckets.getIamPolicy",
            "GCP-STORAGE-PUBLIC",
        ),
        (
            "SQL unclear defaults",
            {"settings": {"ipConfiguration": {}}},
            ["/settings/ipConfiguration"],
            "gcp:expanded.cloud_sql",
            "GCP-SQL-PUBLIC-IP",
        ),
        (
            "Organization policy",
            {
                "spec": {"inheritFromParent": True, "rules": [{"enforce": False}]},
                "effective": {"rules": [{"enforce": True}]},
            },
            ["/spec", "/effective"],
            "gcp:orgpolicy.getEffectivePolicy",
            "",
        ),
        (
            "Logging central sinks",
            {
                "auditConfigs": [
                    {
                        "service": "allServices",
                        "auditLogConfigs": [{"logType": "DATA_READ", "exemptedMembers": []}],
                    }
                ],
                "sinks": [
                    {
                        "name": "central-sink",
                        "destination": (
                            "logging.googleapis.com/projects/synthetic-log/"
                            "locations/global/buckets/security"
                        ),
                        "filter": "logName:cloudaudit.googleapis.com",
                    }
                ],
                "exclusions": [],
            },
            ["/auditConfigs", "/sinks", "/exclusions"],
            "gcp:logging.sinks.list",
            "GCP-LOG-AUDIT",
        ),
    ]
    # Exact material duplicate with a different source record ID; do not lose either row.
    cases.append(cases[0])
    cases += [
        ("Missing evidence", None, [], "gcp:projects.getIamPolicy", ""),
        ("Missing review", {"enabled": False}, ["/enabled"], "gcp:synthetic.configuration", ""),
    ]
    book = Workbook()
    book.remove(book.active)
    rows, index, details, records = [], [], [], {}
    for i, (label, config, pointers, operation, control) in enumerate(cases, 1):
        record = f"synthetic-observation-{i:03}"
        resource = f"projects/{project}/resources/item-{i}"
        if label == "IAM":
            resource = "projects/" + project
        mapping = f"synthetic-evidence-{i:03}"
        reviewed = label != "Missing review"
        row = [None] * 40
        for n, value in {
            0: record,
            1: label + " configuration review",
            2: project,
            3: resource + "; " + operation,
            4: "HIGH" if label == "IAM" else "MEDIUM",
            5: (
                "Potential impact from the recorded configuration; prerequisites "
                "require verification."
            ),
            6: "MEDIUM",
            8: (
                "Restrict the specific setting after validating the documented design "
                "and blocking controls."
            ),
            9: "Potential exposure",
            10: json.dumps(config)[:600] if config else "Original configuration unavailable",
            11: "synthetic-config.json#",
            12: "Existing source control: " + (control or "manual verification"),
            13: "COMPLETE" if config else "NOT_COLLECTED",
            14: "" if config else "Collection evidence missing",
            16: "Effective permissions and execution are unverified",
            17: "FAIL",
            18: label + "; GCP; " + operation,
            19: control or "Manual source control",
            20: "CLI/API configuration",
            21: "" if config else "Missing source configuration",
            22: (
                "Source criterion flags the recorded setting; review the applicable "
                "effective controls."
            ),
            24: "CONDITIONAL",
            25: "FAIL",
            26: i,
            27: resource,
            28: "Synthetic root cause " + label,
            29: "Pending — no decision recorded",
            30: True,
            31: True,
            32: 1,
            34: "GAP_CONFIRMED" if reviewed else "Pending — no decision recorded",
            35: "Synthetic SRE review retained" if reviewed else "",
            37: "Synthetic assessor" if reviewed else "",
            38: datetime(2026, 10, 21) if reviewed else None,
            39: "CONFIRMED; frozen Batch 1"
            if label == "IAM"
            else "CONFIRMED; Batch 2"
            if reviewed
            else "Pending — no decision recorded",
        }.items():
            row[n] = value
        if i == 12:  # make the duplicate material context equal, including frozen review
            row = list(rows[0])
            row[0], row[26] = record, i
        if not control and i != 12:
            row[17] = row[25] = "MANUAL_VALIDATION_REQUIRED"
            row[28] = "Manual observation; source conclusion requires assessor review"
        rows.append(row)
        raw = json.dumps(config, ensure_ascii=True, sort_keys=True) if config is not None else ""
        index.append(
            [
                mapping,
                record,
                project,
                resource,
                operation,
                "synthetic",
                "synthetic-config.json",
                "",
                "2026-10-21T10:00:00Z",
                "synthetic-artifact",
                hashlib.sha256(raw.encode()).hexdigest(),
                "COMPLETE" if config else "NOT_COLLECTED",
                "Available" if config else "Unavailable",
                "See detail parts",
            ]
        )
        for part, text in enumerate(chunks(raw), 1):
            details.append(
                [
                    mapping,
                    record,
                    project,
                    resource,
                    "synthetic-config.json",
                    "",
                    part,
                    text,
                    "" if config else "Missing exported evidence",
                    "Evidence index",
                ]
            )
        records[record] = {
            "selectors": {mapping: pointers},
            "control_id": control,
            "parameters": {
                "project": project,
                "rule": "synthetic-firewall",
                "bucket": "synthetic-bucket",
                "instance": "synthetic-sql",
            },
            "batch": "Batch 1" if label == "IAM" else "Batch 2",
            "canonical": {
                "project": project,
                "resource": resource,
                "root_cause": label,
                "principal_or_workload": "explicit-synthetic-principal",
                "evidence_scope": "synthetic-export",
                "mechanism": "configuration-review",
            },
        }
    add_table(book, REGISTER, COLUMNS, rows, [27, 44, 25] + [42] * 37)
    add_table(book, "GCP_Evidence_Index", INDEX, index)
    add_table(
        book, "GCP_Evidence_Details", DETAILS, details, [25, 25, 25, 40, 35, 28, 9, 100, 40, 25]
    )
    source = output / "synthetic-CloGCP-export.xlsx"
    book.save(source)
    book.close()
    guide = output / "synthetic-guidance.json"
    guide.write_text(
        json.dumps({"schema_version": "clogcp-finalizer-guidance-1", "records": records}, indent=2)
    )
    paths = output / "synthetic-attack-paths.json"
    paths.write_text(
        json.dumps(
            {
                "paths": [
                    {
                        "path_id": "synthetic-supported",
                        "classification": "SUPPORTED",
                        "exploitation_verification": "UNVERIFIED",
                        "prerequisites": [
                            {
                                "name": "observed grant",
                                "state": "AVAILABLE",
                                "observation_ids": ["synthetic-evidence"],
                            }
                        ],
                    },
                    {
                        "path_id": "synthetic-conditional",
                        "classification": "CONDITIONAL",
                        "exploitation_verification": "UNVERIFIED",
                        "prerequisites": [{"name": "binding condition", "state": "CONDITIONAL"}],
                    },
                    {
                        "path_id": "synthetic-blocked",
                        "classification": "BLOCKED",
                        "exploitation_verification": "UNVERIFIED",
                        "prerequisites": [
                            {
                                "name": "internal ingress",
                                "state": "CONTRADICTED",
                                "observation_ids": ["synthetic-ingress"],
                            }
                        ],
                    },
                ]
            },
            indent=2,
        )
    )
    history = output / "synthetic-history.json"
    history.write_text(
        json.dumps(
            {
                "reconciliation": [
                    {
                        "legacy_path_id": "synthetic-batch1",
                        "match_status": "MATCHED_FOR_REVIEW",
                        "new_paths": [{"path_id": "synthetic-conditional"}],
                        "prior_status": "PROVEN",
                        "approval_inherited": False,
                    }
                ]
            },
            indent=2,
        )
    )
    return {"source": source, "guidance": guide, "attack_paths": paths, "history": history}
