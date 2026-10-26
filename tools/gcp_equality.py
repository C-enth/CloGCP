"""GCP behaviour fingerprint for baseline-versus-candidate equality proofs (synthetic only).

Builds the synthetic GCP estate (tools/fixtures/gcp_estate.py), imports one package per
project, assesses each import, builds and assesses the consolidated snapshot, records
deterministic analyst decisions and one false-positive exclusion, then writes a canonical
JSON fingerprint: per-run rule outcome counts, finding IDs with rule, severity and status,
collection-scope statuses, dataset coverage rows, dashboard, report model, Excel cell
values and Word text. Volatile values (timestamps, durations, temporary paths) are masked.

Run it twice with different source trees and compare the outputs byte for byte:

    PYTHONPATH=<baseline>/src:. python -m tools.gcp_equality --output base.json --work W1
    PYTHONPATH=src:. python -m tools.gcp_equality --output cand.json --work W2

The fixture code always comes from this repository (`tools/`), so both runs see identical
packages; only `offline_analyzer` differs. No real evidence is read.
"""

import argparse
import hashlib
import json
import re
import shutil
import zipfile
from pathlib import Path

STAMP = re.compile(
    r"\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}:\d{2}(?:\.\d+)?(?:Z|[+-]\d{2}:?\d{2})?|\d{8}T\d{6}Z"
)
VOLATILE_KEYS = {
    "decided_at",
    "recorded_at",
    "generated_at",
    "created_at",
    "updated_at",
    "assessed_at",
    "imported_at",
    "exported_at",
    "elapsed_seconds",
    "duration_seconds",
}


def scrub(value, work: str):
    if isinstance(value, dict):
        items = {
            ("|".join(map(str, k)) if isinstance(k, tuple) else str(k)): v for k, v in value.items()
        }
        return {
            k: ("<volatile>" if k in VOLATILE_KEYS else scrub(v, work))
            for k, v in sorted(items.items())
        }
    if isinstance(value, (list, tuple)):
        return [scrub(v, work) for v in value]
    if isinstance(value, str):
        return STAMP.sub("<time>", value.replace(work, "<work>"))
    return value


def workbook_cells(path: Path) -> dict:
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True)
    return {
        ws.title: [[c for c in row] for row in ws.iter_rows(values_only=True)]
        for ws in book.worksheets
    }


def docx_text(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    return [t for t in re.findall(r"<w:t[^>]*>([^<]*)</w:t>", xml)]


def fingerprint(work: Path, projects: int) -> dict:
    import offline_analyzer
    from offline_analyzer.contracts.exclusion import ExclusionRequest
    from offline_analyzer.enterprise.snapshot import assess_enterprise
    from offline_analyzer.findings.exclusions import record
    from offline_analyzer.findings.generate import record_decision
    from offline_analyzer.ingestion.importer import import_package
    from offline_analyzer.inventory.environment import set_environment
    from offline_analyzer.reporting.assessment_report import export_assessment, model
    from offline_analyzer.reporting.cloud import dashboard
    from offline_analyzer.rules.engine import assess
    from offline_analyzer.storage.store import Store
    from tools.fixtures.gcp_estate import estate

    shutil.rmtree(work, ignore_errors=True)
    work.mkdir(parents=True)
    packages = estate(work / "packages", projects)
    set_environment(work / "workspace", "SYNTHETIC")
    store = Store(work / "workspace")
    runs = []
    for package in packages:
        run = import_package(store, package)
        assess(store, run)
        runs.append(run)
    snapshot = assess_enterprise(store, runs)
    findings = [
        fid
        for (fid,) in store.rows(
            "SELECT finding_id FROM findings WHERE run_id=? ORDER BY finding_id", [snapshot]
        )
    ]
    for fid in findings[:5]:
        record_decision(store, snapshot, fid, "CONFIRMED", "Synthetic Analyst", "synthetic")
    for fid in findings[5:7]:
        record_decision(store, snapshot, fid, "REJECTED", "Synthetic Analyst", "synthetic")
    if len(findings) > 8:
        record(
            store,
            snapshot,
            findings[8],
            ExclusionRequest(
                action="EXCLUDE",
                reason="FALSE_POSITIVE",
                note="synthetic false positive",
                analyst="Synthetic Analyst",
            ),
        )
    print("offline_analyzer from", Path(offline_analyzer.__file__).parent)
    result = {"runs": {}}
    for run in [*runs, snapshot]:
        outcomes = {}
        for rule, outcome, scope, count in store.rows(
            "SELECT json_extract_string(data,'$.rule_id'), outcome, "
            "json_extract_string(data,'$.evaluation_scope'), count(*) FROM evaluations "
            "WHERE run_id=? GROUP BY 1,2,3 ORDER BY 1,2,3",
            [run],
        ):
            outcomes[f"{rule}|{scope}|{outcome}"] = count
        rows = [
            json.loads(d)
            for (d,) in store.rows(
                "SELECT data FROM findings WHERE run_id=? ORDER BY finding_id", [run]
            )
        ]
        result["runs"][run] = {
            "outcomes": outcomes,
            "findings": [
                [f["finding_id"], f["rule_id"], f.get("suggested_severity"), f.get("status")]
                for f in rows
            ],
        }
    view = scrub(dashboard(store, snapshot), str(work))
    result["dashboard"] = view
    result["report_model"] = {
        include: scrub(model(store, snapshot, include), str(work))
        for include in ("confirmed", "valid", "all")
        if include
        in __import__("offline_analyzer.reporting.assessment_report", fromlist=["INCLUDE"]).INCLUDE
    }
    target = work / "report.zip"
    export_assessment(store, snapshot, target, include="all")
    with zipfile.ZipFile(target) as archive:
        archive.extractall(work / "report")
    for path in sorted((work / "report").iterdir()):
        if path.suffix == ".xlsx":
            result["excel"] = scrub(workbook_cells(path), str(work))
        elif path.suffix == ".docx":
            result["word"] = scrub(docx_text(path), str(work))
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--work", type=Path, required=True)
    parser.add_argument("--projects", type=int, default=10)
    args = parser.parse_args()
    data = fingerprint(args.work, args.projects)
    text = json.dumps(data, indent=1, sort_keys=True, default=str) + "\n"
    args.output.write_text(text, encoding="utf-8")
    print(
        json.dumps(
            {
                "sha256": hashlib.sha256(text.encode()).hexdigest(),
                "runs": len(data["runs"]),
                "findings": sum(len(r["findings"]) for r in data["runs"].values()),
            }
        )
    )


if __name__ == "__main__":
    main()
