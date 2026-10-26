"""Bounded passive inputs. Imported formulas and links are never executed."""

import hashlib
import json
import re
import zipfile
from pathlib import Path

from defusedxml import ElementTree
from openpyxl import load_workbook

MAX_BYTES = 64 * 1024 * 1024
MAX_CELLS = 2_000_000
SECRET = re.compile(
    "-----BEGIN (?:RSA |EC |OPENSSH )?PRIVATE "
    "KEY|(?:AKIA|ASIA)[A-Z0-9]{16}|ya29\\.[A-Za-z0-9_-]{30,}|"
    "AIza[A-Za-z0-9_-]{35}|gh[pousr]_[A-Za-z0-9]{30,}"
)


class FinalizerError(ValueError):
    """Fixed error codes never include customer values."""


def digest(path: Path) -> str:
    result = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            result.update(chunk)
    return result.hexdigest()


def unique_object(pairs: list) -> dict:
    result = {}
    for key, value in pairs:
        if key in result:
            raise FinalizerError("DUPLICATE_JSON_KEY")
        result[key] = value
    return result


def parse_json(text: str):
    if len(text.encode("utf-8")) > MAX_BYTES:
        raise FinalizerError("JSON_BYTE_LIMIT")
    try:
        data = json.loads(text, object_pairs_hook=unique_object)
    except (json.JSONDecodeError, RecursionError) as exc:
        raise FinalizerError("INVALID_JSON") from exc
    pending, count = [(data, 0)], 0
    while pending:
        item, depth = pending.pop()
        count += 1
        if depth > 40 or count > 500000:
            raise FinalizerError("JSON_COMPLEXITY_LIMIT")
        if isinstance(item, dict):
            for key, value in item.items():
                if re.fullmatch(
                    r"(?i)(password|private_?key(?:_?data)?|client_?secret|access_?token|"
                    r"refresh_?token|keyString|secretAccessKey|aws_secret_access_key|"
                    r"aws_session_token|secretString|secretValue)",
                    key,
                ):
                    if value and str(value) not in {"[REDACTED]", "[redacted]", "***"}:
                        raise FinalizerError("CREDENTIAL_MATERIAL_REFUSED")
            pending.extend((v, depth + 1) for v in item.values())
        elif isinstance(item, list):
            pending.extend((v, depth + 1) for v in item)
        elif isinstance(item, str) and SECRET.search(item):
            raise FinalizerError("CREDENTIAL_MATERIAL_REFUSED")
    return data


def read_json(path: Path):
    if path.stat().st_size > MAX_BYTES:
        raise FinalizerError("INPUT_BYTE_LIMIT")
    return parse_json(path.read_text(encoding="utf-8-sig"))


def workbook(path: Path):
    if path.suffix.lower() != ".xlsx" or path.stat().st_size > MAX_BYTES:
        raise FinalizerError("XLSX_REQUIRED_OR_BYTE_LIMIT")
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > 4096 or sum(m.file_size for m in members) > 256 * 1024 * 1024:
                raise FinalizerError("XLSX_EXPANSION_LIMIT")
            for member in members:
                name = member.filename.lower()
                if (
                    member.flag_bits & 1
                    or member.file_size > MAX_BYTES
                    or member.file_size > max(1, member.compress_size) * 300
                    or name.startswith("/")
                    or ".." in Path(name).parts
                ):
                    raise FinalizerError("UNSAFE_XLSX_MEMBER")
                if any(
                    part in name
                    for part in (
                        "vbaproject",
                        "externallinks/",
                        "embeddings/",
                        "connections.xml",
                        "querytables/",
                    )
                ):
                    raise FinalizerError("ACTIVE_WORKBOOK_CONTENT_REFUSED")
                if name.endswith(".rels"):
                    root = ElementTree.fromstring(archive.read(member))
                    if any(
                        r.get("TargetMode") == "External"
                        and not r.get("Target", "").startswith("#")
                        for r in root
                    ):
                        raise FinalizerError("EXTERNAL_WORKBOOK_LINK_REFUSED")
    except zipfile.BadZipFile as exc:
        raise FinalizerError("INVALID_XLSX") from exc
    book = load_workbook(path, data_only=False, keep_links=False)
    if sum(s.max_row * s.max_column for s in book) > MAX_CELLS:
        book.close()
        raise FinalizerError("WORKBOOK_CELL_LIMIT")
    for sheet in book:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str) and SECRET.search(cell.value):
                    book.close()
                    raise FinalizerError("CREDENTIAL_MATERIAL_REFUSED")
    return book


def literal(cell, value) -> None:
    """Typed Excel strings preserve the text without activating '=' or DDE."""
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"


def pointer(document, path: str):
    if path == "":
        return document
    if not path.startswith("/"):
        raise FinalizerError("JSON_POINTER_REQUIRED")
    current = document
    try:
        for token in path.split("/")[1:]:
            if re.search(r"~(?![01])", token):
                raise FinalizerError("INVALID_JSON_POINTER_ESCAPE")
            key = token.replace("~1", "/").replace("~0", "~")
            if isinstance(current, list):
                if not re.fullmatch(r"0|[1-9][0-9]{0,9}", key):
                    raise FinalizerError("INVALID_JSON_ARRAY_POINTER")
                current = current[int(key)]
            else:
                current = current[key]
    except (KeyError, IndexError, ValueError, TypeError) as exc:
        raise FinalizerError("EVIDENCE_SELECTOR_UNRESOLVED") from exc
    return current
