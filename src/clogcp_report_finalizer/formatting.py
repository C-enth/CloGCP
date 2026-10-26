"""Reusable restrained Excel formatting, with bounded readable detail chunks."""

import math
import re

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .safety import literal

INK, BLUE, PALE = "142D40", "254F64", "EDF3F6"


def chunks(text: str, length: int = 1000) -> list[str]:
    result, start = [], 0
    while start < len(text):
        end = min(start + length, len(text))
        breaks = [m.end() for m in re.finditer("\n", text[start:end])]
        if len(breaks) > 10:
            end = start + breaks[9]
        result.append(text[start:end])
        start = end
    return result or [""]


def add_table(book, name: str, headers: list | tuple, rows: list, widths: list | None = None):
    sheet = book.create_sheet(name)
    for r, row in enumerate([list(headers), *rows], 1):
        for c, value in enumerate(row, 1):
            literal(sheet.cell(r, c), value)
    style_sheet(sheet, widths)
    return sheet


def style_sheet(sheet, widths: list | None = None) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "C2" if sheet.max_column > 3 else "A2"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3 if sheet.max_column > 12 else sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = (
        0 if sheet.max_column > 8 else 1,
        0,
    )
    sheet.print_title_rows = "1:1"
    sheet.print_title_cols = "A:B" if sheet.max_column > 8 else None
    sheet.print_options.horizontalCentered = True
    sheet.sheet_properties.outlinePr.summaryRight = False
    sheet.print_area = sheet.dimensions
    for c in range(1, sheet.max_column + 1):
        sheet.column_dimensions[get_column_letter(c)].width = (
            widths[c - 1] if widths and c <= len(widths) else (24 if c <= 2 else 48)
        )
    sheet.row_dimensions[1].height = 46
    for row in sheet:
        for cell in row:
            cell.font = Font(name="Aptos", size=11, color=INK)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.row == 1:
                cell.fill = PatternFill("solid", fgColor=BLUE)
                cell.font = Font(name="Aptos", size=11, bold=True, color="FFFFFF")
            elif cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=PALE)
            if "date" in str(sheet.cell(1, cell.column).value).lower():
                cell.number_format = "yyyy-mm-dd hh:mm"
        if row[0].row > 1:
            lines = max(
                (
                    sum(
                        max(
                            1,
                            math.ceil(
                                len(line) / max(25, sheet.column_dimensions[c.column_letter].width)
                            ),
                        )
                        for line in str(c.value or "").splitlines()
                    )
                    for c in row
                ),
                default=1,
            )
            sheet.row_dimensions[row[0].row].height = min(300, max(44, lines * 14 + 10))
    if sheet.max_row >= 2:
        name = "T_" + re.sub(r"[^A-Za-z0-9_]", "_", sheet.title)
        table = Table(displayName=name, ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)
    sheet.auto_filter.ref = sheet.dimensions
    for header in ("Criticality", "Qualified state", "Resolution state"):
        columns = [c.column for c in sheet[1] if c.value == header]
        for col in columns:
            for row in range(2, sheet.max_row + 1):
                cell = sheet.cell(row, col)
                text = str(cell.value).upper()
                color = (
                    "FBE4E2"
                    if any(w in text for w in ("CRITICAL", "HIGH", "OPEN"))
                    else "FFF0CC"
                    if any(w in text for w in ("MEDIUM", "PENDING", "CONDITIONAL"))
                    else None
                )
                if color:
                    cell.fill = PatternFill("solid", fgColor=color)


def link(cell, sheet: str, row: int, column: str = "A") -> None:
    from openpyxl.worksheet.hyperlink import Hyperlink

    cell.hyperlink = Hyperlink(
        ref=cell.coordinate, location=f"'{sheet.replace(chr(39), chr(39) * 2)}'!{column}{row}"
    )
    cell.font = Font(name="Aptos", size=11, color="145D7A", underline="single")
