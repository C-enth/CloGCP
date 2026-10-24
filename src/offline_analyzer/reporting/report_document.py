"""Assessment report document: Executive summary, Technical findings, Attack paths.

One template for the whole run or a single account/project, in Word (RTF) and Excel.
Restrained formatting: grey table rules, a light header band, bold labels, page numbers.
All text is escaped; spreadsheet text is written as literals, never formulas.
"""

from collections import Counter
from pathlib import Path

from offline_analyzer.ingestion.safety import excel_literal

SEVERITIES = ("CRITICAL", "HIGH", "MEDIUM", "LOW", "INFO")
CONFIRMED = "Validated by analyst"
NOT_COLLECTED = "not collected (seen through organization-wide IAM search)"
PATH_LIMIT = 100
BASIS = (
    "Findings are based on read-only configuration evidence collected by the CloSec "
    "collectors. Exploitation was not performed; misuse scenarios describe what a "
    "configuration would allow if it is not intended. Business impact depends on asset "
    "criticality, which was not provided."
)


def _rtf(text) -> str:
    from offline_analyzer.reporting.approved import rtf_escape

    return rtf_escape(str(text if text is not None else ""))


def summary(data: dict) -> dict:
    """Shared executive-summary facts for Word and Excel."""
    from offline_analyzer.reporting.assessment_report import SCOPE_NOTICE

    findings = data["findings"]
    totals = Counter(f["severity"] for f in findings)
    per_scope: dict[tuple, Counter] = {}
    for f in findings:
        per_scope.setdefault((f["provider"], f["scope"], f["scope_name"]), Counter())[
            f["severity"]
        ] += 1
    collected = {(s["provider"], s["scope"]) for s in data["scopes"] if s["sources"]}
    for s in data["scopes"]:
        if s["sources"]:
            per_scope.setdefault((s["provider"], s["scope"], s["name"]), Counter())
    accounts = [
        {
            "provider": provider,
            "scope": scope,
            "name": name
            if (provider, scope) in collected
            else (f"{name} — {NOT_COLLECTED}" if name else NOT_COLLECTED),
            "collected": (provider, scope) in collected,
            "severity": counts,
            "total": sum(counts.values()),
        }
        for (provider, scope, name), counts in per_scope.items()
    ]
    accounts.sort(key=lambda a: (not a["collected"], a["provider"], a["scope"]))
    feasible = [p for p in data["paths"] if p["feasibility"] == "FEASIBLE"]
    proof = Counter(p["proof_status"] for p in data["paths"] if "proof_status" in p)
    confirmed = sum(f["status"] == CONFIRMED for f in findings)
    scope_names = [
        a["scope"] + (f" ({a['name']})" if a["name"] else "") for a in accounts if a["collected"]
    ]
    overview = (
        f"This report covers {len(scope_names) or len(accounts)} assessed account(s)/project(s). "
        f"{len(findings)} findings are reported: "
        + ", ".join(f"{totals.get(s, 0)} {s.lower()}" for s in SEVERITIES)
        + f". {confirmed} have been confirmed by the assessor; the others are pending review. "
        f"{len(data['paths'])} attack paths were evaluated, of which {len(feasible)} have every "
        "evidence-checkable step observed"
        + (
            " (hop proof status: "
            + ", ".join(
                f"{proof.get(k, 0)} {k.lower()}" for k in ("PROVEN", "HYPOTHESIS", "BLOCKED")
            )
            + ")"
            if proof
            else ""
        )
        + "."
    )
    key = findings[:10]
    return {
        "title": f"{data.get('report_title', 'Cloud Security Assessment Report')}: {_title(data)}",
        "control": [
            ("Cloud scope", provider_scope(data)),
            ("Scope (collected)", ", ".join(scope_names) or "See findings by account"),
            ("Assessment run", data["run_id"]),
            ("Generated (UTC)", data["generated_at"]),
            ("Findings included", data["include"]),
            *(
                [
                    (
                        "Grouped findings",
                        f"{len(data['groups'])} grouped findings / "
                        f"{sum(g['affected_resources'] for g in data['groups'])} affected "
                        f"resources ({len(findings)} findings counted in totals)",
                    )
                ]
                if data.get("grouped") and data.get("groups") is not None
                else []
            ),
            *(
                [("Report scope notice", SCOPE_NOTICE)]
                if data.get("include_key", "confirmed") != "confirmed"
                else []
            ),
            # Editions frozen before exclusions existed carry no count; render them unchanged.
            *(
                [("Excluded by the assessor", f"{data['excluded']} finding(s) not reported")]
                if "excluded" in data
                else []
            ),
            ("Evidence basis", "Read-only configuration evidence; exploitation not performed"),
        ],
        "overview": overview,
        "totals": totals,
        "accounts": accounts,
        "key": key,
        "feasible": feasible,
        "proof": proof,
    }


def _title(data: dict) -> str:
    from offline_analyzer.reporting.assessment_report import scope_title

    return scope_title(data)


def _finding_rows(f: dict) -> list[tuple[str, str]]:
    return [
        ("Finding ID", f["id"]),
        ("Severity", f"{f['severity']} — {f['rationale']}"),
        (
            "Account / project",
            f"{f['provider']} {f['scope']}" + (f" ({f['scope_name']})" if f["scope_name"] else ""),
        ),
        ("Region / resource", f"{f['region']} / {f['resource']}"),
        ("Control", f"{f['control']} {f['checklist_title']}".strip()),
        ("Status", f"{f['status']}; {f['stage']}"),
        ("Why this is flagged", f["why"]),
        ("How it can be misused", f["misuse"]),
        ("Impact", f["impact"]),
        ("Recommendation", f["recommendation"]),
        *((("Immediate mitigation", f["immediate"]),) if f.get("immediate") else ()),
        ("Remediation steps", f["remediation"]),
        ("How to validate", f["validation_procedure"] or "Recollect and confirm the secure state."),
        ("Evidence", f["evidence"]),
        ("STRIDE mapping", f.get("stride") or "Unmapped"),
        ("Related attack paths", f.get("methodology_paths") or "No mapped path"),
        ("Analyst confirmed path interpretations", f.get("evaluated_paths") or "None confirmed"),
        (
            "Report approval",
            f.get("report_approval", "Generated technical report; analyst review required"),
        ),
        *(
            (("Confirmed by", f["analyst"]),)
            if f["status"] == CONFIRMED and f.get("analyst")
            else ()
        ),
        *((("SRE response", f["sre"]),) if f.get("sre") else ()),
    ]


def _path_rows(p: dict) -> list[tuple[str, str]]:
    return [
        ("Path", p["diagram"]),
        (
            "Account / project",
            f"{p['provider']} {p.get('project_id') or p.get('account_id') or ''}",
        ),
        ("Scope", p.get("scope", "single resource")),
        ("Assessment", p["feasibility_reason"]),
        (
            "Unresolved prerequisites",
            ", ".join(p.get("unresolved") or p.get("missing_evidence") or []) or "None",
        ),
        ("Linked findings", ", ".join(p.get("finding_ids", [])) or "None"),
        ("STRIDE categories", ", ".join(p.get("stride", [])) or "Unmapped"),
        ("STRIDE scenarios", ", ".join(p.get("stride_scenarios", [])) or "Unmapped"),
        ("Mapping basis", p.get("stride_mapping_basis", "Unmapped")),
        ("Analyst path interpretation", p.get("analyst_review", "UNREVIEWED")),
    ]


def _ordered_paths(paths: list[dict]) -> list[dict]:
    order = {"FEASIBLE": 0, "INSUFFICIENT_EVIDENCE": 1}
    proof = {"PROVEN": 0, "HYPOTHESIS": 1, "BLOCKED": 2}
    return sorted(
        paths,
        key=lambda p: (order.get(p["feasibility"], 2), proof.get(p.get("proof_status", ""), 3)),
    )


# ---------------------------------------------------------------- shared template sections
STRIDE = (
    "Spoofing",
    "Tampering",
    "Repudiation",
    "Information Disclosure",
    "Denial of Service",
    "Elevation of Privilege",
)
EDGE_MEANING = {
    "network-reachable": "traffic can reach the endpoint",
    "can-invoke-anonymously": "unauthenticated callers are authorized to invoke",
    "can-read-anonymously": "unauthenticated callers can read",
    "can-impersonate": "holder can act as or mint tokens for a service account",
    "authenticates-as": "credential or workload authenticates as an identity",
    "holds-role": "identity is a member of a role binding",
    "can-access": "binding or credential grants a role on a resource",
    "can-call-api": "API key can call services in a scope",
    "reads-credentials": "workload can read credentials (instance metadata)",
    "can-assume": "principal can assume a role through its trust policy",
    "federates-as": "external identity federates into a cloud identity",
    "can-escalate": "permissions enable privilege escalation",
    "not-blocked-by": "no guardrail (SCP, boundary) is evidenced to stop it",
    "stores": "resource holds data of unknown classification",
}
GLOSSARY = (
    (
        "PROVEN (hop)",
        "The configuration that creates the edge is in collected evidence and is "
        "cited (observation or source locator). Not proof of exploitation.",
    ),
    (
        "HYPOTHESIS (hop)",
        "Plausible, but evidence is missing or UNKNOWN; the missing evidence is named.",
    ),
    ("BLOCKED (hop)", "Collected evidence shows the edge does not exist."),
    (
        "Proof status (path)",
        "BLOCKED if any hop is blocked, PROVEN only if every hop is proven, otherwise HYPOTHESIS.",
    ),
    (
        "FEASIBLE",
        "Every evidence-checkable prerequisite is verified; attacker-side steps are "
        "never verified.",
    ),
    ("INSUFFICIENT_EVIDENCE", "At least one evidence-checkable prerequisite is unverified."),
    (
        "INSUFFICIENT EVIDENCE (control)",
        "UNKNOWN or collection error: evidence could not "
        "support an outcome. Never reported as passing.",
    ),
    ("Direction justification", "Why access flows from the hop's source to its target."),
    (
        "Assumption",
        "Attacker-side prerequisite (credential possession, workload compromise, "
        "exploitation); listed, never an edge, never verified.",
    ),
)
LIMITATIONS = (
    "Read-only configuration evidence at collection time; no exploitation, scanning or "
    "runtime testing was performed.",
    "UNKNOWN areas are listed below; missing evidence never produces a passing outcome.",
    "Attack paths restate observed configuration; effective authorization depends on context "
    "named as missing evidence on each hop.",
    "STRIDE mappings are proposed engineering mappings, not verified threat outcomes.",
    "Business impact and asset criticality were not provided.",
)


def provider_scope(data: dict) -> str:
    providers = {s["provider"] for s in data.get("scopes", []) if s.get("sources")}
    return (
        data.get("provider_scope")
        or " + ".join(p for p in ("GCP", "AWS") if p in providers)
        or "Not stated"
    )


def _path_name(p: dict) -> str:
    return f"{p.get('attack_path') or ''} {p.get('title', '')}".strip()


def _scope_of(p: dict) -> str:
    return f"{p['provider']} {p.get('project_id') or p.get('account_id') or ''}".strip()


def hop_evidence(h: dict) -> str:
    parts = []
    if h.get("rule_ids"):
        parts.append("Rules: " + ", ".join(h["rule_ids"]))
    if h.get("finding_ids"):
        parts.append("Findings: " + ", ".join(h["finding_ids"][:5]))
    if h.get("observation_ids"):
        extra = len(h["observation_ids"]) - 5
        parts.append(
            "Observations: "
            + ", ".join(h["observation_ids"][:5])
            + (f" (+{extra})" if extra > 0 else "")
        )
    if h.get("evidence_refs"):
        parts.append("Locators: " + "; ".join(r.get("locator", "") for r in h["evidence_refs"][:5]))
    return "; ".join(parts) or "No source record cited"


HOP_HEADERS = [
    "Path",
    "Title",
    "Account / project",
    "Path proof status",
    "Feasibility",
    "Hop",
    "From",
    "Edge type",
    "To",
    "Hop status",
    "STRIDE",
    *STRIDE,
    "Direction",
    "Direction justification",
    "Missing evidence",
    "Evidence",
]


def hop_rows(paths: list[dict]) -> list[list]:
    """One row per hop, with one STRIDE column per category (1 when mapped)."""
    return [
        [
            p.get("attack_path") or "GCP investigation",
            p.get("title", ""),
            _scope_of(p),
            p.get("proof_status", ""),
            p["feasibility"],
            h["index"],
            h["source"],
            h["edge_type"],
            h["target"],
            h["status"],
            ", ".join(h["stride"]),
            *(1 if s in h["stride"] else 0 for s in STRIDE),
            h["direction"],
            h["justification"],
            "; ".join(h.get("missing_evidence", [])),
            hop_evidence(h),
        ]
        for p in _ordered_paths(paths)
        for h in p.get("hops") or []
    ]


def stride_surface(paths: list[dict]) -> list[dict]:
    """Per STRIDE category: paths, hop counts by status, distinct entry points."""
    rows = []
    for category in STRIDE:
        hops = [(p, h) for p in paths for h in p.get("hops") or [] if category in h["stride"]]
        status = Counter(h["status"] for _p, h in hops)
        rows.append(
            {
                "category": category,
                "paths": len({p["path_id"] for p, _h in hops}),
                "proven": status["PROVEN"],
                "hypothesis": status["HYPOTHESIS"],
                "blocked": status["BLOCKED"],
                "entries": sorted({p.get("entry_point", "") for p, _h in hops} - {""}),
            }
        )
    return rows


def stride_matrix(paths: list[dict]) -> list[list]:
    """One row per path: hops mapped to each STRIDE category."""
    return [
        [
            p.get("attack_path") or "GCP investigation",
            p.get("title", ""),
            _scope_of(p),
            p.get("proof_status", ""),
            len(p.get("hops") or []),
            *(sum(s in h["stride"] for h in p.get("hops") or []) for s in STRIDE),
        ]
        for p in _ordered_paths(paths)
    ]


def path_detail_rows(p: dict) -> list[tuple[str, str]]:
    """Path summary for the report: proof, entry/objective/impact, assumptions, gaps."""
    rows = []
    if "proof_status" in p:
        rows += [
            ("Proof status", p["proof_status"]),
            ("Entry point", p.get("entry_point", "")),
            ("Objective", p.get("objective", "")),
            ("Impact", p.get("impact", "") or "Not stated"),
            ("Attacker-side assumptions", "; ".join(p.get("assumptions", [])) or "None"),
            (
                "Missing evidence",
                "; ".join(p.get("missing_evidence_named", [])) or "None (every hop proven)",
            ),
        ]
    return rows + _path_rows(p)


def unknown_rows(data: dict) -> list[list]:
    return [
        [
            f"{u['provider']} {u['scope']}".strip(),
            u.get("name", ""),
            u["rule"],
            u["insufficient"],
            u["manual"],
        ]
        for u in data.get("unknown_areas", [])
    ]


def coverage_rows(data: dict) -> list[list]:
    labels = ("FAIL", "PASS", "NOT APPLICABLE", "INSUFFICIENT EVIDENCE", "MANUAL VALIDATION")
    return [
        [
            f"{s['provider']} {s['scope']}".strip(),
            s.get("name", ""),
            s.get("sources", 0),
            s.get("gaps", 0),
            *(Counter(s.get("outcomes", {})).get(label, 0) for label in labels),
        ]
        for s in data.get("scopes", [])
    ]


COVERAGE_HEADERS = [
    "Account / project",
    "Name",
    "Collected sources",
    "Collection gaps (error / not collected)",
    "FAIL",
    "PASS",
    "NOT APPLICABLE",
    "INSUFFICIENT EVIDENCE",
    "MANUAL VALIDATION",
]


NOTE_HEADERS = ["Account / project", "Name", "Code", "Reason", "Datasets"]


def note_rows(data: dict) -> list[list]:
    """Collection gaps and limitations (rate limits, expired sign-in, detail limits, IAM
    target-list fallbacks); never a security result."""
    return [
        [
            f"{n['provider']} {n['scope']}".strip(),
            n.get("name", ""),
            n["code"],
            n["reason"],
            n["datasets"],
        ]
        for n in data.get("coverage_notes", [])
    ]


def custom_lines(data: dict) -> list[str]:
    """Analyst scope-limitation text (after the fixed constants), one line per paragraph."""
    text = data.get("custom_limitations") or ""
    return [
        f"Assessor scope limitation: {line.strip()}" for line in text.splitlines() if line.strip()
    ]


def excluded_rows(data: dict, ids: bool = False) -> list[list]:
    """Excluded findings grouped by rule and reason (appendix), never in finding sections."""
    return [
        [g["rule"], g["reason"], g["detail"] or "", g["count"]]
        + (["; ".join(g["finding_ids"])] if ids else [])
        for g in data.get("excluded_findings") or []
    ]


def checklist_rows(data: dict) -> list[list]:
    return [
        [
            row["control_id"],
            row.get("checklist_title", ""),
            row.get("aws", {}).get("outcome", ""),
            row.get("gcp", {}).get("outcome", ""),
        ]
        for row in data.get("checklist", [])
    ]


# ---------------------------------------------------------------- Word (RTF)
TEXT_WIDTH = 9360  # US Letter with 1-inch margins, in twips.
BORDER = "\\brdrs\\brdrw6\\brdrcf1"


def _table(rows: list[list], widths: list[int], header: bool = True) -> str:
    out = []
    for index, row in enumerate(rows):
        head = header and index == 0
        out.append("\\trowd\\trgaph90\\trleft0" + ("\\trhdr" if head else ""))
        position = 0
        for width in widths:
            position += width
            out.append(
                ("\\clcbpat2" if head else "")
                + f"\\clbrdrt{BORDER}\\clbrdrl{BORDER}\\clbrdrb{BORDER}\\clbrdrr{BORDER}"
                + f"\\clvertalt\\cellx{position}"
            )
        cells = "".join(
            "\\pard\\intbl\\sa40 "
            + ("{\\b " if head else "")
            + _rtf(v)
            + ("}" if head else "")
            + "\\cell "
            for v in row
        )
        out.append(cells + "\\row\n")
    return "".join(out) + "\\pard\\sa120\\par\n"


def _labelled(rows: list[tuple[str, str]]) -> str:
    out = []
    for label, value in rows:
        out.append("\\trowd\\trgaph90\\trleft0")
        out.append(
            f"\\clcbpat2\\clbrdrt{BORDER}\\clbrdrl{BORDER}\\clbrdrb{BORDER}\\clbrdrr{BORDER}\\clvertalt\\cellx2500"
        )
        out.append(
            f"\\clbrdrt{BORDER}\\clbrdrl{BORDER}\\clbrdrb{BORDER}\\clbrdrr{BORDER}\\clvertalt\\cellx{TEXT_WIDTH}"
        )
        out.append(
            "\\pard\\intbl\\sa40 {\\b " + _rtf(label) + "}\\cell "
            "\\pard\\intbl\\sa40 " + _rtf(value) + "\\cell \\row\n"
        )
    return "".join(out) + "\\pard\\sa160\\par\n"


def _heading(text: str, size: int = 28, before: int = 360) -> str:
    return f"{{\\pard\\sb{before}\\sa120\\keepn\\b\\fs{size} " + _rtf(text) + "\\par}\n"


def _para(text: str) -> str:
    return "{\\pard\\sa120 " + _rtf(text) + "\\par}\n"


def write_rtf(data: dict, path: Path) -> None:
    s = summary(data)
    out = [
        "{\\rtf1\\ansi\\deff0\\uc1{\\fonttbl{\\f0 Calibri;}}"
        "{\\colortbl;\\red191\\green191\\blue191;\\red234\\green236\\blue239;}"
        "\\paperw12240\\paperh15840\\margl1440\\margr1440\\margt1440\\margb1440\\f0\\fs21\n",
        "{\\footer\\pard\\qc\\fs16 " + _rtf("CloSec assessment report — page ") + "\\chpgn\\par}\n",
        "{\\pard\\sb2400\\sa240\\b\\fs40 " + _rtf(s["title"]) + "\\par}\n",
        _labelled(s["control"]),
        "\\page\n",
        _heading("1. Executive summary", 32, 0),
        _para(s["overview"]),
        _heading("Findings by severity", 24, 200),
        _table(
            [["Severity", "Findings"]]
            + [[x.title(), s["totals"].get(x, 0)] for x in SEVERITIES]
            + [["Total", sum(s["totals"].values())]],
            [4680, 4680],
        ),
        _heading("Findings by account / project", 24, 200),
        _table(
            [["Account / project", "Name", *(x.title() for x in SEVERITIES), "Total"]]
            + [
                [
                    f"{a['provider']} {a['scope']}",
                    a["name"],
                    *(a["severity"].get(x, 0) for x in SEVERITIES),
                    a["total"],
                ]
                for a in s["accounts"]
            ],
            [2100, 2460, 800, 800, 800, 800, 800, 800],
        ),
        _heading("Key risks", 24, 200),
        _table(
            [["#", "Severity", "Finding", "Account / project"]]
            + [
                [i, f["severity"], f["title"], f"{f['provider']} {f['scope']}"]
                for i, f in enumerate(s["key"], 1)
            ]
            if s["key"]
            else [
                ["#", "Severity", "Finding", "Account / project"],
                ["-", "-", "No findings", "-"],
            ],
            [500, 1300, 4960, 2600],
        ),
        _heading("Basis and limitations", 24, 200),
        _para(BASIS),
        *[_para(text) for text in (*LIMITATIONS, *custom_lines(data))],
        "\\page\n",
        _heading("2. Technical findings", 32, 0),
    ]
    if not data["findings"]:
        out.append(_para("No findings are included under the selected marking."))
    for severity in SEVERITIES:
        group = [f for f in data["findings"] if f["severity"] == severity]
        if not group:
            continue
        out.append(_heading(f"{severity.title()} findings ({len(group)})", 26, 240))
        for f in group:
            out.append(_heading(f"{f['title']}", 22, 200))
            out.append(_labelled(_finding_rows(f)))
    out += ["\\page\n", _heading("3. Attack paths", 32, 0)]
    out.append(
        _para(
            "Each path reads left to right: the starting point, each resource or prerequisite, "
            "and what allows the step. FEASIBLE means every evidence-checkable step is observed; "
            "INSUFFICIENT_EVIDENCE lists the prerequisites still unverified. Paths describe "
            "configuration prerequisites only; exploitation was not performed."
        )
    )
    paths = _ordered_paths(data["paths"])
    if not paths:
        out.append(_para("No attack paths were evaluated for this scope."))
    for p in paths[:PATH_LIMIT]:
        out.append(
            _heading(
                f"[{p['feasibility']}] {p.get('attack_path') or ''} {p.get('title', '')}".strip(),
                22,
                200,
            )
        )
        out.append(_labelled(_path_rows(p)))
    if len(paths) > PATH_LIMIT:
        out.append(
            _para(f"{len(paths) - PATH_LIMIT} further paths are listed in the Excel workbook.")
        )
    out.append(_heading("Appendix: Excluded findings"))
    out.append(_para(EXCLUDED_NOTE))
    out.append(
        _table(
            [
                ["Rule", "Reason", "Detail", "Count"],
                *(excluded_rows(data) or [["None", "", "", 0]]),
            ],
            [2600, 2800, 3200, 1000],
        )
    )
    out.append("}")
    path.write_text("".join(out), encoding="ascii")


# ---------------------------------------------------------------- Excel
EXCLUDED_NOTE = (
    "Findings the assessor excluded from this report (for example false positives), grouped "
    "by rule and reason. They are not counted in severity totals or finding sections; the "
    "Excel workbook lists their finding IDs."
)
EXCEL_SHEETS = (
    "Cover",
    "Executive Summary",
    "Coverage and Limitations",
    "Attack Surface",
    "STRIDE Matrix",
    "Attack Paths",
    "Attack Path Hops",
    "Technical Findings",
    "Grouped Findings",
    "Findings Validation",
    "Edge STRIDE Mapping",
    "Appendix",
    "Excluded Findings",
)


def write_excel(data: dict, path: Path, *, gcp_register: dict | None = None) -> None:
    """Cover, summary, coverage, attack surface, STRIDE matrix, paths, one row per hop,
    findings and appendix. Tables have frozen headers and filters; text is literal."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook(write_only=True)
    s = summary(data)
    line = Side(style="thin", color="BFBFBF")
    band = PatternFill("solid", fgColor="EAECEF")
    wrap = Alignment(wrap_text=True, vertical="top")

    def cell(ws, value, *, bold=False, head=False, size=None):
        c = WriteOnlyCell(ws, value=excel_literal(value) if isinstance(value, str) else value)
        if bold or head or size:
            c.font = Font(bold=bold or head, size=size)
        if head:
            c.fill = band
            c.border = Border(bottom=line)
        c.alignment = wrap
        return c

    def sheet_of_lines(name: str, widths: tuple[int, ...]):
        ws = wb.create_sheet(name)
        for column, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(column)].width = width
        return ws

    def table(name: str, headers: list[str], rows: list[list], widths: dict[str, int]):
        sheet = wb.create_sheet(name)
        for column, header in enumerate(headers, 1):
            sheet.column_dimensions[get_column_letter(column)].width = widths.get(header, 16)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(len(rows), 1) + 1}"
        sheet.append([cell(sheet, h, head=True) for h in headers])
        for row in rows:
            sheet.append([cell(sheet, v) for v in row])
        if not rows:
            sheet.append([cell(sheet, "No rows for this report scope.")])

    # Cover
    ws = sheet_of_lines("Cover", (34, 90))
    ws.append([cell(ws, s["title"], bold=True, size=16)])
    ws.append([cell(ws, f"Cloud scope: {provider_scope(data)}", bold=True)])
    ws.append([])
    for label, value in s["control"]:
        ws.append([cell(ws, label, bold=True), cell(ws, value)])
    ws.append([])
    ws.append([cell(ws, "Qualification", bold=True), cell(ws, BASIS)])

    # Executive summary
    ws = sheet_of_lines("Executive Summary", (34, 40, 12, 12, 12, 12, 12, 12))
    ws.append([cell(ws, "1. Executive summary", bold=True, size=12)])
    ws.append([cell(ws, s["overview"])])
    ws.append([])
    ws.append([cell(ws, "Findings by severity", bold=True)])
    ws.append([cell(ws, "Severity", head=True), cell(ws, "Findings", head=True)])
    for x in SEVERITIES:
        ws.append([cell(ws, x.title()), s["totals"].get(x, 0)])
    ws.append([cell(ws, "Total", bold=True), sum(s["totals"].values())])
    ws.append([])
    ws.append([cell(ws, "Attack paths by proof status", bold=True)])
    ws.append([cell(ws, "Proof status", head=True), cell(ws, "Paths", head=True)])
    for x in ("PROVEN", "HYPOTHESIS", "BLOCKED"):
        ws.append([cell(ws, x), s["proof"].get(x, 0)])
    ws.append([])
    ws.append([cell(ws, "Findings by account / project", bold=True)])
    ws.append(
        [
            cell(ws, h, head=True)
            for h in ("Account / project", "Name", *(x.title() for x in SEVERITIES), "Total")
        ]
    )
    for a in s["accounts"]:
        ws.append(
            [
                cell(ws, f"{a['provider']} {a['scope']}"),
                cell(ws, a["name"]),
                *(a["severity"].get(x, 0) for x in SEVERITIES),
                a["total"],
            ]
        )
    ws.append([])
    ws.append([cell(ws, "Key risks", bold=True)])
    ws.append([cell(ws, h, head=True) for h in ("Finding", "Account / project", "Severity")])
    for f in s["key"]:
        ws.append(
            [
                cell(ws, f["title"]),
                cell(ws, f"{f['provider']} {f['scope']}"),
                cell(ws, f["severity"]),
            ]
        )

    # Coverage and limitations (scope table, UNKNOWN areas, path coverage, limitations)
    ws = sheet_of_lines("Coverage and Limitations", (30, 28, 34, 16, 16, 16, 16, 16, 16))
    ws.append([cell(ws, "2. Coverage and limitations", bold=True, size=12)])
    ws.append([cell(ws, "Evidence coverage by account / project", bold=True)])
    ws.append([cell(ws, h, head=True) for h in COVERAGE_HEADERS])
    for row in coverage_rows(data):
        ws.append([cell(ws, v) for v in row])
    ws.append([])
    ws.append([cell(ws, "UNKNOWN and manual areas (never reported as passing)", bold=True)])
    ws.append(
        [
            cell(ws, h, head=True)
            for h in ("Account / project", "Name", "Rule", "Insufficient evidence", "Manual")
        ]
    )
    for row in unknown_rows(data) or [["None", "", "", 0, 0]]:
        ws.append([cell(ws, v) for v in row])
    ws.append([])
    ws.append([cell(ws, "Collection gaps and limitations (reasons)", bold=True)])
    ws.append([cell(ws, h, head=True) for h in NOTE_HEADERS])
    for row in note_rows(data) or [["None", "", "", "", 0]]:
        ws.append([cell(ws, v) for v in row])
    ws.append([])
    ws.append(
        [cell(ws, "Attack-path coverage (why path kinds were or were not formed)", bold=True)]
    )
    ws.append([cell(ws, h, head=True) for h in ("Provider", "Path kind", "Formed", "Reason")])
    for r in data.get("path_coverage", []):
        ws.append(
            [cell(ws, r["provider"]), cell(ws, r["kind"]), r["formed"], cell(ws, r["reason"])]
        )
    ws.append([])
    ws.append([cell(ws, "Limitations", bold=True)])
    for text in (*LIMITATIONS, *custom_lines(data)):
        ws.append([cell(ws, text)])

    paths = _ordered_paths(data["paths"])
    surface = stride_surface(paths)
    table(
        "Attack Surface",
        [
            "STRIDE category",
            "Paths",
            "Proven hops",
            "Hypothesis hops",
            "Blocked hops",
            "Entry points",
        ],
        [
            [
                r["category"],
                r["paths"],
                r["proven"],
                r["hypothesis"],
                r["blocked"],
                "; ".join(r["entries"])
                or ("No edge type maps to Repudiation" if r["category"] == "Repudiation" else ""),
            ]
            for r in surface
        ],
        {"STRIDE category": 26, "Entry points": 90},
    )
    table(
        "STRIDE Matrix",
        ["Path", "Title", "Account / project", "Proof status", "Hops", *STRIDE],
        stride_matrix(paths),
        {"Title": 40, "Account / project": 26},
    )
    table(
        "Attack Paths",
        [
            "Feasibility",
            "Proof status",
            "Path ID",
            "Title",
            "Account / project",
            "Scope",
            "Entry point",
            "Objective",
            "Impact",
            "Hops",
            "Path (A -> B -> C)",
            "Assessment",
            "Attacker-side assumptions",
            "Missing evidence",
            "Unresolved prerequisites",
            "Linked findings",
            "STRIDE (hops)",
            "STRIDE (linked findings)",
            "STRIDE scenarios",
            "Mapping basis",
            "Analyst path interpretation",
        ],
        [
            [
                p["feasibility"],
                p.get("proof_status", "Not modelled (edition before hop model)"),
                p.get("attack_path") or "",
                p.get("title", ""),
                _scope_of(p),
                p.get("scope", "single resource"),
                p.get("entry_point", ""),
                p.get("objective", ""),
                p.get("impact", ""),
                len(p.get("hops") or []),
                p["diagram"],
                p["feasibility_reason"],
                "; ".join(p.get("assumptions", [])),
                "; ".join(p.get("missing_evidence_named", [])),
                ", ".join(p.get("unresolved") or p.get("missing_evidence") or []),
                ", ".join(p.get("finding_ids", [])),
                ", ".join(p.get("hop_stride", [])),
                ", ".join(p.get("stride", [])),
                ", ".join(p.get("stride_scenarios", [])),
                p.get("stride_mapping_basis", "Unmapped"),
                p.get("analyst_review", "UNREVIEWED"),
            ]
            for p in paths
        ],
        {
            "Title": 36,
            "Account / project": 26,
            "Entry point": 40,
            "Objective": 36,
            "Impact": 40,
            "Path (A -> B -> C)": 70,
            "Assessment": 50,
            "Attacker-side assumptions": 40,
            "Missing evidence": 50,
            "Unresolved prerequisites": 40,
            "Linked findings": 30,
        },
    )
    table(
        "Attack Path Hops",
        HOP_HEADERS,
        hop_rows(paths),
        {
            "Title": 32,
            "Account / project": 24,
            "From": 40,
            "To": 40,
            "Edge type": 22,
            "STRIDE": 30,
            "Direction": 50,
            "Direction justification": 70,
            "Missing evidence": 50,
            "Evidence": 60,
            **{x: 9 for x in STRIDE},
        },
    )
    finding_headers = [
        "Finding ID",
        "Severity",
        "Title",
        "Account / project",
        "Name",
        "Region",
        "Resource",
        "Control",
        "Why this is flagged",
        "How it can be misused",
        "Impact",
        "Severity rationale",
        "Recommendation",
        "Remediation steps",
        "How to validate",
        "Evidence",
        "Status",
        "Stage",
        "Confirmed by",
        "SRE response",
        "STRIDE",
        "Related attack paths",
        "Report approval",
    ]
    table(
        "Technical Findings",
        finding_headers,
        [
            [
                f["id"],
                f["severity"],
                f["title"],
                f"{f['provider']} {f['scope']}",
                f["scope_name"],
                f["region"],
                f["resource"],
                f"{f['control']} {f['checklist_title']}".strip(),
                f["why"],
                f["misuse"],
                f["impact"],
                f["rationale"],
                f["recommendation"],
                f["remediation"],
                f["validation_procedure"],
                f["evidence"],
                f["status"],
                f["stage"],
                f.get("analyst", "") if f["status"] == CONFIRMED else "",
                f.get("sre", ""),
                f.get("stride", ""),
                f.get("methodology_paths", ""),
                f.get("report_approval", "Generated technical report"),
            ]
            for f in data["findings"]
        ],
        {
            "Finding ID": 22,
            "Title": 40,
            "Account / project": 26,
            "Resource": 40,
            "Control": 30,
            "Why this is flagged": 55,
            "How it can be misused": 55,
            "Impact": 50,
            "Severity rationale": 40,
            "Recommendation": 50,
            "Remediation steps": 60,
            "How to validate": 50,
            "Evidence": 45,
            "Status": 26,
            "Stage": 26,
        },
    )
    table(
        "Grouped Findings",
        [
            "Group",
            "Rule",
            "Root-cause key",
            "Severity (highest member)",
            "Findings",
            "Affected resources",
            "Member finding IDs",
            "Accounts / projects",
        ],
        [
            [
                g["title"],
                g["rule"],
                g["key"],
                g["severity"],
                len(g["members"]),
                g["affected_resources"],
                "; ".join(m["id"] for m in g["members"]),
                "; ".join(sorted({f"{m['provider']} {m['scope']}" for m in g["members"]})),
            ]
            for g in (data.get("groups") or [])
        ]
        or [["None", "", "", "", 0, 0, "", ""]],
        {"Group": 40, "Root-cause key": 50, "Member finding IDs": 90, "Accounts / projects": 40},
    )
    table(
        "Findings Validation",
        [
            "Finding ID",
            "Provider",
            "Account / project",
            "Title",
            "Validation status",
            "Analyst confirmed path interpretation",
            "STRIDE",
            "Evidence references",
        ],
        [
            [
                f["id"],
                f["provider"],
                f["scope"],
                f["title"],
                f["status"],
                f.get("evaluated_paths", "") if f["status"] == CONFIRMED else "",
                f.get("stride", ""),
                f["evidence"],
            ]
            for f in data["findings"]
        ],
        {
            "Title": 45,
            "Account / project": 30,
            "Validation status": 28,
            "Analyst confirmed path interpretation": 80,
            "STRIDE": 36,
            "Evidence references": 50,
        },
    )
    from offline_analyzer.attack_paths.model import EDGE_STRIDE

    table(
        "Edge STRIDE Mapping",
        ["Edge type", "Meaning (direction source -> target)", "STRIDE", *STRIDE],
        [
            [edge, EDGE_MEANING.get(edge, ""), ", ".join(cats), *(int(x in cats) for x in STRIDE)]
            for edge, cats in EDGE_STRIDE.items()
        ],
        {"Edge type": 24, "Meaning (direction source -> target)": 55, "STRIDE": 45},
    )
    ws = sheet_of_lines("Appendix", (34, 100))
    ws.append([cell(ws, "Appendix", bold=True, size=12)])
    ws.append([cell(ws, "Evidence basis", bold=True), cell(ws, BASIS)])
    ws.append([])
    ws.append([cell(ws, "Glossary", bold=True)])
    for term, meaning in GLOSSARY:
        ws.append([cell(ws, term, bold=True), cell(ws, meaning)])
    ws.append([])
    ws.append([cell(ws, "Control coverage", bold=True)])
    ws.append([cell(ws, h, head=True) for h in ("Control", "Title", "AWS", "GCP")])
    for row in checklist_rows(data):
        ws.append([cell(ws, v) for v in row])
    table(
        "Excluded Findings",
        ["Rule", "Reason", "Detail", "Count", "Finding IDs"],
        excluded_rows(data, ids=True) or [["None", "", "", 0, ""]],
        {"Rule": 26, "Reason": 30, "Detail": 50, "Finding IDs": 90},
    )
    if gcp_register:
        from offline_analyzer.reporting.gcp_register import add_sheets

        add_sheets(wb, gcp_register)
    wb.save(path)
