"""Read-only GCP Excel projection. No rules, relationships or review decisions are created.

The owner supplied the reference's forty headers on 2026-10-23. Layout whitespace is
normalized; no FortiGate records, hint rows or example values are part of this contract.
"""

import json
import math
import re
from collections import defaultdict
from typing import Any

from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo

from offline_analyzer.contracts.r0 import stable_id
from offline_analyzer.findings.detail import _masked
from offline_analyzer.findings.exclusions import excluded
from offline_analyzer.findings.explanations import explain
from offline_analyzer.findings.groups import group_key
from offline_analyzer.findings.recommendations import recommendation
from offline_analyzer.findings.sre import latest_sre, review_state
from offline_analyzer.ingestion.json_input import MAX_BYTES, InputError
from offline_analyzer.ingestion.safety import excel_literal, resolve_under
from offline_analyzer.provenance.execution import checkpoint
from offline_analyzer.reporting.coverage_reasons import reason_text, source_category

COLUMNS = (
    "Record ID (internal)",
    "Observation / gap",
    "Affected asset / scope",
    "Affected objects / services",
    "Criticality",
    "Why this criticality?",
    "Rating confidence (automatic)",
    "Possible business impact",
    "Recommended action",
    "Assessor record type",
    "What the evidence shows",
    "Evidence reference",
    "Expected control / requirement",
    "Evidence at assessment",
    "What the evidence does not establish",
    "Possible attack scenario",
    "Conditions needed for the scenario",
    "Source-reported result",
    "Domain / tool",
    "Control area",
    "Evidence type",
    "Evidence limits / data quality",
    "Assessor conclusion / evidence basis",
    "Related observations",
    "Attack-path evidence level",
    "Assessor evidence result (automatic)",
    "Report sequence (automatic)",
    "Common asset reference (hostname / ID)",
    "Issue group / summary title",
    "Risk rating decision",
    "Show in Overall?",
    "Shown in Overall (automatic)",
    "Record count (automatic)",
    "Applicability (SPOC)",
    "SPOC assessment",
    "SPOC justification / comments",
    "Supporting evidence / reference",
    "Reviewed by",
    "Review date",
    "Final review conclusion (automatic)",
)
SHEETS = ("GCP_Observation_Register", "GCP_Evidence_Index", "GCP_Evidence_Details")
# Analyst and SPOC columns stay blank for the reviewer; recorded tool decisions are kept here.
AUDIT_SHEET = "GCP_Review_Audit"
AUDIT_COLUMNS = (
    "Record ID",
    "Recorded analyst status",
    "Recorded analyst",
    "Analyst decision time",
    "Analyst note",
    "Recorded SRE verdict",
    "SRE reason",
    "SRE contact",
    "SRE recorded time",
)
CITED = re.compile(r"(roles/[A-Za-z0-9_.]+) -> ([^\s;,()]+)")
INDEX_COLUMNS = (
    "Evidence mapping ID",
    "Record ID",
    "GCP project",
    "Resource identifier",
    "Source operation",
    "Dataset",
    "Original evidence reference",
    "JSON path",
    "Collection time",
    "Artifact ID",
    "SHA-256",
    "Collection status",
    "Evidence availability",
    "Configuration excerpt",
)
DETAIL_COLUMNS = (
    "Evidence mapping ID",
    "Record ID",
    "GCP project",
    "Resource identifier",
    "Original evidence reference",
    "JSON path",
    "Excerpt part",
    "Redacted original configuration",
    "Evidence limitation",
    "Back to evidence index",
)
PENDING = "Pending — no decision recorded"
ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
CHUNK = 1200


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (list, tuple, set)):
        return "; ".join(str(v) for v in value if v is not None)
    return str(value)


def _pointer(document: Any, locator: str) -> Any:
    """RFC 6901, including the empty root pointer and literal escaped keys."""
    if locator == "":
        return document
    if not locator.startswith("/"):
        raise ValueError("unsupported evidence locator")
    node = document
    for raw in locator.split("/")[1:]:
        part = raw.replace("~1", "/").replace("~0", "~")
        node = node[int(part)] if isinstance(node, list) and part.isdigit() else node[part]
    return node


def _redact(value: Any, depth: int = 0) -> Any:
    """Use existing secret masking, also after decoding embedded policy JSON."""
    if depth > 40:
        return "[redacted: depth limit]"
    if isinstance(value, str):
        from urllib.parse import unquote

        decoded = unquote(value) if value[:3].lower() in {"%7b", "%5b"} else value
        if decoded.lstrip().startswith(("{", "[")):
            try:
                return _redact(json.loads(decoded), depth + 1)
            except (ValueError, RecursionError):
                pass
        return _masked(value)
    if isinstance(value, dict):
        masked = _masked(value)
        # Environment and metadata payloads can hold credentials under arbitrary names.
        for key in masked:
            if str(key).lower() in {"env", "environment", "environmentvariables", "userdata"}:
                masked[key] = "***MASKED***"
        if re.search(
            r"secret|password|token|credential|private.?key|api.?key",
            str(value.get("name", "")),
            re.I,
        ):
            if "value" in masked:
                masked["value"] = "***MASKED***"
        return {k: _redact(v, depth + 1) for k, v in masked.items()}
    if isinstance(value, list):
        return [_redact(v, depth + 1) for v in value]
    return value


class EvidenceLookup:
    """One bounded raw file cached at a time, verified before any excerpt is returned."""

    def __init__(self, store: Any, run: str):
        self.workspace = store.workspace
        self.hashes = dict(
            store.rows("SELECT artifact_id,sha256 FROM artifacts WHERE run_id=?", [run])
        )
        self.sources = {s["artifact_id"]: s for s in store.source_data(run) if s.get("artifact_id")}
        self.cached_sha: str | None = None
        self.document: Any = None

    def excerpt(self, artifact: str, locator: str) -> tuple[str, str]:
        digest = self.hashes.get(artifact)
        if not digest or not re.fullmatch(r"[0-9a-f]{64}", digest):
            return "", "Source artifact reference is unavailable in the selected assessment."
        try:
            path = resolve_under(self.workspace, f"raw/{digest}.json")
            if digest != self.cached_sha:
                self.document, self.cached_sha = None, None
                if path.stat().st_size > MAX_BYTES:
                    return "", "Source exceeds the bounded evidence lookup size limit."
                with path.open("rb") as stream:
                    content = stream.read(MAX_BYTES + 1)
                if len(content) > MAX_BYTES:
                    return "", "Source exceeds the bounded evidence lookup size limit."
                import hashlib

                if hashlib.sha256(content).hexdigest() != digest:
                    return "", "Source integrity mismatch; configuration excerpt withheld."
                self.document = json.loads(content)
                self.cached_sha = digest
            node = _pointer(self.document, locator)
            return json.dumps(_redact(node), ensure_ascii=True, indent=2, sort_keys=True), ""
        except (OSError, ValueError, TypeError, KeyError, IndexError, RecursionError):
            return "", "Source file or JSON path is unavailable or unreadable."


def build(
    store: Any, run: str, data: dict | None = None, *, provider: str = "", scope: str = ""
) -> dict:
    """Only SELECTs and local file reads. `data` supplies existing report inclusion/edits.

    One row per stored candidate; evaluations without a candidate remain observations or
    evidence limitations. Excluded findings and their evaluations stay excluded.
    """
    empty: dict = {"register": [], "index": [], "details": [], "audit": []}
    if provider == "AWS":
        return empty
    store.require_run(run)
    sources = [s for s in store.source_data(run) if s.get("format") == "gcp-json-v1"]
    selected = [s for s in sources if not scope or s.get("project_id") == scope]
    reported = {f["id"]: f for f in (data or {}).get("findings", []) if f["provider"] == "GCP"}
    if not selected and not reported:
        return empty
    hidden = excluded(store, run)
    candidates = [f for f in store.candidates(run) if f.provider == "GCP"]
    visible = [
        f for f in candidates if f.finding_id not in hidden and (not scope or f.project_id == scope)
    ]
    candidate_keys = {(f.rule_id, f.resource_id) for f in candidates}
    results = [
        r
        for r in store.results(run)
        if r.provider == "GCP" and (not scope or r.project_id == scope)
    ]
    by_key = defaultdict(list)
    for result in results:
        by_key[(result.rule_id, result.resource_id)].append(result)
    source_results = _source_results(store, run, results)
    groups = defaultdict(list)
    for finding in visible:
        groups[group_key(finding)].append(finding.finding_id)
    decisions, sre = store.latest_decisions(run), latest_sre(store, run)
    revisions = _revisions(store, run, decisions)
    engine = _engine_levels(store, run)
    titles = {}
    for finding in visible:
        titles.setdefault(group_key(finding), finding.title)
    records = []
    for finding in visible:
        evaluations = by_key[(finding.rule_id, finding.resource_id)]
        status, stage = review_state(store, finding, decisions, sre)
        decision, response = decisions.get(finding.finding_id), sre.get(finding.finding_id)
        report = reported.get(finding.finding_id, {})
        revision = revisions.get(finding.finding_id)
        explanation = explain(finding)
        members = groups[group_key(finding)]
        records.append(
            {
                "id": finding.finding_id,
                "project": finding.project_id,
                "resource": finding.native_id,
                "title": report.get("title", finding.title),
                "severity": report.get("severity", finding.suggested_severity),
                "rationale": report.get("rationale", finding.severity_rationale),
                "confidence": finding.evidence_confidence,
                "action": report.get("remediation", finding.remediation),
                "type": finding.finding_class or "Observation",
                "domain": finding.security_domain,
                "control": _text([finding.methodology_id, *finding.methodology_criteria]),
                "reason": finding.technical_observation,
                "limits": [
                    *finding.missing_evidence,
                    *(x for r in evaluations for x in r.limitations),
                ],
                "conditions": finding.validation_requirements,
                "refs": finding.evidence,
                "outcome": _text(sorted({str(r.outcome) for r in evaluations})),
                "source_result": _text(
                    sorted(
                        {v for r in evaluations for v in source_results.get(r.evaluation_id, [])}
                    )
                ),
                "related": [i for i in members if i != finding.finding_id],
                "group": stable_id("finding-group", list(group_key(finding))),
                "group_title": f"{titles[group_key(finding)]} ({len(members)} records)",
                "group_size": len(members),
                "expected": recommendation(finding.rule_id).get("expected_posture"),
                "scenario": explanation["misuse"],
                "engine": engine.get(finding.finding_id, []),
                "shown": bool(report) if data is not None else True,
                "status": status,
                "stage": stage,
                "decision": decision,
                "sre": response,
                "revision": revision,
            }
        )
    for result in results:
        if (result.rule_id, result.resource_id) in candidate_keys:
            continue
        limited = str(result.outcome) in {"UNKNOWN", "COLLECTION_ERROR"}
        records.append(
            {
                "id": result.evaluation_id,
                "project": result.project_id,
                "resource": result.native_id,
                "title": result.reasoning,
                "type": "Evidence limitation"
                if limited
                else (result.finding_class or "Gap")
                if str(result.outcome) == "FAIL"
                else "Observation",
                "confidence": result.confidence,
                "control": _text([result.methodology_id, *result.methodology_criteria]),
                "reason": result.reasoning,
                "limits": [
                    *result.missing_dependencies,
                    *result.evidence_needed,
                    *result.limitations,
                ],
                "expected": recommendation(result.rule_id).get("expected_posture"),
                "refs": result.evidence,
                "outcome": str(result.outcome),
                "source_result": _text(source_results.get(result.evaluation_id, [])),
                "shown": False,
            }
        )
    referenced = {ref.artifact_id for record in records for ref in record["refs"]}
    for source in selected:
        if (
            source.get("effective_coverage") not in {"ERROR", "NOT_COLLECTED"}
            or source.get("artifact_id") in referenced
        ):
            continue
        from offline_analyzer.contracts.models import EvidenceRef

        artifact = source.get("artifact_id")
        records.append(
            {
                "id": stable_id("gcp-evidence-limitation", [run, source]),
                "project": source.get("project_id"),
                "resource": "",
                "type": "Evidence limitation",
                "title": f"Collection evidence unavailable: {source['operation']}",
                "reason": _text(source.get("errors")),
                "limits": source.get("limitations") or [],
                "refs": (EvidenceRef(artifact_id=artifact, locator=""),) if artifact else (),
                "outcome": source["effective_coverage"],
                "shown": False,
                "collection_only": True,
                "operations": [source["operation"]],
                "domain": (source.get("gcp_metadata") or {}).get("domain"),
            }
        )
    lookup = EvidenceLookup(store, run)
    for sequence, record in enumerate(sorted(records, key=lambda r: r["id"]), 1):
        _record(empty, record, sequence, lookup, (data or {}).get("paths", []))
    # Classification-only frozen feeds do not carry raw evidence. Never attach AWS files.
    known = {r[0] for r in empty["register"]}
    for fid, finding in sorted(reported.items()):
        if fid not in known:
            _record(
                empty,
                {
                    "id": fid,
                    "project": finding.get("scope"),
                    "resource": finding.get("resource"),
                    "title": finding.get("title"),
                    "severity": finding.get("severity"),
                    "type": "Observation",
                    "reason": "Frozen edition carries no source configuration.",
                    "limits": ["Original source evidence unavailable in this assessment."],
                    "refs": (),
                    "shown": True,
                },
                len(empty["register"]) + 1,
                lookup,
                [],
            )
    return empty


def _revisions(store: Any, run: str, decisions: dict) -> dict:
    """Reuse evidence-current review validity; never turn a draft flag into approval."""
    from offline_analyzer.contracts.report_finding import ReportRevision
    from offline_analyzer.findings.report_ready import current_findings, validity

    found = {}
    for (raw,) in store.rows(
        "SELECT data FROM report_revisions QUALIFY row_number() OVER "
        "(PARTITION BY report_id ORDER BY revision DESC)=1 ORDER BY report_id"
    ):
        revision = ReportRevision.model_validate_json(raw)
        state = validity(
            revision,
            current_findings(store, run, [s.finding_id for s in revision.snapshots]),
            decisions,
            store,
        )
        for snapshot in revision.snapshots:
            if state["member_states"].get(snapshot.finding_id) == "CURRENT":
                if (
                    snapshot.finding_id not in found
                    or revision.primary_finding == snapshot.finding_id
                ):
                    found[snapshot.finding_id] = revision
    return found


def _source_results(store: Any, run: str, results: list) -> dict:
    """Only source-native finding statuses; CLI transport success is not a security result."""
    ids = sorted({oid for r in results for oid in r.supporting_observations})
    if not ids:
        return {}
    observed = {}
    for oid, raw in store.rows(
        "SELECT observation_id,data FROM observations WHERE run_id=? AND tool='gcp' "
        "AND observation_id IN (SELECT unnest(?))",
        [run, ids],
    ):
        item = json.loads(raw)
        if item.get("source_finding_id") or item.get("resource_type") == "GCP::SCCFinding":
            observed[oid] = _text([item.get("source_status"), item.get("source_status_caption")])
    return {
        r.evaluation_id: sorted(
            {observed[oid] for oid in r.supporting_observations if observed.get(oid)}
        )
        for r in results
    }


def _engine_levels(store: Any, run: str) -> dict[str, list[str]]:
    """Latest engine 2.0.1 class per original finding; empty before any recalculation."""
    from offline_analyzer.attack_paths.versioned import latest

    try:
        paths = latest(store.workspace, run)["paths"]
    except (OSError, ValueError, KeyError, TypeError, InputError):
        return {}
    found = defaultdict(list)
    for path in paths:
        for fid in path.get("original_finding_ids") or []:
            found[fid].append(f"{path.get('classification')}: {path.get('path_id')}")
    return {k: sorted(v) for k, v in found.items()}


def _focus(node: Any, text: str) -> Any:
    """Only the allow bindings a record cites (role -> member); others are counted, not shown.

    Uncited or non-policy evidence is returned unchanged, so nothing is hidden by guesswork.
    The complete redacted excerpt always remains in GCP_Evidence_Details.
    """
    pairs = {(role, member.split("?", 1)[0]) for role, member in CITED.findall(text or "")}
    if not pairs or not isinstance(node, dict):
        return node
    holder = node.get("policy") if isinstance(node.get("policy"), dict) else node
    bindings = holder.get("bindings")
    if not isinstance(bindings, list):
        return node
    kept, total = [], 0
    for binding in bindings:
        if not isinstance(binding, dict) or not isinstance(binding.get("members"), list):
            continue
        total += len(binding["members"])
        cited = [
            m
            for m in binding["members"]
            if isinstance(m, str) and (binding.get("role"), m.split("?", 1)[0]) in pairs
        ]
        if cited:
            kept.append({**binding, "members": cited})
    if not kept:
        return node
    focused = {k: node[k] for k in ("resource", "assetType", "name", "subject") if k in node}
    focused["bindings"] = kept
    shown = sum(len(b["members"]) for b in kept)
    if total > shown:
        focused["omitted"] = f"{total - shown} other members; see {SHEETS[2]}"
    return focused


def _unique(values) -> list[str]:
    return list(dict.fromkeys(str(v) for v in values if v))


def _record(output: dict, record: dict, sequence: int, lookup: EvidenceLookup, paths: list) -> None:
    references, snippets, statuses, quality, operations = (
        [],
        [],
        [],
        [],
        list(record.get("operations", [])),
    )
    # Rule-declared unresolved dependencies: what the evidence does not establish.
    unestablished = _unique(record.get("limits", []))
    first_index = len(output["index"]) + 2
    for ref in sorted(set(record["refs"]), key=lambda r: (r.artifact_id, r.locator)):
        checkpoint()
        source = lookup.sources.get(ref.artifact_id, {})
        metadata = source.get("gcp_metadata") or {}
        digest = lookup.hashes.get(ref.artifact_id, "")
        original = (
            metadata.get("original_file") or source.get("path") or metadata.get("error_file") or ""
        )
        reference = f"{original or ref.artifact_id}#{ref.locator}"
        mapping = stable_id("gcp-evidence", [record["id"], ref.artifact_id, ref.locator])
        failed = source.get("effective_coverage") in {"ERROR", "NOT_COLLECTED"}
        snippet, limitation = (
            ("", "Collection failed or was not collected; no usable configuration.")
            if failed
            else lookup.excerpt(ref.artifact_id, ref.locator)
        )
        # Collection coverage reasons: evidence limits and data quality.
        quality.append(limitation)
        quality.append(
            reason_text(source_category(source), metadata.get("completeness_limitations"))
        )
        quality.extend(source.get("limitations") or [])
        status = source.get("effective_coverage") or "Unrecorded collection status"
        statuses.append(status)
        operations.append(source.get("operation") or "")
        references.append(reference)
        if snippet:
            focused = _focus(
                json.loads(snippet), " ".join(_text(record.get(k)) for k in ("reason", "title"))
            )
            compact = json.dumps(focused, ensure_ascii=True, separators=(",", ":"))
            snippets.append(f"{ref.locator or '(root)'}: {compact}")
        detail_row, index_row = len(output["details"]) + 2, len(output["index"]) + 2
        output["index"].append(
            [
                mapping,
                record["id"],
                record.get("project"),
                record.get("resource"),
                source.get("operation"),
                metadata.get("dataset"),
                original,
                ref.locator,
                metadata.get("completed_at") or source.get("collected_at"),
                ref.artifact_id,
                digest,
                metadata.get("status") or status,
                limitation or "Available — SHA-256 verified; redacted excerpt",
                (f"{SHEETS[2]}!H{detail_row}", f"'{SHEETS[2]}'!H{detail_row}"),
            ]
        )
        chunks = _parts(snippet)
        for part, chunk in enumerate(chunks, 1):
            output["details"].append(
                [
                    mapping,
                    record["id"],
                    record.get("project"),
                    record.get("resource"),
                    original,
                    ref.locator,
                    part,
                    chunk,
                    limitation,
                    (f"{SHEETS[1]}!A{index_row}", f"'{SHEETS[1]}'!A{index_row}"),
                ]
            )
    if not record["refs"]:
        quality.append("No source evidence reference recorded; original configuration unavailable.")
    quality = _unique(quality)
    # Main-sheet snippets are explicitly excerpts, never evidence-completeness status.
    actual = "\n".join(snippets)
    if len(actual) > 600:
        actual = (
            actual[:600]
            + "\n[Concise excerpt; full redacted configuration in GCP_Evidence_Details.]"
        )
    response, decision = record.get("sre"), record.get("decision")
    revision = record.get("revision")
    risk = revision.risk_acceptance if revision else None
    outcome = record.get("outcome", "")
    # Engine 2.0.1 classes when a recalculation exists; legacy projections are labelled.
    path_levels = [f"Engine 2.0.1 {level}" for level in record.get("engine", [])] or [
        f"Legacy projection (not engine 2.0.1) {p.get('id', '')}: {p.get('feasibility', '')}"
        for p in paths
        if record["id"] in p.get("finding_ids", [])
    ]
    row = [
        record["id"],
        record.get("title"),
        record.get("project"),
        _text([record.get("resource"), *sorted(set(operations))]),
        record.get("severity"),
        record.get("rationale"),
        record.get("confidence"),
        "",
        record.get("action"),
        record["type"],
        actual or "Original configuration unavailable; see evidence limitations.",
        ("; ".join(references), f"'{SHEETS[1]}'!B{first_index}") if references else "",
        revision.sections.expected_posture
        if revision
        else record.get("expected") or record.get("control"),
        _text(sorted(set(statuses))) or "Source evidence unavailable",
        _text(unestablished) or "No unresolved dependency recorded by the rule.",
        revision.sections.attack_scenario if revision else record.get("scenario", ""),
        _text(record.get("conditions")),
        record.get("source_result"),
        _text([record.get("domain"), "GCP", *sorted(set(operations))]),
        record.get("control"),
        "CLI/API configuration" if snippets else "Evidence unavailable",
        _text(quality) or "No collection limitation recorded for the cited sources.",
        _text([record.get("reason"), decision.note if decision else ""]),
        _text(record.get("related")),
        _text(path_levels),
        outcome,
        sequence,
        record.get("resource"),
        record.get("group_title") or record.get("group"),
        f"Approved by {risk.reviewer} at {risk.approved_at.isoformat()}: {risk.rationale}"
        if risk
        else PENDING,
        revision.include_in_report if revision else PENDING,
        bool(record.get("shown")),
        record.get("group_size", 1),
        "",
        "",
        "",
        _text(revision.sections.references) if revision else "",
        "",
        "",
        _text([record.get("status"), record.get("stage")]) if decision else PENDING,
    ]
    assert len(row) == len(COLUMNS) == 40
    output["register"].append(row)
    if decision or response:
        output["audit"].append(
            [
                record["id"],
                record.get("status") or "",
                decision.analyst if decision else "",
                decision.decided_at.isoformat() if decision else "",
                decision.note if decision else "",
                response.verdict if response else "",
                response.reason if response else "",
                getattr(response, "sre_contact", "") if response else "",
                response.recorded_at.isoformat() if response else "",
            ]
        )


def _parts(text: str) -> list[str]:
    """Keep excerpts below both Excel's cell limit and its visible row-height limit."""
    parts = []
    start = 0
    while start < len(text):
        end = min(start + CHUNK, len(text))
        breaks = [m.end() for m in re.finditer("\n", text[start:end])]
        if len(breaks) > 18:
            end = start + breaks[17]
        parts.append(text[start:end])
        start = end
    return parts or [""]


def add_sheets(workbook: Any, payload: dict) -> None:
    """Native tables, literal strings and internal hyperlinks; existing sheets untouched."""
    if not payload["register"]:
        return
    audit = payload.get("audit") or []
    for name, headers, rows in zip(
        (*SHEETS, AUDIT_SHEET)[: 4 if audit else 3],
        (COLUMNS, INDEX_COLUMNS, DETAIL_COLUMNS, AUDIT_COLUMNS),
        (payload["register"], payload["index"], payload["details"], audit),
        strict=False,
    ):
        if len(rows) > 1_048_575:
            raise InputError("GCP_REGISTER_EXCEL_ROW_LIMIT")
        ws = workbook.create_sheet(name)
        ws.freeze_panes = "D2"
        ws.sheet_view.zoomScale = 85
        ws.row_dimensions[1].height = 64
        ws.sheet_format.defaultRowHeight = 105
        for column, header in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(column)].width = 32 if "ID" in header else 44
        if name == SHEETS[0]:
            ws.column_dimensions["K"].width = 76
        if name == SHEETS[2]:
            ws.column_dimensions["H"].width = 100

        def cells(values: Any, header: bool = False, ws: Any = ws) -> list:
            result = []
            for column, value in enumerate(values, 1):
                link = value if isinstance(value, tuple) else None
                value = value[0] if link else value
                if isinstance(value, str):
                    value = ILLEGAL.sub("", value)
                    if len(value) > 32000:
                        value = (
                            value[:31800]
                            + " [Display shortened; consult evidence details/source reference.]"
                        )
                    value = excel_literal(value)
                cell = WriteOnlyCell(ws, value=value)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.font = Font(
                    name="Calibri", size=11, bold=header, color="FFFFFF" if header else "17365D"
                )
                if header:
                    cell.fill = PatternFill("solid", fgColor="17365D")
                if link:
                    cell.hyperlink = Hyperlink(
                        ref=f"{get_column_letter(column)}{ws.r0_row_count + 1}",
                        location=link[1],
                        display=str(value),
                    )
                    cell.font = Font(name="Calibri", size=11, color="0563C1", underline="single")
                result.append(cell)
            return result

        ws.r0_row_count = 0
        ws.append(cells(headers, True))
        ws.r0_row_count = 1
        for row in rows:
            checkpoint()
            wrapped = max(
                sum(
                    max(
                        1,
                        math.ceil(
                            len(line) / max(1, ws.column_dimensions[get_column_letter(i)].width - 4)
                        ),
                    )
                    for line in str(v[0] if isinstance(v, tuple) else v or "").split("\n")
                )
                for i, v in enumerate(row, 1)
            )
            ws.row_dimensions[ws.r0_row_count + 1].height = min(409, max(60, wrapped * 15 + 10))
            ws.append(cells(row))
            ws.r0_row_count += 1
        reference = f"A1:{get_column_letter(len(headers))}{ws.r0_row_count}"
        ws.auto_filter.ref = reference
        table = Table(displayName=name, ref=reference)
        table.tableColumns = [TableColumn(id=i, name=h) for i, h in enumerate(headers, 1)]
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        from openpyxl.worksheet.filters import AutoFilter

        table.autoFilter = AutoFilter(ref=reference)
        # Explicit TableColumns are required for openpyxl's streaming writer.
        import warnings

        with warnings.catch_warnings():
            warnings.filterwarnings(
                "ignore", message="In write-only mode you must add table columns manually"
            )
            ws.add_table(table)
