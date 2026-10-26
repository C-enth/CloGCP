"""Standalone assessment workbook; never opens, integrates with or overwrites customer trackers."""

import json
import os
import re
from collections import Counter
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.writer.excel import ExcelWriter
from openpyxl.xml.functions import Element, SubElement

from offline_analyzer import PRODUCT_NAME
from offline_analyzer.attack_paths.prerequisites import path_records
from offline_analyzer.correlation.prowler import MAPPINGS
from offline_analyzer.findings.explanations import explain
from offline_analyzer.findings.generate import validation_status
from offline_analyzer.ingestion.json_input import InputError
from offline_analyzer.ingestion.safety import excel_literal
from offline_analyzer.inventory.accounts import workspace_registry
from offline_analyzer.provenance.execution import CheckpointWriter
from offline_analyzer.reporting.metrics import assessment_metrics
from offline_analyzer.reporting.tracker import REFERENCE_COLUMNS, TRACKER_COLUMNS, tracker_rows
from offline_analyzer.rules.engine import verify_raw
from offline_analyzer.rules.registry import ALL_RULES, BY_ID, primary_operations

NOTICE = (
    "LIMITED ASSESSMENT: rules evaluate only the named methodology criteria; mappings are "
    "PROPOSED pending methodology-owner approval and no control is fully covered. "
    "Only listed evidence scopes/resources assessed. Partial/uncollected scopes are not clean. "
    "Candidates are not analyst-confirmed. Configuration is not proof of exploitation. "
    "Display text is sanitized and may be shortened; raw evidence remains unchanged."
)


def mapping_text(item) -> str:
    if not item.methodology_id:
        return "UNRESOLVED: provisional rule"
    return f"{item.methodology_id} {item.mapping_status}: {', '.join(item.methodology_criteria)}"


ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def safe_text(value: object) -> str:
    if value is None:
        return "UNRESOLVED"
    if not isinstance(value, str):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True)
    if ILLEGAL.search(value):
        value = ILLEGAL.sub("", value)
    value = excel_literal(value)
    if len(value) > 32000:
        value = value[:31900] + " [SHORTENED: consult raw evidence/Evidence Register]"
    return value


REPORT_SHEETS = 28  # Stable workbook contract; tests assert the generated sheet count.
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
LETTERS = [get_column_letter(i) for i in range(1, 129)]


def style_id(ws, **styles) -> str:
    """Register one cell style with the workbook once; rows reuse its index."""
    cell = WriteOnlyCell(ws)
    for name, value in styles.items():
        setattr(cell, name, value)
    return str(cell.style_id)


def row_element(ws, values, index):
    """One row of typed cells; strings are always inline strings, never formulas.

    Replaces openpyxl's per-cell Cell objects, style hashing and element writes, which
    dominated large exports (~5M cells), with the same XML built once per row.
    """
    attributes = {"r": str(index)}
    attributes.update(ws.row_dimensions.get(index, {}))
    row = Element("row", attributes)
    for column, (value, style) in enumerate(values):
        reference = LETTERS[column] + attributes["r"]
        if type(value) is bool:
            cell = SubElement(row, "c", {"r": reference, "s": style, "t": "b"})
            SubElement(cell, "v").text = "1" if value else "0"
        elif type(value) is int:
            cell = SubElement(row, "c", {"r": reference, "s": style, "t": "n"})
            SubElement(cell, "v").text = str(value)
        else:
            cell = SubElement(row, "c", {"r": reference, "s": style, "t": "inlineStr"})
            if value == "":
                continue  # openpyxl writes an empty styled cell without content.
            text = SubElement(SubElement(cell, "is"), "t")
            if value != value.strip():
                text.set(XML_SPACE, "preserve")
            text.text = value
    return row


def sheet(workbook, name, headers):
    if len(headers) > len(LETTERS):
        raise ValueError("too many report columns")
    ws = workbook.create_sheet(name)
    if getattr(workbook, "r0_progress", None):
        workbook.r0_progress(f"Writing sheet: {name}", len(workbook.worksheets) - 1, None)
    ws.freeze_panes = "D4"
    ws.r0_row_count = 0  # Report-owned counter; no library internals are read.
    ws.merged_cells.add(f"A1:{get_column_letter(len(headers))}1")
    ws.row_dimensions[1].height = 80
    ws.row_dimensions[2].height = 26
    ws.row_dimensions[3].height = 36
    ws.sheet_format.defaultRowHeight = 90
    ws.sheet_view.zoomScale = 85
    for index, _ in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(index)].width = 26 if index < 4 else 42
    ws.r0_body_style = style_id(ws, alignment=Alignment(vertical="top", wrap_text=True))
    header_style = style_id(
        ws,
        font=Font(bold=True, color="FFFFFF"),
        fill=PatternFill("solid", fgColor="17365D"),
        alignment=Alignment(wrap_text=True),
    )
    # Stream rows through the write-only sheet's own XML writer, one element per row.
    ws._get_writer()
    ws._values_to_row = lambda values, _index: values
    ws._writer.write_row = lambda xf, values, index: xf.write(row_element(ws, values, index))
    append(ws, [NOTICE])
    append(ws, ["Assessment run", workbook.properties.identifier])
    ws.merged_cells.add("B2:D2")
    ws.append([(safe_text(header), header_style) for header in headers])
    ws.r0_row_count = 3
    return ws


def append(ws, values):
    from offline_analyzer.provenance.execution import checkpoint

    checkpoint()
    style = ws.r0_body_style
    ws.append(
        [
            # Computed counts remain useful Excel numbers; untrusted text is sanitized.
            (value if type(value) in (int, bool) else safe_text(value), style)
            for value in values
        ]
    )
    ws.r0_row_count += 1


def methodology_sheet(wb, store, run_id: str) -> None:
    """Per-rule mapping and whether this run supplied the rule's evidence at all."""
    ws = sheet(
        wb,
        "Methodology Coverage",
        [
            "Rule ID",
            "Rule version",
            "Methodology control",
            "Mapping status",
            "Criteria evaluated",
            "Criteria not evaluated",
            "Source-linked attack paths",
            "Evidence operations",
            "Evidence in this run",
            "Outcomes in this run",
        ],
    )
    supplied = {s["operation"] for s in store.source_data(run_id)}
    outcomes = Counter((r.rule_id, str(r.outcome)) for r in store.results(run_id))
    for rule in ALL_RULES:
        link = rule.methodology
        operations = primary_operations(rule) | set(rule.context_operations)
        append(
            ws,
            [
                rule.rule_id,
                rule.version,
                link.control_id if link else None,
                link.status if link else "UNRESOLVED",
                link.criteria if link else (),
                link.not_evaluated if link else (),
                link.attack_paths if link else (),
                sorted(operations),
                "SUPPLIED"
                if primary_operations(rule) & supplied
                else "NOT_SUPPLIED: control criteria not assessed in this run",
                {o: n for (rid, o), n in sorted(outcomes.items()) if rid == rule.rule_id},
            ],
        )


def account_coverage(store, run_id: str, registry) -> list[list]:
    """Per account and evidence family; registry accounts without evidence are listed."""
    sources = list(store.source_data(run_id))
    families = sorted({s["operation"] for s in sources if s["operation"] != "prowler:Scan"})
    accounts = {s["account_id"] for s in sources if s["account_id"]}
    registered = registry.accounts if registry else {}
    accounts |= set(registry.in_scope if registry else ())
    rows = []
    for account in sorted(accounts):
        entry = registered.get(account)
        status = (
            ("IN_SCOPE" if entry.in_scope else "OUT_OF_SCOPE")
            if entry
            else ("NOT_IN_REGISTRY" if registry else "NO_REGISTRY")
        )
        for family in families or ["(no evidence families in run)"]:
            matching = [
                s for s in sources if s["account_id"] == account and s["operation"] == family
            ]
            rows.append(
                [
                    account,
                    entry.tracker_key if entry else None,
                    status,
                    family,
                    sorted({s["region"] or "unscoped" for s in matching}),
                    sorted({s["effective_coverage"] for s in matching}) or ["NOT_COLLECTED"],
                    bool(matching),
                ]
            )
    return rows


def iam_sheets(wb, store, run_id: str, registry) -> None:
    """IAM investigations are analyzer conclusions with evidence trails, not AWS answers."""
    from offline_analyzer.iam.analysis import IamAnalysis

    analysis = IamAnalysis(store, run_id, list(store.source_data(run_id)), registry)
    investigations = sheet(
        wb,
        "IAM Investigations",
        [
            "Type",
            "Principal",
            "Target",
            "Service",
            "Authorization state",
            "Classification",
            "Reasoning",
            "Supporting statements",
            "Blocking controls",
            "Missing dependencies",
            "Preconditions and notes",
        ],
    )
    for record in analysis.transitions() + analysis.passrole():
        append(
            investigations,
            [
                record.kind,
                record.principal,
                record.target,
                record.service,
                record.state,
                record.classification,
                record.reasoning,
                record.supporting,
                record.blocking,
                record.missing,
                record.preconditions,
            ],
        )
    relationships = sheet(
        wb, "IAM Relationships", ["Relationship", "Source", "Target", "Basis", "Detail", "Evidence"]
    )
    for item in analysis.relationships():
        append(
            relationships,
            [item.kind, item.source, item.target, item.basis, item.detail, item.evidence],
        )
    permissions = sheet(
        wb,
        "Effective Permissions",
        ["Role", "Account", "Privilege level", "Probe decisions", "Missing evidence"],
    )
    for role in sorted(analysis.inventory.roles.values(), key=lambda r: r.arn):
        level, notes = analysis.privilege(role)
        append(
            permissions,
            [role.arn, role.account, level, notes, analysis.privilege_missing(role)],
        )
    gaps = sheet(wb, "IAM Evidence Gaps", ["Category", "Scope", "Missing evidence"])
    for gap in analysis.evidence_gaps():
        append(gaps, list(gap))
    # Every other limit that dropped records (rows only when something was dropped).
    from offline_analyzer.attack_paths.feasibility import view_limits
    from offline_analyzer.provenance.truncation import text

    entries = view_limits(store, run_id)
    cached = store.db.execute(
        "SELECT data FROM enterprise_projection WHERE run_id=?", [run_id]
    ).fetchone()
    if cached:
        entries += json.loads(cached[0]).get("truncation_detail") or []
    for entry in entries:
        append(gaps, ["Truncated view", entry["scope"], text(entry)])


def export_excel(
    store,
    run_id: str,
    output: Path,
    account_keys: dict[str, str] | None = None,
    progress=None,
) -> Path:
    wb = Workbook(write_only=True)
    try:
        return _export_excel(store, run_id, output, wb, account_keys, progress)
    except BaseException:
        if getattr(wb, "r0_output_owned", False):
            output.unlink(missing_ok=True)
        raise
    finally:
        for worksheet in wb.worksheets:
            if not worksheet.closed:
                worksheet.close()
            writer = worksheet._writer
            if writer is not None and Path(writer.out).exists():
                writer.cleanup()


def _export_excel(store, run_id, output, wb, account_keys, progress) -> Path:
    store.require_run(run_id)
    if output.exists():
        raise InputError("REPORT_ALREADY_EXISTS")
    if not store.db.execute("SELECT 1 FROM assessments WHERE run_id=?", [run_id]).fetchone():
        raise InputError("RUN_NOT_ASSESSED")
    verify_raw(store, run_id)
    wb.r0_progress = progress  # Optional (stage, completed, total) observer; no data passed.
    wb.properties.identifier = run_id
    wb.properties.title = f"{PRODUCT_NAME} assessment"
    wb.properties.creator = PRODUCT_NAME
    metrics_sheet = sheet(wb, "Assessment Summary", ["Metric", "Value", "Meaning"])
    for metric, value in assessment_metrics(store, run_id).items():
        append(
            metrics_sheet,
            [
                metric,
                value,
                "Distinct metric; counts are not additive"
                if metric != "outcomes"
                else "Resource and collection-scope records combined; see separate outcome totals",
            ],
        )
    headers = [
        "Finding ID",
        "Title",
        "Domain",
        "Account",
        "Region",
        "Resource ID",
        "Native ID",
        "Rule",
        "Methodology mapping",
        "Why this is flagged",
        "How it can be misused (potential; not demonstrated)",
        "Technical security impact",
        "Recommendation",
        "Suggested severity",
        "Severity rationale",
        "Evidence confidence",
        "Evidence refs",
        "Contradictory observations",
        "Missing evidence",
        "Remediation",
        "Engine status",
        "Analyst validation status",
        "Analyst",
        "Analyst note",
        "Technical validation requirements",
        "Methodology control",
        "Mapping status",
        "Criteria evaluated",
        "Source-linked attack paths",
        "STRIDE category (source)",
        "Methodology inherent severity (source; not technical severity)",
        "Finding class",
        "Exposure status",
        "Compensating controls",
    ]
    findings = sheet(wb, "Findings Register", headers)
    confirmed = sheet(wb, "Confirmed Findings", headers)
    # Separate record sets per class; all remain unconfirmed candidates.
    by_class = {
        "SUPPORTED_SECURITY_GAP": sheet(wb, "Security Gaps", headers),
        "POTENTIAL_EXPOSURE": sheet(wb, "Potential Exposures", headers),
        "BASELINE_DEVIATION": sheet(wb, "Baseline Deviations", headers),
    }
    review = sheet(
        wb,
        "Review Queue",
        [
            "Kind",
            "ID",
            "Account",
            "Region",
            "Resource",
            "Status",
            "Reason",
            "Missing dependencies",
            "Evidence",
        ],
    )
    from offline_analyzer.findings.exclusions import excluded

    latest = store.latest_decisions(run_id)
    hidden, skipped = excluded(store, run_id), 0
    for finding in store.candidates(run_id):
        if finding.finding_id in hidden:
            skipped += 1  # Excluded from client reporting by the assessor.
            continue
        status, decision = validation_status(store, finding, latest)
        explanation = explain(finding)
        row = [
            finding.finding_id,
            finding.title,
            finding.security_domain,
            finding.account_id,
            finding.region,
            finding.resource_id,
            finding.native_id,
            finding.rule_id,
            mapping_text(finding),
            explanation["why_flagged"],
            explanation["misuse"],
            explanation["impact"],
            explanation["recommendation"],
            finding.suggested_severity,
            finding.severity_rationale,
            finding.evidence_confidence,
            [ref.model_dump() for ref in finding.evidence],
            finding.contradictory_observations,
            finding.missing_evidence,
            finding.remediation,
            "CANDIDATE",
            status,
            decision.analyst if decision else None,
            decision.note if decision else None,
            finding.validation_requirements,
            finding.methodology_id,
            finding.mapping_status,
            finding.methodology_criteria,
            finding.attack_paths,
            finding.stride_category,
            finding.methodology_inherent_severity,
            finding.finding_class,
            finding.exposure_status,
            finding.compensating_controls,
        ]
        append(findings, row)
        if finding.finding_class in by_class:
            append(by_class[finding.finding_class], row)
        if status == "CONFIRMED":
            append(confirmed, row)
        if status in {"CANDIDATE", "NEEDS_VALIDATION", "NEEDS_REVALIDATION"}:
            append(
                review,
                [
                    "candidate",
                    finding.finding_id,
                    finding.account_id,
                    finding.region,
                    finding.native_id,
                    status,
                    finding.technical_observation,
                    finding.missing_evidence,
                    [r.model_dump() for r in finding.evidence],
                ],
            )
    append(
        metrics_sheet,
        ["excluded_from_reporting", skipped, "Findings the assessor excluded from reports"],
    )
    from offline_analyzer.findings.exclusions import listing

    # Transparency appendix: every excluded finding with its ID, rule, reason and state.
    excluded_sheet = sheet(
        wb,
        "Excluded Findings",
        [
            "Finding ID",
            "Rule ID",
            "Reason",
            "Detail",
            "False positive",
            "Previous state",
            "Carried over from",
            "Analyst",
            "Recorded at",
        ],
    )
    for item in listing(store, run_id)["items"]:
        append(
            excluded_sheet,
            [
                item["id"],
                item["rule"],
                item["reason_label"],
                item["note"],
                item["false_positive"],
                item["previous_state"] or "",
                item["carried_from"] or "",
                item["analyst"],
                item["excluded_at"],
            ],
        )
    controls = sheet(
        wb,
        "Control Assessment",
        [
            "Evaluation ID",
            "Rule ID",
            "Rule version",
            "Mapping",
            "Account",
            "Region",
            "Resource",
            "Outcome",
            "Reasoning",
            "Evidence confidence",
            "Missing prerequisites",
            "Limitations",
            "Evidence refs",
            "Supporting observations",
            "Contradictory observations",
            "Assessment time",
            "Evaluation scope",
            "Methodology control",
            "Criteria evaluated",
            "Excluded (reason)",
        ],
    )
    from offline_analyzer.contracts.r0 import stable_id
    from offline_analyzer.findings.exclusions import evaluation_exclusions

    excluded_by_finding = evaluation_exclusions(store, run_id)
    for result in store.results(run_id):
        append(
            controls,
            [
                result.evaluation_id,
                result.rule_id,
                result.rule_version,
                mapping_text(result),
                result.account_id,
                result.region,
                result.native_id,
                result.outcome,
                result.reasoning,
                result.confidence,
                result.missing_dependencies,
                result.limitations,
                [r.model_dump() for r in result.evidence],
                result.supporting_observations,
                result.contradictory_observations,
                result.assessed_at.isoformat(),
                result.evaluation_scope,
                result.methodology_id,
                result.methodology_criteria,
                excluded_by_finding.get(
                    stable_id("finding", [result.rule_id, result.resource_id or ""]), ""
                ),
            ],
        )
        if (
            result.outcome in {"UNKNOWN", "COLLECTION_ERROR", "MANUAL_VALIDATION_REQUIRED"}
            or result.contradictory_observations
        ):
            append(
                review,
                [
                    "evaluation",
                    result.evaluation_id,
                    result.account_id,
                    result.region,
                    result.native_id,
                    "NEEDS_VALIDATION" if result.contradictory_observations else result.outcome,
                    result.reasoning
                    + (
                        " Prowler/native contradiction requires timestamp/scope reconciliation."
                        if result.contradictory_observations
                        else ""
                    ),
                    result.missing_dependencies,
                    [r.model_dump() for r in result.evidence],
                ],
            )
    coverage = sheet(
        wb,
        "Account Coverage",
        [
            "Account",
            "Region",
            "Operation",
            "Declared coverage",
            "Effective coverage",
            "Errors",
            "Records received",
            "Limitations",
            "Artifact ID",
        ],
    )
    represented = set()
    for source in store.source_data(run_id):
        represented.add((source["account_id"], source["region"]))
        append(
            coverage,
            [
                source["account_id"],
                source["region"],
                source["operation"],
                source["coverage"],
                source["effective_coverage"],
                source["errors"],
                source["record_count"],
                source["limitations"],
                source["artifact_id"],
            ],
        )
    package = json.loads(
        store.db.execute("SELECT package FROM runs WHERE run_id=?", [run_id]).fetchone()[0]
    )
    for target in package["targets"]:
        if (target["account_id"], target["region"]) not in represented:
            append(
                coverage,
                [
                    target["account_id"],
                    target["region"],
                    "ALL R0",
                    "NOT_COLLECTED",
                    "NOT_COLLECTED",
                    [],
                    0,
                    "Target declared but evidence absent",
                    None,
                ],
            )
    for result in store.results(run_id):
        if result.evaluation_scope == "COLLECTION_SCOPE":
            append(
                coverage,
                [
                    result.account_id,
                    result.region,
                    BY_ID[result.rule_id].operation,
                    "SCOPE_EVALUATION",
                    str(result.outcome),
                    [],
                    None,
                    result.reasoning,
                    [r.artifact_id for r in result.evidence],
                ],
            )
    evidence = sheet(
        wb,
        "Evidence Register",
        [
            "Observation ID",
            "Artifact ID",
            "Raw locator",
            "SHA-256",
            "Account",
            "Region",
            "Resource",
            "Resource type",
            "Source tool",
            "Source version",
            "Source check",
            "Source status",
            "Source severity",
            "Collection time",
            "Normalizer",
            "Mapping support",
            "Unresolved inference",
            "Raw storage path",
            "Scanner finding UID",
            "Source lifecycle ID",
            "Source lifecycle caption",
            "Collection account",
            "Collection Region",
            "Trail relationship",
            "Source evidence details",
            "Parent source record",
        ],
    )
    run_refs = store.run_refs(run_id)  # One bulk read instead of ~2 queries per observation.
    for obs in store.observations(run_id):
        located = run_refs.get(obs.observation_id, [])
        refs = tuple(ref for ref, _ in located)
        for ref, digest in located:
            append(
                evidence,
                [
                    obs.observation_id,
                    ref.artifact_id,
                    ref.locator,
                    digest,
                    obs.account_id,
                    obs.region,
                    obs.native_id,
                    obs.resource_type,
                    obs.source_tool,
                    obs.source_version,
                    obs.source_check_id,
                    obs.source_status,
                    obs.source_severity,
                    obs.collected_at.isoformat() if obs.collected_at else None,
                    obs.normalization_version,
                    "SUPPORTED"
                    if obs.source_tool == "aws" or obs.source_check_id in MAPPINGS
                    else "UNSUPPORTED_SCANNER_CHECK",
                    obs.unresolved,
                    f"raw/{digest}.json",
                    obs.source_finding_id,
                    obs.source_status_id,
                    obs.source_status_caption,
                    obs.collection_account_id,
                    obs.collection_region,
                    obs.source_details.get("trail_relationship"),
                    {
                        "status_detail": obs.configuration.get("status_detail"),
                        "metadata": obs.source_details,
                    },
                    ref.locator.rsplit("/resources/", 1)[0]
                    if obs.source_tool == "prowler"
                    else ref.locator,
                ],
            )
        if obs.source_tool == "prowler" and (
            obs.source_check_id not in MAPPINGS
            or obs.unresolved
            or obs.configuration.get("inactive")
        ):
            append(
                review,
                [
                    "scanner observation",
                    obs.observation_id,
                    obs.account_id,
                    obs.region,
                    obs.native_id,
                    "UNRESOLVED",
                    "Unsupported check or insufficient source fields",
                    obs.unresolved,
                    [r.model_dump() for r in refs],
                ],
            )
    # Include artifacts even when listings are empty or no resource can be normalized.
    for artifact, digest, size in store.rows(
        "SELECT artifact_id,sha256,size_bytes FROM artifacts WHERE run_id=? ORDER BY artifact_id",
        [run_id],
    ):
        append(
            evidence,
            [
                "ARTIFACT",
                artifact,
                "/",
                digest,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                "size_bytes=" + str(size),
                None,
                f"raw/{digest}.json",
            ],
        )
    incomplete = sheet(
        wb,
        "Incomplete Investigations",
        [
            "Evaluation ID",
            "Rule",
            "Methodology control",
            "Account",
            "Region",
            "Resource",
            "Outcome",
            "Reasoning",
            "Evidence required",
            "Scope",
        ],
    )
    for result in store.results(run_id):
        if str(result.outcome) in {"UNKNOWN", "COLLECTION_ERROR", "MANUAL_VALIDATION_REQUIRED"}:
            append(
                incomplete,
                [
                    result.evaluation_id,
                    result.rule_id,
                    result.methodology_id,
                    result.account_id,
                    result.region,
                    result.native_id,
                    result.outcome,
                    result.reasoning,
                    result.missing_dependencies,
                    result.evaluation_scope,
                ],
            )
    methodology_sheet(wb, store, run_id)
    registry = workspace_registry(store.workspace)
    keys = account_keys or (registry.tracker_keys() if registry else {})
    tracker = sheet(wb, "Tracker Staging", list(TRACKER_COLUMNS + REFERENCE_COLUMNS))
    for row in tracker_rows(store, run_id, keys, registry):
        append(tracker, row)
    coverage_rows = account_coverage(store, run_id, registry)
    accounts = sheet(
        wb,
        "Assessment Coverage",
        [
            "Account",
            "Tracker key",
            "Registry status",
            "Evidence family",
            "Regions with evidence",
            "Coverage states",
            "Assessed at all",
        ],
    )
    for row in coverage_rows:
        append(accounts, row)
    paths = sheet(
        wb,
        "Attack Path Prerequisites",
        [
            "Record ID",
            "Attack path",
            "Account",
            "Region",
            "Anchor resource",
            "State",
            "Prerequisite status",
            "Unresolved prerequisites",
            "Verified blockers",
            "Linked threats",
            "Supporting evaluations",
            "Reasoning",
            "Definition version",
        ],
    )
    for record in path_records(store, run_id):
        append(
            paths,
            [
                record.record_id,
                record.path_id,
                record.account_id,
                record.region,
                record.anchor,
                record.state,
                record.prerequisites,
                record.missing,
                record.blockers,
                record.threats,
                record.evaluations,
                record.reasoning,
                record.definition_version,
            ],
        )
    iam_sheets(wb, store, run_id, registry)
    from offline_analyzer.reporting.enterprise import enterprise_sheets

    enterprise_sheets(wb, store, run_id)
    if store.db.execute(
        "SELECT 1 FROM observations WHERE run_id=? AND tool='gcp' LIMIT 1", [run_id]
    ).fetchone():
        from offline_analyzer.reporting.multicloud import sheets

        sheets(wb, store, run_id)
    if store.db.execute(
        "SELECT 1 FROM sources WHERE run_id=? AND json_extract_string(data,'$.format')="
        "'gcp-json-v1' LIMIT 1",
        [run_id],
    ).fetchone():
        from offline_analyzer.reporting.gcp_register import add_sheets, build

        add_sheets(wb, build(store, run_id))
    for worksheet in wb.worksheets:
        if worksheet.title in {
            "GCP_Observation_Register",
            "GCP_Evidence_Index",
            "GCP_Evidence_Details",
        }:
            continue
        worksheet.auto_filter.ref = (
            f"A3:{get_column_letter(len(worksheet.column_dimensions))}{worksheet.r0_row_count}"
        )
    # Write directly through an exclusive handle; existing remediation work cannot be overwritten.
    # Write-only workbook and ZIP writer stream to disk, keeping report memory bounded.
    output.parent.mkdir(parents=True, exist_ok=True)
    if progress:
        progress("Saving workbook", len(wb.worksheets), len(wb.worksheets))
    descriptor = os.open(output, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
    wb.r0_output_owned = True
    with os.fdopen(descriptor, "wb") as stream:
        with ZipFile(CheckpointWriter(stream), "w", ZIP_DEFLATED, allowZip64=True) as archive:
            ExcelWriter(wb, archive).write_data()
    return output
