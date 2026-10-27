"""Per-account SRE validation workbooks and the Finalized (confirmed vulnerability) report.

Every output uses the same plain-language columns as the dashboard (why flagged, potential
misuse, impact, recommendation) and the stage vocabulary of ADR-008. Text cells are written
as literals, so finding or analyst text can never become a spreadsheet formula.
"""

import io
import re
import zipfile
from collections import Counter
from datetime import UTC, datetime

from offline_analyzer.attack_paths.model import HOP_KEYS
from offline_analyzer.findings.explanations import MISUSE_QUALIFIER, explain
from offline_analyzer.findings.location import display_region
from offline_analyzer.findings.recommendations import recommendation
from offline_analyzer.findings.sre import STAGES, VERDICTS, latest_sre, review_state
from offline_analyzer.ingestion.json_input import InputError
from offline_analyzer.ingestion.safety import excel_literal

SEVERITIES = ("CRITICAL", "HIGH", "MEDIUM", "LOW", "INFO")
OPEN_FOR_SRE = {"TRIAGE", "WITH_SRE", "SRE_REVALIDATION"}
SRE_CHOICES = ("Gap confirmed", "No gap (BAU)", "Needs info")
DEFINITIONS = (
    ("Why this is flagged", "The check that failed and the configuration observed for it."),
    (
        "How it can be misused",
        "What the configuration would allow someone to do if it is not intended. "
        + MISUSE_QUALIFIER,
    ),
    ("Impact", "The security consequence of the configuration."),
    ("Recommendation", "The recommended corrective action."),
    ("How to validate", "How the owner can confirm the configuration or the fix."),
    *(
        (f"Stage: {label}", meaning)
        for label, meaning in (
            ("Analyst triage", "Not yet sent to SRE."),
            ("With SRE", "Sent to SRE, or SRE asked for more information."),
            ("Evidence changed: revalidate with SRE", "New evidence differs from what SRE saw."),
            ("No gap (BAU)", "SRE confirmed the configuration is intended business-as-usual."),
            ("Gap confirmed by SRE: analyst to confirm", "SRE agreed it is a gap."),
            ("Confirmed vulnerability", "The assessor marked it valid (SRE response optional)."),
            ("Rejected by analyst", "The analyst concluded it is not a valid finding."),
        )
    ),
)


def _scope(finding) -> tuple[str, str]:
    return finding.provider or "AWS", finding.project_id or finding.account_id or ""


def rows(store, run: str, provider: str = "") -> list[dict]:
    """One plain-language record per candidate with its analyst status, SRE verdict and stage.

    Findings excluded from client reporting are omitted (see findings/exclusions.py).
    """
    from offline_analyzer.findings.exclusions import excluded
    from offline_analyzer.ingestion.labels import scope_names
    from offline_analyzer.methodology.cloud_checklist import describe

    store.require_run(run)
    names = scope_names(store)
    decisions = store.latest_decisions(run)
    sre = latest_sre(store, run)
    collected = {s["artifact_id"]: s for s in store.source_data(run) if s.get("artifact_id")}
    hidden = excluded(store, run)
    result = []
    for finding in store.candidates(run):
        finding_provider, scope = _scope(finding)
        if (provider and finding_provider != provider) or finding.finding_id in hidden:
            continue
        status, stage = review_state(store, finding, decisions, sre)
        explanation = explain(finding)
        verdict = sre.get(finding.finding_id)
        decision = decisions.get(finding.finding_id)
        evidence = [
            f"{collected.get(ref.artifact_id, {}).get('operation', 'source')} {ref.locator}"
            for ref in finding.evidence
        ]
        result.append(
            {
                "id": finding.finding_id,
                "severity": finding.suggested_severity,
                "title": finding.title,
                "provider": finding_provider,
                "scope": scope,
                "scope_name": names.get((finding_provider, scope), ""),
                "region": display_region(finding.region),
                "resource": finding.native_id,
                "control": finding.methodology_id or "Not mapped",
                "control_subject": describe(finding.methodology_id)["checklist_title"],
                "why": explanation["why_flagged"],
                "misuse": explanation["misuse"],
                "impact": explanation["impact"],
                "recommendation": explanation["recommendation"],
                "validate": recommendation(finding.rule_id).get("validation_procedure")
                or "Confirm the configuration with the resource owner and recollect evidence.",
                "evidence": "; ".join(evidence) or "No evidence reference",
                "analyst_status": status,
                "analyst": decision.analyst if decision else "",
                "analyst_note": decision.note if decision else "",
                "stage": stage,
                "stage_label": STAGES[stage],
                "sre_verdict": VERDICTS[verdict.verdict] if verdict else "",
                "sre_contact": verdict.sre_contact if verdict else "",
                "sre_date": verdict.recorded_at.date().isoformat() if verdict else "",
                "sre_reason": verdict.reason if verdict else "",
            }
        )
    order = {s: i for i, s in enumerate(SEVERITIES)}
    result.sort(key=lambda r: (order.get(r["severity"], 9), r["provider"], r["scope"], r["title"]))
    return result


def _label(row: dict) -> str:
    return f"{row['scope_name']} ({row['scope']})" if row["scope_name"] else row["scope"]


def _safe(text: str) -> str:
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", text)[:80].strip("_") or "scope"


def _sheet(wb, title: str, headers: list[str], values: list[list], widths: dict | None = None):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet(title)
    ws.append([excel_literal(h) for h in headers])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E7E9EC")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in values:
        ws.append([excel_literal(v) if isinstance(v, str) else v for v in row])
    for column, header in enumerate(headers, 1):
        width = (widths or {}).get(header, 18)
        ws.column_dimensions[get_column_letter(column)].width = width
        if width >= 30:
            for (cell,) in ws.iter_rows(min_row=2, min_col=column, max_col=column):
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    if values:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(values) + 1}"
    return ws


WIDE = {
    "Title": 40,
    "Why this is flagged": 50,
    "How it can be misused (potential; not demonstrated)": 50,
    "Impact": 45,
    "Recommendation": 45,
    "How to validate": 45,
    "Evidence references": 40,
    "SRE BAU justification / comments": 45,
    "SRE comments": 40,
    "Analyst note": 40,
    "Related attack paths": 45,
}


def _definitions(wb, title: str, facts: list[tuple[str, str]]):
    from openpyxl.styles import Alignment, Font

    ws = wb.active
    ws.title = "Read me"
    ws.append([excel_literal(title)])
    ws["A1"].font = Font(bold=True, size=14)
    for key, value in facts:
        ws.append([excel_literal(key), excel_literal(value)])
    ws.append([])
    ws.append([excel_literal("Definitions")])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    for key, value in DEFINITIONS:
        ws.append([excel_literal(key), excel_literal(value)])
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 100
    for _, cell in ws.iter_rows(min_row=2, max_col=2):
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def sre_workbook(store, run: str, provider: str, scope: str) -> tuple[str, bytes]:
    """Workbook for one account/project listing its findings that still need an SRE verdict."""
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation

    if provider not in {"AWS", "GCP"} or not scope or len(scope) > 200:
        raise InputError("INVALID_SRE_SCOPE")
    selected = [
        r for r in rows(store, run, provider) if r["scope"] == scope and r["stage"] in OPEN_FOR_SRE
    ]
    label = _label(selected[0]) if selected else scope
    wb = Workbook()
    _definitions(
        wb,
        f"SRE validation request: {provider} {label}",
        [
            ("Account / project", f"{provider} {label}"),
            ("Assessment run", run),
            ("Generated (UTC)", datetime.now(UTC).isoformat(timespec="seconds")),
            ("Findings to validate", str(len(selected))),
            (
                "What to do",
                "For each finding, check it against the intended (BAU) configuration and fill "
                "SRE verdict, SRE name, Date and comments. Choose 'No gap (BAU)' only with a "
                "justification. Return the workbook to the security analyst, who records the "
                "verdicts in CloSec.",
            ),
        ],
    )
    headers = [
        "Finding ID",
        "Severity",
        "Title",
        "Checklist control",
        "Region",
        "Resource",
        "Why this is flagged",
        "How it can be misused (potential; not demonstrated)",
        "Impact",
        "Recommendation",
        "How to validate",
        "Evidence references",
        "Current stage",
        "SRE verdict",
        "SRE name",
        "Date",
        "SRE BAU justification / comments",
    ]
    ws = _sheet(
        wb,
        "Findings for SRE",
        headers,
        [
            [
                r["id"],
                r["severity"],
                r["title"],
                r["control"],
                r["region"],
                r["resource"],
                r["why"],
                r["misuse"],
                r["impact"],
                r["recommendation"],
                r["validate"],
                r["evidence"],
                r["stage_label"],
                "",
                r["sre_contact"],
                "",
                r["sre_reason"],
            ]
            for r in selected
        ],
        WIDE,
    )
    if selected:
        choice = DataValidation(
            type="list", formula1='"' + ",".join(SRE_CHOICES) + '"', allow_blank=True
        )
        ws.add_data_validation(choice)
        choice.add(f"N2:N{len(selected) + 1}")
    buffer = io.BytesIO()
    wb.save(buffer)
    return f"CloSec-SRE-{provider}-{_safe(label)}.xlsx", buffer.getvalue()


def sre_workbooks_zip(store, run: str, provider: str = "") -> bytes:
    """One SRE workbook per account/project that has findings open for SRE."""
    scopes = sorted(
        {
            (r["provider"], r["scope"])
            for r in rows(store, run, provider)
            if r["stage"] in OPEN_FOR_SRE
        }
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for scope_provider, scope in scopes:
            if scope:
                name, data = sre_workbook(store, run, scope_provider, scope)
                archive.writestr(name, data)
    return buffer.getvalue()


def finalized(store, run: str, provider: str = "", scope: str = "") -> dict:
    """Confirmed vulnerabilities, consolidated, with the attack paths that involve them."""
    from offline_analyzer.attack_paths.diagram import diagram, diagram_text
    from offline_analyzer.attack_paths.feasibility import attack_paths
    from offline_analyzer.findings.finalisation import filtered, frozen

    edition = frozen(store, run)
    if edition:
        return filtered(edition, "view", provider, scope)

    confirmed = [
        r
        for r in rows(store, run, provider)
        if r["stage"] == "CONFIRMED_VULNERABILITY" and (not scope or r["scope"] == scope)
    ]
    ids = {r["id"] for r in confirmed}
    paths = [
        p
        for p in attack_paths(store, run)["items"]
        if ids & set(p.get("finding_ids") or [])
        and not p.get("excluded_dependency")
        and (not provider or p["provider"] == provider)
    ]
    cross = {}
    for fid, source, target, kind in store.rows(
        "SELECT json_extract_string(f.value,'$'),json_extract_string(e.data,'$.source_account'),"
        "json_extract_string(e.data,'$.destination_account'),"
        "json_extract_string(e.data,'$.relationship_type') "
        "FROM enterprise_edges e, json_each(e.data,'$.related_findings') f WHERE e.run_id=? "
        "AND json_extract_string(f.value,'$') IN (SELECT unnest(?))",
        [run, sorted(ids)],
    ):
        if source and target and source != target:
            cross.setdefault(fid, set()).add(f"{kind or 'relationship'}: {source} -> {target}")
    for r in confirmed:
        related = [
            f"[{p['feasibility']}] {p.get('attack_path') or ''} {p.get('title', '')}: "
            f"{diagram_text(p)}"
            for p in paths
            if r["id"] in (p.get("finding_ids") or [])
        ]
        r["paths"] = "; ".join(related + sorted(cross.get(r["id"], ())))
    by_account: dict[tuple, dict] = {}
    for r in confirmed:
        entry = by_account.setdefault(
            (r["provider"], r["scope"]),
            {
                "provider": r["provider"],
                "scope": r["scope"],
                "name": r["scope_name"],
                "severity": Counter(),
                "total": 0,
            },
        )
        entry["severity"][r["severity"]] += 1
        entry["total"] += 1
    from offline_analyzer.findings.exclusions import listing

    return {
        "run_id": run,
        "generated_at": datetime.now(UTC).isoformat(timespec="seconds"),
        "findings": confirmed,
        "excluded": sum(
            not scope or i["scope"] == scope for i in listing(store, run, provider)["items"]
        ),
        "severity": dict(Counter(r["severity"] for r in confirmed)),
        "accounts": [
            {**a, "severity": dict(a["severity"])}
            for a in sorted(by_account.values(), key=lambda a: (a["provider"], a["scope"]))
        ],
        "controls": [
            {"control": c, "count": n}
            for c, n in sorted(Counter(r["control"] for r in confirmed).items())
        ],
        "paths": [
            {
                "feasibility": p["feasibility"],
                "provider": p["provider"],
                "scope": p.get("project_id") or p.get("account_id") or "",
                "path": p.get("attack_path") or "",
                "title": p.get("title", ""),
                "kind": p.get("scope", "single resource"),
                "steps": diagram_text(p),
                "diagram": diagram(p),
                "reason": p.get("feasibility_reason", ""),
                "finding_ids": sorted(ids & set(p.get("finding_ids") or [])),
                "stride": p.get("stride", []),
                "stride_scenarios": p.get("stride_scenarios", []),
                **{k: p[k] for k in HOP_KEYS if k in p},
            }
            for p in paths
        ],
    }


def finalized_document(store, run: str, provider: str = "", scope: str = "") -> dict:
    """Confirmed vulnerabilities in the assessment report template, with any SRE response."""
    from offline_analyzer.findings.finalisation import filtered, frozen
    from offline_analyzer.reporting.assessment_report import model

    edition = frozen(store, run)
    if edition:
        return filtered(edition, "document", provider, scope)

    data = model(store, run, "confirmed", provider, scope)
    sre = latest_sre(store, run)
    for f in data["findings"]:
        record = sre.get(f["id"])
        if record:
            f["sre"] = (
                f"{VERDICTS[record.verdict]} by {record.sre_contact or 'unrecorded'} on "
                f"{record.recorded_at.date().isoformat()}. {record.reason}".strip()
            )
    data["report_title"] = "Confirmed Vulnerabilities Report"
    data["include"] = "Only findings the assessor marked valid (SRE response optional)"
    return data


def finalized_files(data: dict, kind: str) -> bytes:
    """Excel or Word bytes of `finalized_document` output in the shared template."""
    import tempfile
    from pathlib import Path

    from offline_analyzer.reporting.docx import write_docx
    from offline_analyzer.reporting.report_document import write_excel, write_rtf

    with tempfile.TemporaryDirectory() as folder:
        target = Path(folder) / ("report." + kind)
        {"xlsx": write_excel, "rtf": write_rtf, "docx": write_docx}[kind](data, target)
        return target.read_bytes()
