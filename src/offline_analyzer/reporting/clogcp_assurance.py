"""Additional read-only coverage, comparison and review workbook for CloGCP.

Existing assessment exports, their forty-column register and decisions are unchanged.
Unsupported families are explicit rows, never inferred as passing controls.
"""

import json
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from offline_analyzer.contracts.gcp_expanded_catalog import DATASETS, EXTENSIONS_22
from offline_analyzer.findings.clogcp_reconcile import reconcile, snapshots
from offline_analyzer.findings.groups import group_key
from offline_analyzer.ingestion.safety import excel_literal
from offline_analyzer.rules.registry import PROVIDER_RULES

CONTEXT = {
    "GCP::ServiceAccount": "Key subject identity and attached-identity context",
    "GCP::CustomRole": "Public custom-role / dataset-role permission resolution",
    "GCP::Project": "Collected project identity, hierarchy and Compute metadata if present",
    "GCP::AssetIAMSearchResult": "Organization/resource IAM and federation analysis",
    "GCP::AssetResourceSnapshot": (
        "Explicit target-proxy SSL attachment and Compute project metadata"
    ),
    "GCP::WorkloadIdentityPool": "Federation provider and pool-binding context",
    "GCP::ForwardingRule": "Explicit external target-proxy attachment for TLS analysis",
    "GCP::WorkloadIdentityProvider": "Federation issuer, condition and pool-binding context",
    "GCP::OrgPolicy": "Effective Storage public-access prevention context",
    "GCP::LoggingSink": "Exclusion derivation and audit-routing review context",
}

# These are consumers of the collected family, not additional controls on that family.
CONTEXT_RULES = {
    "GCP::ForwardingRule": {"GCP-LB-TLS"},
    "GCP::WorkloadIdentityProvider": {"GCP-ORG-FEDERATION"},
    "GCP::OrgPolicy": {"GCP-STORAGE-PUBLIC", "GCP-STORAGE-PAP", "GCP-ORG-PUBLIC-DATA"},
    "GCP::ServiceAccount": {"GCP-KEY-PRIVILEGED"},
    "GCP::CustomRole": {"GCP-IAM-CUSTOM-PUBLIC", "GCP-BQ-PUBLIC-DATA", "GCP-ORG-PUBLIC-DATA"},
    "GCP::Project": {"GCP-VM-OSLOGIN", "GCP-VM-SERIAL-PORT", "GCP-VM-PROJECT-SSH-KEYS"},
    "GCP::AssetResourceSnapshot": {
        "GCP-LB-TLS",
        "GCP-VM-OSLOGIN",
        "GCP-VM-SERIAL-PORT",
        "GCP-VM-PROJECT-SSH-KEYS",
    },
    "GCP::WorkloadIdentityPool": {"GCP-ORG-FEDERATION", "GCP-WIF-SHARED-ISSUER"},
    "GCP::LoggingSink": {"GCP-LOG-EXCLUSION"},
    "GCP::Bucket": {"GCP-STORAGE-PUBLIC", "GCP-ORG-PUBLIC-DATA"},
    "GCP::RunService": {"GCP-ORG-PUBLIC-SERVERLESS", "GCP-FUNCTION-PUBLIC"},
    "GCP::ComputeInstance": {
        "GCP-NET-FW-ADMIN-INTERNET",
        "GCP-NET-FW-DB-INTERNET",
        "GCP-NET-FW-ANY-INTERNET",
    },
    "GCP::IAMPolicy": {
        "GCP-RUN-PUBLIC",
        "GCP-FUNCTION-PUBLIC",
        "GCP-IAM-CUSTOM-PUBLIC",
        "GCP-LOG-AUDIT-COVERAGE",
    },
    "GCP::IAMBinding": {
        "GCP-RUN-RUNTIME-PRIVILEGE",
        "GCP-FUNCTION-RUNTIME-PRIVILEGE",
        "GCP-GKE-NODE-PRIVILEGE",
    },
}

UNSUPPORTED = {
    "GCP::IAMDenyPolicy": (
        "Deny-policy condition/exception and effective-authorization evaluator not implemented."
    ),
    "GCP::PrincipalAccessBoundaryBinding": (
        "Effective principal access boundary evaluation not implemented."
    ),
    "GCP::Disk": (
        "Encryption metadata retained; CMEK/CSEK requirement, key availability "
        "and resource IAM are not established."
    ),
    "GCP::Snapshot": (
        "Snapshot encryption/IAM and approved retention/copy policy require additional evidence."
    ),
    "GCP::Route": (
        "Routes alone do not establish effective reachability, listeners or "
        "approved routing intent."
    ),
    "GCP::SSLPolicy": (
        "Target-proxy attachment and external forwarding chain required for live TLS findings."
    ),
    "GCP::LogMetric": (
        "Approved detection event catalog, filter semantics and alert "
        "references require manual verification."
    ),
    "GCP::NotificationChannel": (
        "Delivery, ownership and escalation routing cannot be established from channel metadata."
    ),
    "GCP::SCCSource": (
        "Source listing does not establish detector enablement, entitlement or "
        "complete finding visibility."
    ),
    "GCP::KMSKeyVersion": (
        "Key-version state alone does not establish insecure use or "
        "availability of protected resources."
    ),
}


def coverage_matrix(store: Any = None, run: str | None = None) -> list[dict]:
    rules = [r for r in PROVIDER_RULES if r.provider == "GCP"]
    observations = list(store.observations(run, provider="GCP")) if store else []
    results = list(store.results(run)) if store else []
    sources = list(store.source_data(run)) if store else []
    kinds = defaultdict(list)
    evaluations = defaultdict(list)
    source_families = defaultdict(list)
    for o in observations:
        kinds[o.resource_type].append(o)
    for r in results:
        evaluations[r.rule_id].append(r)
    for s in sources:
        family = (s.get("gcp_metadata") or {}).get("dataset_contract", {}).get("family")
        source_families[family].append(s)
    rows = []
    for d in DATASETS + EXTENSIONS_22:
        kind = d["resource_type"]
        relevant = [
            r
            for r in rules
            if r.resource_type == kind
            or kind == "GCP::IAMPolicy"
            and r.resource_type == "GCP::IAMBinding"
        ]
        # Org search consumers have resource-specific executable selectors, not catalog-kind rules.
        if kind == "GCP::AssetIAMSearchResult":
            relevant = [r for r in rules if r.rule_id.startswith("GCP-ORG-")]
        ss = source_families.get(d["family"], [])
        counts = Counter(str(e.outcome) for r in relevant for e in evaluations[r.rule_id])
        rows.append(
            {
                "dataset_family": d["family"],
                "domain": d["domain"],
                "operation": d["operation"],
                "normalized_object": kind,
                "declared_support": d["analysis_support"],
                "security_checks": "; ".join(r.rule_id for r in relevant),
                "context_consumers": "; ".join(sorted(CONTEXT_RULES.get(kind, set()))),
                "applicable_criteria": "; ".join(r.title for r in relevant),
                "analysis_status": "CRITERIA_IMPLEMENTED"
                if relevant
                else "CONTEXT_CONSUMER"
                if kind in CONTEXT
                else "UNSUPPORTED_AUTOMATED_ANALYSIS",
                "context_use": CONTEXT.get(kind, ""),
                "limitation": UNSUPPORTED.get(
                    kind,
                    (
                        "Only the named criteria are implemented; remaining service controls "
                        "require review."
                    )
                    if relevant
                    else (
                        "No defensible automatic control has been implemented for this metadata"
                        " family; approved requirements and effective configuration context are"
                        " needed."
                    ),
                ),
                "datasets_imported": len(ss),
                "collection_states": dict(Counter(s["effective_coverage"] for s in ss)),
                "resource_observations": len(kinds[kind]),
                "evaluator_executions": sum(counts.values()),
                "outcomes": dict(counts),
                "findings": counts.get("FAIL", 0),
                "report_output": (
                    "GCP_Observation_Register → GCP_Evidence_Index → GCP_Evidence_Details"
                )
                if relevant
                else "Unassessed Families / Dataset Coverage",
                "qualification": (
                    "Execution/resource counts are type/criterion totals; repeated catalog "
                    "aliases are not independent assessments."
                ),
            }
        )
    review_only = {
        "GCP-SECRET-LIFECYCLE",
        "GCP-ARTIFACT-TAGS",
        "GCP-BUILD-APPROVAL",
        "GCP-PUBSUB-PUSH-AUTH",
        "GCP-GKE-SECRETS-ENCRYPTION",
        "GCP-LB-ARMOR",
        "GCP-LOG-RETENTION",
        "GCP-LOG-AUDIT-COVERAGE",
    }
    for row in rows:
        ids = set(row["security_checks"].split("; ")) - {""}
        if ids and ids <= review_only:
            row["analysis_status"] = "REVIEW_ONLY"
    if store is None:
        return rows
    # Scope counts to the selected project and this exact family's source artifacts.
    by_project_family, by_artifact, by_project_rule = (
        defaultdict(list),
        defaultdict(list),
        defaultdict(list),
    )
    for s in sources:
        family = (s.get("gcp_metadata") or {}).get("dataset_contract", {}).get("family")
        by_project_family[(s.get("project_id"), family)].append(s)
    for o in observations:
        by_artifact[o.evidence.artifact_id].append(o)
    for result in results:
        by_project_rule[(result.project_id, result.rule_id)].append(result)
    finding_keys = {(f.project_id, f.resource_id, f.rule_id) for f in store.candidates(run)}
    scoped = []
    for project in sorted({s.get("project_id") for s in sources} - {None}):
        for row in rows:
            family_sources = by_project_family[(project, row["dataset_family"])]
            ids = set(row["security_checks"].split("; ")) - {""}
            observed = [o for s in family_sources for o in by_artifact[s.get("artifact_id")]]
            resources = {o.resource_id for o in observed if o.resource_id}
            artifacts = {s.get("artifact_id") for s in family_sources} - {None}
            consumers = set(row["context_consumers"].split("; ")) - {""}
            executions = [
                r
                for rid in ids
                for r in by_project_rule[(project, rid)]
                if r.resource_id in resources or any(e.artifact_id in artifacts for e in r.evidence)
            ]
            context_executions = [
                r
                for rid in consumers
                for r in by_project_rule[(project, rid)]
                if any(e.artifact_id in artifacts for e in r.evidence)
            ]
            counts = dict(Counter(str(r.outcome) for r in executions))
            scoped.append(
                {
                    **row,
                    "project": project,
                    "datasets_imported": len(family_sources),
                    "collection_states": dict(
                        Counter(s["effective_coverage"] for s in family_sources)
                    ),
                    "resource_observations": len(observed),
                    "evaluator_executions": len(executions),
                    "context_evidence_executions": len(context_executions),
                    "outcomes": counts,
                    "findings": sum(
                        (project, r.resource_id, r.rule_id) in finding_keys for r in executions
                    ),
                    "qualification": (
                        "Counts use this project and family source artifacts/canonical "
                        "subjects; aliases can reference the same evaluation. Scope gaps "
                        "appear separately."
                    ),
                }
            )
    return scoped


def write_sheet(wb: Workbook, title: str, rows: list[dict]) -> None:
    if not hasattr(wb, "clogcp_rows"):
        wb.clogcp_rows = {}
    wb.clogcp_rows[title] = rows
    ws = wb.create_sheet(title)
    if not rows:
        rows = [
            {
                "status": "NO_RECORDS",
                "limitation": "No records in this selected assessment; no assurance is implied.",
            }
        ]
    headers = list(rows[0])
    ws.append(headers)
    for row in rows:
        cells = []
        for h in headers:
            value = row.get(h)
            if isinstance(value, (dict, list, tuple)):
                value = json.dumps(value, sort_keys=True, ensure_ascii=False)
            if isinstance(value, str):
                import re

                value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)
                if len(value) > 32000:
                    value = value[:31800] + " [SHORTENED; full value retained in companion JSON]"
                value = excel_literal(value)
            cells.append(value)
        ws.append(cells)
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
    for col in ws.columns:
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(
            65, max(18, len(str(col[0].value)) + 2)
        )
        for c in col[1:]:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    table = Table(displayName="CloGCP_" + str(len(wb.worksheets)), ref=ws.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)


def export_assurance(store: Any, run: str, output: Path, history=None) -> dict:
    if output.exists() or output.with_suffix(".json").exists():
        raise FileExistsError("CloGCP assurance output must be a new file")
    matrix = coverage_matrix(store, run)
    current = snapshots(store, run)
    comparison = reconcile(current, history)
    findings = list(store.candidates(run))
    results = list(store.results(run))
    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "Dataset Coverage", matrix)
    write_sheet(
        wb,
        "Unassessed Families",
        [r for r in matrix if r["analysis_status"] != "CRITERIA_IMPLEMENTED"],
    )
    write_sheet(wb, "Finding Comparison", comparison)
    distribution = Counter(
        (f.project_id, f.security_domain, f.suggested_severity)
        for f in findings
        if f.provider == "GCP"
    )
    write_sheet(
        wb,
        "Domain Distribution",
        [
            dict(project=p, domain=d, criticality=s, findings=n)
            for (p, d, s), n in sorted(distribution.items())
        ],
    )
    write_sheet(
        wb,
        "Insufficient Evidence",
        [
            dict(
                project=r.project_id,
                resource=r.native_id,
                rule=r.rule_id,
                status=str(r.outcome),
                reason=r.reasoning,
                missing=r.missing_dependencies,
                limitations=r.limitations,
                evidence=[e.model_dump() for e in r.evidence],
            )
            for r in results
            if r.provider == "GCP"
            and r.outcome in {"UNKNOWN", "COLLECTION_ERROR", "MANUAL_VALIDATION_REQUIRED"}
        ],
    )
    write_sheet(
        wb,
        "High Critical Review",
        [
            dict(
                project=f.project_id,
                resource=f.native_id,
                rule=f.rule_id,
                criticality=f.suggested_severity,
                finding_id=f.finding_id,
                status="PENDING_ASSESSOR_REVIEW",
                observed=f.technical_observation,
                evidence=[e.model_dump() for e in f.evidence],
            )
            for f in findings
            if f.provider == "GCP" and f.suggested_severity in {"HIGH", "CRITICAL"}
        ],
    )
    grouped = defaultdict(list)
    for f in findings:
        if f.provider == "GCP":
            grouped[group_key(f)].append(f)
    write_sheet(
        wb,
        "Unique Observations",
        [
            dict(
                rule=rid,
                observation=[f.technical_observation for f in members],
                member_count=len(members),
                projects=sorted({f.project_id for f in members}),
                finding_ids=[f.finding_id for f in members],
                resources=[f.native_id for f in members],
                qualification=(
                    "Grouping preserves all resource-level finding IDs; it does not infer "
                    "relationships."
                ),
            )
            for (rid, _kind, _value), members in sorted(grouped.items())
        ],
    )
    with output.with_suffix(".json").open("x") as file:
        json.dump(wb.clogcp_rows, file, indent=2, ensure_ascii=False)
    wb.save(output)
    return {
        "coverage_families": len({r["dataset_family"] for r in matrix}),
        "project_family_rows": len(matrix),
        "comparison_records": len(comparison),
        "findings": len(findings),
    }
