"""Local-only lossless cell inventory. Outputs are sensitive, not sanitized."""

import argparse
import json
import tempfile
from datetime import date, datetime
from pathlib import Path
from zipfile import BadZipFile, ZipFile

import defusedxml  # noqa: F401 -- dependency required for openpyxl XML hardening
from openpyxl import load_workbook
from openpyxl.xml import DEFUSEDXML

from offline_analyzer.ingestion.safety import sha256_file, validate_relative_path

MAX_FILE_BYTES = 64 * 1024 * 1024
MAX_EXPANDED_BYTES = 256 * 1024 * 1024
MAX_CELLS = 2_000_000


def preflight(path: Path) -> None:
    if path.suffix.lower() != ".xlsx":
        raise ValueError("only XLSX accepted")
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ValueError("workbook exceeds byte budget")
    if not DEFUSEDXML:
        raise RuntimeError("XML hardening unavailable")
    try:
        with ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > 4096:
                raise ValueError("archive member budget exceeded")
            names = [member.filename for member in members]
            if len(names) != len(set(names)):
                raise ValueError("duplicate ZIP members")
            total = 0
            for member in members:
                validate_relative_path(member.filename)
                lower = member.filename.lower()
                if (
                    "vbaproject" in lower
                    or "externallinks/" in lower
                    or "embeddings/" in lower
                    or member.flag_bits & 1
                ):
                    raise ValueError("active/external/encrypted workbook content rejected")
                total += member.file_size
                if total > MAX_EXPANDED_BYTES or member.file_size > MAX_EXPANDED_BYTES:
                    raise ValueError("expanded byte budget exceeded")
                if member.file_size > max(1, member.compress_size) * 1000:
                    raise ValueError("compression ratio budget exceeded")
    except BadZipFile as exc:
        raise ValueError("invalid XLSX archive") from exc


def cell_value(value):
    if isinstance(value, (date, datetime)):
        return {"iso_date": value.isoformat()}
    return value


def extract(source: Path, output: Path) -> None:
    """Snapshot source; bound cells; stage outputs; never print source content."""
    if output.exists():
        raise FileExistsError("output directory already exists")
    original_hash = sha256_file(source)
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=output.parent) as temp:
        staging = Path(temp)
        snapshot = staging / "source.xlsx"
        with source.open("rb") as incoming, snapshot.open("wb") as outgoing:
            copied = 0
            while chunk := incoming.read(1024 * 1024):
                copied += len(chunk)
                if copied > MAX_FILE_BYTES:
                    raise ValueError("snapshot byte budget exceeded")
                outgoing.write(chunk)
        if sha256_file(snapshot) != original_hash:
            raise ValueError("source changed during snapshot")
        preflight(snapshot)
        result = staging / "result"
        result.mkdir()
        workbook = load_workbook(snapshot, read_only=True, data_only=False, keep_links=False)
        summary = []
        visited = 0
        try:
            with (
                (result / "cells.jsonl").open("w", encoding="utf-8") as stream,
                (result / "coverage-template.jsonl").open("w", encoding="utf-8") as coverage,
            ):
                for sheet in workbook.worksheets:
                    # Untrusted dimension metadata may hide populated cells: scan actual XML.
                    sheet.reset_dimensions()
                    nonempty = 0
                    rows = 0
                    for row_number, row in enumerate(sheet.iter_rows(), start=1):
                        rows = row_number
                        visited += max(1, len(row))
                        if visited > MAX_CELLS:
                            raise ValueError("cell/row budget exceeded")
                        coordinates = []
                        for col_number, cell in enumerate(row, start=1):
                            if cell.value is None:
                                continue
                            nonempty += 1
                            coordinates.append(cell.coordinate)
                            record = {
                                "sheet": sheet.title,
                                "row": row_number,
                                "column": col_number,
                                "coordinate": cell.coordinate,
                                "type": cell.data_type,
                                "value": cell_value(cell.value),
                            }
                            stream.write(json.dumps(record, ensure_ascii=False) + "\n")
                        if coordinates:
                            coverage.write(
                                json.dumps(
                                    {
                                        "source_sha256": original_hash,
                                        "sheet": sheet.title,
                                        "row": row_number,
                                        "cells": coordinates,
                                        "mapping_status": "UNMAPPED",
                                        "record_kind": None,
                                        "original_id": None,
                                        "requirement": None,
                                        "domain": None,
                                        "subdomain": None,
                                        "evidence": None,
                                        "collection_procedure": None,
                                        "validation": None,
                                        "pass_fail": None,
                                        "source_automation": None,
                                        "dependencies": None,
                                        "cross_references": None,
                                        "severity_criteria": None,
                                        "reporting_fields": None,
                                        "aws_scope": None,
                                        "feasibility": "UNREVIEWED",
                                        "limitations": ["Source semantics require review"],
                                        "implementation_status": "NOT_IMPLEMENTED",
                                        "rule_version": None,
                                        "reviewer": None,
                                        "quality_flags": [],
                                    },
                                    ensure_ascii=False,
                                )
                                + "\n"
                            )
                    summary.append(
                        {"sheet": sheet.title, "rows_scanned": rows, "nonempty_cells": nonempty}
                    )
        finally:
            workbook.close()
        if sha256_file(source) != original_hash:
            raise ValueError("source changed during extraction")
        inventory = {
            "schema_version": "1.0",
            "parser_version": "0.0.1",
            "source_sha256": original_hash,
            "worksheets": summary,
            "semantic_counts": None,
            "limitations": [
                "Semantic table mapping and coverage audit require review",
                "Formulas preserved as inert text, never evaluated",
            ],
        }
        (result / "inventory.json").write_text(json.dumps(inventory, indent=2), encoding="utf-8")
        result.rename(output)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    try:
        extract(args.source, args.output)
    except (ValueError, OSError, RuntimeError):
        # Source exceptions can contain customer sheet names/paths. Keep CLI silent.
        parser.exit(1, "Extraction failed; inspect locally without sharing private contents.\n")
    print("Local extraction complete. Output is private and has not been sanitized.")


if __name__ == "__main__":
    main()
