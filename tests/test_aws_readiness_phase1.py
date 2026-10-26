"""AWS readiness phase 1 (shared code): truncation honesty, per-scope snapshot membership,
exclusion-aware enterprise counts and account-qualified group keys. Synthetic data only."""

import json

import pytest
from openpyxl import load_workbook

from offline_analyzer.attack_paths import feasibility
from offline_analyzer.attack_paths.feasibility import attack_paths
from offline_analyzer.contracts.exclusion import ExclusionRequest
from offline_analyzer.enterprise import correlate as correlation
from offline_analyzer.enterprise.correlate import correlate
from offline_analyzer.enterprise.queries import EnterpriseViews
from offline_analyzer.enterprise.snapshot import (
    assess_enterprise,
    members,
    partially_superseded,
    superseded_imports,
)
from offline_analyzer.enterprise.validation import validate_snapshot
from offline_analyzer.findings.exclusions import record
from offline_analyzer.findings.groups import groups
from offline_analyzer.iam import analysis as iam_analysis
from offline_analyzer.ingestion.importer import import_package
from offline_analyzer.reporting.assessment_report import model
from offline_analyzer.reporting.excel import export_excel
from offline_analyzer.rules.engine import assess
from offline_analyzer.storage.store import Store
from tools.fixtures.aws_readiness import R1, R2, A, B, iam, network, write
from tools.fixtures.r12_demo import build as r12_build
from tools.gcp_fixture import build as gcp_build

T1, T2, T3 = "2026-10-01T00:00:00+00:00", "2026-10-05T00:00:00+00:00", "2026-10-08T00:00:00+00:00"


def cells(path, sheet):
    book = load_workbook(path, read_only=True)
    values = [str(v) for row in book[sheet].iter_rows(values_only=True) for v in row if v]
    book.close()
    return values


@pytest.fixture
def r12(tmp_path):
    r12_build(tmp_path / "r12")
    store = Store(tmp_path / "work")
    run = import_package(store, tmp_path / "r12")
    assess(store, run)
    yield store, run, tmp_path
    store.__exit__(None, None, None)


# --- 1. truncation visibility -------------------------------------------------------------


def test_untruncated_views_carry_no_truncation_record(r12):
    store, run, tmp_path = r12
    paths = attack_paths(store, run)
    assert paths["truncated"] is False and "truncation" not in paths
    assert not any(n["code"].startswith("TRUNCATED_") for n in model(store, run)["coverage_notes"])
    export_excel(store, run, tmp_path / "w.xlsx")
    assert "Truncated view" not in cells(tmp_path / "w.xlsx", "IAM Evidence Gaps")


def test_iam_budgets_are_reported_with_shown_and_total(r12, monkeypatch):
    store, run, tmp_path = r12
    monkeypatch.setattr(iam_analysis, "MAX_TRANSITIONS", 2)
    monkeypatch.setattr(iam_analysis, "MAX_PASSROLE", 2)
    paths = attack_paths(store, run)
    assert paths["truncated"] is True
    (passrole,) = [e for e in paths["truncation"] if e["budget"] == "PassRole investigation budget"]
    assert passrole["shown"] == 2 and passrole["total"] > passrole["shown"]
    assert passrole["dropped"] == passrole["total"] - passrole["shown"]
    assert passrole["provider"] == "AWS" and "AP-06" in passrole["effect"]
    notes = [n for n in model(store, run, "all")["coverage_notes"] if n["provider"] == "AWS"]
    assert any(n["code"] == "TRUNCATED_PASSROLE_INVESTIGATION_BUDGET" for n in notes)
    export_excel(store, run, tmp_path / "w.xlsx")
    gaps = cells(tmp_path / "w.xlsx", "IAM Evidence Gaps")
    assert any("AssumeRole investigation budget: 2 of" in g for g in gaps)
    assert any("PassRole investigation budget: 2 of" in g for g in gaps)


def test_transition_budget_counts_every_unevaluated_pair(r12, monkeypatch):
    store, run, _ = r12
    full = iam_analysis.IamAnalysis(store, run, list(store.source_data(run))).transitions()
    monkeypatch.setattr(iam_analysis, "MAX_TRANSITIONS", 3)
    limited = iam_analysis.IamAnalysis(store, run, list(store.source_data(run)))
    assert len(limited.transitions()) == 3
    (entry,) = limited.truncation
    assert entry["total"] == len(full) and entry["total_basis"].startswith("exact")


def test_gcp_investigation_display_limit_is_counted(tmp_path, monkeypatch):
    with Store(tmp_path / "work") as store:
        run = import_package(store, gcp_build(tmp_path / "gcp", "demo-alpha", "123456789012"))
        assess(store, run)
        everything = attack_paths(store, run)
        investigations = [p for p in everything["items"] if p["provider"] == "GCP"]
        assert len(investigations) >= 2 and "truncation" not in everything
        monkeypatch.setattr(feasibility, "INVESTIGATION_LIMIT", 1)
        limited = attack_paths(store, run)
        (entry,) = limited["truncation"]
        assert limited["truncated"] and entry["shown"] == 1
        assert entry["total"] == len(investigations) and entry["provider"] == "GCP"


def test_enterprise_relationship_budget_is_reported(r12, monkeypatch):
    store, run, tmp_path = r12
    monkeypatch.setattr(correlation, "MAX_EDGES", 3)
    summary = correlate(store, run)
    (entry,) = [e for e in summary["truncation_detail"] if e["provider"] == "Enterprise"]
    assert entry["shown"] == 3 and entry["total"] > 3
    assert (
        store.db.execute("SELECT count(*) FROM enterprise_edges WHERE run_id=?", [run]).fetchone()[
            0
        ]
        == 3
    )
    notes = model(store, run, "all")["coverage_notes"]
    assert any(n["code"] == "TRUNCATED_ENTERPRISE_RELATIONSHIP_BUDGET" for n in notes)
    export_excel(store, run, tmp_path / "w.xlsx")
    assert any(
        "Enterprise relationship budget: 3 of" in g
        for g in cells(tmp_path / "w.xlsx", "IAM Evidence Gaps")
    )


# --- 2. snapshot membership -----------------------------------------------------------------


def stale_estate(tmp_path):
    """Old import of A+B; newer import of A (SSH restricted); newest import of A Region R1
    only (SSH restricted, IMDSv2 required)."""
    old = iam(A, B) + iam(B, A) + network(A, R1) + network(B, R1) + network(A, R2)
    newer = iam(A, B) + network(A, R1, open_ssh=False) + network(A, R2)
    newest = network(A, R1, open_ssh=False, tokens="required")
    store = Store(tmp_path / "work")
    runs = [
        import_package(store, write(tmp_path / name, items, when))
        for name, items, when in (("old", old, T1), ("newer", newer, T2), ("newest", newest, T3))
    ]
    for run in runs:
        assess(store, run)
    return store, runs


def sg_results(store, run):
    return {
        (r.account_id, r.region): str(r.outcome)
        for r in store.results(run)
        if r.rule_id == "R1-EC2-SG-SENSITIVE-INGRESS" and r.evaluation_scope == "RESOURCE"
    }


def test_multi_account_import_never_downgrades_a_newer_result(tmp_path):
    store, (old, newer, newest) = stale_estate(tmp_path)
    try:
        # The old import is still the latest for B: it stays a member, partly superseded.
        assert superseded_imports(store) == {}
        assert set(members(store)) == {old, newer, newest}
        lost = partially_superseded(store)[old]
        # Every dataset of A is newer elsewhere; none of B is.
        assert all(label.startswith(A + " ") for label in lost)
        assert lost[f"{A} global iam:GetAccountAuthorizationDetails"] == newer
        assert lost[f"{A} {R1} ec2:DescribeSecurityGroups"] == newest
        assert lost[f"{A} {R2} ec2:DescribeSecurityGroups"] == newer
        snapshot = assess_enterprise(store)
        results = sg_results(store, snapshot)
        # Previously "Native snapshots disagree" (UNKNOWN) from the old A snapshot.
        assert results[(A, R1)] == sg_results(store, newest)[(A, R1)]
        assert results[(A, R1)] == "MANUAL_VALIDATION_REQUIRED"
        assert results[(B, R1)] == "FAIL" and results[(A, R2)] == "FAIL"
        imds = {
            (r.account_id, r.region): str(r.outcome)
            for r in store.results(snapshot)
            if r.rule_id == "R0-EC2-IMDSV2" and r.evaluation_scope == "RESOURCE"
        }
        assert imds[(A, R1)] == "PASS" and imds[(B, R1)] == "FAIL"
        check = validate_snapshot(store, snapshot)
        assert check["status"] in {"PASS", "WARNING"} and check["superseded_sources"][old] > 0
        assert all(c["status"] == "PASS" for c in check["checks"])
        assert not any(
            r.outcome == "UNKNOWN" and "disagree" in r.reasoning for r in store.results(snapshot)
        )
    finally:
        store.__exit__(None, None, None)


def test_runs_view_lists_partly_superseded_scopes(tmp_path):
    from offline_analyzer.api.services import Views

    store, (old, newer, newest) = stale_estate(tmp_path)
    try:
        listed = {r["run_id"]: r for r in Views(store.workspace).runs(store)}
        scopes = listed[old]["superseded_scopes"]
        assert scopes[f"{A} {R1} ec2:DescribeSecurityGroups"] == newest
        assert scopes[f"{A} global iam:GetAccountAuthorizationDetails"] == newer
        assert "superseded_scopes" not in listed[newest] and not listed[old]["superseded_by"]
    finally:
        store.__exit__(None, None, None)


def test_single_account_and_per_region_imports(tmp_path):
    with Store(tmp_path / "work") as store:
        first = import_package(store, write(tmp_path / "a1", iam(A, B) + network(A, R1), T1))
        region = import_package(store, write(tmp_path / "a2", network(A, R2), T2))
        again = import_package(
            store, write(tmp_path / "a3", iam(A, B) + network(A, R1, open_ssh=False), T3)
        )
        for run in (first, region, again):
            assess(store, run)
        # Same account, newer full import: the old import is superseded, not included.
        assert superseded_imports(store) == {first: [again]}
        snapshot = assess_enterprise(store)
        # Per-Region imports of the same account merge: both Regions are assessed.
        assert set(sg_results(store, snapshot)) == {(A, R1), (A, R2)}
        assert sg_results(store, snapshot)[(A, R1)] == "MANUAL_VALIDATION_REQUIRED"


# --- 3. exclusions and 4. group keys ----------------------------------------------------------


def two_accounts(tmp_path):
    items = iam(A, B) + iam(B, A) + network(A, R1) + network(B, R1)
    store = Store(tmp_path / "work")
    run = import_package(store, write(tmp_path / "p", items, T1))
    assess(store, run)
    return store, assess_enterprise(store, [run])


def test_excluded_finding_leaves_enterprise_counts_and_references(tmp_path):
    store, snap = two_accounts(tmp_path)
    try:
        target = next(
            c
            for c in store.candidates(snap)
            if c.rule_id == "R1-EC2-SG-SENSITIVE-INGRESS" and c.account_id == B
        )
        views = EnterpriseViews()
        before = {a["account"]: a["findings"] for a in views.accounts(store, snap)}
        edges = views.relationships(store, snap, limit=200)["items"]
        assert any(target.finding_id in e["related_findings"] for e in edges)
        evaluations = store.db.execute(
            "SELECT count(*), sum(hash(data)) FROM evaluations WHERE run_id=?", [snap]
        ).fetchone()
        record(
            store,
            snap,
            target.finding_id,
            ExclusionRequest(action="EXCLUDE", reason="FALSE_POSITIVE", note="synthetic",
                             analyst="synthetic"),
        )  # fmt: skip
        after = {a["account"]: a["findings"] for a in views.accounts(store, snap)}
        assert after[B] == before[B] - 1 and after[A] == before[A]
        edges = views.relationships(store, snap, limit=200)["items"]
        assert not any(target.finding_id in e["related_findings"] for e in edges)
        assert any(e.get("excluded_related_findings") == 1 for e in edges)
        export_excel(store, snap, tmp_path / "w.xlsx")
        book = load_workbook(tmp_path / "w.xlsx", read_only=True)
        holding = [
            ws.title
            for ws in book
            if any(
                target.finding_id in str(v) for row in ws.iter_rows(values_only=True) for v in row
            )
        ]
        book.close()
        assert holding == ["Excluded Findings"]
        assert (
            store.db.execute(
                "SELECT count(*), sum(hash(data)) FROM evaluations WHERE run_id=?", [snap]
            ).fetchone()
            == evaluations
        )
    finally:
        store.__exit__(None, None, None)


def test_same_role_name_and_group_id_in_two_accounts_form_separate_groups(tmp_path):
    store, snap = two_accounts(tmp_path)
    try:
        found = groups(store, snap)
        sg = [g for g in found if g["rule_id"] == "R1-EC2-SG-SENSITIVE-INGRESS"]
        assert sorted(g["key"] for g in sg) == [
            f"{A}/{R1}/sg-00000000000000001",
            f"{B}/{R1}/sg-00000000000000001",
        ]
        assert all(g["member_count"] == 1 and g["affected_scopes"] == 1 for g in sg)
        trust = [g for g in found if g["rule_id"] == "R1-IAM-ROLE-TRUST-CONDITIONS"]
        assert sorted(g["key"] for g in trust) == [
            f"arn:aws:iam::{A}:role/shared-name-trust",
            f"arn:aws:iam::{B}:role/shared-name-trust",
        ]
        assert len({g["group_id"] for g in sg + trust}) == 4
        assert json.dumps(found)  # serialisable for the API
    finally:
        store.__exit__(None, None, None)
