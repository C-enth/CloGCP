"""Golden regression datasets for high-risk false-positive conditions (synthetic inputs).

Each scenario pins the source-criterion outcome, finding class, technical severity and
report rows. Outcomes follow the methodology clause; class/severity follow evidence.
"""

import json
from datetime import UTC, datetime

import pytest
from openpyxl import load_workbook

from offline_analyzer.cli import main, summary
from offline_analyzer.contracts.r0 import Package, Source
from offline_analyzer.findings.generate import record_decision, validation_status
from offline_analyzer.ingestion.importer import import_package
from offline_analyzer.ingestion.json_input import InputError
from offline_analyzer.ingestion.safety import sha256_file
from offline_analyzer.inventory.accounts import install_registry, workspace_registry
from offline_analyzer.reporting.excel import export_excel
from offline_analyzer.rules.engine import assess
from offline_analyzer.storage.store import Store
from tools.fixtures.r1_demo import database, group, permission, rid
from tools.fixtures.r11_demo import OK, SCP, TIME, A, C, build, registry_file


@pytest.fixture(scope="module")
def golden(tmp_path_factory):
    root = tmp_path_factory.mktemp("golden")
    build(root / "input")
    workspace = root / "workspace"
    with Store(workspace) as store:
        install_registry(workspace, registry_file(root / "registry.json"))
        run = import_package(store, root / "input")
        assess(store, run)
        report = root / "golden.xlsx"
        export_excel(store, run, report)
        yield store, run, report


def find(store, run, rule, name):
    for result in store.results(run):
        if result.rule_id != rule or not result.resource_id:
            continue
        obs = next(store.observations(run, resource_id=result.resource_id))
        if name in {
            obs.native_id,
            obs.configuration.get("RoleName"),
            obs.configuration.get("GroupId"),
        }:
            return result
    raise AssertionError(f"no result for {rule} {name}")


def candidate(store, run, result):
    return next((c for c in store.candidates(run) if c.resource_id == result.resource_id
                 and c.rule_id == result.rule_id), None)  # fmt: skip


IAM, SG, RDS = "R1-IAM-ROLE-TRUST-CONDITIONS", "R1-EC2-SG-SENSITIVE-INGRESS", "R1-RDS-PUBLIC-ACCESS"
BUCKET, BPA = "R1-S3-BUCKET-PUBLIC-ACCESS", "R1-S3-ACCOUNT-BPA"
GOLDEN = [
    # Cross-account trust: narrow internal principal, registry and permission context.
    (IAM, "golden-narrow-internal", "MANUAL_VALIDATION_REQUIRED", None, None),
    (IAM, "golden-vendor-admin", "FAIL", "POTENTIAL_EXPOSURE", "HIGH"),
    (IAM, "golden-internal-root", "FAIL", "BASELINE_DEVIATION", "LOW"),
    # R1.2: no RCP evidence, so outside-principal assumption is potential, never CRITICAL.
    (IAM, "golden-public-admin", "FAIL", "POTENTIAL_EXPOSURE", "HIGH"),
    (IAM, "golden-vendor-customer-admin", "FAIL", "POTENTIAL_EXPOSURE", "HIGH"),
    (IAM, "golden-vendor-unknown-policy", "FAIL", "POTENTIAL_EXPOSURE", "MEDIUM"),
    (IAM, "golden-listroles-vendor", "FAIL", "POTENTIAL_EXPOSURE", "MEDIUM"),
    # Security-group exposure: routed, blocked by routing, main-table fallback, no routes.
    (SG, rid("sg", 21), "FAIL", "POTENTIAL_EXPOSURE", "HIGH"),
    (SG, rid("sg", 22), "FAIL", "BASELINE_DEVIATION", "LOW"),
    (SG, rid("sg", 23), "FAIL", "POTENTIAL_EXPOSURE", "HIGH"),
    (SG, rid("sg", 41), "FAIL", "POTENTIAL_EXPOSURE", "MEDIUM"),
    (SG, rid("sg", 31), "PASS", None, None),
    # Public RDS: restricted SG, open+routed, SG evidence missing, open but private routing.
    (RDS, f"arn:aws:rds:us-east-1:{A}:db:golden-restricted", "FAIL", "BASELINE_DEVIATION", "LOW"),
    (RDS, f"arn:aws:rds:us-east-1:{A}:db:golden-open-routed", "FAIL", "POTENTIAL_EXPOSURE",
     "HIGH"),
    (RDS, f"arn:aws:rds:us-east-1:{A}:db:golden-sg-missing", "FAIL", "POTENTIAL_EXPOSURE",
     "MEDIUM"),
    (RDS, f"arn:aws:rds:us-east-1:{A}:db:golden-open-private", "FAIL", "BASELINE_DEVIATION", "LOW"),
    # IMDS: container hop limit, stopped instance, pending metadata change.
    ("R1-EC2-IMDS-HOP-LIMIT", rid("i", 22), "FAIL", "BASELINE_DEVIATION", "LOW"),
    ("R0-EC2-IMDSV2", rid("i", 22), "PASS", None, None),
    ("R0-EC2-IMDSV2", rid("i", 23), "FAIL", "BASELINE_DEVIATION", "MEDIUM"),
    ("R0-EC2-IMDSV2", rid("i", 24), "UNKNOWN", None, None),
    ("R1-EC2-IMDS-HOP-LIMIT", rid("i", 24), "UNKNOWN", None, None),
    # S3: account baseline vs effective bucket exposure.
    (BPA, A, "FAIL", "BASELINE_DEVIATION", "MEDIUM"),
    (BPA, C, "PASS", None, None),
    (BUCKET, "golden-restricted-bucket", "PASS", None, None),
    (BUCKET, "golden-public-policy", "FAIL", "POTENTIAL_EXPOSURE", "HIGH"),
    (BUCKET, "golden-public-acl", "FAIL", "POTENTIAL_EXPOSURE", "HIGH"),
    (BUCKET, "golden-object-acls", "MANUAL_VALIDATION_REQUIRED", None, None),
    (BUCKET, "golden-missing-bucket-bpa", "UNKNOWN", None, None),
    (BUCKET, "golden-account-restricted", "PASS", None, None),
]  # fmt: skip


@pytest.mark.parametrize("rule,name,outcome,finding_class,severity", GOLDEN)
def test_golden_conclusion(golden, rule, name, outcome, finding_class, severity):
    store, run, _ = golden
    result = find(store, run, rule, name)
    assert str(result.outcome) == outcome, result.reasoning
    found = candidate(store, run, result)
    if outcome != "FAIL":
        assert found is None  # no suggested finding without a failed criterion
        return
    assert (found.finding_class, found.suggested_severity) == (finding_class, severity)
    assert found.review_status == "CANDIDATE"  # never auto-confirmed
    assert found.methodology_id and found.methodology_criteria
    for claim in ("exploited", "compromised", "is reachable", "accessed by"):
        assert claim not in (found.technical_observation + (found.exposure_status or "")).lower()


def test_trust_findings_preserve_missing_scp_and_boundary(golden):
    store, run, _ = golden
    for name in ("golden-vendor-admin", "golden-internal-root", "golden-public-admin"):
        found = candidate(store, run, find(store, run, IAM, name))
        assert SCP in found.missing_evidence
        assert "does not establish effective access" in found.exposure_status
    found = candidate(store, run, find(store, run, IAM, "golden-vendor-customer-admin"))
    assert any("Permissions boundary" in c for c in found.compensating_controls)
    # Customer-managed admin policy, but the boundary document is missing: unresolved.
    assert "POTENTIALLY_PRIVILEGED" in found.exposure_status
    assert any("permissions boundary document" in e for e in found.missing_evidence)
    unknown = candidate(store, run, find(store, run, IAM, "golden-vendor-unknown-policy"))
    assert any("not-exported" in e for e in unknown.missing_evidence)
    narrow = find(store, run, IAM, "golden-narrow-internal")
    assert "registered assessment account" in narrow.reasoning


def test_routing_and_container_context(golden):
    store, run, _ = golden
    blocked = candidate(store, run, find(store, run, SG, rid("sg", 22)))
    assert "internet gateway" in blocked.exposure_status
    assert blocked.compensating_controls
    missing = candidate(store, run, find(store, run, SG, rid("sg", 41)))
    assert any("ec2:DescribeRouteTables" in e for e in missing.missing_evidence)
    routed = candidate(store, run, find(store, run, SG, rid("sg", 21)))
    assert "not demonstrated" in routed.exposure_status
    assert any("NACL" in e for e in routed.missing_evidence)
    hop = candidate(store, run, find(store, run, "R1-EC2-IMDS-HOP-LIMIT", rid("i", 22)))
    assert any("eks:cluster-name" in c for c in hop.compensating_controls)
    stopped = candidate(store, run, find(store, run, "R0-EC2-IMDSV2", rid("i", 23)))
    assert stopped.exposure_status.startswith("Instance state stopped")


def test_bucket_inventory_never_complete_and_account_baseline_is_not_exposure(golden):
    store, run, _ = golden
    scopes = [
        r
        for r in store.results(run)
        if r.rule_id == BUCKET and r.evaluation_scope == "COLLECTION_SCOPE"
    ]
    assert scopes and all(str(r.outcome) == "UNKNOWN" for r in scopes)
    baseline = candidate(store, run, find(store, run, BPA, A))
    assert "no bucket is shown to be public" in baseline.exposure_status


def test_report_rows_match_candidates(golden):
    store, run, report = golden
    wb = load_workbook(report)
    header = [c.value for c in wb["Findings Register"][3]]
    index = {
        name: header.index(name) for name in ("Finding ID", "Finding class", "Suggested severity")
    }
    rows = {
        r[index["Finding ID"]]: r
        for r in wb["Findings Register"].iter_rows(min_row=4, values_only=True)
    }
    candidates = list(store.candidates(run))
    assert set(rows) == {c.finding_id for c in candidates}
    for found in candidates:
        row = rows[found.finding_id]
        assert row[index["Finding class"]] == found.finding_class
        assert row[index["Suggested severity"]] == found.suggested_severity
    for sheet, finding_class in (("Security Gaps", "SUPPORTED_SECURITY_GAP"),
                                 ("Potential Exposures", "POTENTIAL_EXPOSURE"),
                                 ("Baseline Deviations", "BASELINE_DEVIATION")):  # fmt: skip
        ids = {r[0] for r in wb[sheet].iter_rows(min_row=4, values_only=True)}
        assert ids == {c.finding_id for c in candidates if c.finding_class == finding_class}
    incomplete = {
        r[0] for r in wb["Incomplete Investigations"].iter_rows(min_row=4, values_only=True)
    }
    assert incomplete == {
        r.evaluation_id
        for r in store.results(run)
        if str(r.outcome) in {"UNKNOWN", "COLLECTION_ERROR", "MANUAL_VALIDATION_REQUIRED"}
    }
    # Tracker staging: registry accounts without evidence are explicit, never Pass.
    tracker = list(wb["Tracker Staging"].iter_rows(min_row=4, values_only=True))
    assert "Pass" not in {r[3] for r in tracker}
    unassessed = [r for r in tracker if r[15] == "777777777777"]
    assert unassessed and {r[3] for r in unassessed} == {"Not Assessed"}
    assert {r[1] for r in unassessed} == {"ACCT-004"}
    assert {r[16] for r in tracker if r[15] == C} == {"NOT_IN_REGISTRY_MAPPING_REQUIRED"}
    assert all(cell.data_type != "f" for ws in wb for row in ws for cell in row)
    metrics = summary(store, run)
    assert metrics["registry_accounts_without_evidence"] == 2  # INTERNAL and 777...
    assert metrics["evidence_accounts_not_in_registry"] == 1  # account C
    wb.close()


def test_missing_collection_metadata_never_passes_scope(tmp_path):
    path = tmp_path / "groups.json"
    path.write_text(
        json.dumps({"SecurityGroups": [group(A, 1, "x", [permission(443, "0.0.0.0/0")])]})
    )
    declared = Source(path=path.name, sha256=sha256_file(path), format="aws-json-v1",
                      operation="ec2:DescribeSecurityGroups", account_id=A, region="us-east-1",
                      collected_at=TIME, coverage="COMPLETE", request_parameters=None)  # fmt: skip
    unscoped = declared.model_copy(update={"account_id": None, "region": None})
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, tmp_path, Package(sources=(declared,)))
        assess(store, run)
        assert next(store.source_data(run))["effective_coverage"] == "PARTIAL"
        scope = [
            r
            for r in store.results(run)
            if r.evaluation_scope == "COLLECTION_SCOPE" and r.rule_id == SG
        ]
        assert scope and all(str(r.outcome) == "UNKNOWN" for r in scope)
        run2 = import_package(store, tmp_path, Package(sources=(unscoped,)))
        assess(store, run2)
        assert any(
            r.evaluation_id.startswith("unassigned-source") and str(r.outcome) == "UNKNOWN"
            for r in store.results(run2)
        )


def prowler_row(status):
    return {
        "class_uid": 2004,
        "metadata": {"product": {"name": "Prowler", "version": "5.10.0"},
                     "event_code": "rds_instance_storage_encrypted"},
        "status_code": status, "status_id": 1, "status": "New", "severity": "Medium",
        "time_dt": TIME, "finding_info": {"uid": "golden-prowler-1"},
        "cloud": {"provider": "aws", "account": {"uid": A}, "region": "us-east-1"},
        "resources": [{"uid": f"arn:aws:rds:us-east-1:{A}:db:golden-contradiction",
                       "type": "AwsRdsDbInstance", "region": "us-east-1"}],
    }  # fmt: skip


def test_contradictory_prowler_and_native_observations(tmp_path):
    native = tmp_path / "rds.json"
    native.write_text(
        json.dumps(
            {"DBInstances": [database(A, "golden-contradiction", "postgres", False, False)], **OK}
        )
    )
    scanner = tmp_path / "prowler.json"
    scanner.write_text(json.dumps([prowler_row("PASS")]))
    sources = (
        Source(path=native.name, sha256=sha256_file(native), format="aws-json-v1",
               operation="rds:DescribeDBInstances", account_id=A, region="us-east-1",
               collected_at=TIME, coverage="COMPLETE", request_parameters={}),
        Source(path=scanner.name, sha256=sha256_file(scanner), format="prowler-ocsf-5.10",
               operation="prowler:Scan", account_id=A, region="us-east-1", collected_at=TIME),
    )  # fmt: skip
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, tmp_path, Package(sources=sources))
        assess(store, run)
        result = next(
            r for r in store.results(run) if r.rule_id == "R0-RDS-ENCRYPTION" and r.resource_id
        )
        assert str(result.outcome) == "FAIL"  # native configuration decides
        assert result.contradictory_observations and result.confidence == "MEDIUM"
        found = candidate(store, run, result)
        assert found.evidence_confidence == "MEDIUM" and found.suggested_severity == "MEDIUM"
        export_excel(store, run, tmp_path / "c.xlsx")
        wb = load_workbook(tmp_path / "c.xlsx")
        queue = [r for r in wb["Review Queue"].iter_rows(min_row=4, values_only=True)
                 if r[1] == result.evaluation_id]  # fmt: skip
        assert queue and queue[0][5] == "NEEDS_VALIDATION"
        wb.close()


def test_repeated_import_preserves_analyst_decisions(tmp_path):
    build(tmp_path / "input")
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, tmp_path / "input")
        assess(store, run)
        target = candidate(store, run, find(store, run, IAM, "golden-public-admin"))
        record_decision(store, run, target.finding_id, "CONFIRMED", "golden-analyst", "validated")
        before = [c.model_dump_json() for c in store.candidates(run)]
        assert import_package(store, tmp_path / "input") == run
        assess(store, run)
        assert [c.model_dump_json() for c in store.candidates(run)] == before
        assert validation_status(store, target)[0] == "CONFIRMED"
        others = [c for c in store.candidates(run) if c.finding_id != target.finding_id]
        assert all(validation_status(store, c)[0] == "CANDIDATE" for c in others)
        export_excel(store, run, tmp_path / "decided.xlsx")
        wb = load_workbook(tmp_path / "decided.xlsx")
        assert [r[0] for r in wb["Confirmed Findings"].iter_rows(min_row=4, values_only=True)] == [
            target.finding_id
        ]
        wb.close()
        package = Package.model_validate_json((tmp_path / "input" / "manifest.json").read_bytes())
        changed = package.model_copy(update={"sources": tuple(
            s.model_copy(update={"collected_at": datetime(2026, 10, 11, tzinfo=UTC)})
            for s in package.sources)})  # fmt: skip
        run2 = import_package(store, tmp_path / "input", changed)
        assess(store, run2)
        updated = next(c for c in store.candidates(run2) if c.finding_id == target.finding_id)
        assert validation_status(store, updated)[0] == "NEEDS_REVALIDATION"


@pytest.mark.parametrize(
    "mutate,code",
    [
        (lambda d: d["entries"].append(dict(d["entries"][0])), "INVALID_ACCOUNT_REGISTRY"),
        (lambda d: d["entries"][1].update(tracker_key="ACCT-001"), "INVALID_ACCOUNT_REGISTRY"),
        (lambda d: d.pop("source"), "INVALID_ACCOUNT_REGISTRY"),
        (lambda d: d["entries"][0].update(account_id="12345"), "INVALID_ACCOUNT_REGISTRY"),
        (lambda d: d["entries"][0].update(name_guess="row 1"), "INVALID_ACCOUNT_REGISTRY"),
    ],
)
def test_registry_rejects_invalid_or_inferred_mappings(tmp_path, mutate, code):
    data = json.loads(registry_file(tmp_path / "r.json").read_text())
    mutate(data)
    (tmp_path / "r.json").write_text(json.dumps(data))
    with pytest.raises(InputError, match=code):
        install_registry(tmp_path, tmp_path / "r.json")
    assert workspace_registry(tmp_path) is None


def test_registry_cli_history_and_no_ids_printed(tmp_path, capsys):
    workspace = tmp_path / "ws"
    first = registry_file(tmp_path / "one.json")
    assert main(["--workspace", str(workspace), "accounts", "load", str(first)]) == 0
    second = registry_file(tmp_path / "two.json", ((A, "ACCT-009"),))
    assert main(["--workspace", str(workspace), "accounts", "load", str(second)]) == 0
    output = capsys.readouterr().out
    assert A not in output  # account IDs are never printed
    assert workspace_registry(workspace).tracker_keys() == {A: "ACCT-009"}
    assert len(list((workspace / "account-registry-history").iterdir())) == 1
    assert main(["--workspace", str(workspace), "accounts", "show"]) == 0
    shown = json.loads(capsys.readouterr().out)
    assert shown["accounts"] == 1 and shown["with_tracker_key"] == 1


def test_verifier_classifies_origin_and_covers_every_profile(tmp_path):
    from offline_analyzer.contracts.r0 import AWS_OPERATIONS
    from offline_analyzer.verification import verify_package

    package = build(tmp_path / "input")
    synthetic = verify_package(tmp_path / "input", package, origin="synthetic")
    assert synthetic["real_export_validation"] == "NOT_ESTABLISHED"
    assert {s["validation_class"] for s in synthetic["sources"]} == {
        "SYNTHETIC_NOT_REAL_EXPORT_VALIDATION"
    }
    operations = {p["operation"] for p in synthetic["profiles"]}
    assert operations == {source.operation for source in package.sources}
    assert operations <= set(AWS_OPERATIONS)
    assert all(p["parser_accepted"] == p["sources"] for p in synthetic["profiles"])
    declared = verify_package(tmp_path / "input", package, origin="tool-generated")
    # The verifier proves structure only; origin is the operator's recorded declaration.
    assert declared["declared_origin"] == "tool-generated"
    assert declared["real_export_validation"] == "STRUCTURE_COMPATIBLE_OPERATOR_DECLARED"
    dump = json.dumps(declared)
    for private in ("golden-vendor-admin", "203.0.113.50", A):
        assert private not in dump


def test_verifier_counts_duplicate_resource_records(tmp_path):
    from offline_analyzer.verification import verify_package

    path = tmp_path / "sg.json"
    item = group(A, 1, "dup", [permission(22, "10.0.0.0/8")])
    path.write_text(json.dumps({"SecurityGroups": [item, item], **OK}))
    source = Source(path=path.name, sha256=sha256_file(path), format="aws-json-v1",
                    operation="ec2:DescribeSecurityGroups", account_id=A, region="us-east-1",
                    collected_at=TIME, coverage="COMPLETE", request_parameters={})  # fmt: skip
    report = verify_package(tmp_path, Package(sources=(source,)), origin="tool-generated")
    assert report["sources"][0]["duplicate_resource_records"] == 1


def test_generated_r11_fixture_preserves_historical_bytes(tmp_path):
    from fixture_fingerprints import package_fingerprint

    build(tmp_path / "golden")
    assert len(list((tmp_path / "golden").iterdir())) == 27
    assert package_fingerprint(tmp_path / "golden") == (
        "1e2cdb3b76cce0df6449305870b9eca861c42d63222438b0fd4aa12ccd0de853"
    )
    first = registry_file(tmp_path / "registry.json")
    second = registry_file(tmp_path / "registry-again.json")
    assert first.read_bytes() == second.read_bytes()
