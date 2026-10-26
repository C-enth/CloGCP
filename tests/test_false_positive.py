"""False-positive disposition (analyst layer) on synthetic data: rule results unchanged, findings
excluded from confirmed lists, reports and totals, attack paths flagged, consolidated and
carry-over behaviour. All projects, identities and packages are synthetic."""

import json
import zipfile
from pathlib import Path

import pytest
from pydantic import ValidationError

from offline_analyzer.attack_paths.feasibility import attack_paths
from offline_analyzer.contracts.exclusion import BulkExclusionRequest, ExclusionRequest
from offline_analyzer.enterprise.snapshot import assess_enterprise
from offline_analyzer.findings.exclusions import (
    FALSE_POSITIVE_PRESETS,
    appendix,
    excluded,
    false_positives,
    listing,
    record,
    record_many,
)
from offline_analyzer.ingestion.importer import import_package
from offline_analyzer.ingestion.json_input import InputError
from offline_analyzer.ingestion.safety import sha256_file
from offline_analyzer.inventory.environment import set_environment
from offline_analyzer.reporting.assessment_report import model
from offline_analyzer.reporting.cloud import dashboard
from offline_analyzer.reporting.sre_reports import rows
from offline_analyzer.rules.engine import assess
from offline_analyzer.storage.store import Store
from tools.gcp_fixture import build

ANALYST = "Synthetic analyst"
PRESET = FALSE_POSITIVE_PRESETS[0]


def fp(note=PRESET):
    return ExclusionRequest(action="EXCLUDE", reason="FALSE_POSITIVE", note=note, analyst=ANALYST)


RESTORE = ExclusionRequest(action="RESTORE", analyst=ANALYST)


def variant(directory: Path, project: str, number: str, extra: bool = False) -> Path:
    """Synthetic package; `extra` adds a binding so the import differs (a new import)."""
    package = build(directory, project, number)
    if extra:
        manifest = json.loads((package / "manifest.json").read_text())
        for dataset in manifest["datasets"]:
            if dataset["operation"] == "gcp:projects.getIamPolicy":
                path = package / dataset["path"]
                policy = json.loads(path.read_text())
                policy["bindings"].append(
                    {"role": "roles/viewer", "members": ["user:extra@example.invalid"]}
                )
                path.write_text(json.dumps(policy))
                dataset["sha256"] = sha256_file(path)
        (package / "manifest.json").write_text(json.dumps(manifest))
    return package


@pytest.fixture
def workspace(tmp_path):
    set_environment(tmp_path / "work", "SYNTHETIC")
    with Store(tmp_path / "work") as store:
        alpha = import_package(store, variant(tmp_path / "alpha", "demo-alpha", "123456789012"))
        beta = import_package(store, variant(tmp_path / "beta", "demo-beta", "222222222222"))
        for run in (alpha, beta):
            assess(store, run)
        yield store, alpha, beta, tmp_path


def cross(store, run):
    return next(f for f in store.candidates(run) if f.rule_id == "GCP-IAM-CROSSPROJECT")


def test_reason_is_mandatory():
    with pytest.raises(ValidationError):
        fp("")
    with pytest.raises(ValidationError):
        BulkExclusionRequest(
            action="EXCLUDE", reason="FALSE_POSITIVE", analyst=ANALYST, finding_ids=["f-1"]
        )
    assert fp().note == PRESET


def test_single_mark_and_restore_keep_the_rule_result(workspace):
    store, alpha, _beta, _ = workspace
    finding = cross(store, alpha)
    outcomes = sorted(str(r.outcome) for r in store.results(alpha))
    entry = record(store, alpha, finding.finding_id, fp())
    assert entry.reason == "FALSE_POSITIVE" and entry.note == PRESET
    assert entry.previous_state == "CANDIDATE" and entry.recorded_at
    assert any("worker@demo-beta" in p for p in entry.principals)
    assert finding.finding_id in false_positives(store, alpha)
    # The rule evaluation is unchanged.
    assert sorted(str(r.outcome) for r in store.results(alpha)) == outcomes
    assert finding.finding_id not in {r["id"] for r in rows(store, alpha)}
    item = next(i for i in listing(store, alpha)["items"] if i["id"] == finding.finding_id)
    assert item["false_positive"] and item["previous_state"] == "CANDIDATE"
    record(store, alpha, finding.finding_id, RESTORE)
    assert finding.finding_id not in excluded(store, alpha)
    assert finding.finding_id in {r["id"] for r in rows(store, alpha)}


def test_bulk_mark_restore_and_limit(workspace):
    store, alpha, _beta, _ = workspace
    ids = [f.finding_id for f in store.candidates(alpha)][:3]
    request = BulkExclusionRequest(
        action="EXCLUDE",
        reason="FALSE_POSITIVE",
        note="Synthetic bulk",
        analyst=ANALYST,
        finding_ids=ids,
    )
    assert len(record_many(store, alpha, request)) == 3
    assert set(ids) <= set(false_positives(store, alpha))
    with pytest.raises(InputError, match="FINDING_ALREADY_EXCLUDED"):  # all-or-nothing
        record_many(store, alpha, request)
    restore = BulkExclusionRequest(action="RESTORE", analyst=ANALYST, finding_ids=ids)
    record_many(store, alpha, restore)
    assert not set(ids) & set(excluded(store, alpha))
    with pytest.raises(ValidationError):  # bounded selection
        BulkExclusionRequest(
            action="RESTORE", analyst=ANALYST, finding_ids=[f"f-{i}" for i in range(2001)]
        )


def test_excluded_from_totals_reports_and_appendix(workspace):
    from openpyxl import load_workbook

    from offline_analyzer.reporting.docx import write_docx
    from offline_analyzer.reporting.excel import REPORT_SHEETS, export_excel
    from offline_analyzer.reporting.report_document import write_excel

    store, alpha, _beta, tmp_path = workspace
    finding = cross(store, alpha)
    before = dashboard(store, alpha, "GCP")["providers"]["GCP"]["severity"]
    record(store, alpha, finding.finding_id, fp())
    after = dashboard(store, alpha, "GCP")["providers"]["GCP"]["severity"]
    assert sum(after.values()) == sum(before.values()) - 1
    data = model(store, alpha)
    assert finding.finding_id not in {f["id"] for f in data["findings"]}
    (group,) = [g for g in data["excluded_findings"] if g["rule"] == "GCP-IAM-CROSSPROJECT"]
    assert group == {
        "rule": "GCP-IAM-CROSSPROJECT",
        "reason": "False positive",
        "detail": PRESET,
        "count": 1,
        "finding_ids": [finding.finding_id],
    }
    assert appendix(store, alpha) == data["excluded_findings"]
    word, book = tmp_path / "r.docx", tmp_path / "r.xlsx"
    write_docx(data, word)
    xml = zipfile.ZipFile(word).read("word/document.xml").decode()
    assert "Appendix C Excluded findings" in xml and PRESET in xml
    assert finding.finding_id not in xml.split("Appendix C Excluded findings")[0]
    write_excel(data, book)
    sheet = load_workbook(book)["Excluded Findings"]
    values = [c for row in sheet.iter_rows(values_only=True) for c in row]
    assert "GCP-IAM-CROSSPROJECT" in values and 1 in values
    assert any(finding.finding_id in str(v) for v in values)
    workbook = tmp_path / "w.xlsx"
    export_excel(store, alpha, workbook)
    loaded = load_workbook(workbook)
    assert len(loaded.sheetnames) >= REPORT_SHEETS and "Excluded Findings" in loaded.sheetnames
    ids = [r[0] for r in loaded["Excluded Findings"].iter_rows(min_row=2, values_only=True)]
    assert finding.finding_id in ids
    for name in ("Findings Register", "Confirmed Findings"):
        cells = [str(v) for row in loaded[name].iter_rows(values_only=True) for v in row]
        assert not any(finding.finding_id in v for v in cells), name


def test_finalized_list_never_contains_false_positives(workspace):
    from offline_analyzer.findings.generate import record_decision
    from offline_analyzer.reporting.sre_reports import finalized

    store, alpha, _beta, _ = workspace
    finding = cross(store, alpha)
    record_decision(store, alpha, finding.finding_id, "CONFIRMED", ANALYST, "")
    assert finding.finding_id in {r["id"] for r in finalized(store, alpha)["findings"]}
    record(store, alpha, finding.finding_id, fp())
    assert finding.finding_id not in {r["id"] for r in finalized(store, alpha)["findings"]}
    assert finalized(store, alpha)["excluded"] == 1


def test_attack_path_depending_on_a_false_positive_is_flagged_and_left_out(workspace):
    store, alpha, _beta, _ = workspace
    paths = [p for p in attack_paths(store, alpha)["items"] if p.get("finding_ids")]
    assert paths
    target = paths[0]
    linked = target["finding_ids"][0]
    record(store, alpha, linked, fp("Synthetic path dependency"))
    flagged = {p["path_id"]: p for p in attack_paths(store, alpha)["items"]}[target["path_id"]]
    assert flagged["false_positive_dependency"] == {
        "flag": "depends on a false-positive finding",
        "finding_ids": [linked],
    }
    assert target["path_id"] not in {p["path_id"] for p in model(store, alpha)["paths"]}
    # The path identity and an analyst review fingerprint are unchanged by the flag.
    from offline_analyzer.attack_paths.review import fingerprint

    assert fingerprint(flagged) == fingerprint(target)


def test_consolidated_snapshot_reflects_dispositions_both_ways(workspace):
    store, alpha, beta, _ = workspace
    snapshot = assess_enterprise(store, [alpha, beta])
    # In the snapshot demo-beta is collected too, so alpha's cross-project binding is an
    # informational link there; use a finding present in both runs.
    shared = {f.finding_id for f in store.candidates(snapshot)}
    in_alpha = next(f for f in store.candidates(alpha) if f.finding_id in shared)
    record(store, alpha, in_alpha.finding_id, fp())
    assert in_alpha.finding_id in false_positives(store, snapshot)
    assert in_alpha.finding_id not in {r["id"] for r in rows(store, snapshot)}
    in_beta = next(f for f in store.candidates(beta) if f.rule_id == "GCP-IAM-BASIC")
    record(store, snapshot, in_beta.finding_id, fp("Marked in the consolidated view"))
    assert in_beta.finding_id in false_positives(store, beta)
    snap_totals = dashboard(store, snapshot, "GCP")["providers"]["GCP"]["excluded_findings"]
    assert snap_totals == 2


def test_carry_over_to_a_new_import_only_on_an_exact_match(workspace):
    store, alpha, beta, tmp_path = workspace
    finding = cross(store, alpha)
    record(store, alpha, finding.finding_id, fp())
    newer = import_package(
        store, variant(tmp_path / "alpha-2", "demo-alpha", "123456789012", extra=True)
    )
    assess(store, newer)
    assert newer != alpha
    carried = false_positives(store, newer)[finding.finding_id]
    assert carried.carried_from == alpha
    item = next(i for i in listing(store, newer)["items"] if i["id"] == finding.finding_id)
    assert item["carried_from"] == alpha
    # Another project never shares the finding ID, so nothing is carried across projects.
    assert finding.finding_id not in {f.finding_id for f in store.candidates(beta)}
    assert not false_positives(store, beta)
    # Same finding ID but different principals: not carried over.
    tampered = carried.model_copy(
        update={
            "principals": ("serviceAccount:other@demo-gamma.iam.gserviceaccount.com",),
            "carried_from": None,
        }
    )
    store.db.execute(
        "UPDATE finding_exclusions SET data=? WHERE finding_id=?",
        [tampered.model_dump_json(), finding.finding_id],
    )
    assert finding.finding_id not in false_positives(store, newer)
    assert finding.finding_id in false_positives(store, alpha)  # still where it was made


def test_workbench_hides_false_positives_by_default_with_count_and_tag(workspace):
    from offline_analyzer.api.services import Views

    store, alpha, _beta, tmp_path = workspace
    views = Views(store.workspace)
    finding = cross(store, alpha)
    record(store, alpha, finding.finding_id, fp())
    default = views.records(store, alpha, "findings", "", "", "", 0, 100, {})
    assert finding.finding_id not in {i["finding_id"] for i in default["items"]}
    assert default["false_positive_excluded"] == 1
    shown = views.records(store, alpha, "findings", "", "", "", 0, 100, {"false_positive": "show"})
    tagged = next(i for i in shown["items"] if i["finding_id"] == finding.finding_id)
    assert tagged["false_positive"]["reason"] == PRESET
    only = views.records(store, alpha, "findings", "", "", "", 0, 100, {"false_positive": "only"})
    assert [i["finding_id"] for i in only["items"]] == [finding.finding_id]
    facets = views.finding_facets(store, alpha)
    assert facets["false_positive_excluded"] == 1
