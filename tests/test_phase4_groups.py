"""Root-cause grouping (owner brief Phase 4): group keys, grouped view and bulk actions,
group-level report edits with per-member review state, grouped Word and report Excel.
Synthetic AWS + GCP data only."""

import zipfile
from types import SimpleNamespace

import pytest
from pydantic import ValidationError
from test_theme_browser import workspace as estate_workspace

from offline_analyzer.contracts.exclusion import ExclusionRequest
from offline_analyzer.contracts.report_finding import ReportEdit
from offline_analyzer.findings.exclusions import record as exclude
from offline_analyzer.findings.generate import record_decisions
from offline_analyzer.findings.groups import group_id, group_key, groups
from offline_analyzer.findings.report_ready import (
    current_findings,
    latest_revision,
    save_revision,
    sections_for,
    validity,
)
from offline_analyzer.reporting.assessment_report import model
from offline_analyzer.storage.store import Store


@pytest.fixture(scope="module")
def estate(tmp_path_factory):
    root = tmp_path_factory.mktemp("phase4")
    work = estate_workspace(root)
    with Store(work) as store:
        snapshot = next(r for (r,) in store.rows("SELECT run_id FROM enterprise_snapshots"))
        runs = sorted({r for (r,) in store.rows("SELECT run_id FROM runs")} - {snapshot})
        yield store, runs, snapshot, root


def fake(rule, native="", text=""):
    return SimpleNamespace(rule_id=rule, native_id=native, technical_observation=text)


def test_group_keys_by_kind():
    sa = "serviceAccount:worker@demo-beta.iam.gserviceaccount.com"
    assert group_key(fake("GCP-IAM-CROSSPROJECT", "projects/p/bindings/x", f"hold {sa}.")) == (
        "GCP-IAM-CROSSPROJECT",
        "principal",
        sa,
    )
    assert group_key(fake("GCP-NET-FW-ADMIN-INTERNET", "projects/p/global/firewalls/ssh")) == (
        "GCP-NET-FW-ADMIN-INTERNET",
        "resource",
        "projects/p/global/firewalls/ssh",
    )
    assert group_key(fake("GCP-ORG-BASIC-ROLE", "//cloudresourcemanager/folders/700")) == (
        "GCP-ORG-BASIC-ROLE",
        "org_node",
        "folders/700",
    )
    # No principal or node: falls back to the rule; other rules group by rule ID.
    assert group_key(fake("GCP-ORG-BASIC-ROLE", "//x/projects/p")) == (
        "GCP-ORG-BASIC-ROLE",
        "rule",
        "",
    )
    assert group_key(fake("GCP-VM-OSLOGIN", "projects/p/instances/a"))[1] == "rule"
    assert group_id(fake("GCP-VM-OSLOGIN", "a")) == group_id(fake("GCP-VM-OSLOGIN", "b"))


def test_groups_span_runs_exclude_excluded_members_and_take_bulk_actions(estate):
    store, runs, snapshot, _ = estate
    grouped = groups(store, snapshot)
    assert grouped and all(g["member_count"] == len(g["members"]) for g in grouped)
    member_groups = {g["group_id"] for run in runs for g in groups(store, run)}
    assert {g["group_id"] for g in grouped} <= member_groups  # same IDs in member runs
    target = max(grouped, key=lambda g: g["member_count"])
    assert target["severity"] in {"CRITICAL", "HIGH", "MEDIUM", "LOW", "INFO"}
    first = target["members"][0]
    exclude(
        store,
        snapshot,
        first,
        ExclusionRequest(action="EXCLUDE", reason="ACCEPTED_RISK", analyst="analyst"),
    )
    try:
        after = next(g for g in groups(store, snapshot) if g["group_id"] == target["group_id"])
        assert first not in after["members"] and after["excluded_members"] == 1
        record_decisions(store, snapshot, after["members"], "CONFIRMED", "analyst", "")
        again = next(g for g in groups(store, snapshot) if g["group_id"] == target["group_id"])
        assert set(again["status_mix"]) == {"CONFIRMED_VULNERABILITY"}
    finally:
        exclude(store, snapshot, first, ExclusionRequest(action="RESTORE", analyst="analyst"))


def test_group_report_edit_member_limit_and_per_member_review(estate):
    store, _runs, snapshot, _ = estate
    findings = sorted(store.candidates(snapshot), key=lambda f: f.finding_id)
    primary = findings[0]
    edit = dict(
        expected_revision=0,
        analyst="analyst",
        sections=sections_for(primary).model_copy(
            update={"immediate_mitigation": "Synthetic immediate step."}
        ),
        include_in_report=True,
    )
    with pytest.raises(ValidationError):
        ReportEdit(**edit, member_ids=[f"finding-{i}" for i in range(2001)])
    assert ReportEdit(**edit, member_ids=[f"finding-{i}" for i in range(2000)])
    members = [f.finding_id for f in findings[:3]]
    revision = save_revision(
        store, snapshot, primary.finding_id, ReportEdit(**edit, member_ids=members)
    )
    changed = members[2]
    store.db.execute(
        "UPDATE findings SET data=json_merge_patch(data, ?) WHERE run_id=? AND finding_id=?",
        ['{"technical_observation": "Synthetic changed evidence."}', snapshot, changed],
    )
    state = validity(
        latest_revision(store, primary.finding_id),
        current_findings(store, snapshot, members),
        store.latest_decisions(snapshot),
        store,
    )
    assert state["member_states"][changed] == "NEEDS_REVIEW"
    assert {state["member_states"][m] for m in members[:2]} == {"CURRENT"}
    data = model(store, snapshot, include="all")
    rows = {f["id"]: f for f in data["findings"]}
    for member in members[:2]:
        assert rows[member]["immediate"] == "Synthetic immediate step."
        assert rows[member]["report_revision"] == revision.revision
    assert rows[changed].get("immediate") is None


def test_grouped_and_ungrouped_word_and_excel(estate, tmp_path):
    from openpyxl import load_workbook

    from offline_analyzer.reporting.docx import write_docx
    from offline_analyzer.reporting.report_document import summary, write_excel

    store, _runs, snapshot, _ = estate
    data = model(store, snapshot, include="all")
    assert data["grouped"] and data["groups"]
    total = sum(len(g["members"]) for g in data["groups"])
    assert total == len(data["findings"])  # totals still count member findings
    for group in data["groups"]:
        assert group["severity"] == group["representative"]["severity"]
    cover = dict(summary(data)["control"])
    assert cover["Grouped findings"].startswith(f"{len(data['groups'])} grouped findings / ")
    word = tmp_path / "grouped.docx"
    write_docx(data, word)
    xml = zipfile.ZipFile(word).read("word/document.xml").decode()
    assert "affected resource(s))" in xml and "groups)" in xml
    book = tmp_path / "grouped.xlsx"
    write_excel(data, book)
    sheet = load_workbook(book)["Grouped Findings"]
    assert sheet.max_row - 1 == len(data["groups"])
    flat = model(store, snapshot, include="all", grouped=False)
    flat_word = tmp_path / "flat.docx"
    write_docx(flat, flat_word)
    flat_xml = zipfile.ZipFile(flat_word).read("word/document.xml").decode()
    assert "affected resource(s))" not in flat_xml
    assert "Grouped findings" not in dict(summary(flat)["control"])
