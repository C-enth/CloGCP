"""2.2.6 manifest fields are consumed, not only accepted: detail limits, IAM target-list
fallbacks and RATE_LIMITED / AUTH_ERROR become explicit coverage reasons (synthetic data)."""

import hashlib
import json
import zipfile
from pathlib import Path
from xml.etree.ElementTree import fromstring

import pytest
from openpyxl import load_workbook
from test_gcp_collector_226 import LIMIT_KEYS, package

from offline_analyzer.ingestion.gcp import inspect_package
from offline_analyzer.ingestion.importer import import_package
from offline_analyzer.reporting.assessment_report import export_assessment, model
from offline_analyzer.reporting.cloud import dashboard
from offline_analyzer.rules.engine import assess
from offline_analyzer.storage.store import Store

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def rewrite(folder: Path, update) -> Path:
    path = folder / "collection_manifest.json"
    manifest = json.loads(path.read_text())
    update(manifest, folder)
    path.write_text(json.dumps(manifest))
    lines = [
        f"{hashlib.sha256(p.read_bytes()).hexdigest()}  {p.relative_to(folder).as_posix()}"
        for p in sorted(
            p for p in folder.rglob("*") if p.is_file() and p.name != "checksums.sha256"
        )
    ]
    (folder / "checksums.sha256").write_text("\n".join(lines) + "\n")
    return folder


def fail(item: dict, folder: Path, status: str, text: str) -> None:
    if item.get("file"):
        (folder / item["file"]).unlink(missing_ok=True)
    item.update(status=status, file=None, record_count=None)
    item.pop("pages", None)
    item["exit_code"] = 1 if item.get("exit_code") is not None else None
    item["error_file"] = f"{item['domain']}/{item['dataset']}.stderr.txt"
    (folder / item["error_file"]).write_text(text)


def sources(folder: Path) -> dict[str, dict]:
    package_, _ = inspect_package(folder)
    return {
        s.gcp_metadata["dataset"]: {
            "coverage": s.coverage,
            "notes": s.gcp_metadata["completeness_limitations"],
            "error": s.gcp_metadata["error_category"],
        }
        for s in package_.sources
    }


def imported(tmp_path: Path, folder: Path):
    store = Store(tmp_path / "workspace")
    run = import_package(store, folder)
    assess(store, run)
    return store, run


def test_detail_limit_marks_the_family_partial_with_a_reason(tmp_path):
    limits = {key: key == "buckets" for key in LIMIT_KEYS}
    folder = package(tmp_path, buckets=2, manifest_extra={"detail_limits_reached": limits})
    found = sources(folder)
    for name in ("buckets", "bucket_iam_1", "bucket_details_2"):
        assert "DETAIL_LIMIT_REACHED_BUCKETS" in found[name]["notes"]
        assert found[name]["coverage"] == "PARTIAL"
    assert "DETAIL_LIMIT_REACHED_BUCKETS" not in found["sinks"]["notes"]
    store, run = imported(tmp_path, folder)
    with store:
        gaps = {g["dataset"]: g for g in dashboard(store, run, "GCP")["gaps"]}
        assert "remaining buckets not described" in gaps["bucket_iam_1"]["coverage_reason"]
        assert gaps["bucket_iam_1"]["effective_coverage"] == "PARTIAL"
        # No collection-scope PASS can come from a limited family.
        assert not [
            r
            for r in store.results(run)
            if r.rule_id.startswith("GCP-STORAGE")
            and r.evaluation_scope == "COLLECTION_SCOPE"
            and str(r.outcome) == "PASS"
        ]
        notes = {n["code"] for n in model(store, run)["coverage_notes"]}
        assert "DETAIL_LIMIT_REACHED_BUCKETS" in notes


def test_partial_rest_list_without_fallback_leaves_iam_partial_with_reason(tmp_path):
    def partial(manifest, folder):
        for item in manifest["datasets"]:
            if item["dataset"] == "run_services_rest":
                assert item["status"] == "SUCCESS" and item["file"]
                item["status"] = "PARTIAL"

    folder = rewrite(package(tmp_path, buckets=1), partial)
    found = sources(folder)
    iam = [n for n in found if n.startswith("run_iam_")]
    assert iam
    for name in ("run_services_rest", *iam):
        assert "IAM_TARGET_LIST_INCOMPLETE" in found[name]["notes"]
        assert found[name]["coverage"] == "PARTIAL"  # never Effective Coverage (COMPLETE)
    store, run = imported(tmp_path, folder)
    with store:
        gaps = {g["dataset"]: g for g in dashboard(store, run, "GCP")["gaps"]}
        assert "IAM was collected only for the listed targets" in gaps[iam[0]]["coverage_reason"]
        target = tmp_path / "report.zip"
        export_assessment(store, run, target, provider="GCP")
    with zipfile.ZipFile(target) as bundle:
        word = next(n for n in bundle.namelist() if n.endswith(".docx"))
        book = next(n for n in bundle.namelist() if n.endswith(".xlsx"))
        with zipfile.ZipFile(bundle.open(word)) as docx:
            text = " ".join(fromstring(docx.read("word/document.xml")).itertext())
        assert "Collection gaps and limitations" in text and "IAM_TARGET_LIST_INCOMPLETE" in text
        sheet = load_workbook(bundle.open(book))["Coverage and Limitations"]
        values = [c for row in sheet.iter_rows(values_only=True) for c in row if c]
        assert "IAM_TARGET_LIST_INCOMPLETE" in values


def test_fallback_used_is_a_coverage_note(tmp_path):
    found = sources(package(tmp_path, fallback=True))
    assert "IAM_TARGETS_FALLBACK_USED" in found["run_services_rest"]["notes"]
    assert found["run_services_rest"]["coverage"] == "ERROR"
    assert "IAM_TARGETS_FROM_GCLOUD_LIST" in found["run_iam_1"]["notes"]


def test_fallback_that_also_failed_is_unknown_never_pass(tmp_path):
    def both_failed(manifest, folder):
        kept = []
        for item in manifest["datasets"]:
            if item["dataset"].startswith("run_iam_"):
                (folder / item["file"]).unlink()
                continue
            if item["dataset"] == "cloud_run":
                fail(item, folder, "RATE_LIMITED", "HTTPError 429: synthetic")
            kept.append(item)
        manifest["datasets"] = kept
        manifest["dependency_fallbacks"][0].update(fallback_status="RATE_LIMITED", targets=0)

    folder = rewrite(package(tmp_path, fallback=True), both_failed)
    found = sources(folder)
    for name in ("run_services_rest", "cloud_run"):
        assert "IAM_TARGETS_UNAVAILABLE" in found[name]["notes"]
        assert found[name]["coverage"] == "ERROR"
    assert not any(n.startswith("run_iam_") for n in found)
    store, run = imported(tmp_path, folder)
    with store:
        failed = {
            s["artifact_id"]
            for s in store.source_data(run)
            if (s.get("gcp_metadata") or {}).get("dataset") in {"run_services_rest", "cloud_run"}
        }
        assert failed
        assert not [
            r
            for r in store.results(run)
            if str(r.outcome) == "PASS" and failed & {e.artifact_id for e in r.evidence}
        ]
        gaps = {g["dataset"]: g for g in dashboard(store, run, "GCP")["gaps"]}
        assert "UNKNOWN, never a pass" in gaps["cloud_run"]["coverage_reason"]


@pytest.mark.parametrize(
    ("status", "code", "phrase"),
    [
        ("RATE_LIMITED", "COLLECTION_RATE_LIMITED", "rate-limited"),
        ("AUTH_ERROR", "COLLECTION_AUTH_FAILURE", "sign-in expired"),
    ],
)
def test_rate_limit_and_expired_sign_in_are_error_coverage_with_reason(
    tmp_path, status, code, phrase
):
    def failed(manifest, folder):
        item = next(i for i in manifest["datasets"] if i["dataset"] == "sinks")
        fail(item, folder, status, "synthetic stderr")

    folder = rewrite(package(tmp_path, buckets=1), failed)
    found = sources(folder)
    assert found["sinks"]["coverage"] == "ERROR" and found["sinks"]["error"] == code
    store, run = imported(tmp_path, folder)
    with store:
        gaps = {g["dataset"]: g for g in dashboard(store, run, "GCP")["gaps"]}
        assert gaps["sinks"]["effective_coverage"] == "ERROR"
        assert phrase in gaps["sinks"]["coverage_reason"]
        assert any(n["code"] == code for n in model(store, run)["coverage_notes"])
