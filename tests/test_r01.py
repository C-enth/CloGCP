"""Source-format synthetic fixtures only: no real AWS/Prowler export is represented."""

import json
from collections import Counter

import pytest
from openpyxl import load_workbook

from offline_analyzer.cli import summary
from offline_analyzer.contracts.r0 import NORMALIZATION_VERSION, Package, Source
from offline_analyzer.findings.generate import record_decision, validation_status
from offline_analyzer.ingestion.importer import import_package, single_file_package
from offline_analyzer.ingestion.safety import sha256_file
from offline_analyzer.reporting.excel import export_excel
from offline_analyzer.rules.engine import assess
from offline_analyzer.storage.store import Store
from offline_analyzer.verification import main as verify_main
from offline_analyzer.verification import verify_package
from tools.fixtures.demo import TIME, A, B, build_package


def aws_source(directory, name, data, operation, account=A, region="us-east-1", **context):
    path = directory / name
    path.write_text(json.dumps(data), encoding="utf-8")
    return Source(
        path=name,
        sha256=sha256_file(path),
        format="aws-json-v1",
        operation=operation,
        account_id=account,
        region=region,
        collected_at=TIME,
        coverage="COMPLETE",
        request_parameters={},
        **context,
    )


def scanner_source(directory, mutate=None):
    staging = directory / "generator"
    base = build_package(staging)
    records = json.loads((staging / "prowler.json").read_text())
    row = records[0]
    row["finding_info"]["uid"] = "customer-finding-private"
    row["metadata"]["version"] = "1.3.0"
    row["metadata"]["profiles"] = ["cloud", "datetime"]
    row["message"] = "customer-secret-observation-do-not-return"
    row["resources"][0]["data"] = {"details": "private data", "metadata": {"encrypted": False}}
    if mutate:
        mutate(row)
    path = directory / "scanner.json"
    path.write_text(json.dumps([row]))
    return base.sources[-1].model_copy(update={"path": path.name, "sha256": sha256_file(path)})


def db_response(encrypted=True):
    return {
        "DBInstances": [
            {
                "DBInstanceArn": f"arn:aws:rds:us-east-1:{A}:db:db",
                "Engine": "postgres",
                "StorageEncrypted": encrypted,
            }
        ],
        "ResponseMetadata": {"HTTPStatusCode": 200, "RequestId": "private-request"},
    }


def test_pinned_source_fields_are_preserved(tmp_path):
    source = scanner_source(tmp_path)
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, tmp_path, Package(sources=(source,)))
        obs = next(store.observations(run))
        assert obs.source_check_id == "rds_instance_storage_encrypted"
        assert obs.source_finding_id == "customer-finding-private"
        assert obs.source_status == "PASS" and obs.source_status_id == 1
        assert obs.source_status_caption == "New"
        assert obs.source_details["metadata_version"] == "1.3.0"
        assert obs.source_details["resource_data"]["metadata"]["encrypted"] is False
        assert obs.normalization_version == NORMALIZATION_VERSION
        raw = store.workspace / "raw" / f"{source.sha256}.json"
        assert sha256_file(raw) == source.sha256
        assert obs.native_id != obs.source_finding_id


@pytest.mark.parametrize(
    "code,lifecycle,caption",
    [
        ("FAIL", 3, "Suppressed"),
        ("PASS", 4, "Resolved"),
        ("PASS", 5, "Archived"),
        ("PASS", 1, "Suppressed"),
        ("PASS", None, "New"),
        ("PASS", True, "New"),
        ("PASS", [3], "New"),
        ("FAIL", 99, "Other"),
    ],
)
def test_lifecycle_does_not_become_control_status(tmp_path, code, lifecycle, caption):
    def mutate(row):
        row["status_code"] = code
        row["status_id"] = lifecycle
        row["status"] = caption

    source = scanner_source(tmp_path, mutate)
    # Same native resource as scanner, with an independent explicit configuration failure.
    native = aws_source(
        tmp_path, "db.json", db_response(False), "rds:DescribeDBInstances", account=A
    )
    # Change scanner owner to native owner using the full resource ARN.
    path = tmp_path / source.path
    row = json.loads(path.read_text())[0]
    row["cloud"]["account"]["uid"] = A
    row["cloud"]["region"] = "us-east-1"
    row["resources"][0]["uid"] = f"arn:aws:rds:us-east-1:{A}:db:db"
    row["resources"][0]["region"] = "us-east-1"
    path.write_text(json.dumps([row]))
    source = source.model_copy(update={"sha256": sha256_file(path)})
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, tmp_path, Package(sources=(native, source)))
        assess(store, run)
        result = next(
            r for r in store.results(run) if r.resource_id and r.rule_id.startswith("R0-")
        )
        assert result.outcome == "FAIL"
        assert len(result.supporting_observations) == 1
        assert not result.contradictory_observations
        obs = next(o for o in store.observations(run) if o.source_tool == "prowler")
        assert obs.source_status == code
        assert next(store.candidates(run)).review_status == "CANDIDATE"


@pytest.mark.parametrize("field", ["metadata", "class_uid", "cloud", "status_code", "finding_info"])
def test_missing_scanner_metadata_is_retained(tmp_path, field):
    source = scanner_source(tmp_path, lambda row: row.pop(field))
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, tmp_path, Package(sources=(source,)))
        assess(store, run)
        assert list(store.observations(run))
        assert not list(store.candidates(run))
        assert all(r.outcome != "PASS" for r in store.results(run))


@pytest.mark.parametrize(
    "params,extra,expected",
    [
        (None, {}, "PARTIAL"),
        ({"DBInstanceIdentifier": "db"}, {}, "PARTIAL"),
        ({"Filters": [{"Name": "db-instance-id", "Values": ["db"]}]}, {}, "PARTIAL"),
        ({"Filters": False}, {}, "PARTIAL"),
        ({"Marker": "private-token"}, {}, "PARTIAL"),
        ({}, {"Marker": {}}, "PARTIAL"),
        ({}, {"Marker": "private-token"}, "PARTIAL"),
        ({}, {"IsTruncated": True}, "PARTIAL"),
        ({}, {"ResponseMetadata": {"HTTPStatusCode": 403}}, "ERROR"),
        ({}, {}, "COMPLETE"),
    ],
)
def test_collection_metadata_not_inferred_from_resource_pass(tmp_path, params, extra, expected):
    data = db_response()
    data.update(extra)
    source = aws_source(tmp_path, "db.json", data, "rds:DescribeDBInstances")
    source = source.model_copy(update={"request_parameters": params})
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, tmp_path, Package(sources=(source,)))
        assess(store, run)
        assert next(store.source_data(run))["effective_coverage"] == expected
        result = next(
            r for r in store.results(run) if r.resource_id and r.rule_id.startswith("R0-")
        )
        assert result.outcome == ("COLLECTION_ERROR" if expected == "ERROR" else "PASS")
        scope = [
            r
            for r in store.results(run)
            if r.evaluation_scope == "COLLECTION_SCOPE" and r.rule_id == "R0-RDS-ENCRYPTION"
        ]
        if expected != "COMPLETE":
            assert scope and scope[0].outcome != "PASS"


def test_undeclared_aws_response_error_is_retained(tmp_path):
    source = aws_source(
        tmp_path,
        "error.json",
        {"Error": {"Code": "AccessDenied", "Message": "private"}},
        "rds:DescribeDBInstances",
    )
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, tmp_path, Package(sources=(source,)))
        assess(store, run)
        assert summary(store, run)["source_collections_with_errors"] == 1
        assert any(r.outcome == "COLLECTION_ERROR" for r in store.results(run))


def test_nested_reservation_owner_without_inventing_region(tmp_path):
    source = aws_source(
        tmp_path,
        "ec2.json",
        {
            "Reservations": [
                {
                    "OwnerId": B,
                    "Instances": [
                        {
                            "InstanceId": "i-12345678",
                            "State": {"Name": "running"},
                            "MetadataOptions": {
                                "State": "applied",
                                "HttpTokens": "required",
                                "HttpEndpoint": "enabled",
                            },
                        }
                    ],
                }
            ]
        },
        "ec2:DescribeInstances",
        account=None,
        region=None,
    )
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, tmp_path, Package(sources=(source,)))
        assess(store, run)
        obs = next(store.observations(run))
        assert obs.account_id == B and obs.region is None and obs.resource_id is None
        assert summary(store, run)["unique_accounts"] == 1
        assert all(r.outcome == "UNKNOWN" for r in store.results(run))


def test_shadow_copies_share_resource_but_keep_collection_context(tmp_path):
    trail = {
        "TrailARN": f"arn:aws:cloudtrail:us-east-1:{A}:trail/organization",
        "HomeRegion": "us-east-1",
        "IsOrganizationTrail": True,
        "IsMultiRegionTrail": True,
        "LogFileValidationEnabled": False,
    }
    sources = tuple(
        aws_source(
            tmp_path,
            f"trail-{index}.json",
            {"trailList": [trail]},
            "cloudtrail:DescribeTrails",
            account=account,
            region=region,
        )
        for index, (account, region) in enumerate(
            [(A, "us-east-1"), (A, "eu-west-1"), (B, "eu-west-1")]
        )
    )
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, tmp_path, Package(sources=sources))
        assess(store, run)
        observations = list(store.observations(run))
        assert len(observations) == 3
        assert len({o.resource_id for o in observations}) == 1
        assert all(o.account_id == A and o.region == "us-east-1" for o in observations)
        assert {o.source_details["trail_relationship"] for o in observations} == {
            "HOME_REGION",
            "REGIONAL_SHADOW",
            "ORGANIZATION_SHADOW",
        }
        findings = list(store.candidates(run))
        assert len(findings) == 1
        assert len(findings[0].evidence) == 3
        assert import_package(store, tmp_path, Package(sources=sources)) == run
        assert summary(store, run)["observed_owner_accounts"] == 1
        assert summary(store, run)["unique_accounts"] == 2
        assert summary(store, run)["unique_identified_resources"] == 1


@pytest.mark.parametrize("home", [None, "eu-west-1", 12])
def test_cloudtrail_home_region_must_agree(tmp_path, home):
    source = aws_source(
        tmp_path,
        "trail.json",
        {
            "trailList": [
                {
                    "TrailARN": f"arn:aws:cloudtrail:us-east-1:{A}:trail/t",
                    "HomeRegion": home,
                    "LogFileValidationEnabled": True,
                }
            ]
        },
        "cloudtrail:DescribeTrails",
    )
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, tmp_path, Package(sources=(source,)))
        assess(store, run)
        result = next(
            r for r in store.results(run) if r.resource_id and r.rule_id.startswith("R0-")
        )
        assert result.outcome == "UNKNOWN" and result.missing_dependencies


def test_configuration_change_preserves_issue_id_and_analyst_history(tmp_path):
    source = aws_source(tmp_path, "db.json", db_response(False), "rds:DescribeDBInstances")
    with Store(tmp_path / "workspace") as store:
        run1 = import_package(store, tmp_path, Package(sources=(source,)))
        assess(store, run1)
        finding = next(store.candidates(run1))
        record_decision(store, run1, finding.finding_id, "CONFIRMED", "analyst", "retained")
        fixed = aws_source(tmp_path, "fixed.json", db_response(True), "rds:DescribeDBInstances")
        run2 = import_package(store, tmp_path, Package(sources=(fixed,)))
        assess(store, run2)
        assert not list(store.candidates(run2))
        assert store.latest_decision(finding.finding_id).note == "retained"
        # Recurrence of the same configuration on newly collected evidence is the same issue.
        recurrence = source.model_copy(update={"collected_at": source.collected_at.replace(hour=9)})
        run3 = import_package(store, tmp_path, Package(sources=(recurrence,)))
        assess(store, run3)
        recurring = next(store.candidates(run3))
        assert recurring.finding_id == finding.finding_id
        assert validation_status(store, recurring)[0] == "NEEDS_REVALIDATION"


def test_metrics_and_excel_have_distinct_record_kinds(tmp_path):
    package = build_package(tmp_path / "input")
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, tmp_path / "input", package)
        assess(store, run)
        metrics = summary(store, run)
        assert metrics["unique_accounts"] == 2 and metrics["unique_regions"] == 3
        assert metrics["raw_observation_occurrences"] == 10
        assert metrics["normalized_observations"] == 8
        # 10 R0 records (unchanged) plus 4 R1 resource rows lacking fields and 2 scope rows.
        assert metrics["control_evaluation_records"] == 16
        assert metrics["resource_evaluation_outcomes"] == {
            "FAIL": 3,
            "MANUAL_VALIDATION_REQUIRED": 1,
            "PASS": 2,
            "UNKNOWN": 4,
        }
        assert metrics["collection_scope_outcomes"] == {"COLLECTION_ERROR": 2, "UNKNOWN": 4}
        assert metrics["source_collections_with_errors"] == 1
        assert metrics["evaluations_missing_evidence"] == 11
        output = tmp_path / "report.xlsx"
        export_excel(store, run, output)
        wb = load_workbook(output, data_only=False)
        reported = {
            r[0]: json.loads(r[1]) if isinstance(r[1], str) else r[1]
            for r in wb["Assessment Summary"].iter_rows(min_row=4, values_only=True)
            if r[0]
            in {"unique_accounts", "candidate_findings", "normalized_observations", "outcomes"}
        }
        for key, value in reported.items():
            assert value == metrics[key]
        assert Counter(
            row[16] for row in wb["Control Assessment"].iter_rows(min_row=4, values_only=True)
        ) == {"RESOURCE": 10, "COLLECTION_SCOPE": 6}
        assert all(cell.data_type != "f" for sheet in wb for row in sheet for cell in row)
        wb.close()


def test_local_verifier_does_not_expose_values_paths_or_salt(tmp_path):
    source = scanner_source(tmp_path)
    path = tmp_path / source.path
    data = json.loads(path.read_text())
    data[0]["customer-private-field-name@example.test"] = "private-customer-string"
    path.write_text(json.dumps(data))
    source = source.model_copy(update={"sha256": sha256_file(path)})
    package = Package(sources=(source,))
    report = verify_package(tmp_path, package, include_pseudonyms=True)
    dumped = json.dumps(report)
    for private in (
        A,
        B,
        "scanner.json",
        str(tmp_path),
        "customer-finding-private",
        "customer-private-field-name@example.test",
        "private-customer-string",
        "customer-secret-observation-do-not-return",
        "private data",
    ):
        assert private not in dumped
    assert "_key" not in dumped and "salt" not in dumped and "mapping" not in dumped
    diag = report["sources"][0]
    assert diag["parser_accepted"] and diag["tool_versions"] == {"5.10.0": 1}
    assert diag["structure_types"]["item.status_id"] == {"number": 1}
    assert diag["array_sizes"][""]["total_items"] == 1
    assert report["observed_identity_totals"]["unique_observed_accounts"] == 1
    report2 = verify_package(tmp_path, package, include_pseudonyms=True)
    assert (
        report["observed_identity_totals"]["identity_relationships"]
        != report2["observed_identity_totals"]["identity_relationships"]
    )


def test_pseudonyms_preserve_cross_file_identity_links(tmp_path):
    source = aws_source(tmp_path, "db.json", db_response(), "rds:DescribeDBInstances")
    second = aws_source(tmp_path, "second.json", db_response(False), "rds:DescribeDBInstances")
    report = verify_package(tmp_path, Package(sources=(source, second)), include_pseudonyms=True)
    assert report["observed_identity_totals"]["unique_identified_resources"] == 1
    assert len(report["observed_identity_totals"]["identity_relationships"]) == 1


def test_local_verifier_unsupported_version_and_malformed_never_leak(tmp_path, capsys):
    source = scanner_source(
        tmp_path, lambda row: row["metadata"]["product"].update(version="5.19.0")
    )
    report = verify_package(tmp_path, Package(sources=(source,)))
    assert report["sources"][0]["rejection_code"] == "UNSUPPORTED_PROWLER_VERSION_OR_CLASS"
    assert report["observed_identity_totals"]["unique_identified_resources"] == 0
    path = tmp_path / "private-file-name.json"
    path.write_text('{"private-email@example.test": "sensitive",')
    output = tmp_path / "diagnostics.json"
    assert (
        verify_main(
            [
                "--file",
                str(path),
                "--format",
                "prowler-ocsf-5.10",
                "--operation",
                "prowler:Scan",
                "--origin",
                "synthetic",
                "--collected-at",
                TIME,
                "--output",
                str(output),
            ]
        )
        == 2
    )
    displayed = capsys.readouterr().out
    assert "private-file-name" not in displayed and "sensitive" not in displayed
    assert "sensitive" not in output.read_text()
    original = output.read_bytes()
    assert (
        verify_main(
            [
                "--file",
                str(path),
                "--format",
                "prowler-ocsf-5.10",
                "--operation",
                "prowler:Scan",
                "--origin",
                "synthetic",
                "--collected-at",
                TIME,
                "--output",
                str(output),
            ]
        )
        == 1
    )
    assert output.read_bytes() == original


def test_large_array_verification_is_incremental(tmp_path):
    rows = db_response()["DBInstances"] * 2000
    source = aws_source(tmp_path, "large.json", {"DBInstances": rows}, "rds:DescribeDBInstances")
    report = verify_package(tmp_path, Package(sources=(source,)))
    assert report["sources"][0]["array_sizes"]["DBInstances"]["total_items"] == 2000
    assert report["sources"][0]["normalized_resource_records"] == 2000
    assert report["observed_identity_totals"]["unique_identified_resources"] == 1


def test_absent_resource_slots_have_real_provenance(tmp_path):
    source = scanner_source(tmp_path, lambda row: row.pop("resources"))
    native = aws_source(
        tmp_path, "ec2.json", {"Reservations": [{"OwnerId": A}]}, "ec2:DescribeInstances"
    )
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, tmp_path, Package(sources=(source, native)))
        assess(store, run)
        for obs in store.observations(run):
            raw = json.loads((store.workspace / "raw" / f"{obs.integrity_sha256}.json").read_text())
            for segment in obs.evidence.locator.lstrip("/").split("/"):
                raw = raw[int(segment)] if isinstance(raw, list) else raw[segment]
            assert raw is not None
        assert not list(store.candidates(run))


def test_disagreeing_scanner_timestamps_do_not_corroborate(tmp_path):
    source = scanner_source(tmp_path, lambda row: row.update(time=1))
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, tmp_path, Package(sources=(source,)))
        obs = next(store.observations(run))
        assert "scanner timestamps conflict or are malformed" in obs.unresolved


def test_verifier_reports_unsupported_tool_version_without_raw_values(tmp_path):
    source = scanner_source(
        tmp_path, lambda row: row["metadata"]["product"].update(version="5.19.0")
    )
    report = verify_package(tmp_path, Package(sources=(source,)))
    assert report["sources"][0]["tool_versions"] == {"5.19.0": 1}
    assert report["sources"][0]["compatibility_state"] == "REJECTED"
    assert "customer-secret-observation-do-not-return" not in json.dumps(report)


def test_missing_identity_rows_not_counted_as_collection_gaps(tmp_path):
    source = aws_source(
        tmp_path,
        "db.json",
        {"DBInstances": [{"StorageEncrypted": True}, None]},
        "rds:DescribeDBInstances",
    )
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, tmp_path, Package(sources=(source,)))
        assess(store, run)
        metrics = summary(store, run)
        # Two unresolved rows for each of the two RDS rules; no resource is evaluated.
        assert metrics["resource_evaluation_outcomes"] == {"UNKNOWN": 4}
        # Three R0 scope gaps plus one for the RDS-family R1 rule (EC2 family not supplied).
        assert metrics["unresolved_collection_gaps"] == 4
        assert metrics["outcomes"]["UNKNOWN"] == 8
        assert metrics["resources_evaluated"] == 0


def test_legacy_records_remain_readable_and_get_explicit_scope(tmp_path):
    from offline_analyzer.contracts.r0 import Normalized, Result

    source = aws_source(tmp_path, "db.json", db_response(), "rds:DescribeDBInstances")
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, tmp_path, Package(sources=(source,)))
        assess(store, run)
        obs = next(store.observations(run)).model_dump(mode="json")
        obs["normalization_version"] = "r0.1"
        for key in (
            "source_finding_id",
            "source_status_id",
            "source_status_caption",
            "source_details",
            "collection_account_id",
            "collection_region",
        ):
            obs.pop(key)
        assert Normalized.model_validate(obs).normalization_version == "r0.1"
        scope = next(
            r for r in store.results(run) if r.evaluation_scope == "COLLECTION_SCOPE"
        ).model_dump(mode="json")
        scope.pop("evaluation_scope")
        assert Result.model_validate(scope).evaluation_scope == "COLLECTION_SCOPE"


def test_known_native_account_survives_invalid_owner_as_unresolved(tmp_path):
    source = aws_source(
        tmp_path,
        "ec2.json",
        {
            "Reservations": [
                {
                    "OwnerId": "bad",
                    "Instances": [
                        {
                            "InstanceId": "i-12345678",
                            "State": {"Name": "running"},
                            "MetadataOptions": {
                                "State": "applied",
                                "HttpEndpoint": "enabled",
                                "HttpTokens": "required",
                            },
                        }
                    ],
                }
            ]
        },
        "ec2:DescribeInstances",
    )
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, tmp_path, Package(sources=(source,)))
        assess(store, run)
        assert next(store.observations(run)).account_id is None
        assert all(r.outcome == "UNKNOWN" for r in store.results(run))


def test_mixed_good_bad_profile_verification_excludes_partial_identities(tmp_path):
    source = scanner_source(tmp_path)
    path = tmp_path / source.path
    records = json.loads(path.read_text())
    bad = json.loads(json.dumps(records[0]))
    bad["metadata"]["product"]["version"] = "5.19.0"
    records.append(bad)
    path.write_text(json.dumps(records))
    source = source.model_copy(update={"sha256": sha256_file(path)})
    report = verify_package(tmp_path, Package(sources=(source,)))
    assert not report["sources"][0]["parser_accepted"]
    assert report["observed_identity_totals"]["unique_identified_resources"] == 0


def test_empty_scanner_is_not_version_verification(tmp_path):
    path = tmp_path / "empty.json"
    path.write_text("[]")
    package = single_file_package(
        path,
        format="prowler-ocsf-5.10",
        operation="prowler:Scan",
        account_id=None,
        region=None,
        collected_at=TIME,
    )
    report = verify_package(tmp_path, package)
    assert report["sources"][0]["compatibility_state"] == "ACCEPTED_EMPTY_PROFILE_UNVERIFIED"
    assert report["sources"][0]["tool_versions"] == {}


def test_shadow_configuration_disagreement_remains_unknown(tmp_path):
    body = {
        "trailList": [
            {
                "TrailARN": f"arn:aws:cloudtrail:us-east-1:{A}:trail/t",
                "HomeRegion": "us-east-1",
                "IsOrganizationTrail": True,
                "LogFileValidationEnabled": True,
            }
        ]
    }
    first = aws_source(tmp_path, "home.json", body, "cloudtrail:DescribeTrails")
    body["trailList"][0]["LogFileValidationEnabled"] = False
    second = aws_source(
        tmp_path, "shadow.json", body, "cloudtrail:DescribeTrails", region="eu-west-1"
    )
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, tmp_path, Package(sources=(first, second)))
        assess(store, run)
        result = next(
            r for r in store.results(run) if r.resource_id and r.rule_id.startswith("R0-")
        )
        assert result.outcome == "UNKNOWN"
        assert len(result.evidence) == 2
        assert not list(store.candidates(run))


def test_numeric_report_counts_and_string_injection_safety(tmp_path):
    source = aws_source(tmp_path, "db.json", db_response(False), "rds:DescribeDBInstances")
    source = source.model_copy(update={"collector_version": "=UNTRUSTED()"})
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, tmp_path, Package(sources=(source,)))
        assess(store, run)
        export_excel(store, run, tmp_path / "report.xlsx")
    wb = load_workbook(tmp_path / "report.xlsx", data_only=False)
    counts = {
        row[0]: row[1] for row in wb["Assessment Summary"].iter_rows(min_row=4, values_only=True)
    }
    assert counts["candidate_findings"] == 1 and type(counts["candidate_findings"]) is int
    assert any(
        row[9] == "'=UNTRUSTED()"
        for row in wb["Evidence Register"].iter_rows(min_row=4, values_only=True)
    )
    assert all(cell.data_type != "f" for sheet in wb for row in sheet for cell in row)
    wb.close()


def test_verification_failure_does_not_mutate_evidence(tmp_path):
    source = scanner_source(tmp_path)
    path = tmp_path / source.path
    before = path.read_bytes()
    report = verify_package(tmp_path, Package(sources=(source,)))
    assert report["sources"][0]["integrity_unchanged"]
    assert path.read_bytes() == before


@pytest.mark.parametrize("name", [[], {}, 1, False, None])
def test_malformed_ec2_state_is_unknown_not_application_crash(tmp_path, name):
    source = aws_source(
        tmp_path,
        "ec2.json",
        {
            "Reservations": [
                {
                    "OwnerId": A,
                    "Instances": [
                        {
                            "InstanceId": "i-12345678",
                            "State": {"Name": name},
                            "MetadataOptions": {
                                "State": "applied",
                                "HttpTokens": "required",
                                "HttpEndpoint": "enabled",
                            },
                        }
                    ],
                }
            ]
        },
        "ec2:DescribeInstances",
    )
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, tmp_path, Package(sources=(source,)))
        assess(store, run)
        result = next(
            r for r in store.results(run) if r.resource_id and r.rule_id.startswith("R0-")
        )
        assert result.outcome == "UNKNOWN"
        assert result.missing_dependencies == ("State.Name",)


def test_failed_native_responses_do_not_support_successful_snapshot(tmp_path):
    source = aws_source(tmp_path, "good.json", db_response(True), "rds:DescribeDBInstances")
    bad_data = db_response(False)
    bad_data["ResponseMetadata"]["HTTPStatusCode"] = 403
    bad = aws_source(tmp_path, "bad.json", bad_data, "rds:DescribeDBInstances")
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, tmp_path, Package(sources=(source, bad)))
        assess(store, run)
        result = next(
            r for r in store.results(run) if r.resource_id and r.rule_id.startswith("R0-")
        )
        assert result.outcome == "PASS"
        assert len(result.supporting_observations) == 1
        assert len(result.evidence) == 2
        assert summary(store, run)["source_collections_with_errors"] == 1
