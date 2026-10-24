"""Opt-in practical workflow in Chromium against the production bundle (synthetic only)."""

import json
import os
import socket
import threading
import time
from pathlib import Path

import pytest

from tests.browser_diagnostics import capture_browser, word_text

pytestmark = pytest.mark.skipif(
    os.environ.get("RUN_GUI_BROWSER") != "1", reason="opt-in browser runtime required"
)


def test_operational_gui_workflow(tmp_path):
    import uvicorn
    from openpyxl import load_workbook
    from playwright.sync_api import expect, sync_playwright

    from offline_analyzer.api.app import create_app
    from tools.fixtures.r12_demo import build

    package = tmp_path / "package"
    build(package)
    files = [str(p) for p in sorted(package.iterdir())]
    listener = socket.socket()
    listener.bind(("127.0.0.1", 0))
    port = listener.getsockname()[1]
    listener.listen(128)
    origin = f"http://127.0.0.1:{port}"
    app = create_app(tmp_path / "work", origin=origin, bootstrap="browser-test")
    server = uvicorn.Server(
        uvicorn.Config(app, host="127.0.0.1", port=port, log_level="critical", access_log=False)
    )
    thread = threading.Thread(target=lambda: server.run(sockets=[listener]), daemon=True)
    thread.start()
    for _ in range(100):
        if server.started:
            break
        time.sleep(0.05)
    assert server.started
    screenshots = Path("artifacts/gui-screenshots")
    screenshots.mkdir(parents=True, exist_ok=True)
    try:
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(headless=True)
            context = browser.new_context(
                accept_downloads=True, viewport={"width": 1440, "height": 1000}
            )
            context.grant_permissions(["clipboard-read", "clipboard-write"], origin=origin)
            with capture_browser(context, Path("artifacts/recovery-browser/operational")):
                external, errors = [], []

                def route(request):
                    if not request.request.url.startswith(origin):
                        external.append(request.request.url)
                        request.abort()
                    else:
                        request.continue_()

                context.route("**/*", route)
                page = context.new_page()
                page.on("pageerror", lambda exc: errors.append(str(exc)))
                page.goto(origin + "/#token=browser-test")
                # Unlabeled workspaces are conspicuously not approved for customer reporting.
                expect(page.get_by_label("Workspace environment")).to_contain_text("UNLABELED")
                expect(page.get_by_role("note").first).to_contain_text("UNLABELED WORKSPACE")
                page.get_by_role("button", name="Evidence import", exact=True).click()
                page.get_by_label("Select package folder").set_input_files(str(package))
                page.get_by_role("button", name="Validate evidence", exact=True).click()
                page.get_by_role("heading", name="Validation preview").wait_for()
                expect(page.get_by_text("NOT_ESTABLISHED", exact=True)).to_be_visible()
                with page.expect_download() as compatibility:
                    page.get_by_role(
                        "link", name="Download sanitized compatibility summary"
                    ).click()
                summary = json.loads(Path(compatibility.value.path()).read_text())
                assert summary["real_export_validation"] == "NOT_ESTABLISHED"
                assert summary["raw_values_included"] is False
                page.get_by_role("button", name="Import validated evidence").click()
                page.get_by_role("button", name="Run assessment", exact=True).wait_for()
                expect(page.get_by_text("Active run origin: synthetic")).to_be_visible()
                expect(
                    page.get_by_role("button", name="Run assessment", exact=True)
                ).to_be_enabled()
                page.get_by_role("button", name="Run assessment", exact=True).click()
                status = page.get_by_label("Local operation details").get_by_role("status")
                expect(status).to_contain_text("Assessment")
                expect(status).to_contain_text("SUCCEEDED")
                page.get_by_role("button", name="Dashboard", exact=True).click()
                page.get_by_text("Detailed metrics and collection records", exact=True).click()
                page.get_by_text("Accounts represented", exact=True).wait_for()
                assert (
                    page.locator(".metric")
                    .filter(has_text="Accounts represented")
                    .locator("strong")
                    .inner_text()
                    == "5"
                )
                assert (
                    page.locator(".metric")
                    .filter(has_text="Candidate findings")
                    .locator("strong")
                    .inner_text()
                    == "8"
                )
                expect(
                    page.get_by_label("Active evidence run").locator("option:checked")
                ).to_contain_text("Assessed")
                page.screenshot(path=str(screenshots / "dashboard.png"), full_page=True)
                page.get_by_text("Evidence understood", exact=True).wait_for()
                page.screenshot(path=str(screenshots / "coverage-0.4.2.png"), full_page=True)
                page.get_by_role("button", name="Findings workbench", exact=True).click()
                page.get_by_label("Filter findings").select_option("HIGH")
                expect(page.locator("tbody tr").first).to_contain_text("HIGH")
                page.get_by_label("Filter findings").select_option("")
                page.get_by_label("Finding account").select_option("AWS:111111111111")
                expect(page.locator("tbody tr").first).to_contain_text("111111111111")
                page.get_by_label("Finding account").select_option("")
                page.get_by_role("button", name="View", exact=True).first.click()
                page.get_by_role("tab", name="Client Reporting", exact=True).click()
                page.get_by_label("Analyst identity").fill("browser-reviewer")
                page.get_by_label("Analyst remarks").fill("Validate against original evidence")
                # Analyst-triggered transitions only; technical conclusions remain fixed.
                page.get_by_role("tab", name="Client Reporting", exact=True).click()
                page.get_by_role("button", name="Confirm", exact=True).click()
                expect(
                    page.locator(".records tbody").get_by_text("CONFIRMED", exact=True)
                ).to_be_visible()
                page.get_by_role("button", name="View", exact=True).first.click()
                page.get_by_role("tab", name="Client Reporting", exact=True).click()
                editor = page.get_by_label("Report finding editor")
                editor.get_by_label("Report reviewer").fill("browser-report-reviewer")
                editor.get_by_label("Report finding title").fill(
                    "Approved synthetic client finding"
                )
                editor.get_by_label("Include in client report").check()
                editor.get_by_role("button", name="Save report draft", exact=True).click()
                expect(editor.get_by_role("status")).to_contain_text("Report draft saved")
                editor.get_by_role("button", name="Approve recommendation", exact=True).click()
                expect(editor.get_by_role("status")).to_contain_text(
                    "recommendation approval recorded"
                )
                editor.get_by_role("button", name="Approve report narrative", exact=True).click()
                expect(editor.get_by_role("button", name="Copy Report Finding")).to_be_enabled()
                editor.get_by_role("button", name="Copy Report Finding").click()
                expect(editor.get_by_role("status")).to_contain_text("copied")
                clipboard = page.evaluate("navigator.clipboard.readText()")
                assert clipboard.startswith("1. Finding Title\nApproved synthetic client finding")
                assert "13. Applicable References" in clipboard
                page.screenshot(path=str(screenshots / "report-finding-0.4.2.png"), full_page=True)
                page.get_by_role("button", name="Close details", exact=True).click()
                page.get_by_role("button", name="Report center", exact=True).click()
                page.get_by_role("button", name="Export all approved findings (Word DOCX)").click()
                report_row = page.locator("tr").filter(
                    has_text="CloSec-assessment-approved-docx.docx"
                )
                report_row.wait_for()
                with page.expect_download() as approved_report:
                    report_row.get_by_role("link", name="Download").click()
                approved_rtf = word_text(approved_report.value.path())
                assert (
                    "1. Finding Title" in approved_rtf
                    and "Approved synthetic client finding" in approved_rtf
                )
                page.screenshot(path=str(screenshots / "report-center-0.4.2.png"), full_page=True)
                page.get_by_role("button", name="Findings workbench", exact=True).click()
                page.get_by_role("button", name="View", exact=True).first.click()
                page.get_by_role("tab", name="Client Reporting", exact=True).click()
                page.get_by_role("button", name="Reject", exact=True).click()
                expect(
                    page.locator(".records tbody").get_by_text("REJECTED", exact=True)
                ).to_be_visible()
                page.get_by_role("button", name="View", exact=True).first.click()
                page.get_by_role("tab", name="Client Reporting", exact=True).click()
                page.get_by_role("button", name="Record exception", exact=True).click()
                expect(
                    page.locator(".records tbody").get_by_text("NEEDS_VALIDATION", exact=True)
                ).to_be_visible()
                page.screenshot(path=str(screenshots / "findings.png"), full_page=True)
                page.reload()
                page.get_by_role("button", name="Findings workbench", exact=True).click()
                expect(
                    page.locator(".records tbody").get_by_text("NEEDS_VALIDATION", exact=True)
                ).to_be_visible()
                page.get_by_role("button", name="Evidence import", exact=True).click()
                page.get_by_label("Select JSON files").set_input_files(files)
                page.get_by_role("button", name="Validate evidence", exact=True).click()
                page.get_by_role("button", name="Import validated evidence").wait_for()
                page.get_by_role("button", name="Import validated evidence").click()
                page.get_by_text("SUCCEEDED", exact=True).wait_for()
                page.get_by_role("button", name="Findings workbench", exact=True).click()
                expect(
                    page.locator(".records tbody").get_by_text("NEEDS_VALIDATION", exact=True)
                ).to_be_visible()
                for screen in [
                    "Account explorer",
                    "IAM intelligence",
                    "Evidence explorer",
                    "Network analysis",
                    "Security investigations",
                ]:
                    page.get_by_role("button", name=screen, exact=True).click()
                    page.locator("tbody tr").first.wait_for()
                    if screen == "Security investigations":
                        page.screenshot(
                            path=str(screenshots / "security-investigations-0.4.2.png"),
                            full_page=True,
                        )
                    if screen == "Account explorer":
                        page.get_by_role("button", name="View", exact=True).first.click()
                        page.get_by_role("heading", name="Account 111111111111").wait_for()
                    if screen == "IAM intelligence":
                        page.get_by_role("button", name="relationships", exact=True).click()
                        page.locator(".records").first.get_by_role(
                            "button", name="View", exact=True
                        ).first.click()
                        page.get_by_text(
                            "Existing analyzer relationship; no inferred edges.", exact=False
                        ).wait_for()
                        page.get_by_role("button", name="Close details").click()
                        page.screenshot(path=str(screenshots / "iam.png"), full_page=True)
                    if screen == "Evidence explorer":
                        page.get_by_role("button", name="View", exact=True).first.click()
                        page.get_by_role("heading", name="Evidence traceability").wait_for()
                        page.get_by_role("heading", name="Integrity and source locators").wait_for()
                        page.get_by_role("button", name="Close details").click()
                page.get_by_role("button", name="Report center", exact=True).click()
                page.get_by_role("button", name="Generate Excel", exact=True).click()
                page.locator("tr", has_text="assessment-excel.xlsx").get_by_role(
                    "link", name="Download"
                ).wait_for()
                with page.expect_download() as download:
                    page.locator("tr", has_text="assessment-excel.xlsx").get_by_role(
                        "link", name="Download"
                    ).click()
                saved = tmp_path / "export.xlsx"
                download.value.save_as(saved)
                workbook = load_workbook(saved)
                assert workbook["Findings Register"].max_row - 3 == 8
                assert workbook["Control Assessment"].max_row - 3 == 97
                workbook.close()
                page.screenshot(path=str(screenshots / "reports.png"), full_page=True)
                assert not external, external
                assert not errors, errors
            browser.close()
    finally:
        server.should_exit = True
        thread.join(timeout=20)
        listener.close()
        assert not thread.is_alive()


def test_one_command_launcher_lifecycle(tmp_path, monkeypatch):
    """Exercise initialization, browser-open URL and authenticated clean shutdown."""
    from urllib.parse import urlsplit

    import httpx

    from offline_analyzer import launcher

    opened = []
    monkeypatch.setattr(launcher.webbrowser, "open", lambda url: opened.append(url))
    monkeypatch.setattr(
        "sys.argv",
        ["offline-analyzer-gui", "--workspace", str(tmp_path / "launcher"), "--port", "0"],
    )
    result = []
    thread = threading.Thread(target=lambda: result.append(launcher.main()), daemon=True)
    thread.start()
    for _ in range(100):
        if opened:
            break
        time.sleep(0.05)
    assert opened
    parts = urlsplit(opened[0])
    origin = f"{parts.scheme}://{parts.netloc}"
    token = parts.fragment.removeprefix("token=")
    with httpx.Client(base_url=origin, trust_env=False) as client:
        assert client.get("/api/health").json()["ready"]
        session = client.post(
            "/api/session", json={}, headers={"Origin": origin, "X-Bootstrap-Token": token}
        )
        assert session.status_code == 200
        assert client.get("/").status_code == 200
        assert (
            client.post(
                "/api/shutdown",
                json={},
                headers={"Origin": origin, "X-CSRF-Token": session.json()["csrf"]},
            ).status_code
            == 200
        )
    thread.join(timeout=15)
    assert result == [0] and not thread.is_alive()
    launcher.WorkspaceLock(tmp_path / "launcher").close()


def test_launcher_reports_database_startup_failure(tmp_path, monkeypatch):
    from offline_analyzer import launcher

    workspace = tmp_path / "invalid-database"
    workspace.mkdir()
    (workspace / "assessment.duckdb").write_bytes(b"not a database")
    monkeypatch.setattr(
        "sys.argv",
        ["offline-analyzer-gui", "--workspace", str(workspace), "--port", "0", "--no-browser"],
    )
    result = []
    thread = threading.Thread(target=lambda: result.append(launcher.main()), daemon=True)
    thread.start()
    thread.join(timeout=15)
    assert result == [1] and not thread.is_alive()
    launcher.WorkspaceLock(workspace).close()


@pytest.mark.parametrize("profile,transport", [("canonical", "folder"), ("v030", "folder")])
def test_collector_package_gui_workflow(tmp_path, profile, transport):
    """Collector package → validate → import → assess → collection view → Excel (synthetic)."""
    import importlib.util

    import uvicorn
    from openpyxl import load_workbook
    from playwright.sync_api import expect, sync_playwright

    from offline_analyzer.api.app import create_app
    from tests.test_recovery import archive_bytes
    from tools.fixtures.r12_demo import build

    spec = importlib.util.spec_from_file_location("fixture", "tools/collector_fixture.py")
    fixture = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(fixture)
    build(tmp_path / "r12")
    package = fixture.convert(tmp_path / "r12", tmp_path / "collector")[0]
    if profile == "v030":
        from tests.test_activation import v030, zipped

        package = v030.convert(package, tmp_path / "v030")
    archive = tmp_path / "collector.zip"
    archive.write_bytes(zipped(package) if profile == "v030" else archive_bytes(package))
    listener = socket.socket()
    listener.bind(("127.0.0.1", 0))
    port = listener.getsockname()[1]
    listener.listen(128)
    origin = f"http://127.0.0.1:{port}"
    app = create_app(tmp_path / "work", origin=origin, bootstrap="browser-test")
    server = uvicorn.Server(
        uvicorn.Config(app, host="127.0.0.1", port=port, log_level="critical", access_log=False)
    )
    thread = threading.Thread(target=lambda: server.run(sockets=[listener]), daemon=True)
    thread.start()
    for _ in range(100):
        if server.started:
            break
        time.sleep(0.05)
    try:
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(headless=True)
            context = browser.new_context(accept_downloads=True)
            external, errors = [], []

            def route(request):
                if not request.request.url.startswith(origin):
                    external.append(request.request.url)
                    request.abort()
                else:
                    request.continue_()

            context.route("**/*", route)
            page = context.new_page()
            page.on("pageerror", lambda exc: errors.append(str(exc)))
            page.goto(origin + "/#token=browser-test")
            expect(page).to_have_title("CloSec")
            expect(page.locator(".brand")).to_contain_text("CloSec")
            page.get_by_role("button", name="Evidence import", exact=True).click()
            expect(
                page.get_by_role("button", name="Import AWS Evidence ZIP", exact=True)
            ).to_have_count(0)
            expect(
                page.get_by_role("button", name="Import AWS Evidence Folder", exact=True)
            ).to_be_visible()
            with page.expect_file_chooser() as chooser:
                page.get_by_role(
                    "button",
                    name="Import AWS Evidence ZIP"
                    if transport == "zip"
                    else "Import AWS Evidence Folder",
                    exact=True,
                ).click()
            chooser.value.set_files(str(archive if transport == "zip" else package))
            page.get_by_role("button", name="Validate evidence", exact=True).click()
            collector = page.get_by_label("Collector package")
            expect(collector).to_contain_text(
                "Original collector outcomes"
                if profile == "v030"
                else "synthetic-collector-fixture"
            )
            expect(collector).to_contain_text("1 unsupported (retained, not parsed)")
            expect(collector.get_by_text("UNSUPPORTED_OPERATION_RETAINED")).to_be_visible()
            page.get_by_role("button", name="Import validated evidence").click()
            page.get_by_role("button", name="Run assessment", exact=True).wait_for()
            page.get_by_role("button", name="Run assessment", exact=True).click()
            status = page.get_by_label("Local operation details").get_by_role("status")
            expect(status).to_contain_text("Assessment")
            expect(status).to_contain_text("SUCCEEDED")
            page.get_by_role("button", name="Dashboard", exact=True).click()
            page.get_by_text("Detailed metrics and collection records", exact=True).click()
            operations = page.locator("section", has_text="Collection by operation")
            expect(
                operations.locator("tr", has_text="cloudtrail:DescribeTrails").first
            ).to_contain_text("AccessDenied")
            requests = page.locator("section", has_text="Targeted evidence requests")
            expect(
                requests.locator("tr", has_text="cloudtrail:DescribeTrails").first
            ).to_contain_text("Collect the registered operation")
            with page.expect_download() as csv_download:
                page.get_by_role("link", name="Download evidence request list (CSV)").click()
            text = Path(csv_download.value.path()).read_text()
            assert text.startswith("account_id,region,operation,state,reason,request")
            expect(
                page.get_by_role("heading", name="Evidence collection readiness")
            ).to_be_visible()
            page.get_by_label("Collection status", exact=True).select_option("ERROR")
            expect(
                page.locator(
                    "section", has=page.get_by_role("heading", name="Evidence collection readiness")
                )
            ).to_contain_text("AccessDenied")
            with page.expect_download() as request_download:
                page.get_by_role("link", name="Download collector requests (JSON)").click()
            request = json.loads(Path(request_download.value.path()).read_text())
            assert request["schema"] == "closec-evidence-request-1"
            assert request["requests"] and request["manual_validation"]
            page.get_by_role("button", name="Findings workbench", exact=True).click()
            page.get_by_role("button", name="Priority review queue", exact=True).click()
            expect(page.get_by_label("Validation status", exact=True)).to_have_value("CANDIDATE")
            expect(page.get_by_label("Finding class", exact=True)).to_be_visible()
            expect(page.get_by_label("Evidence confidence", exact=True)).to_be_visible()
            expect(page.locator(".records tbody tr").first).to_be_visible()
            if profile == "v030":
                screenshots = Path("artifacts/gui-screenshots")
                screenshots.mkdir(parents=True, exist_ok=True)
                page.screenshot(
                    path=str(screenshots / "activation-v030-findings.png"), full_page=True
                )
            page.get_by_role("button", name="Report center", exact=True).click()
            page.get_by_role("button", name="Generate Excel", exact=True).click()
            row = page.locator("tr", has_text="CloSec-assessment-excel.xlsx")
            row.get_by_role("link", name="Download").wait_for()
            expect(page.get_by_role("button", name="Generate Excel", exact=True)).to_be_enabled()
            with page.expect_download() as download:
                row.get_by_role("link", name="Download").click()
            assert download.value.suggested_filename == "CloSec-assessment-excel.xlsx"
            saved = tmp_path / "export.xlsx"
            download.value.save_as(saved)
            workbook = load_workbook(saved)
            assert workbook.properties.title == "CloSec assessment"
            assert len(workbook.sheetnames) == 27
            workbook.close()
            assert not external, external
            assert not errors, errors
            browser.close()
    finally:
        server.should_exit = True
        thread.join(timeout=20)
        listener.close()
