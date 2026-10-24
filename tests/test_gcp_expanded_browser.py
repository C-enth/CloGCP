"""Real launcher, compiled frontend, local API, DuckDB and security engine acceptance."""

import json
import os
import time
from pathlib import Path

import pytest

from tests.browser_diagnostics import word_text
from tests.launcher_support import LauncherSession
from tests.test_gcp_legacy import zip_pilot
from tools.fixtures.gcp_expanded import expanded

pytestmark = pytest.mark.skipif(os.environ.get("RUN_GUI_BROWSER") != "1", reason="Opt-in browser")


@pytest.mark.parametrize(
    "transport,browser_name",
    [
        ("Folder", "chromium"),
        ("Folder", "webkit"),
    ],
)
def test_expanded_gui_modules_review_exports_and_restart(tmp_path, transport, browser_name):
    from openpyxl import load_workbook
    from playwright.sync_api import expect, sync_playwright

    from offline_analyzer.storage.store import Store

    folder = expanded(tmp_path / "collection", mixed=True)
    archive = zip_pilot(folder, tmp_path / "collection.zip")
    work = tmp_path / "workspace"
    measured = {
        "origin": "synthetic",
        "browser": browser_name,
        "transport": transport,
        "modules": [],
        "jobs": [],
    }
    with LauncherSession(work) as session, sync_playwright() as playwright:
        browser = getattr(playwright, browser_name).launch(headless=True)
        context = browser.new_context(
            accept_downloads=True, viewport={"width": 1600, "height": 1050}
        )
        if browser_name == "chromium":
            context.grant_permissions(["clipboard-read", "clipboard-write"], origin=session.origin)
        errors, outbound = [], []
        context.route(
            "**/*",
            lambda route: (
                route.continue_()
                if route.request.url.startswith(session.origin)
                else (outbound.append("EXTERNAL_REQUEST"), route.abort())
            ),
        )
        page = context.new_page()
        page.on("pageerror", lambda error: errors.append(str(error)))
        page.goto(session.url)
        page.get_by_role("button", name="Evidence import", exact=True).wait_for()
        health = context.request.get(session.origin + "/api/health").json()
        assert health["ready"]
        assert context.request.get(session.origin + "/api/runtime").json()["compatible"]
        page.get_by_role("tab", name="GCP", exact=True).click()
        page.get_by_role("button", name="Evidence import", exact=True).click()
        with page.expect_file_chooser() as chooser:
            page.get_by_role("button", name=f"Import GCP Evidence {transport}", exact=True).click()
        chooser.value.set_files(str(folder if transport == "Folder" else archive))
        if transport == "Folder":
            paths = page.get_by_label("Select package folder").evaluate(
                "input => Array.from(input.files).map(f => f.webkitRelativePath)"
            )
            assert len(paths) == sum(f.is_file() for f in folder.rglob("*"))
            assert all("/" in path for path in paths)
            assert any(path.endswith("/collection_manifest.json") for path in paths)

        def wait_job(job_id):
            deadline = time.monotonic() + 60
            while time.monotonic() < deadline:
                job = context.request.get(session.origin + "/api/jobs/" + job_id).json()
                if job["state"] not in {"PENDING", "RUNNING"}:
                    break
                page.wait_for_timeout(50)
            assert job["state"] == "COMPLETED", job.get("error")
            measured["jobs"].append(
                {
                    "id": job_id,
                    "state": job["state"],
                    "type": job["kind"],
                    "duration_seconds": job["finished_at"] - job["started_at"],
                    "stage": job.get("stage"),
                }
            )
            return job

        page.get_by_role("button", name="Validate evidence", exact=True).click()
        button = page.get_by_role("button", name="Import validated evidence", exact=True)
        expect(button).to_be_visible(timeout=30000)
        expect(page.get_by_label("Manifest detection details")).to_contain_text(
            "assessment-expanded-2.1"
        )
        page.get_by_role(
            "checkbox", name="I acknowledge that this legacy package", exact=False
        ).check()
        with page.expect_response(
            lambda response: response.url.endswith("/api/import")
        ) as submitted:
            button.click()
        wait_job(submitted.value.json()["id"])
        assess = page.get_by_role("button", name="Run assessment", exact=True)
        expect(assess).to_be_enabled(timeout=30000)
        run = page.get_by_label("Active evidence run").input_value()
        measured["run_id"] = run
        with page.expect_response(
            lambda response: response.url.endswith(f"/api/runs/{run}/assess")
        ) as submitted:
            assess.click()
        wait_job(submitted.value.json()["id"])
        expect(assess).to_be_enabled(timeout=30000)
        expect(page.get_by_label("Active evidence run").locator("option:checked")).to_contain_text(
            "Assessed"
        )

        modules = {
            "Account explorer": ["accounts"],
            "IAM intelligence": ["roles", "policies", "relationships", "gaps", "iam_findings"],
            "Evidence explorer": ["evidence", "coverage"],
            "Network analysis": ["network", "network_findings"],
            "Security investigations": ["security_investigations"],
            "Findings workbench": ["findings"],
            "Cross-account explorer": [
                "roles",
                "policies",
                "relationships",
                "gaps",
                "iam_findings",
            ],
        }
        for module, kinds in modules.items():
            page.get_by_role("button", name=module, exact=True).click()
            expect(page.locator(".records")).to_have_count(len(kinds))
            for index, kind in enumerate(kinds):
                data = context.request.get(
                    session.origin + f"/api/runs/{run}/records/{kind}?provider=GCP&limit=30"
                ).json()
                record = page.locator(".records").nth(index)
                expect(record.locator(".toolbar .muted").first).to_have_text(
                    f"{data['total']} records", timeout=30000
                )
                expect(record.locator("tbody tr")).to_have_count(min(30, data["total"]))
                measured["modules"].append(
                    {
                        "module": module,
                        "endpoint": f"/api/runs/{run}/records/{kind}",
                        "run_id": run,
                        "backend_count": data["total"],
                        "displayed_count": min(30, data["total"]),
                        "status": "PASS",
                        "limitation": "Unavailable/coverage-limited" if data["total"] == 0 else "",
                    }
                )
            assert page.get_by_label("Active evidence run").input_value() == run

        page.get_by_role("button", name="Dashboard", exact=True).click()
        page.get_by_role("tab", name="Assessment Coverage", exact=True).click()
        page.get_by_text("Collector-to-assessment coverage", exact=True).click()
        expect(page.get_by_text("NOT_EVALUATED", exact=True).first).to_be_visible()
        page.get_by_role("button", name="Findings workbench", exact=True).click()
        finding_data = context.request.get(
            session.origin + f"/api/runs/{run}/records/findings?provider=GCP&limit=30"
        ).json()
        assert finding_data["total"] > 0
        page.locator(".records tbody tr").first.get_by_role(
            "button", name="View", exact=True
        ).click()
        page.get_by_role("tab", name="Client Reporting", exact=True).click()
        page.get_by_label("Analyst identity").fill("synthetic-ar01-reviewer")
        page.get_by_label("Analyst remarks").fill("Synthetic evidence reviewed locally")
        page.get_by_role("tab", name="Client Reporting", exact=True).click()
        page.get_by_role("button", name="Confirm", exact=True).click()
        expect(page.locator(".records tbody")).to_contain_text("CONFIRMED")
        page.locator(".records tbody tr").first.get_by_role(
            "button", name="View", exact=True
        ).click()
        page.get_by_role("tab", name="Client Reporting", exact=True).click()
        editor = page.get_by_label("Report finding editor")
        editor.get_by_label("Report reviewer").fill("synthetic-ar01-reviewer")
        editor.get_by_label("Include in client report").check()
        editor.get_by_label("Recommended remediation").fill(
            "Synthetic reviewed remediation: use scoped roles and rotate keys "
            "after authorized validation."
        )
        editor.get_by_role("button", name="Save report draft", exact=True).click()
        expect(editor.get_by_role("status")).to_contain_text("Report draft saved")
        editor.get_by_role("button", name="Approve recommendation", exact=True).click()
        expect(editor.get_by_role("status")).to_contain_text("recommendation approval recorded")
        editor.get_by_role("button", name="Approve report narrative", exact=True).click()
        expect(editor.get_by_role("button", name="Copy Report Finding")).to_be_enabled()
        if browser_name == "chromium":
            editor.get_by_role("button", name="Copy Report Finding").click()
            expect(editor.get_by_role("status")).to_contain_text("copied")
            assert page.evaluate("navigator.clipboard.readText()")
        page.get_by_role("button", name="Close details", exact=True).click()
        page.get_by_role("button", name="Report center", exact=True).click()
        for format, suffix in (("Excel", "xlsx"), ("DOCX", "docx")):
            page.get_by_role(
                "button",
                name="Generate Excel"
                if format == "Excel"
                else "Export all approved findings (Word DOCX)",
                exact=True,
            ).click()
            report = (
                page.locator("tr")
                .filter(
                    has_text="CloSec-assessment-"
                    + ("excel" if suffix == "xlsx" else "approved-docx")
                    + f".{suffix}"
                )
                .filter(has_text=run[:20])
                .first
            )
            link = report.get_by_role("link", name="Download", exact=True)
            expect(link).to_be_visible(timeout=30000)
            with page.expect_download() as download:
                link.click()
            destination = tmp_path / f"result.{suffix}"
            download.value.save_as(destination)
            assert destination.stat().st_size > 0
            if suffix == "xlsx":
                workbook = load_workbook(destination, read_only=True)
                assert "Cloud Collection Coverage" in workbook.sheetnames
                workbook.close()
            else:
                assert "CloSec" in word_text(destination)
        page.get_by_role("button", name="Manage Assessment Data", exact=True).click()
        expect(
            page.get_by_role("table", name="Assessed scopes").get_by_role(
                "button", name="sample-project", exact=True
            )
        ).to_be_visible()
        page.reload()
        expect(page.get_by_label("Active evidence run")).to_have_value(run)
        assert not errors and not outbound
        browser.close()
    with Store(work) as store:
        assert len(list(store.source_data(run))) == 115
        assert (
            list(store.observations(run))
            and list(store.results(run))
            and list(store.candidates(run))
        )
        assert any(d.status == "CONFIRMED" for d in store.latest_decisions(run).values())
    # Actual launcher restart verifies persisted assessment selection/data remain available.
    with LauncherSession(work) as session:
        assert any(
            r["run_id"] == run and r["assessed"] for r in session.client.get("/api/runs").json()
        )
    output = Path("artifacts/ar01-browser")
    output.mkdir(parents=True, exist_ok=True)
    (output / f"{browser_name}-{transport}.json").write_text(json.dumps(measured, indent=2))
