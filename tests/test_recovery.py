"""Synthetic operational recovery, import boundaries and semantic reporting regression."""

import io
import json
import stat
import threading
import time
import zipfile
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from offline_analyzer.api.app import create_app
from offline_analyzer.enterprise.refresh import refresh_enterprise, refresh_plan
from offline_analyzer.enterprise.snapshot import assess_enterprise
from offline_analyzer.findings.generate import record_decision, validation_status
from offline_analyzer.ingestion.archive import unpack_package
from offline_analyzer.ingestion.collector import COLLECTOR_MANIFEST, Operation, inspect_collector
from offline_analyzer.ingestion.importer import import_package
from offline_analyzer.ingestion.json_input import InputError
from offline_analyzer.ingestion.safety import sha256_file
from offline_analyzer.inventory.environment import record_run_origin, set_environment
from offline_analyzer.normalization.adapters import configuration_for
from offline_analyzer.reporting.collection import collection_readiness
from offline_analyzer.rules.engine import assess
from offline_analyzer.storage.store import Store
from tests.test_api import ORIGIN, TOKEN, authenticated, await_job
from tests.test_collector import fixture


def archive_bytes(package: Path) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for path in package.rglob("*.json"):
            archive.writestr("package/" + path.relative_to(package).as_posix(), path.read_bytes())
    return output.getvalue()


@pytest.mark.parametrize(
    "name",
    [
        "../escape.json",
        "/root.json",
        "C:/x.json",
        "a\\..\\b.json",
        "CON.json",
        "a./x.json",
        "a//x.json",
        "execute.py",
    ],
)
def test_archive_refuses_unsafe_members_before_extracting(tmp_path, name):
    path = tmp_path / "test.zip"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("collector-manifest.json", "{}")
        archive.writestr(name, "{}")
    destination = tmp_path / "expanded"
    destination.mkdir()
    with pytest.raises(InputError):
        unpack_package(path, destination)
    assert not list(destination.iterdir())


@pytest.mark.parametrize("case", ["symlink", "duplicate", "oversize", "corrupt", "unknown-root"])
def test_archive_rejects_invalid_packages(tmp_path, case):
    path = tmp_path / "test.zip"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("collector-manifest.json", "{}")
        if case == "symlink":
            info = zipfile.ZipInfo("link.json")
            info.external_attr = (stat.S_IFLNK | 0o777) << 16
            archive.writestr(info, "outside")
        elif case == "duplicate":
            archive.writestr("COLLECTOR-MANIFEST.json", "{}")
        elif case == "oversize":
            archive.writestr("large.json", " " * (65 * 1024 * 1024))
        elif case == "unknown-root":
            archive.writestr("other/manifest.json", "{}")
    if case == "corrupt":
        path.write_bytes(b"corrupt zip")
    destination = tmp_path / "expanded"
    destination.mkdir()
    with pytest.raises(InputError):
        unpack_package(path, destination)


def add_identity(package, account="111111111111"):
    data = json.loads((package / COLLECTOR_MANIFEST).read_text())
    path = package / "responses/identity.json"
    path.write_text(
        json.dumps(
            {
                "Account": account,
                "Arn": f"arn:aws:iam::{account}:user/synthetic",
                "UserId": "synthetic",
            }
        )
    )
    data["operations"] = [
        o for o in data["operations"] if o["operation"] != "sts:GetCallerIdentity"
    ]
    data["operations"].append(
        {
            "operation": "sts:GetCallerIdentity",
            "region": "global",
            "status": "SUCCEEDED",
            "pages": 1,
            "path": "responses/identity.json",
            "sha256": sha256_file(path),
            "started_at": data["started_at"],
            "completed_at": data["completed_at"],
        }
    )
    (package / COLLECTOR_MANIFEST).write_text(json.dumps(data))


def test_sts_identity_metadata_is_validated_without_inventing_resources(tmp_path):
    from tools.fixtures.r12_demo import build

    build(tmp_path / "native")
    package = fixture.convert(tmp_path / "native", tmp_path / "collector")[0]
    add_identity(package)
    with Store(tmp_path / "work") as store:
        run = import_package(store, package)
        assert any(s["operation"] == "sts:GetCallerIdentity" for s in store.source_data(run))
        assert not any(o.operation == "sts:GetCallerIdentity" for o in store.observations(run))
    add_identity(package, "222222222222")
    with pytest.raises(InputError, match="IDENTITY_MISMATCH"):
        inspect_collector(package)


def test_global_collector_partition_preserves_account_bpa_identity(tmp_path):
    from tools.fixtures.r12_demo import build

    build(tmp_path / "native")
    package = fixture.convert(tmp_path / "native", tmp_path / "collector")[0]
    data = json.loads((package / COLLECTOR_MANIFEST).read_text())
    path = package / "responses/bpa.json"
    path.write_text(
        json.dumps(
            {
                "PublicAccessBlockConfiguration": {
                    "BlockPublicAcls": True,
                    "IgnorePublicAcls": True,
                    "BlockPublicPolicy": True,
                    "RestrictPublicBuckets": True,
                }
            }
        )
    )
    data["operations"].append(
        {
            "operation": "s3control:GetPublicAccessBlock",
            "region": "global",
            "status": "SUCCEEDED",
            "pages": 1,
            "path": "responses/bpa.json",
            "sha256": sha256_file(path),
            "started_at": data["started_at"],
            "completed_at": data["completed_at"],
        }
    )
    (package / COLLECTOR_MANIFEST).write_text(json.dumps(data))
    with Store(tmp_path / "work") as store:
        run = import_package(store, package)
        bpa = next(
            o for o in store.observations(run) if o.operation == "s3control:GetPublicAccessBlock"
        )
        assert bpa.partition == "aws" and bpa.resource_id and bpa.region == "global"
        assert not bpa.unresolved


@pytest.mark.parametrize("interface", ["api", "importer"])
def test_declared_synthetic_collector_cannot_be_mislabelled_into_assessment(tmp_path, interface):
    from tools.fixtures.r12_demo import build

    build(tmp_path / "native")
    package = fixture.convert(tmp_path / "native", tmp_path / "collector")[0]
    work = tmp_path / "work"
    set_environment(work, "ASSESSMENT", "isolation-test")
    if interface == "api":
        app = create_app(work, origin=ORIGIN, bootstrap=TOKEN)
        with TestClient(app, base_url=ORIGIN) as client:
            headers = authenticated(client)
            response = client.post(
                "/api/uploads",
                headers=headers,
                files=[("files", ("collector.zip", archive_bytes(package), "application/zip"))],
                data={"paths": '["collector.zip"]', "context": '{"origin":"tool-generated"}'},
            )
            assert response.status_code == 400
            assert response.json()["error"] == "ASSESSMENT_WORKSPACE_REJECTS_SYNTHETIC"
            assert not (work / "uploads").exists()
    else:
        with Store(work) as store, pytest.raises(InputError, match="REJECTS_SYNTHETIC"):
            import_package(store, package)
    with Store(work) as store:
        assert store.db.execute("SELECT count(*) FROM runs").fetchone()[0] == 0


def test_readiness_separates_declared_regions_and_unsupported_access_denial(tmp_path):
    package = tmp_path / "package"
    package.mkdir()
    path = package / "instances.json"
    path.write_text('{"Reservations": []}')
    stamp = "2026-10-10T07:00:00Z"
    data = {
        "schema": "closec-collector-1",
        "package_id": "synthetic-readiness",
        "collector": {"name": "synthetic", "version": "1.0"},
        "account_id": "111111111111",
        "dataset_origin": "SYNTHETIC",
        "started_at": stamp,
        "completed_at": stamp,
        "regions": ["us-east-1", "eu-west-1"],
        "operations": [
            {
                "operation": "ec2:DescribeInstances",
                "region": "us-east-1",
                "status": "SUCCEEDED",
                "pages": 1,
                "path": "instances.json",
                "sha256": sha256_file(path),
                "started_at": stamp,
                "completed_at": stamp,
            },
            {
                "operation": "securityhub:GetFindings",
                "region": "eu-west-1",
                "status": "FAILED",
                "error_code": "AccessDenied",
                "started_at": stamp,
                "completed_at": stamp,
            },
        ],
    }
    (package / COLLECTOR_MANIFEST).write_text(json.dumps(data))
    with Store(tmp_path / "work") as store:
        run = import_package(store, package)
        result = collection_readiness(store, run)
    row = result["accounts"][0]
    assert row["regions"] == ["us-east-1"]
    assert row["declared_regions"] == ["eu-west-1", "us-east-1"]
    assert row["access_denied"] == 1 and row["unsupported_evidence"] == 1
    assert any(op.startswith("eu-west-1:") for op in row["missing_required_operations"])


@pytest.mark.parametrize(
    "versions,declared",
    [
        ([], "v1"),
        ([{"IsDefaultVersion": True, "VersionId": "v2", "Document": {"Statement": []}}], "v1"),
        ([{"IsDefaultVersion": True, "Document": {"Statement": []}}] * 2, None),
    ],
)
def test_aggregate_iam_missing_or_ambiguous_default_never_resolves(versions, declared):
    config, _ = configuration_for(
        "iam:GetAccountAuthorizationDetails#policy",
        {"DefaultVersionId": declared, "PolicyVersionList": versions},
    )
    assert config["DefaultDocument"] is None


def test_pagination_and_unknown_format_fail_closed(tmp_path):
    base = dict(
        operation="ec2:DescribeInstances",
        region="us-east-1",
        status="SUCCEEDED",
        path="p.json",
        sha256="0" * 64,
        started_at="2026-10-10T07:00:00Z",
        completed_at="2026-10-10T07:00:00Z",
    )
    with pytest.raises(ValueError):
        Operation(**base, format="unregistered")


@pytest.mark.parametrize(
    "version,profile,accepted",
    [
        ("5.26.1", "prowler-ocsf-5.26.1", True),
        ("5.10.0", "prowler-ocsf-5.10", True),
        ("5.26.1", "prowler-ocsf-5.10", False),
        ("5.27.0", "prowler-ocsf-5.26.1", False),
    ],
)
def test_prowler_exact_profiles_are_synthetic_only(tmp_path, version, profile, accepted):
    from offline_analyzer.verification import verify_package
    from tools.fixtures.demo import build_package

    package_path = tmp_path / "package"
    build_package(package_path)
    path = package_path / "prowler.json"
    records = json.loads(path.read_text())
    for record in records:
        record["metadata"]["product"]["version"] = version
        record.setdefault("unmapped", {})["scan_id"] = "synthetic-scan"
    path.write_text(json.dumps(records))
    manifest = json.loads((package_path / "manifest.json").read_text())
    for source in manifest["sources"]:
        if source["operation"] == "prowler:Scan":
            source["format"] = profile
            source["sha256"] = sha256_file(path)
    (package_path / "manifest.json").write_text(json.dumps(manifest))
    diagnostic = verify_package(package_path, origin="synthetic")
    scanner = next(s for s in diagnostic["sources"] if s["operation"] == "prowler:Scan")
    assert scanner["parser_accepted"] is accepted
    assert diagnostic["real_export_validation"] == "NOT_ESTABLISHED"


def test_zip_multaccount_requests_updated_evidence_review_restart(tmp_path):
    from offline_analyzer.release_audit import audit
    from offline_analyzer.storage.backup import backup, restore
    from tools.fixtures.r12_demo import build

    build(tmp_path / "native")
    packages = fixture.convert(tmp_path / "native", tmp_path / "collector")[:2]
    work = tmp_path / "work"
    set_environment(work, "SYNTHETIC", "recovery")
    app = create_app(work, origin=ORIGIN, bootstrap=TOKEN)
    with TestClient(app, base_url=ORIGIN) as client:
        headers = authenticated(client)
        runs = []
        for package in packages:
            preview = client.post(
                "/api/uploads",
                headers=headers,
                files=[("files", ("collector.zip", archive_bytes(package), "application/zip"))],
                data={"paths": '["collector.zip"]', "context": '{"origin":"synthetic"}'},
            ).json()
            assert preview["can_import"]
            run = await_job(
                client,
                client.post(
                    "/api/import", headers=headers, json={"upload_id": preview["upload_id"]}
                ),
            )["run_id"]
            runs.append(run)
        enterprise = await_job(
            client, client.post("/api/enterprise/assess", headers=headers, json={"members": runs})
        )["run_id"]
        readiness = client.get(f"/api/runs/{enterprise}/collection").json()
        assert len(readiness["accounts"]) == 2
        assert any(a["access_denied"] and a["unsupported_evidence"] for a in readiness["accounts"])
        request = client.get(f"/api/runs/{enterprise}/evidence-requests.json").json()
        from offline_analyzer.contracts.r0 import AWS_OPERATIONS

        assert request["requests"] and request["manual_validation"]
        assert all(r["operation"] in AWS_OPERATIONS for r in request["requests"])
        assert any(
            r["related_controls"] and r["related_investigations"] for r in request["requests"]
        )
        assert client.get(f"/api/runs/{enterprise}/evidence-requests.csv").status_code == 200
    with Store(work) as store:
        finding = next(store.candidates(enterprise))
        record_decision(
            store, enterprise, finding.finding_id, "CONFIRMED", "synthetic-analyst", "Reviewed"
        )
        data = json.loads((packages[0] / COLLECTOR_MANIFEST).read_text())
        data["package_id"] += "-updated"
        (packages[0] / COLLECTOR_MANIFEST).write_text(json.dumps(data))
        updated = import_package(store, packages[0])
        record_run_origin(work, updated, "synthetic")
        assert updated != runs[0]
        plan = refresh_plan(store, enterprise, [updated])
        assert runs[1] in plan["retained_members"] and runs[0] in plan["replaced_members"]
        assert plan["changed_resource_facts"] == 0
        assert plan["previous_affected_resources"] > 0
        refreshed = refresh_enterprise(store, enterprise, [updated])["run_id"]
        same = next(c for c in store.candidates(refreshed) if c.finding_id == finding.finding_id)
        assert validation_status(store, same)[0] == "CONFIRMED"
        assert import_package(store, packages[0]) == updated
        assert len(collection_readiness(store, refreshed)["accounts"]) == 2
        assert store.require_run(enterprise) is None
    with TestClient(create_app(work, origin=ORIGIN, bootstrap=TOKEN), base_url=ORIGIN) as client:
        headers = authenticated(client)
        preview = client.post(
            "/api/enterprise/refresh-preview",
            headers=headers,
            json={"baseline": enterprise, "replacements": [updated]},
        ).json()
        assert preview["mode"] == "SAFE_FULL_RECOMPUTATION"
        result = await_job(
            client,
            client.post(
                "/api/enterprise/refresh",
                headers=headers,
                json={"baseline": enterprise, "replacements": [updated]},
            ),
        )
        assert result["run_id"] == refreshed
    assert audit(work, refreshed)["result"] == "PASS"
    backup(work, tmp_path / "backup")
    restore(tmp_path / "backup", tmp_path / "restored")
    with Store(tmp_path / "restored") as store:
        assert collection_readiness(store, refreshed)["collector_packages"]
        same = next(c for c in store.candidates(refreshed) if c.finding_id == finding.finding_id)
        assert validation_status(store, same)[0] == "CONFIRMED"


def test_report_status_is_responsive_while_single_writer_is_busy(tmp_path):
    app = create_app(tmp_path / "work", origin=ORIGIN, bootstrap=TOKEN)
    entered, release = threading.Event(), threading.Event()
    with TestClient(app, base_url=ORIGIN) as client:
        authenticated(client)

        def blocked(store):
            entered.set()
            release.wait(5)

        job = app.state.coordinator.submit("Report", blocked, key="export")
        try:
            assert entered.wait(2)
            start = time.perf_counter()
            response = client.get(f"/api/jobs/{job['id']}")
            assert response.json()["status"] == "RUNNING"
            assert time.perf_counter() - start < 1
        finally:
            release.set()


def test_optimized_workbook_semantics_match_reference_writer_and_database(tmp_path, monkeypatch):
    from copy import copy

    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    from offline_analyzer.reporting import excel
    from tools.fixtures.r12_demo import build

    def append(ws, values):
        cells = []
        for value in values:
            cell = WriteOnlyCell(
                ws, value if type(value) in (int, bool) else excel.safe_text(value)
            )
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cells.append(cell)
        ws.append(cells)
        ws.r0_row_count += 1

    def sheet(wb, name, headers):
        ws = wb.create_sheet(name)
        ws.freeze_panes = "D4"
        ws.r0_row_count = 0
        ws.merged_cells.add(f"A1:{get_column_letter(len(headers))}1")
        for i, _ in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(i)].width = 26 if i < 4 else 42
        append(ws, [excel.NOTICE])
        append(ws, ["Assessment run", wb.properties.identifier])
        ws.merged_cells.add("B2:D2")
        row = []
        for value in headers:
            cell = WriteOnlyCell(ws, excel.safe_text(value))
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.alignment = Alignment(wrap_text=True)
            row.append(cell)
        ws.append(row)
        ws.r0_row_count = 3
        return ws

    build(tmp_path / "package")
    with Store(tmp_path / "work") as store:
        run = import_package(store, tmp_path / "package")
        assess(store, run)
        enterprise = assess_enterprise(store, [run])
        excel.export_excel(store, enterprise, tmp_path / "optimized.xlsx")
        monkeypatch.setattr(excel, "sheet", sheet)
        monkeypatch.setattr(excel, "append", append)
        excel.export_excel(store, enterprise, tmp_path / "reference.xlsx")
        expected_findings = {c.finding_id for c in store.candidates(enterprise)}
        expected_evaluations = {r.evaluation_id for r in store.results(enterprise)}
    optimized = load_workbook(tmp_path / "optimized.xlsx")
    reference = load_workbook(tmp_path / "reference.xlsx")
    assert optimized.sheetnames == reference.sheetnames
    for ws in optimized:
        old = reference[ws.title]
        assert list(ws.values) == list(old.values)
        assert ws.auto_filter.ref == old.auto_filter.ref
        assert ws.freeze_panes == old.freeze_panes
        for row in ws:
            for cell in row:
                other = old[cell.coordinate]
                assert cell.data_type == other.data_type != "f"
                if cell.value is not None:
                    assert copy(cell.alignment) == copy(other.alignment)
                    assert copy(cell.font) == copy(other.font)
                    assert copy(cell.fill) == copy(other.fill)
    assert {
        row[0].value for row in optimized["Findings Register"].iter_rows(min_row=4)
    } == expected_findings
    assert {
        row[0].value for row in optimized["Control Assessment"].iter_rows(min_row=4)
    } == expected_evaluations
    optimized.close()
    reference.close()
