"""IAM evaluator semantics and R1.2 golden scenarios (synthetic evidence only)."""

import json
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from offline_analyzer.findings.generate import record_decision, validation_status
from offline_analyzer.iam.analysis import IamAnalysis
from offline_analyzer.iam.evaluator import (
    ALLOWED,
    DENIED,
    POTENTIAL,
    UNKNOWN_DECISION,
    Request,
    evaluate_identity,
)
from offline_analyzer.iam.organizations import OrganizationModel
from offline_analyzer.iam.policy import (
    FALSE,
    TRUE,
    UNKNOWN,
    Condition,
    evaluate_condition,
    parse_policy,
)
from offline_analyzer.ingestion.importer import import_package
from offline_analyzer.inventory.accounts import install_registry, workspace_registry
from offline_analyzer.reporting.excel import export_excel
from offline_analyzer.rules.engine import assess
from offline_analyzer.storage.store import Store
from tools.fixtures.r12_demo import M, N, S, T, U, X, arn, build, registry_file

ROLE = SimpleNamespace(
    arn=arn(T, "probe"), account=T, boundary_arn=None, service_linked=False, partition="aws"
)
NO_ORG = OrganizationModel([], [])


def inventory(*documents, missing=(), boundary=None):
    docs = [parse_policy(d, f"p{i}") for i, d in enumerate(documents)]
    return SimpleNamespace(
        identity_documents=lambda role: (docs, list(missing)),
        managed_document=lambda arn_: parse_policy(boundary, "boundary") if boundary else None,
    )


def statement(effect="Allow", action="iam:PassRole", resource="*", **extra):
    return {"Statement": [{"Effect": effect, "Action": action, "Resource": resource, **extra}]}


TARGET = arn(T, "x")


def decide(
    *documents, action="iam:PassRole", resource=TARGET, context=None, org=NO_ORG, role=ROLE, **kw
):
    return evaluate_identity(
        inventory(*documents, **kw), org, role, Request(action, resource, context or {})
    )


# --- policy semantics -------------------------------------------------------------------


@pytest.mark.parametrize(
    "document,action,resource,expected",
    [
        (statement(action="IAM:passrole"), "iam:PassRole", arn(T, "x"), POTENTIAL),
        (statement(action="iam:Pass*"), "iam:PassRole", arn(T, "x"), POTENTIAL),
        (statement(action="iam:Get*"), "iam:PassRole", arn(T, "x"), DENIED),
        (
            {"Statement": [{"Effect": "Allow", "NotAction": "iam:*", "Resource": "*"}]},
            "iam:PassRole",
            arn(T, "x"),
            DENIED,
        ),
        (
            {"Statement": [{"Effect": "Allow", "NotAction": "iam:*", "Resource": "*"}]},
            "s3:GetObject",
            "arn:aws:s3:::b/k",
            POTENTIAL,
        ),
        (
            {
                "Statement": [
                    {"Effect": "Allow", "Action": "iam:PassRole", "NotResource": arn(T, "x")}
                ]
            },
            "iam:PassRole",
            arn(T, "x"),
            DENIED,
        ),
        (
            {
                "Statement": [
                    {"Effect": "Allow", "Action": "iam:PassRole", "NotResource": arn(T, "x")}
                ]
            },
            "iam:PassRole",
            arn(T, "y"),
            POTENTIAL,
        ),
        (
            statement(resource=f"arn:aws:iam::{T}:role/${{aws:username}}"),
            "iam:PassRole",
            arn(T, "x"),
            POTENTIAL,
        ),
        (
            statement(resource=f"arn:aws:ec2:*:{T}:instance/*", action="ec2:RunInstances"),
            "ec2:RunInstances",
            None,
            POTENTIAL,
        ),
    ],
)
def test_action_and_resource_semantics(document, action, resource, expected):
    # No organization evidence: the best attainable state is POTENTIALLY_ALLOWED.
    assert decide(document, action=action, resource=resource).state == expected


UNSUPPORTED = [
    {"NumericLessThan": {"aws:MultiFactorAuthAge": "3600"}},
    {"DateLessThan": {"aws:CurrentTime": "2030-01-01T00:00:00Z"}},
    {"IpAddress": {"aws:SourceIp": "203.0.113.0/24"}},
    {"ForAnyValue:StringEquals": {"aws:TagKeys": "x"}},
    {"BinaryEquals": {"aws:x": "AA=="}},
    {"StringEquals": {"aws:PrincipalTag/team": "${aws:username}"}},
    {"StringEquals": {"aws:RequestedRegion": "us-east-1"}},  # key not in request context
]


@pytest.mark.parametrize("condition", UNSUPPORTED)
def test_unsupported_semantics_never_become_allowed(condition):
    full_org = organization_allowing_everything()
    result = decide(statement(Condition=condition), org=full_org)
    assert result.state == POTENTIAL and result.state != ALLOWED
    unsupported_principal = {
        "Statement": [{"Effect": "Allow", "Action": "*", "Resource": "*", "Principal": "*"}]
    }
    assert decide(unsupported_principal, org=full_org).state == UNKNOWN_DECISION
    malformed = {"Statement": [{"Effect": "Allow", "Action": "*", "Resource": "*", "Bogus": 1}]}
    assert decide(malformed, org=full_org).state == UNKNOWN_DECISION
    deny_unsupported = {
        "Statement": [
            statement()["Statement"][0],
            {"Effect": "Deny", "Action": "*", "Resource": "*", "Condition": condition},
        ]
    }
    assert decide(deny_unsupported, org=full_org).state == POTENTIAL


@pytest.mark.parametrize(
    "operator,key,values,context,expected",
    [
        ("StringEquals", "aws:PrincipalOrgID", ("o-1",), {"aws:principalorgid": ("o-1",)}, TRUE),
        ("StringEquals", "aws:PrincipalOrgID", ("o-1",), {"aws:principalorgid": ()}, FALSE),
        ("StringNotEquals", "aws:PrincipalOrgID", ("o-1",), {"aws:principalorgid": ()}, TRUE),
        ("StringEqualsIfExists", "aws:PrincipalOrgID", ("o-1",), {"aws:principalorgid": ()}, TRUE),
        (
            "StringLike",
            "aws:PrincipalArn",
            ("arn:*:role/a*",),
            {"aws:principalarn": ("arn:x:role/ab",)},
            TRUE,
        ),
        ("Bool", "aws:MultiFactorAuthPresent", ("true",), {}, UNKNOWN),
        ("Null", "sts:ExternalId", ("false",), {"sts:externalid": ()}, FALSE),
        (
            "ArnLike",
            "aws:PrincipalArn",
            ("arn:aws:iam::1:role/*",),
            {"aws:principalarn": ("a", "b")},
            UNKNOWN,
        ),
    ],
)
def test_condition_three_valued_logic(operator, key, values, context, expected):
    assert evaluate_condition(Condition(operator, key, values), context)[0] == expected


def test_explicit_deny_and_boundary():
    allow, deny = statement(), statement(effect="Deny")
    assert decide(allow, deny).state == DENIED
    boundary_role = SimpleNamespace(**{**ROLE.__dict__, "boundary_arn": "arn:aws:iam::1:policy/b"})
    assert decide(allow, role=boundary_role).state == POTENTIAL  # boundary document missing
    assert "permissions boundary document" in " ".join(decide(allow, role=boundary_role).missing)
    narrow = statement(action="s3:*")
    assert decide(allow, role=boundary_role, boundary=narrow).state == DENIED
    assert decide(allow, role=boundary_role, boundary=statement(effect="Deny")).state == DENIED
    assert decide(missing=("managed policy document X",)).state == UNKNOWN_DECISION
    assert decide(allow, missing=("managed policy document X",)).state == POTENTIAL


def org_observations(scp_content=None, types=("SERVICE_CONTROL_POLICY",), attach_account=True):
    def obs(kind, native, config):
        return SimpleNamespace(resource_type=kind, native_id=native, configuration=config)

    full = {"Statement": [{"Effect": "Allow", "Action": "*", "Resource": "*"}]}
    items = [
        obs("AWS::Organizations::Organization", "o-abcdefghij", {"MasterAccountId": M}),
        obs(
            "AWS::Organizations::Root",
            "r-ab12",
            {"PolicyTypes": [{"Type": t, "Status": "ENABLED"} for t in types]},
        ),
        obs("AWS::Organizations::Account", T, {"ParentId": "r-ab12"}),
        obs(
            "AWS::Organizations::Policy",
            "p-FullAWSAccess",
            {"Type": "SERVICE_CONTROL_POLICY", "Content": full},
        ),
        obs(
            "AWS::Organizations::TargetPolicyList",
            "r-ab12#SERVICE_CONTROL_POLICY",
            {
                "Filter": "SERVICE_CONTROL_POLICY",
                "PolicyIds": ["p-FullAWSAccess"],
                "Truncated": False,
            },
        ),
    ]
    if attach_account:
        ids = ["p-FullAWSAccess"] + (["p-custom000001"] if scp_content else [])
        items.append(
            obs(
                "AWS::Organizations::TargetPolicyList",
                f"{T}#SERVICE_CONTROL_POLICY",
                {"Filter": "SERVICE_CONTROL_POLICY", "PolicyIds": ids, "Truncated": False},
            )
        )
    if scp_content:
        items.append(
            obs(
                "AWS::Organizations::Policy",
                "p-custom000001",
                {"Type": "SERVICE_CONTROL_POLICY", "Content": scp_content},
            )
        )
    return OrganizationModel(items, [])


def organization_allowing_everything():
    return org_observations()


def test_scp_applicability_and_inheritance():
    allow = statement()
    assert decide(allow, org=organization_allowing_everything()).state == ALLOWED
    blocking = org_observations(statement(effect="Deny"))
    assert decide(allow, org=blocking).state == DENIED
    assert decide(allow, org=org_observations(attach_account=False)).state == POTENTIAL
    assert decide(allow, org=NO_ORG).state == POTENTIAL  # SCPs unknown, not assumed absent
    management = SimpleNamespace(**{**ROLE.__dict__, "account": M})
    model = org_observations(statement(effect="Deny"))
    assert decide(allow, org=model, role=management).state == ALLOWED  # SCPs not applicable
    linked = SimpleNamespace(**{**ROLE.__dict__, "service_linked": True})
    assert decide(allow, org=model, role=linked).state == ALLOWED
    assert decide(allow, org=org_observations(statement(effect="Deny"), types=())).state == ALLOWED
    narrow_scp = org_observations(statement(action="s3:*"))
    # Every level must allow: FullAWSAccess at the root, but both policies at the account
    # level are evaluated together, so the account level still allows via FullAWSAccess.
    assert decide(allow, org=narrow_scp).state == ALLOWED


# --- golden IAM package -----------------------------------------------------------------


@pytest.fixture(scope="module")
def golden(tmp_path_factory):
    root = tmp_path_factory.mktemp("iam")
    build(root / "input")
    with Store(root / "workspace") as store:
        install_registry(store.workspace, registry_file(root / "registry.json"))
        run = import_package(store, root / "input")
        assess(store, run)
        analysis = IamAnalysis(
            store, run, list(store.source_data(run)), workspace_registry(store.workspace)
        )
        report = root / "iam.xlsx"
        export_excel(store, run, report)
        yield store, run, analysis, report


def transition(analysis, source, target):
    return next(
        r
        for r in analysis.transitions()
        if r.principal.endswith(source) and r.target.endswith(target)
    )


def passrole(analysis, principal, target, service="ec2.amazonaws.com"):
    return next(
        r
        for r in analysis.passrole()
        if r.principal.endswith(f"/{principal}")
        and r.target.endswith(f"/{target}")
        and r.service == service
    )


def rule_result(store, run, rule, name):
    return next(
        r
        for r in store.results(run)
        if r.rule_id == rule and r.resource_id and (r.native_id or "").endswith(f"/{name}")
    )


def candidate(store, run, result):
    return next(
        (
            c
            for c in store.candidates(run)
            if c.resource_id == result.resource_id and c.rule_id == result.rule_id
        ),
        None,
    )


TRUST, INSTANCE = "R1-IAM-ROLE-TRUST-CONDITIONS", "R1-IAM-INSTANCE-ROLE-PRIVILEGE"


def test_01_wildcard_trust_with_restrictive_conditions(golden):
    store, run, analysis, _ = golden
    assert str(rule_result(store, run, TRUST, "t-wildcard-conditioned").outcome) == "PASS"
    assert transition(analysis, "role/s-caller", "t-wildcard-conditioned").state == ALLOWED
    outside = transition(analysis, "role/n-caller", "t-wildcard-conditioned")
    assert outside.state == DENIED and "grants no matching principal" in outside.blocking[0]
    assert transition(analysis, "role/u-caller", "t-wildcard-conditioned").state == POTENTIAL


def test_02_specific_external_trust_missing_source_policy(golden):
    store, run, analysis, _ = golden
    assert (
        str(rule_result(store, run, TRUST, "t-specific-external").outcome)
        == "MANUAL_VALIDATION_REQUIRED"
    )
    record = transition(analysis, f"account {X}", "t-specific-external")
    assert record.classification == "UNKNOWN_MISSING_EVIDENCE"
    assert record.missing == [f"IAM evidence for account {X}"]


def test_03_valid_source_permission_and_target_trust(golden):
    _, _, analysis, _ = golden
    record = transition(analysis, "role/s-caller", "t-admin-target")
    assert record.state == ALLOWED and record.classification == "SUPPORTED_TRANSITION"
    assert any(s.startswith("source inline:") for s in record.supporting)
    assert any(s.startswith("trust trust:") for s in record.supporting)
    edges = [r for r in analysis.relationships() if r.kind == "PRINCIPAL_CAN_ASSUME_ROLE_SUPPORTED"]
    assert (arn(S, "s-caller"), arn(T, "t-admin-target")) in {(e.source, e.target) for e in edges}
    assert all(e.basis == "EVALUATED" for e in edges)
    # Factual trust edges are kept separately from evaluated transitions.
    trusts = [r for r in analysis.relationships() if r.kind == "ROLE_TRUSTS_PRINCIPAL"]
    assert all(r.basis == "OBSERVED" for r in trusts)


def test_04_explicit_deny_overrides_allow(golden):
    _, _, analysis, _ = golden
    record = transition(analysis, "role/s-deny-caller", "t-admin-target")
    assert record.state == DENIED and record.blocking[0].endswith("s-deny-caller/assume#1")
    assert transition(analysis, "role/s-deny-caller", "t-wildcard-conditioned").state == ALLOWED
    blocks = [r for r in analysis.relationships() if r.kind == "POLICY_BLOCKS_ACTION"]
    assert any(r.source.endswith("s-deny-caller/assume#1") for r in blocks)


def test_05_missing_boundary_and_session_conditions(golden):
    _, _, analysis, _ = golden
    boundary = transition(analysis, "role/s-boundary-missing", "t-admin-target")
    assert boundary.state == POTENTIAL
    assert any("permissions boundary document" in m for m in boundary.missing)
    session = transition(analysis, "role/s-caller", "t-mfa-target")
    assert session.state == POTENTIAL
    assert any("MultiFactorAuthPresent" in m for m in session.missing)


def test_06_scp_blocks_assume_role(golden):
    _, _, analysis, _ = golden
    record = transition(analysis, "role/scp-restricted-caller", "t-admin-target")
    assert record.state == DENIED and record.blocking[0].startswith(
        "SERVICE_CONTROL_POLICY:p-denyassume01"
    )


def test_07_scp_unavailable_or_not_applicable(golden):
    _, _, analysis, _ = golden
    outside = analysis.inventory.by_arn[arn(N, "n-caller")]
    decision = analysis.identity(outside, "sts:AssumeRole", arn(T, "t-admin-target"))
    assert decision.state == ALLOWED and any("not applicable" in n for n in decision.notes)
    detached = analysis.inventory.by_arn[arn(U, "u-caller")]
    decision = analysis.identity(detached, "sts:AssumeRole", arn(T, "t-admin-target"))
    assert decision.state == POTENTIAL  # IAM conclusion kept; SCP dependency recorded
    assert any("organization hierarchy" in m for m in decision.missing)


def test_08_passrole_without_service_operation(golden):
    _, _, analysis, _ = golden
    record = passrole(analysis, "p-broad-passrole", "t-lambda-admin", "lambda.amazonaws.com")
    assert record.classification == "BLOCKED_COMBINATION"
    assert any(b.startswith("lambda:CreateFunction") for b in record.blocking)


def test_09_passrole_restricted_target(golden):
    _, _, analysis, _ = golden
    record = passrole(analysis, "p-restricted-passrole", "t-ec2-readonly")
    assert record.classification == "NOT_AN_ESCALATION"
    summary = next(
        r
        for r in analysis.passrole()
        if r.principal.endswith("p-restricted-passrole")
        and "other service-trusted roles" in r.target
    )
    assert arn(T, "t-ec2-admin") in summary.reasoning


def test_10_incompatible_service_trust(golden):
    _, _, analysis, _ = golden
    record = passrole(analysis, "p-ecs-mismatch", "t-ecs-only")
    assert record.classification == "BLOCKED_COMBINATION"
    assert "does not admit ec2.amazonaws.com" in record.blocking[0]


def test_11_supported_escalation_prerequisites_not_confirmed(golden):
    store, run, analysis, _ = golden
    record = passrole(analysis, "p-broad-passrole", "t-ec2-admin")
    assert record.classification == "SUPPORTED_ESCALATION_PREREQUISITES"
    assert any("Attacker control" in p for p in record.preconditions)
    result = rule_result(store, run, INSTANCE, "p-broad-passrole")
    found = candidate(store, run, result)
    assert (found.finding_class, found.suggested_severity) == ("SUPPORTED_SECURITY_GAP", "HIGH")
    assert found.review_status == "CANDIDATE" and found.methodology_id == "AWS-19"
    assert found.methodology_criteria == ("AWS-19:P4", "AWS-19:F3")
    assert "not evaluated" in found.exposure_status


def test_12_notaction_and_notresource(golden):
    _, _, analysis, _ = golden
    assert not [r for r in analysis.passrole() if r.principal.endswith("/p-notaction")]
    role = analysis.inventory.by_arn[arn(T, "p-notaction")]
    assert analysis.identity(role, "iam:PassRole", arn(T, "t-ec2-admin")).state == DENIED
    assert analysis.identity(role, "s3:GetObject", "arn:aws:s3:::b/k").state == ALLOWED
    summary = next(
        r
        for r in analysis.passrole()
        if r.principal.endswith("/p-notresource") and "other service-trusted roles" in r.target
    )
    assert arn(T, "t-ec2-admin") in summary.reasoning


def test_13_unsupported_condition_operator(golden):
    _, _, analysis, _ = golden
    record = passrole(analysis, "p-unsupported-condition", "t-ec2-admin")
    assert record.classification == "POTENTIAL_ESCALATION" and record.state != ALLOWED
    assert "condition operator DateLessThan not supported" in record.missing


def test_14_missing_managed_policy_document(golden):
    store, run, analysis, _ = golden
    role = analysis.inventory.by_arn[arn(T, "p-potential")]
    assert analysis.privilege(role)[0] == "UNRESOLVED"
    record = passrole(analysis, "p-potential", "t-ec2-admin")
    assert record.classification == "UNKNOWN_MISSING_EVIDENCE"
    found = candidate(store, run, rule_result(store, run, INSTANCE, "p-potential"))
    assert (found.finding_class, found.suggested_severity) == ("POTENTIAL_EXPOSURE", "MEDIUM")
    assert any("not-exported" in m for m in found.missing_evidence)


def test_15_duplicate_and_incomplete_collections(golden):
    store, run, analysis, _ = golden
    names = [r.name for r in analysis.inventory.roles.values() if r.account == S]
    assert len(names) == len(set(names)) == 7  # two identical ListRoles exports, one role each
    truncated = analysis.inventory.roles[(S, "s-truncated")]
    assert truncated.attached is None
    assert any(
        s["operation"] == "iam:ListAttachedRolePolicies" and s["effective_coverage"] == "PARTIAL"
        for s in store.source_data(run)
    )
    conflict = analysis.inventory.roles[(S, "s-conflict")]
    assert conflict.inline["assume"] is None
    assert transition(analysis, "role/s-conflict", "t-admin-target").state == UNKNOWN_DECISION
    assert analysis.inventory.roles[(S, "s-boundary-missing")].boundary_arn.endswith(
        "missing-boundary"
    )


def test_16_reimport_preserves_analyst_decisions(tmp_path):
    build(tmp_path / "input")
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, tmp_path / "input")
        assess(store, run)
        target = next(c for c in store.candidates(run) if c.native_id.endswith("/p-broad-passrole"))
        record_decision(store, run, target.finding_id, "CONFIRMED", "analyst", "validated")
        before = [c.model_dump_json() for c in store.candidates(run)]
        assert import_package(store, tmp_path / "input") == run
        assess(store, run)
        assert [c.model_dump_json() for c in store.candidates(run)] == before
        assert validation_status(store, target)[0] == "CONFIRMED"
        others = [c for c in store.candidates(run) if c.finding_id != target.finding_id]
        assert all(validation_status(store, c)[0] == "CANDIDATE" for c in others)


def test_17_end_to_end_excel_reconciliation(golden):
    store, run, analysis, report = golden
    wb = load_workbook(report)
    rows = lambda name: list(wb[name].iter_rows(min_row=4, values_only=True))  # noqa: E731
    investigations = analysis.transitions() + analysis.passrole()
    assert len(rows("IAM Investigations")) == len(investigations)
    assert sorted(r[5] for r in rows("IAM Investigations")) == sorted(
        i.classification for i in investigations
    )
    assert len(rows("IAM Relationships")) == len(analysis.relationships())
    assert len(rows("Effective Permissions")) == len(analysis.inventory.roles)
    assert len(rows("IAM Evidence Gaps")) == len(analysis.evidence_gaps())
    findings = rows("Findings Register")
    assert len(findings) == len(list(store.candidates(run)))
    paths = [r for r in rows("Attack Path Prerequisites") if r[1] == "AP-06"]
    assert paths and {r[5] for r in paths} == {"UNKNOWN"}
    assert all(cell.data_type != "f" for ws in wb for row in ws for cell in row)
    wb.close()


def test_trust_candidate_uses_transitions_and_never_critical_without_support(golden):
    store, run, _, _ = golden
    found = candidate(store, run, rule_result(store, run, TRUST, "t-admin-target"))
    assert found.finding_class == "BASELINE_DEVIATION"  # registered internal account
    assert found.suggested_severity == "MEDIUM"  # admin-equivalent target
    assert "SUPPORTED_TRANSITION=1" in found.exposure_status
    assert all(c.suggested_severity != "CRITICAL" for c in store.candidates(run))


def test_generated_r12_fixture_preserves_historical_bytes(tmp_path):
    from fixture_fingerprints import package_fingerprint

    build(tmp_path / "golden")
    assert len(list((tmp_path / "golden").iterdir())) == 49
    assert package_fingerprint(tmp_path / "golden") == (
        "f252938360fa99589ad5f6b06f438f10386ed4ddad245f3c3500f31c9bc2dfc3"
    )


def test_no_aws_or_network_calls_in_iam_package():
    from pathlib import Path

    root = Path(__file__).parents[1] / "src" / "offline_analyzer"
    for path in list((root / "iam").glob("*.py")) + [root / "normalization" / "iam_adapters.py"]:
        text = path.read_text()
        for forbidden in (
            "boto3",
            "botocore",
            "socket",
            "urllib.request",
            "http.client",
            "requests",
        ):
            assert forbidden not in text, (path, forbidden)


def test_url_encoded_role_documents_parse(golden):
    _, _, analysis, _ = golden
    caller = analysis.inventory.roles[(S, "s-caller")]
    assert caller.trust is not None and caller.inline["assume"] is not None
    assert json.dumps(caller.trust_raw)  # decoded from URL-encoded API form


# --- R1.1 classification corrections ------------------------------------------------------


def test_network_and_storage_exposure_never_claim_supported_gap(tmp_path):
    from tools.fixtures.r11_demo import build as build_r11

    build_r11(tmp_path / "input")
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, tmp_path / "input")
        assess(store, run)
        classes = {
            (c.rule_id, c.finding_class)
            for c in store.candidates(run)
            if c.rule_id
            in {"R1-EC2-SG-SENSITIVE-INGRESS", "R1-RDS-PUBLIC-ACCESS", "R1-S3-BUCKET-PUBLIC-ACCESS"}
        }
        assert classes and all(cls != "SUPPORTED_SECURITY_GAP" for _, cls in classes)
        routed = next(c for c in store.candidates(run) if c.native_id == "sg-00000000000000015")
        assert routed.finding_class == "POTENTIAL_EXPOSURE"
        assert "not demonstrated" in routed.exposure_status
        bucket = next(c for c in store.candidates(run) if c.native_id == "golden-public-policy")
        assert "data sensitivity not established" in bucket.exposure_status


def test_outcome_class_severity_confidence_and_status_are_independent(tmp_path):
    from offline_analyzer.contracts.r0 import Package, Source
    from offline_analyzer.ingestion.safety import sha256_file
    from tools.fixtures.r1_demo import database

    native = {
        "DBInstances": [
            database("111111111111", f"db{i}", "postgres", False, False) for i in (1, 2)
        ],
        "ResponseMetadata": {"HTTPStatusCode": 200},
    }
    (tmp_path / "rds.json").write_text(json.dumps(native))
    scanner = [
        {
            "class_uid": 2004,
            "metadata": {
                "product": {"name": "Prowler", "version": "5.10.0"},
                "event_code": "rds_instance_storage_encrypted",
            },
            "status_code": "PASS",
            "status_id": 1,
            "status": "New",
            "severity": "Low",
            "time_dt": "2026-10-10T07:00:00+00:00",
            "finding_info": {"uid": "x"},
            "cloud": {"provider": "aws", "account": {"uid": "111111111111"}, "region": "us-east-1"},
            "resources": [
                {
                    "uid": "arn:aws:rds:us-east-1:111111111111:db:db1",
                    "type": "AwsRdsDbInstance",
                    "region": "us-east-1",
                }
            ],
        }
    ]
    (tmp_path / "prowler.json").write_text(json.dumps(scanner))
    sources = (
        Source(
            path="rds.json",
            sha256=sha256_file(tmp_path / "rds.json"),
            format="aws-json-v1",
            operation="rds:DescribeDBInstances",
            account_id="111111111111",
            region="us-east-1",
            collected_at="2026-10-10T07:00:00+00:00",
            coverage="COMPLETE",
            request_parameters={},
        ),
        Source(
            path="prowler.json",
            sha256=sha256_file(tmp_path / "prowler.json"),
            format="prowler-ocsf-5.10",
            operation="prowler:Scan",
            account_id="111111111111",
            region="us-east-1",
            collected_at="2026-10-10T07:00:00+00:00",
        ),
    )
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, tmp_path, Package(sources=sources))
        assess(store, run)
        found = {
            c.native_id.rsplit(":", 1)[-1]: c
            for c in store.candidates(run)
            if c.rule_id == "R0-RDS-ENCRYPTION"
        }
        contradicted, clean = found["db1"], found["db2"]
        # Contradiction lowers evidence confidence only; outcome, class and severity hold.
        assert (contradicted.evidence_confidence, clean.evidence_confidence) == ("MEDIUM", "HIGH")
        assert contradicted.suggested_severity == clean.suggested_severity == "MEDIUM"
        assert contradicted.finding_class == clean.finding_class == "BASELINE_DEVIATION"
        record_decision(store, run, clean.finding_id, "CONFIRMED", "analyst", "ok")
        assess(store, run)
        after = next(c for c in store.candidates(run) if c.finding_id == clean.finding_id)
        # Analyst status is stored separately and never rewrites the candidate.
        assert after.model_dump() == clean.model_dump()
        assert validation_status(store, after)[0] == "CONFIRMED"
