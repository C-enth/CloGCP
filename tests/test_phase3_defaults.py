"""Safe report defaults, exclusion-aware counts, latest-import snapshots, provider-neutral refresh
and staleness, path-review carry-over, custom limitations, size-limit messages and the additive
migration. Synthetic data only; AWS and GCP both exercised."""

import json
import zipfile

import pytest
from test_false_positive import variant
from test_theme_browser import workspace as estate_workspace

from offline_analyzer.attack_paths.feasibility import attack_paths
from offline_analyzer.attack_paths.review import record as review_path
from offline_analyzer.contracts.exclusion import ExclusionRequest
from offline_analyzer.enterprise.snapshot import assess_enterprise, members, superseded_imports
from offline_analyzer.findings.exclusions import excluded_ids, record
from offline_analyzer.ingestion.importer import import_package
from offline_analyzer.ingestion.json_input import InputError
from offline_analyzer.inventory.environment import record_run_origin, set_environment
from offline_analyzer.reporting.assessment_report import model
from offline_analyzer.reporting.cloud import dashboard
from offline_analyzer.rules.engine import assess
from offline_analyzer.storage.store import Store

ACCEPTED = ExclusionRequest(action="EXCLUDE", reason="ACCEPTED_RISK", analyst="analyst")


@pytest.fixture(scope="module")
def estate(tmp_path_factory):
    root = tmp_path_factory.mktemp("phase3")
    work = estate_workspace(root)
    with Store(work) as store:
        runs = {r: s for r, s in store.rows("SELECT run_id, package FROM runs")}
        snapshot = next(r for (r,) in store.rows("SELECT run_id FROM enterprise_snapshots"))
        yield store, sorted(set(runs) - {snapshot}), snapshot, root


def provider_of(store, run):
    return {"GCP" if s["format"] == "gcp-json-v1" else "AWS" for s in store.source_data(run)}


def test_report_scope_defaults_to_confirmed_with_a_cover_notice(estate):
    from offline_analyzer.reporting.report_document import summary

    store, _members, snapshot, _ = estate
    default = model(store, snapshot)
    assert default["include_key"] == "confirmed"
    from offline_analyzer.reporting.assessment_report import VALIDATION

    assert all(f["status"] == VALIDATION["CONFIRMED"] for f in default["findings"])
    assert not any(k == "Report scope notice" for k, _ in summary(default)["control"])
    wide = model(store, snapshot, include="valid")
    assert any(k == "Report scope notice" for k, _ in summary(wide)["control"])


def test_excluded_findings_leave_every_count(estate):
    from offline_analyzer.api.services import Views
    from offline_analyzer.findings.quality import run_quality
    from offline_analyzer.reporting.metrics import assessment_metrics
    from offline_analyzer.reporting.tracker import tracker_rows

    store, runs, _snapshot, _ = estate
    aws = next(r for r in runs if provider_of(store, r) == {"AWS"})
    views = Views(store.workspace)
    finding = sorted(store.candidates(aws), key=lambda f: f.finding_id)[0]
    before = {
        "metrics": assessment_metrics(store, aws)["candidate_findings"],
        "quality": sum(v for k, v in run_quality(store, aws).items() if k != "items"),
        "runs": next(r for r in views.runs(store) if r["run_id"] == aws)["finding_count"],
        "tracker": sum(len(str(r)) for r in tracker_rows(store, aws, {}, None)),
    }
    record(store, aws, finding.finding_id, ACCEPTED)
    try:
        assert finding.finding_id in excluded_ids(store, aws)
        assert assessment_metrics(store, aws)["candidate_findings"] == before["metrics"] - 1
        assert sum(v for k, v in run_quality(store, aws).items() if k != "items") == (
            before["quality"] - 1
        )
        assert next(r for r in views.runs(store) if r["run_id"] == aws)["finding_count"] == (
            before["runs"] - 1
        )
        evaluations = views.records(store, aws, "evaluations", "", "", "", 0, 100, {})
        marked = [i for i in evaluations["items"] if i.get("excluded")]
        assert marked and marked[0]["excluded"].startswith("excluded (Risk accepted")
        # Paths depending on an excluded finding are flagged and not counted on the dashboard.
        flagged = [p for p in attack_paths(store, aws)["items"] if p.get("excluded_dependency")]
        assert all(finding.finding_id in p["excluded_dependency"]["finding_ids"] for p in flagged)
        unflagged = [
            p for p in attack_paths(store, aws)["items"] if not p.get("excluded_dependency")
        ]
        for row in dashboard(store, aws, "AWS")["stride"]:
            assert row["paths"] == sum(row["category"] in p.get("stride", []) for p in unflagged)
    finally:
        record(
            store,
            aws,
            finding.finding_id,
            ExclusionRequest(action="RESTORE", analyst="analyst"),
        )


def test_enterprise_issues_with_only_excluded_findings_are_not_counted(estate):
    from offline_analyzer.enterprise.queries import EnterpriseViews

    store, _runs, snapshot, _ = estate
    views = EnterpriseViews()
    # A synthetic systemic issue linked to one real finding of the snapshot.
    finding = sorted(store.candidates(snapshot), key=lambda f: f.finding_id)[0]
    store.db.execute(
        "INSERT INTO enterprise_issues SELECT run_id, 'synthetic-issue', "
        "json_merge_patch(data, ?) FROM enterprise_issues WHERE run_id=? LIMIT 1",
        [
            json.dumps(
                {"investigation_id": "synthetic-issue", "finding_ids": [finding.finding_id]}
            ),
            snapshot,
        ],
    )
    listed = views.issues(store, snapshot, multi=False)
    issue = next(i for i in listed["items"] if i["investigation_id"] == "synthetic-issue")
    for fid in issue["finding_ids"]:
        record(store, snapshot, fid, ACCEPTED)
    try:
        after = views.issues(store, snapshot, multi=False)
        assert after["total"] == listed["total"] - 1
        assert issue["investigation_id"] not in {i["investigation_id"] for i in after["items"]}
    finally:
        for fid in issue["finding_ids"]:
            record(store, snapshot, fid, ExclusionRequest(action="RESTORE", analyst="analyst"))


@pytest.mark.parametrize(("kind", "sheets"), [("aws", 28), ("gcp", 43), ("mixed", 43)])
def test_excel_sheet_count_contract(estate, kind, sheets, tmp_path):
    from openpyxl import load_workbook

    from offline_analyzer.reporting.excel import REPORT_SHEETS, export_excel

    store, runs, snapshot, _ = estate
    run = {
        "aws": next(r for r in runs if provider_of(store, r) == {"AWS"}),
        "gcp": next(r for r in runs if provider_of(store, r) == {"GCP"}),
        "mixed": snapshot,
    }[kind]
    path = tmp_path / f"{kind}.xlsx"
    export_excel(store, run, path)
    names = load_workbook(path, read_only=True).sheetnames
    assert len(names) == sheets and REPORT_SHEETS == 28
    assert "Excluded Findings" in names


def test_snapshot_defaults_to_the_latest_import_per_project_and_gcp_refresh(tmp_path):
    from offline_analyzer.api.services import Views
    from offline_analyzer.enterprise.queries import scope_collections
    from offline_analyzer.enterprise.refresh import refresh_plan

    set_environment(tmp_path / "work", "SYNTHETIC")
    with Store(tmp_path / "work") as store:
        older = import_package(store, variant(tmp_path / "a1", "demo-alpha", "123456789012"))
        beta = import_package(store, variant(tmp_path / "b1", "demo-beta", "222222222222"))
        for run in (older, beta):
            record_run_origin(tmp_path / "work", run, "synthetic")
            assess(store, run)
        baseline = assess_enterprise(store)
        newer_package = variant(tmp_path / "a2", "demo-alpha", "123456789012", extra=True)
        manifest = json.loads((newer_package / "manifest.json").read_text())
        for dataset in manifest["datasets"]:
            dataset["completed_at"] = "2026-10-13T00:00:00Z"  # collected one day later
        (newer_package / "manifest.json").write_text(json.dumps(manifest))
        newer = import_package(store, newer_package)
        record_run_origin(tmp_path / "work", newer, "synthetic")
        assess(store, newer)
        assert superseded_imports(store) == {older: [newer]}
        assert set(members(store)) == {newer, beta}
        listed = {r["run_id"]: r for r in Views(store.workspace).runs(store)}
        assert listed[older]["superseded_by"] == [newer] and not listed[newer]["superseded_by"]
        # Provider-neutral refresh: GCP projects are scope keys, not refused.
        plan = refresh_plan(store, baseline, [newer])
        assert plan["affected_accounts"] == ["GCP:demo-alpha"]
        assert set(plan["members"]) == {newer, beta}
        keys = {c["account"] for c in scope_collections(store, baseline)}
        assert keys == {"demo-alpha", "demo-beta"}
        # An explicit selection may still include the superseded import.
        explicit = assess_enterprise(store, [older, newer, beta])
        assert explicit != assess_enterprise(store)


def test_path_reviews_apply_where_the_fingerprint_is_identical(estate):
    store, runs, snapshot, _ = estate
    gcp = next(r for r in runs if provider_of(store, r) == {"GCP"})
    member_paths = {p["path_id"]: p for p in attack_paths(store, gcp)["items"]}
    snapshot_paths = {p["path_id"]: p for p in attack_paths(store, snapshot)["items"]}
    shared = sorted(set(member_paths) & set(snapshot_paths))
    if not shared:
        pytest.skip("no path shared between a member and the snapshot")
    path_id = shared[0]
    review_path(store, gcp, path_id, "CONFIRMED", "analyst")
    from offline_analyzer.attack_paths.review import fingerprint

    same = fingerprint(member_paths[path_id]) == fingerprint(snapshot_paths[path_id])
    carried = {p["path_id"]: p for p in attack_paths(store, snapshot)["items"]}[path_id]
    assert carried["analyst_review"] == ("CONFIRMED" if same else "NEEDS_REVALIDATION")
    # A review recorded for a different fingerprint never applies.
    store.db.execute(
        "UPDATE path_reviews SET data=json_merge_patch(data, ?) WHERE path_id=?",
        [json.dumps({"fingerprint": "f" * 64}), path_id],
    )
    stale = {p["path_id"]: p for p in attack_paths(store, gcp)["items"]}[path_id]
    assert stale["analyst_review"] == "NEEDS_REVALIDATION"


def test_custom_limitations_render_after_the_fixed_constants(estate, tmp_path):
    from openpyxl import load_workbook

    from offline_analyzer.reporting.docx import write_docx
    from offline_analyzer.reporting.limitations import latest, save
    from offline_analyzer.reporting.report_document import LIMITATIONS, write_excel

    store, _runs, snapshot, _ = estate
    save(store, snapshot, "Synthetic scope: two projects only.\nNo data-plane testing.", "a")
    assert latest(store, snapshot).startswith("Synthetic scope")
    data = model(store, snapshot)
    word = tmp_path / "r.docx"
    write_docx(data, word)
    xml = zipfile.ZipFile(word).read("word/document.xml").decode()
    assert xml.index(LIMITATIONS[-1]) < xml.index("Assessor scope limitation: Synthetic scope")
    assert "Assessor scope limitation: No data-plane testing." in xml
    book = tmp_path / "r.xlsx"
    write_excel(data, book)
    cells = [
        str(c)
        for row in load_workbook(book)["Coverage and Limitations"].iter_rows(values_only=True)
        for c in row
        if c
    ]
    assert cells.index(LIMITATIONS[-1]) < cells.index(
        "Assessor scope limitation: Synthetic scope: two projects only."
    )
    with pytest.raises(InputError, match="INVALID_REPORT_LIMITATIONS"):
        save(store, snapshot, "x" * 4001, "a")


def test_size_limits_fail_closed_with_an_explanation():
    from offline_analyzer.ingestion.errors import explanation

    for code in (
        "GCP_ASSESSMENT_RECORD_LIMIT",
        "ENTERPRISE_OBSERVATION_LIMIT",
        "INVALID_ENTERPRISE_MEMBERS",
    ):
        text = explanation(code)
        assert "truncated" in text["message"] or "No snapshot" in text["message"]
        assert text["action"]


def test_additive_migration_takes_a_backup_and_keeps_every_status(tmp_path):
    import duckdb

    from offline_analyzer.storage.location import inspect_workspace, prepare_open
    from tools.upgrade_state import build, read

    build(tmp_path)
    before = json.loads((tmp_path / "before.json").read_text())
    work = tmp_path / "work"
    with duckdb.connect(str(work / "assessment.duckdb")) as db:  # an older workspace
        db.execute("DROP TABLE report_limitations")
    assert inspect_workspace(work)["schema"] == "OLDER"
    opened = prepare_open(work)
    assert opened["migrated"] and opened["integrity"] == "OK" and opened["backup"]["files"]
    backups = list((tmp_path / "work-backups").glob("pre-migration-*/assessment.duckdb"))
    assert len(backups) == 1
    assert read(work) == before
