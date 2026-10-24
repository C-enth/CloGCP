"""Synthetic fixtures only."""

import json
from datetime import UTC, datetime
from zipfile import ZipFile

import pytest
from openpyxl import Workbook
from pydantic import ValidationError

from offline_analyzer.contracts.models import Artifact, Evaluation, Manifest, PathEvaluation
from offline_analyzer.ingestion.safety import excel_literal, resolve_under, sha256_file
from offline_analyzer.methodology.extract import extract, preflight


@pytest.mark.parametrize("path", ["../escape", "/tmp/a", "C:/file", "a\\b", "."])
def test_reject_paths(path):
    with pytest.raises(ValueError):
        Artifact(artifact_id="a", relative_path=path, sha256="0" * 64, size_bytes=0)


def test_symlink_escape(tmp_path):
    (tmp_path / "link").symlink_to(tmp_path.parent, target_is_directory=True)
    with pytest.raises(ValueError):
        resolve_under(tmp_path, "link/outside")


def evaluation(**overrides):
    values = dict(
        evaluation_id="eval",
        run_id="run",
        control_id="SYN-1",
        rule_version="1",
        methodology_sha256="0" * 64,
        accounts=("123456789012",),
        resource_ids=(),
        outcome="UNKNOWN",
        evidence=(),
        reasoning="missing collection",
        missing_dependencies=("config",),
        confidence="UNASSESSED",
        limitations=(),
        assessed_at=datetime.now(UTC),
    )
    values.update(overrides)
    return Evaluation(**values)


def test_unknown_is_not_pass():
    assert evaluation().outcome == "UNKNOWN"
    with pytest.raises(ValidationError):
        evaluation(outcome="PASS")
    with pytest.raises(ValidationError):
        evaluation(outcome="PASS", evidence=({"artifact_id": "a", "locator": "/1"},))
    assert (
        evaluation(
            outcome="PASS",
            missing_dependencies=(),
            evidence=({"artifact_id": "a", "locator": "/1"},),
        ).outcome
        == "PASS"
    )


def test_timezone_and_unknown_fields():
    with pytest.raises(ValidationError):
        evaluation(assessed_at=datetime(2026, 1, 1))
    with pytest.raises(ValidationError):
        evaluation(invented=True)


def test_supported_path_needs_edges():
    with pytest.raises(ValidationError):
        PathEvaluation(
            path_definition_id="p",
            definition_version="1",
            run_id="r",
            state="SUPPORTED",
            relationship_ids=(),
            evidence=(),
            prerequisites=(),
            blockers=(),
            missing_dependencies=(),
            reasoning="no proof",
            limitations=(),
            assessed_at=datetime.now(UTC),
        )


def test_manifest_window_and_duplicates():
    data = dict(
        run_id="r",
        source_id="s",
        collector="synthetic",
        collector_version="1",
        collection_started=datetime(2026, 2, 1, tzinfo=UTC),
        collection_finished=datetime(2026, 1, 1, tzinfo=UTC),
        artifacts=(),
        coverage=(),
    )
    with pytest.raises(ValidationError):
        Manifest(**data)
    data["collection_finished"] = data["collection_started"]
    a = Artifact(artifact_id="a", relative_path="a.json", sha256="0" * 64, size_bytes=0)
    data["artifacts"] = (a, a)
    with pytest.raises(ValidationError):
        Manifest(**data)


def test_lossless_extract(tmp_path):
    source = tmp_path / "synthetic.xlsx"
    workbook = Workbook()
    workbook.active.title = "Controls"
    workbook.active.append(["Control ID", "Requirement"])
    workbook.active.append(["SYN-1", "Synthetic only"])
    workbook.active["C3"] = "=1+1"
    workbook.create_sheet("Empty")
    workbook.save(source)
    before = sha256_file(source)
    output = tmp_path / "private"
    extract(source, output)
    assert sha256_file(source) == before
    inventory = json.loads((output / "inventory.json").read_text())
    assert len(inventory["worksheets"]) == 2
    assert inventory["semantic_counts"] is None
    cells = [json.loads(line) for line in (output / "cells.jsonl").read_text().splitlines()]
    assert any(c["value"] == "=1+1" and c["type"] == "f" for c in cells)
    assert len(cells) == 5
    with pytest.raises(FileExistsError):
        extract(source, output)


@pytest.mark.parametrize("member", ["../escape", "xl/vbaProject.bin", "xl/externalLinks/a.xml"])
def test_hostile_xlsx(tmp_path, member):
    path = tmp_path / "bad.xlsx"
    with ZipFile(path, "w") as archive:
        archive.writestr(member, "test")
    with pytest.raises(ValueError):
        preflight(path)


def test_bounds(tmp_path, monkeypatch):
    from offline_analyzer.methodology import extract as module

    source = tmp_path / "synthetic.xlsx"
    workbook = Workbook()
    workbook.active.append([1, 2, 3])
    workbook.save(source)
    monkeypatch.setattr(module, "MAX_CELLS", 1)
    output = tmp_path / "result"
    with pytest.raises(ValueError):
        extract(source, output)
    assert not output.exists()
    with pytest.raises(ValueError):
        sha256_file(source, max_bytes=1)


@pytest.mark.parametrize("value", ["=cmd", "+cmd", "-cmd", "@cmd", " \t=cmd"])
def test_formula_injection(value):
    assert excel_literal(value).startswith("'")
    assert excel_literal("ordinary") == "ordinary"


def test_artifact_integrity(tmp_path):
    from offline_analyzer.ingestion.safety import verify_artifact

    source = tmp_path / "a.json"
    source.write_text("{}")
    artifact = Artifact(
        artifact_id="a", relative_path="a.json", sha256=sha256_file(source), size_bytes=2
    )
    assert verify_artifact(tmp_path, artifact) == source
    source.write_text("[]")
    with pytest.raises(ValueError, match="checksum"):
        verify_artifact(tmp_path, artifact)
    source.write_text("changed")
    with pytest.raises(ValueError, match="size"):
        verify_artifact(tmp_path, artifact)


def test_coverage_template_preserves_unmapped_rows(tmp_path):
    source = tmp_path / "synthetic.xlsx"
    workbook = Workbook()
    workbook.active.append(["ID", "Requirement"])
    workbook.active.append(["SYN-1", "first; second"])
    workbook.active["A4"] = "Unclassified note"
    workbook.save(source)
    output = tmp_path / "private"
    extract(source, output)
    rows = [
        json.loads(line) for line in (output / "coverage-template.jsonl").read_text().splitlines()
    ]
    assert [row["row"] for row in rows] == [1, 2, 4]
    assert all(row["mapping_status"] == "UNMAPPED" for row in rows)
    assert all(row["feasibility"] == "UNREVIEWED" for row in rows)


def test_misleading_dimensions_do_not_hide_records(tmp_path):
    from io import BytesIO

    source = tmp_path / "synthetic.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "Header"
    workbook.active["A7"] = "SYN-7"
    workbook.save(source)
    original = source.read_bytes()
    with ZipFile(BytesIO(original)) as incoming, ZipFile(source, "w") as outgoing:
        for member in incoming.infolist():
            data = incoming.read(member.filename)
            if member.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b'ref="A1:A7"', b'ref="A1:A1"')
            outgoing.writestr(member, data)
    output = tmp_path / "private"
    extract(source, output)
    cells = [json.loads(line) for line in (output / "cells.jsonl").read_text().splitlines()]
    assert any(cell["coordinate"] == "A7" for cell in cells)
