"""Compiled GUI + real loopback APIs: AWS/GCP import, approval and consolidation."""

import io
import os
import socket
import threading
import time
import zipfile
from pathlib import Path

import pytest

from tests.browser_diagnostics import word_text

pytestmark = pytest.mark.skipif(os.environ.get("RUN_GUI_BROWSER") != "1", reason="Opt-in Chromium")


def test_multicloud_operational_workflow(tmp_path):
    import uvicorn
    from openpyxl import load_workbook
    from playwright.sync_api import expect, sync_playwright

    from offline_analyzer.api.app import create_app
    from offline_analyzer.reporting.report_document import EXCEL_SHEETS
    from tools.fixtures.r12_demo import build as aws_build
    from tools.gcp_fixture import build

    aws_build(tmp_path / "aws")
    package = build(tmp_path / "gcp")
    archive = tmp_path / "gcp.zip"
    with zipfile.ZipFile(archive, "w") as out:
        for file in package.iterdir():
            out.write(file, file.name)
    workspace = tmp_path / "work"
    listener = socket.socket()
    listener.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
    listener.bind(("127.0.0.1", 0))
    port = listener.getsockname()[1]
    origin = f"http://127.0.0.1:{port}"
    listener.listen(128)
    app = create_app(workspace, origin=origin, bootstrap="multicloud-test")
    server = uvicorn.Server(uvicorn.Config(app, log_level="critical", access_log=False))
    thread = threading.Thread(target=lambda: server.run(sockets=[listener]), daemon=True)
    thread.start()
    for _ in range(100):
        if server.started:
            break
        time.sleep(0.05)
    assert server.started
    artifacts = Path("artifacts/gcp-browser")
    artifacts.mkdir(parents=True, exist_ok=True)
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True)
            context = browser.new_context(
                accept_downloads=True, viewport={"width": 1600, "height": 1050}
            )
            context.grant_permissions(["clipboard-read", "clipboard-write"], origin=origin)
            external, errors = [], []

            def route(request):
                if request.request.url.startswith(origin):
                    request.continue_()
                else:
                    external.append(request.request.url)
                    request.abort()

            context.route("**/*", route)
            page = context.new_page()
            page.on("pageerror", lambda error: errors.append(str(error)))
            page.goto(origin + "/#token=multicloud-test")

            def import_evidence(cloud, file, transport):
                page.get_by_role("button", name="Evidence import", exact=True).click()
                with page.expect_file_chooser() as chooser:
                    page.get_by_role(
                        "button", name=f"Import {cloud} Evidence {transport}", exact=True
                    ).click()
                chooser.value.set_files(str(file))
                page.get_by_role("button", name="Validate evidence", exact=True).click()
                page.get_by_role("button", name="Import validated evidence", exact=True).wait_for()
                page.get_by_role("button", name="Import validated evidence", exact=True).click()
                expect(
                    page.get_by_role("button", name="Run assessment", exact=True)
                ).to_be_enabled()
                page.get_by_role("button", name="Run assessment", exact=True).click()
                expect(
                    page.get_by_label("Active evidence run").locator("option:checked")
                ).to_contain_text("Assessed")

            import_evidence("AWS", tmp_path / "aws", "Folder")
            page.get_by_role("tab", name="AWS", exact=True).click()
            expect(page.get_by_label("AWS scope summary")).to_contain_text("Candidate findings8")
            page.screenshot(path=str(artifacts / "aws.png"), full_page=True)
            page.get_by_role("button", name="HIGH findings:", exact=False).click()
            expect(page.get_by_label("Cloud severity", exact=True)).to_have_value("HIGH")
            expect(page.locator(".cloud-records tbody tr").first).to_contain_text("HIGH")
            import_evidence("GCP", package, "Folder")
            page.get_by_role("tab", name="GCP", exact=True).click()
            expect(page.get_by_label("GCP scope summary")).to_contain_text("Scopes expected2")
            expect(page.get_by_label("Security domain comparison")).to_contain_text(
                "Credentials and API Security"
            )
            page.screenshot(path=str(artifacts / "gcp.png"), full_page=True)
            page.get_by_role("tab", name="Security Findings", exact=True).click()
            page.get_by_label("Cloud Search").select_option(
                "API key lacks application and API restrictions"
            )
            table = page.locator(".cloud-records tbody")
            expect(table.locator("tr")).to_have_count(1)
            table.get_by_role("button", name="View").click()
            expect(page.get_by_role("button", name="Confirm", exact=True)).to_have_count(0)
            page.get_by_role("tab", name="Client Reporting", exact=True).click()
            page.get_by_label("Analyst identity").fill("synthetic-gcp-reviewer")
            page.get_by_label("Analyst remarks").fill("Synthetic native metadata reviewed")
            page.get_by_role("tab", name="Client Reporting", exact=True).click()
            page.get_by_role("button", name="Confirm", exact=True).click()
            expect(table).to_contain_text("CONFIRMED")
            table.get_by_role("button", name="View").click()
            page.get_by_role("tab", name="Client Reporting", exact=True).click()
            editor = page.get_by_label("Report finding editor")
            editor.get_by_label("Report reviewer").fill("synthetic-gcp-reviewer")
            editor.get_by_label("Include in client report").check()
            editor.get_by_role("button", name="Save report draft", exact=True).click()
            expect(editor.get_by_role("status")).to_contain_text("Report draft saved")
            editor.get_by_role("button", name="Approve recommendation", exact=True).click()
            expect(editor.get_by_role("status")).to_contain_text("recommendation approval recorded")
            editor.get_by_role("button", name="Approve report narrative", exact=True).click()
            expect(editor.get_by_role("button", name="Copy Report Finding")).to_be_enabled()
            editor.get_by_role("button", name="Copy Report Finding").click()
            expect(editor.get_by_role("status")).to_contain_text("copied")
            assert "demo-alpha" in page.evaluate("navigator.clipboard.readText()")
            page.screenshot(path=str(artifacts / "approved-gcp.png"), full_page=True)
            page.get_by_role("button", name="Close details", exact=True).click()
            page.get_by_role("button", name="Build enterprise assessment", exact=True).click()
            expect(
                page.get_by_label("Active evidence run").locator("option:checked")
            ).to_contain_text("Enterprise")
            page.get_by_role("tab", name="Consolidated", exact=True).click()
            expect(page.get_by_label("AWS scope summary")).to_contain_text("Candidate findings8")
            expect(page.get_by_label("GCP scope summary")).to_contain_text("Approved findings1")
            page.screenshot(path=str(artifacts / "consolidated.png"), full_page=True)
            matrix = page.get_by_label("Security domain comparison")
            matrix.locator("tr").filter(has_text="Credentials and API Security").get_by_role(
                "button", name="Evaluations", exact=True
            ).click()
            expect(page.locator(".cloud-records tbody")).to_contain_text("GCP-API-BOTH")
            page.get_by_role("tab", name="Attack Paths", exact=True).click()
            expect(
                page.get_by_role("heading", name="Attack-path investigations", exact=True)
            ).to_be_visible()
            expect(page.locator(".cloud-dashboard")).to_contain_text("INSUFFICIENT_EVIDENCE")
            diagram = page.get_by_label("GCP attack paths", exact=True).locator("figure").first
            expect(diagram).to_be_visible()
            expect(diagram).to_contain_text("STRIDE:")
            diagram.get_by_role("button", name="Confirm path interpretation", exact=True).click()
            expect(diagram).to_contain_text("Analyst path interpretation: CONFIRMED")
            # Hops: direction, status, STRIDE, justification, table fallback, click-through.
            hops = diagram.get_by_label("Attack path hops")
            expect(hops.locator(".hop-edge").first).to_be_visible()
            expect(hops).to_contain_text("→")
            expect(hops.locator(".badge.stride").first).to_be_visible()
            expect(diagram).to_contain_text("Proof:")
            diagram.get_by_role("button", name="Hop table", exact=True).click()
            expect(
                diagram.get_by_role("columnheader", name="Direction justification")
            ).to_be_visible()
            diagram.get_by_role("button", name="Diagram", exact=True).click()
            diagram.locator(".hop-links button", has_text="Evidence").first.click()
            expect(page.get_by_role("heading", name="Evidence traceability")).to_be_visible()
            page.get_by_role("button", name="Close details", exact=True).click()
            surface = page.get_by_label("Attack surface by STRIDE", exact=True)
            expect(surface.locator("tbody tr")).to_have_count(6)
            expect(surface).to_contain_text("No edge type maps to Repudiation.")
            # A bounded relationship list never claims the attack paths are incomplete.
            expect(page.get_by_text("displayed paths are not a complete")).to_have_count(0)
            page.get_by_role("tab", name="AWS", exact=True).click()
            page.get_by_role("tab", name="Attack Paths", exact=True).click()
            aws_paths = page.get_by_label("AWS attack paths", exact=True)
            # Never a blank canvas: diagrams, or the reasons no path was formed.
            expect(
                aws_paths.locator("figure.path-diagram, .path-empty, .path-coverage-details").first
            ).to_be_visible()
            page.get_by_role("tab", name="Consolidated", exact=True).click()
            page.get_by_role("tab", name="Attack Paths", exact=True).click()
            page.screenshot(path=str(artifacts / "attack-paths.png"), full_page=True)
            # The approved checklist classifies GCP findings by control (API key -> 1.8.5).
            page.get_by_role("button", name="Assessment checklist", exact=True).click()
            controls = page.get_by_label("Checklist controls")
            page.get_by_label("Search controls").select_option("1.8.5")
            row = controls.locator("tbody tr").filter(has_text="1.8.5")
            expect(row).to_contain_text("FAIL")
            row.get_by_role("button", name="findings").click()
            expect(
                page.locator("tbody tr")
                .filter(has_text="1.8.5")
                .filter(has_text="API key lacks application and API restrictions")
            ).to_have_count(1)
            page.screenshot(path=str(artifacts / "checklist.png"), full_page=True)
            page.get_by_role("button", name="Report center", exact=True).click()
            # Native Office files only; PDF output was withdrawn.
            with page.expect_download(timeout=60000) as download:
                page.get_by_role("button", name="Download assessment report (Excel + Word)").click()
            assert download.value.suggested_filename.endswith("-assessment-report.zip")
            with zipfile.ZipFile(download.value.path()) as bundle:
                names = bundle.namelist()
                (book,) = [n for n in names if n.endswith(".xlsx")]
                (word,) = [n for n in names if n.endswith(".docx")]
                assert not any(n.endswith(".pdf") for n in names)
                assert "demo-alpha" in word_text(bundle.read(word))
                report = load_workbook(io.BytesIO(bundle.read(book)), read_only=True)
                assert report.sheetnames == list(EXCEL_SHEETS)
                report.close()
            page.get_by_role("button", name="Export all approved findings (Word DOCX)").click()
            rtf_row = page.locator("tr").filter(has_text="CloSec-assessment-approved-docx.docx")
            with page.expect_download() as download:
                rtf_row.get_by_role("link", name="Download").click()
            assert "demo-alpha" in word_text(download.value.path())
            expect(
                page.get_by_role("button", name="Export all approved findings (PDF)")
            ).to_have_count(0)
            page.get_by_role("button", name="Generate Excel", exact=True).click()
            excel = page.locator("tr").filter(has_text="CloSec-assessment-excel.xlsx")
            with page.expect_download() as download:
                excel.get_by_role("link", name="Download").click()
            with open(download.value.path(), "rb") as incoming:
                wb = load_workbook(incoming)
            assert {r[1] for r in list(wb["Unified Findings Register"].values)[3:]} == {
                "AWS",
                "GCP",
            }
            wb.close()
            page.reload()
            page.get_by_role("tab", name="Consolidated", exact=True).click()
            expect(page.get_by_label("GCP scope summary")).to_contain_text("Approved findings1")
            server.should_exit = True
            thread.join(10)
            assert not thread.is_alive()
            listener.close()
            listener = socket.socket()
            listener.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
            listener.bind(("127.0.0.1", port))
            listener.listen(128)
            app = create_app(workspace, origin=origin, bootstrap="multicloud-restart")
            server = uvicorn.Server(uvicorn.Config(app, log_level="critical", access_log=False))
            thread = threading.Thread(target=lambda: server.run(sockets=[listener]), daemon=True)
            thread.start()
            for _ in range(100):
                if server.started:
                    break
                time.sleep(0.05)
            assert server.started
            page.goto(origin + "/#token=multicloud-restart")
            page.get_by_role("tab", name="Consolidated", exact=True).click()
            expect(page.get_by_label("GCP scope summary")).to_contain_text("Approved findings1")
            expect(page.get_by_label("AWS scope summary")).to_contain_text("Candidate findings8")
            page.get_by_role("button", name="Finalized findings", exact=True).click()
            freeze = page.get_by_role("button", name="Freeze finalised assessment", exact=True)
            expect(freeze).to_be_disabled()
            page.get_by_label("Confirm freeze assessment", exact=True).check()
            freeze.click()
            expect(page.get_by_role("status").filter(has_text="Frozen edition")).to_be_visible()
            expect(page.get_by_role("link", name="GCP PDF", exact=True)).to_have_count(0)
            with page.expect_download() as final_word:
                page.get_by_role("link", name="GCP DOCX", exact=True).click()
            assert Path(final_word.value.path()).read_bytes().startswith(b"PK")
            page.reload()
            page.get_by_role("button", name="Finalized findings", exact=True).click()
            expect(page.get_by_role("status").filter(has_text="Frozen edition")).to_be_visible()
            assert not external and not errors
            (artifacts / "result.json").write_text(
                '{"external_requests":0,"page_errors":0,"workflow":"PASSED"}'
            )
            browser.close()
    finally:
        server.should_exit = True
        thread.join(10)
        listener.close()
