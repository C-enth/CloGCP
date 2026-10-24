"""Local enterprise API and practical browser workflow with separately imported accounts."""

import os
import socket
import threading
import time
from pathlib import Path

import pytest
import uvicorn
from fastapi.testclient import TestClient
from test_api import ORIGIN, TOKEN, authenticated, await_job, upload_package
from test_enterprise import A, B, iam_rows, package

from offline_analyzer.api.app import create_app
from offline_analyzer.enterprise.queries import EnterpriseViews
from offline_analyzer.storage.store import Store
from tests.browser_diagnostics import capture_browser


def test_api_assembles_selected_runs_and_persists_qualified_queries(tmp_path):
    target = package(tmp_path / "target", iam_rows(B), B)
    source = package(tmp_path / "source", iam_rows(A, True), A)
    app = create_app(tmp_path / "work", origin=ORIGIN, bootstrap=TOKEN)
    with TestClient(app, base_url=ORIGIN) as client:
        headers = authenticated(client)
        ids = []
        for folder in (target, source):
            upload = upload_package(client, headers, folder).json()
            ids.append(
                await_job(
                    client,
                    client.post(
                        "/api/import", headers=headers, json={"upload_id": upload["upload_id"]}
                    ),
                )["run_id"]
            )
        run = await_job(
            client, client.post("/api/enterprise/assess", headers=headers, json={"members": ids})
        )["run_id"]
        assert (
            client.get(f"/api/runs/{run}/enterprise/dashboard").json()["represented_accounts"] == 2
        )
        page = client.get(
            f"/api/runs/{run}/enterprise/relationships?source={A}&destination={B}&kind=ASSUME_ROLE"
        ).json()
        assert {e["status"] for e in page["items"]} == {"DERIVED_SUPPORTED", "BLOCKED"}
        missing = client.get(f"/api/runs/{run}/enterprise/relationships?missing_only=true").json()
        assert all(e["missing"] for e in missing["items"])
        privileged = client.get(f"/api/runs/{run}/enterprise/relationships?privileged=true").json()
        assert all(
            e["target_privilege"]
            in {"ADMIN_EQUIVALENT", "ESCALATION_CAPABLE", "POTENTIALLY_PRIVILEGED"}
            for e in privileged["items"]
        )
        graph = client.get(f"/api/runs/{run}/enterprise/graph?node_limit=2&edge_limit=2").json()
        assert len(graph["nodes"]) <= 2 and graph["truncated"]
        assert client.get(f"/api/runs/{run}/enterprise/graph?node_limit=151").status_code == 400
        issue = client.get(f"/api/runs/{run}/enterprise/investigations").json()["items"][0]
        assert (
            client.post(
                f"/api/runs/{run}/enterprise/review/{issue['investigation_id']}",
                headers=headers,
                json={
                    "status": "NEEDS_VALIDATION",
                    "analyst": "synthetic",
                    "note": "Review shared root",
                },
            ).status_code
            == 200
        )
        assert (
            client.get(f"/api/runs/{run}/enterprise/investigations").json()["items"][0][
                "analyst_status"
            ]
            == "NEEDS_VALIDATION"
        )
        evidence = page["items"][0]["observation_ids"][0]
        detail = client.get(f"/api/runs/{run}/evidence/{evidence}").json()
        assert detail["relationships"]
        assert client.post("/api/enterprise/assess", json={"members": ids}).status_code == 403


@pytest.mark.skipif(os.environ.get("RUN_GUI_BROWSER") != "1", reason="opt-in local browser")
def test_enterprise_browser_separate_imports_to_graph_review_report(tmp_path):
    from openpyxl import load_workbook
    from playwright.sync_api import expect, sync_playwright

    target = package(tmp_path / "target", iam_rows(B), B)
    source = package(tmp_path / "source", iam_rows(A, True), A)
    listener = socket.socket()
    listener.bind(("127.0.0.1", 0))
    listener.listen(128)
    origin = f"http://127.0.0.1:{listener.getsockname()[1]}"
    app = create_app(tmp_path / "work", origin=origin, bootstrap="synthetic-enterprise-browser")
    server = uvicorn.Server(uvicorn.Config(app, log_level="critical", access_log=False))
    thread = threading.Thread(target=lambda: server.run(sockets=[listener]), daemon=True)
    thread.start()
    for _ in range(200):
        if server.started:
            break
        time.sleep(0.05)
    assert server.started
    external = []
    errors = []
    try:
        with sync_playwright() as pw:
            browser = pw.chromium.launch(headless=True)
            context = browser.new_context(
                accept_downloads=True, viewport={"width": 1440, "height": 1000}
            )

            with capture_browser(context, Path("artifacts/recovery-browser/enterprise")):

                def route(request):
                    if request.request.url.startswith(origin):
                        request.continue_()
                    else:
                        external.append(request.request.url)
                        request.abort()

                context.route("**/*", route)
                page = context.new_page()
                page.on("pageerror", lambda e: errors.append(str(e)))
                page.goto(origin + "/#token=synthetic-enterprise-browser")
                for folder in (target, source):
                    page.get_by_role("button", name="Evidence import", exact=True).click()
                    page.get_by_label("Select package folder").set_input_files(str(folder))
                    page.get_by_role("button", name="Validate evidence", exact=True).click()
                    page.get_by_role("button", name="Import validated evidence").wait_for()
                    page.get_by_role("button", name="Import validated evidence").click()
                    status = page.get_by_label("Local operation details").get_by_role("status")
                    expect(status).to_contain_text("Import")
                    expect(status).to_contain_text("SUCCEEDED")
                    expect(page.get_by_label("Active evidence run")).to_be_enabled()
                page.get_by_role("button", name="Dashboard", exact=True).click()
                page.get_by_role("button", name="Build enterprise assessment", exact=True).click()
                expect(
                    page.get_by_label("Active evidence run").locator("option:checked")
                ).to_contain_text("Enterprise")
                page.get_by_role("button", name="Cross-account explorer", exact=True).click()
                page.get_by_label("Source account", exact=True).select_option(A)
                page.get_by_label("Destination account", exact=True).select_option(B)
                page.get_by_label("Relationship type", exact=True).select_option("ASSUME_ROLE")
                expect(
                    page.get_by_role("img", name="Enterprise relationship graph")
                ).to_be_visible()
                page.get_by_role("button", name="Inspect relationship", exact=True).first.click()
                page.get_by_role("heading", name="Typed relationship qualification").wait_for()
                page.get_by_role(
                    "button", name="Inspect relationship evidence", exact=False
                ).first.click()
                page.get_by_role("heading", name="Integrity and source locators").wait_for()
                page.get_by_role("button", name="Close details").last.click()
                page.get_by_role("button", name="Close details").last.click()
                page.get_by_role("button", name="Zoom in graph").click()
                page.locator(".graph-node").first.click()
                expect(page.locator(".graph-node.selected")).to_have_count(1)
                page.locator(".graph-edge").first.click()
                page.get_by_role("heading", name="Typed relationship qualification").wait_for()
                page.get_by_role("button", name="Close details").click()
                screenshots = Path("artifacts/sprint3-screenshots")
                screenshots.mkdir(parents=True, exist_ok=True)
                page.screenshot(path=str(screenshots / "cross-account.png"), full_page=True)
                page.get_by_role("button", name="Findings workbench", exact=True).click()
                page.get_by_role("button", name="Review shared investigation").first.click()
                page.get_by_label("Shared investigation analyst").fill(
                    "synthetic-enterprise-reviewer"
                )
                page.get_by_label("Shared investigation remarks").fill(
                    "Validate shared trust against original evidence"
                )
                page.get_by_role("button", name="Validate investigation", exact=True).click()
                expect(
                    page.get_by_role("cell", name="NEEDS_VALIDATION", exact=True)
                ).to_be_visible()
                page.reload()
                page.get_by_role("button", name="Findings workbench", exact=True).click()
                expect(
                    page.get_by_role("cell", name="NEEDS_VALIDATION", exact=True)
                ).to_be_visible()
                page.get_by_role("button", name="Report center", exact=True).click()
                page.get_by_role("button", name="Generate Excel", exact=True).click()
                page.locator("tr", has_text="assessment-excel.xlsx").get_by_role(
                    "link", name="Download"
                ).wait_for()
                with page.expect_download() as download:
                    page.locator("tr", has_text="assessment-excel.xlsx").get_by_role(
                        "link", name="Download"
                    ).click()
                path = tmp_path / "report.xlsx"
                download.value.save_as(path)
                wb = load_workbook(path)
                assert wb["Cross-Account Relationships"].max_row > 3
                assert wb["Enterprise Account Summary"].max_row - 3 == 3
                wb.close()
                assert not external and not errors
            browser.close()
    finally:
        server.should_exit = True
        thread.join(timeout=20)
        listener.close()
        assert not thread.is_alive()
    with Store(tmp_path / "work") as store:
        runs = store.db.execute("SELECT run_id FROM enterprise_snapshots").fetchall()
        assert (
            EnterpriseViews().issues(store, runs[0][0])["items"][0]["analyst_status"]
            == "NEEDS_VALIDATION"
        )
