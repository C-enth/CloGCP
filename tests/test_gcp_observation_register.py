"""Excel-only synthetic regression for the owner's forty-column GCP evidence register."""

import hashlib
import io
import json
import zipfile
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from offline_analyzer.enterprise.snapshot import assess_enterprise
from offline_analyzer.findings.generate import record_decision
from offline_analyzer.findings.sre import record_sre
from offline_analyzer.ingestion.importer import import_package
from offline_analyzer.reporting.assessment_report import export_assessment, model
from offline_analyzer.reporting.excel import export_excel
from offline_analyzer.reporting.gcp_register import (
    AUDIT_SHEET,
    COLUMNS,
    PENDING,
    SHEETS,
    EvidenceLookup,
    _pointer,
    _record,
    _redact,
    add_sheets,
    build,
)
from offline_analyzer.reporting.report_document import write_excel
from offline_analyzer.rules.engine import assess
from offline_analyzer.storage.store import Store
from tools.fixtures.r12_demo import build as aws_build
from tools.gcp_fixture import build as gcp_build


@pytest.fixture(scope="module")
def estate(tmp_path_factory):
    root = tmp_path_factory.mktemp("gcp-register")
    gcp_build(root / "gcp")
    gcp_build(root / "gcp-other", "demo-second", "222222222222")
    aws_build(root / "aws")
    with Store(root / "workspace") as store:
        runs = [import_package(store, root / name) for name in ("gcp", "gcp-other", "aws")]
        for run in runs:
            assess(store, run)
        run = assess_enterprise(store, runs)
        findings = [f for f in store.candidates(run) if f.provider == "GCP"]
        record_decision(
            store, run, findings[0].finding_id, "CONFIRMED", "Reviewer", "Synthetic review"
        )
        record_sre(
            store,
            run,
            [findings[0].finding_id],
            "GAP_CONFIRMED",
            "Reviewer",
            reason="Synthetic response",
        )
        yield store, run, root, runs, findings[0]


def database_digest(store):
    tables = [r[0] for r in store.rows("SHOW TABLES")]
    payload = [(t, list(store.rows(f'SELECT * FROM "{t}" ORDER BY ALL'))) for t in tables]
    return hashlib.sha256(json.dumps(payload, default=str, sort_keys=True).encode()).hexdigest()


def register_tail(names):
    """The register sheets end the workbook; the audit sheet follows only recorded decisions."""
    return names[-3:] == list(SHEETS) or names[-4:] == [*SHEETS, AUDIT_SHEET]


def values(ws):
    return list(ws.iter_rows(values_only=True))


def test_exact_contract_and_traceability_and_read_only(estate, tmp_path):
    store, run, _, _, reviewed = estate
    before = database_digest(store)
    raw_before = {
        p.name: hashlib.sha256(p.read_bytes()).hexdigest()
        for p in (store.workspace / "raw").glob("*.json")
    }
    data = model(store, run, "confirmed", "GCP")
    payload = build(store, run, data, provider="GCP")
    wb = Workbook(write_only=True)
    add_sheets(wb, payload)
    path = tmp_path / "register.xlsx"
    wb.save(path)
    saved = load_workbook(path)
    # Recorded decisions exist, so the review audit sheet follows the three register sheets.
    assert saved.sheetnames == [*SHEETS, AUDIT_SHEET]
    register, index, details = [saved[s] for s in SHEETS]
    # Owner-provided contract independently pinned in exact order.
    expected = (
        "Record ID (internal)|Observation / gap|Affected asset / "
        "scope|Affected objects / services|Criticality|Why this "
        "criticality?|Rating confidence (automatic)|Possible business "
        "impact|Recommended action|Assessor record type|What the evidence "
        "shows|Evidence reference|Expected control / requirement|Evidence at"
        " assessment|What the evidence does not establish|Possible attack "
        "scenario|Conditions needed for the scenario|Source-reported "
        "result|Domain / tool|Control area|Evidence type|Evidence limits / "
        "data quality|Assessor conclusion / evidence basis|Related "
        "observations|Attack-path evidence level|Assessor evidence result "
        "(automatic)|Report sequence (automatic)|Common asset reference "
        "(hostname / ID)|Issue group / summary title|Risk rating "
        "decision|Show in Overall?|Shown in Overall (automatic)|Record count"
        " (automatic)|Applicability (SPOC)|SPOC assessment|SPOC "
        "justification / comments|Supporting evidence / reference|Reviewed "
        "by|Review date|Final review conclusion (automatic)"
    )  # noqa: E501
    assert tuple(expected.split("|")) == COLUMNS == tuple(values(register)[0])
    records = {r[0]: dict(zip(COLUMNS, r, strict=True)) for r in values(register)[1:]}
    assert records[reviewed.finding_id][COLUMNS[4]] == reviewed.suggested_severity
    assert records[reviewed.finding_id][COLUMNS[9]] == reviewed.finding_class
    # Owner rule: SPOC, reviewer and review-date cells stay blank for the analyst; the
    # recorded tool decisions move to the audit sheet unchanged.
    for column in (33, 34, 35, 37, 38):
        assert records[reviewed.finding_id][COLUMNS[column]] in (None, "")
    audit = {r[0]: r for r in values(saved[AUDIT_SHEET])[1:]}
    assert audit[reviewed.finding_id][2] == "Reviewer"
    assert audit[reviewed.finding_id][5] == "GAP_CONFIRMED"
    assert audit[reviewed.finding_id][6] == "Synthetic response"
    assert records[reviewed.finding_id][COLUMNS[31]] is True
    assert any(r[COLUMNS[39]] == PENDING for r in records.values())
    assert any(r[COLUMNS[9]] == "Evidence limitation" for r in records.values())
    assert any(r[COLUMNS[9]] == "Observation" for r in records.values())
    assert all(
        r[COLUMNS[9]] == "Observation"
        for r in records.values()
        if r[COLUMNS[25]] in {"PASS", "NOT_APPLICABLE", "MANUAL_VALIDATION_REQUIRED"}
    )
    # Every actual reference resolves to an index row, then to an original JSON excerpt.
    indexed = {r[1] for r in values(index)[1:]}
    for row in register.iter_rows(min_row=2):
        if row[11].value:
            assert row[0].value in indexed and row[11].hyperlink.location.startswith(
                "'GCP_Evidence_Index'"
            )
    lookup = EvidenceLookup(store, run)
    for row in index.iter_rows(min_row=2):
        assert row[13].hyperlink.location.startswith("'GCP_Evidence_Details'")
        matching = [r for r in values(details)[1:] if r[0] == row[0].value]
        assert matching and all(r[1] == row[1].value for r in matching)
        actual, limitation = lookup.excerpt(row[9].value, row[7].value or "")
        if not limitation:
            assert "".join(r[7] or "" for r in matching) == actual
    for ws in saved:
        assert ws.freeze_panes == "D2" and ws.auto_filter.ref.startswith("A1:")
        assert len(ws.tables) == 1
        assert all(c.data_type != "f" for row in ws for c in row)
    assert database_digest(store) == before
    assert {
        p.name: hashlib.sha256(p.read_bytes()).hexdigest()
        for p in (store.workspace / "raw").glob("*.json")
    } == raw_before
    assert build(store, run, data, provider="GCP") == payload


def test_existing_sheets_and_aws_unchanged(estate, tmp_path):
    store, run, _, runs, _ = estate
    for provider in ("AWS", "GCP", ""):
        data = model(store, run, "all", provider)
        old, new = tmp_path / f"old-{provider}.xlsx", tmp_path / f"new-{provider}.xlsx"
        write_excel(data, old)
        write_excel(data, new, gcp_register=build(store, run, data, provider=provider))
        left, right = load_workbook(old), load_workbook(new)
        assert right.sheetnames[: len(left.sheetnames)] == left.sheetnames
        for name in left.sheetnames:
            assert values(left[name]) == values(right[name])
            assert left[name].auto_filter.ref == right[name].auto_filter.ref
            assert left[name].freeze_panes == right[name].freeze_panes
        if provider == "AWS":
            with zipfile.ZipFile(old) as a, zipfile.ZipFile(new) as b:
                assert a.namelist() == b.namelist()
                # Creation timestamps are the only nondeterministic workbook property.
                assert all(a.read(n) == b.read(n) for n in a.namelist() if n != "docProps/core.xml")
        else:
            assert register_tail(right.sheetnames)
    # Full report-center Excel still works; AWS-only retains the established sheet contract.
    for single, gcp in ((runs[0], True), (runs[2], False)):
        target = tmp_path / f"full-{gcp}.xlsx"
        export_excel(store, single, target)
        workbook = load_workbook(target, read_only=True)
        assert (SHEETS[0] in workbook.sheetnames) is gcp
        if not gcp:
            assert len(workbook.sheetnames) == 28


def test_scoped_assessment_export(estate, tmp_path):
    store, run, _, _, _ = estate
    target = tmp_path / "assessment.zip"
    export_assessment(store, run, target, include="confirmed", provider="GCP", scope="demo-alpha")
    with zipfile.ZipFile(target) as archive:
        name = next(n for n in archive.namelist() if n.endswith(".xlsx"))
        workbook = load_workbook(io.BytesIO(archive.read(name)))
    assert register_tail(workbook.sheetnames)
    assert {r[2] for r in values(workbook[SHEETS[0]])[1:]} == {"demo-alpha"}
    assert {r[2] for r in values(workbook[SHEETS[1]])[1:]} == {"demo-alpha"}


def test_missing_tampered_and_failed_evidence_is_a_limitation(estate, monkeypatch):
    store, run, _, _, _ = estate
    lookup = EvidenceLookup(store, run)
    artifact = next(iter(lookup.hashes))
    assert lookup.excerpt("missing", "/0")[0] == ""
    assert lookup.excerpt(artifact, "/nonexistent/path")[0] == ""
    original_open = type(store.workspace).open
    monkeypatch.setattr(
        type(store.workspace),
        "open",
        lambda p, *a, **k: (
            io.BytesIO(b"{}") if p.parent.name == "raw" else original_open(p, *a, **k)
        ),
    )
    lookup.cached_sha = None
    snippet, limitation = lookup.excerpt(artifact, "")
    assert not snippet and "integrity mismatch" in limitation


def test_redaction_pointer_chunking_and_excel_literals(tmp_path):
    assert _pointer({"a/b": {"~": [False, None]}}, "/a~1b/~0/0") is False
    assert _pointer({"a": 1}, "") == {"a": 1}
    secret = {
        "password": "synthetic-private",
        "policy": '{"token":"synthetic-token"}',
        "env": [{"name": "CUSTOM", "value": "synthetic-env"}],
        "name": "api-key",
        "value": "synthetic-value",
    }
    redacted = json.dumps(_redact(secret))
    assert not any(
        word in redacted
        for word in ("synthetic-private", "synthetic-token", "synthetic-env", "synthetic-value")
    )
    payload = {"register": [["=CMD()", *([""] * 39)]], "index": [], "details": []}
    wb = Workbook(write_only=True)
    add_sheets(wb, payload)
    file = tmp_path / "literal.xlsx"
    wb.save(file)
    saved = load_workbook(file)
    assert saved[SHEETS[0]]["A2"].data_type != "f"
    assert saved[SHEETS[1]].max_row == 1  # no artificial example/default row


def test_long_excerpts_are_complete_and_failures_never_become_configuration():
    from offline_analyzer.contracts.models import EvidenceRef

    excerpt = json.dumps({"configuration": "synthetic " * 7000})
    lookup = SimpleNamespace(
        sources={
            "artifact": {
                "operation": "gcp:compute.instances.list",
                "effective_coverage": "COMPLETE",
                "path": "original.json",
            }
        },
        hashes={"artifact": "a" * 64},
        excerpt=lambda a, p: (excerpt, ""),
    )
    record = {
        "id": "synthetic-record",
        "type": "Observation",
        "refs": (EvidenceRef(artifact_id="artifact", locator="/0"),),
    }
    output = {"register": [], "index": [], "details": []}
    _record(output, record, 1, lookup, [])
    assert len(output["details"]) > 1
    assert "".join(row[7] for row in output["details"]) == excerpt
    assert "Concise excerpt" in output["register"][0][10]
    lookup.sources["artifact"]["effective_coverage"] = "ERROR"
    lookup.excerpt = lambda a, p: pytest.fail("failed collection must not be read as configuration")
    output = {"register": [], "index": [], "details": []}
    _record(output, {**record, "type": "Evidence limitation"}, 1, lookup, [])
    assert output["details"][0][7] == ""
    assert output["register"][0][4] is None
    assert output["register"][0][13] == "ERROR"
    assert "no usable configuration" in output["details"][0][8]


def test_recorded_report_and_risk_decisions_are_preserved(estate):
    from offline_analyzer.contracts.report_finding import ReportApproval, ReportEdit
    from offline_analyzer.findings.report_ready import (
        approve_revision,
        report_detail,
        save_revision,
    )

    store, run, _, _, finding = estate
    sections = report_detail(store, run, finding.finding_id)["sections"]
    sections.update(
        title="Synthetic analyst title", attack_scenario="Synthetic scenario recorded by reviewer"
    )
    revision = save_revision(
        store,
        run,
        finding.finding_id,
        ReportEdit(
            expected_revision=0,
            analyst="Reviewer",
            sections=sections,
            member_ids=[finding.finding_id],
            include_in_report=False,
        ),
    )
    approve_revision(
        store,
        run,
        finding.finding_id,
        ReportApproval(
            expected_revision=revision.revision,
            analyst="Risk owner",
            kind="RISK_ACCEPTANCE",
            rationale="Synthetic time-bound risk decision recorded for regression.",
        ),
    )
    before = database_digest(store)
    data = model(store, run, "confirmed", "GCP")
    payload = build(store, run, data, provider="GCP")
    row = next(r for r in payload["register"] if r[0] == finding.finding_id)
    assert row[1] == "Synthetic analyst title"
    assert row[15] == "Synthetic scenario recorded by reviewer"
    assert "Risk owner" in row[29]
    assert row[30] is False  # the saved draft's toggle, independent of this export's scope
    assert row[31] is True  # the existing confirmed report model includes this finding
    assert row[33] == ""  # applicability has no persisted SPOC decision
    assert database_digest(store) == before


def test_failed_collection_only_still_exports_limitations(tmp_path):
    from offline_analyzer.contracts.gcp import GCPDataset, GCPManifest

    folder = tmp_path / "denied-package"
    folder.mkdir()
    manifest = GCPManifest(
        collector_version="synthetic-gcp-1",
        assessment_id="synthetic-denied",
        expected_projects=["demo-denied"],
        collected_projects=[],
        dataset_origin="synthetic",
        datasets=[
            GCPDataset(
                project_id="demo-denied",
                operation="gcp:compute.instances.list",
                status="ACCESS_DENIED",
                record_count=0,
                errors=["ACCESS_DENIED"],
            )
        ],
    )
    (folder / "manifest.json").write_text(manifest.model_dump_json())
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, folder)
        assess(store, run)
        assert not list(store.candidates(run))
        before = database_digest(store)
        target = tmp_path / "denied.xlsx"
        export_excel(store, run, target)
        workbook = load_workbook(target)
        rows = values(workbook[SHEETS[0]])[1:]
        assert rows and all(r[9] == "Evidence limitation" and r[4] is None for r in rows)
        assert all("unavailable" in r[10] for r in rows)
        assert database_digest(store) == before
