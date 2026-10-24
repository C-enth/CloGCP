"""Enterprise synthetic integration: source imports, conservative edges and persistent review."""

import json

import pytest
from openpyxl import load_workbook

from offline_analyzer.contracts.r0 import Package, Scope, Source
from offline_analyzer.enterprise.queries import EnterpriseViews
from offline_analyzer.enterprise.snapshot import assess_enterprise
from offline_analyzer.ingestion.importer import import_package
from offline_analyzer.ingestion.json_input import InputError
from offline_analyzer.ingestion.safety import sha256_file
from offline_analyzer.inventory.accounts import AccountRegistry, install_registry
from offline_analyzer.normalization.enterprise_adapters import PROFILES
from offline_analyzer.reporting.excel import export_excel
from offline_analyzer.storage.store import Store
from offline_analyzer.verification import verify_package
from tools.fixtures.r12_demo import TIME, allow, deny, doc, gaad_role, trust

A, B, C, X = "111111111111", "222222222222", "333333333333", "444444444444"


def package(folder, rows, account, targets=True):
    folder.mkdir(parents=True)
    sources = []
    for i, (op, body, params) in enumerate(rows):
        path = folder / f"{i}.json"
        path.write_text(json.dumps(body))
        sources.append(
            Source(
                path=path.name,
                sha256=sha256_file(path),
                format="aws-json-v1",
                operation=op,
                account_id=account,
                region="us-east-1",
                collected_at=TIME,
                collector_version="synthetic-enterprise-1",
                request_parameters=params,
                coverage="COMPLETE",
            )
        )
    p = Package(
        sources=tuple(sources),
        targets=(Scope(account_id=account, region="us-east-1"),) if targets else (),
    )
    (folder / "manifest.json").write_text(p.model_dump_json())
    return folder


def iam_rows(account, source=False):
    arn = f"arn:aws:iam::{B}:role/target"
    if source:
        roles = [
            gaad_role(
                A,
                "allowed",
                trust({"Service": "ec2.amazonaws.com"}),
                inline={"grant": doc(allow("sts:AssumeRole", arn))},
            ),
            gaad_role(
                A,
                "denied",
                trust({"Service": "ec2.amazonaws.com"}),
                inline={"grant": doc(allow("sts:AssumeRole", arn), deny("sts:AssumeRole", arn))},
            ),
        ]
    else:
        roles = [
            gaad_role(
                B,
                "target",
                trust({"AWS": [f"arn:aws:iam::{A}:root", f"arn:aws:iam::{C}:root"]}),
                inline={"admin": doc(allow("*"))},
            )
        ]
    return [
        (
            "iam:GetAccountAuthorizationDetails",
            {
                "RoleDetailList": roles,
                "Policies": [],
                "UserDetailList": [],
                "GroupDetailList": [],
                "IsTruncated": False,
            },
            {},
        ),
        (
            "organizations:DescribeOrganization",
            {
                "Error": {"Code": "AWSOrganizationsNotInUseException"},
                "ResponseMetadata": {"HTTPStatusCode": 400},
            },
            {},
        ),
    ]


@pytest.fixture
def workspace(tmp_path):
    with Store(tmp_path / "workspace") as store:
        yield store


def test_cross_account_correlation_across_separate_imports_preserves_history(workspace, tmp_path):
    target = import_package(workspace, package(tmp_path / "target", iam_rows(B), B))
    first = assess_enterprise(workspace)
    v = EnterpriseViews()
    refs = v.relationships(
        workspace, first, source=A, destination=B, kind="TRUST_POLICY_ALLOW_REFERENCE"
    )
    assert refs["total"] == 1 and refs["items"][0]["status"] == "DIRECTLY_OBSERVED"
    unresolved = v.relationships(workspace, first, source=A, kind="ASSUME_ROLE")["items"]
    assert unresolved and unresolved[0]["status"] == "POTENTIAL" and unresolved[0]["missing"]
    source = import_package(workspace, package(tmp_path / "source", iam_rows(A, True), A))
    second = assess_enterprise(workspace)
    assert second != first
    transitions = v.relationships(workspace, second, source=A, destination=B, kind="ASSUME_ROLE")[
        "items"
    ]
    allowed = next(e for e in transitions if e["source"].endswith("/allowed"))
    blocked = next(e for e in transitions if e["source"].endswith("/denied"))
    assert allowed["status"] == "DERIVED_SUPPORTED" and not allowed["missing"]
    assert blocked["status"] == "BLOCKED" and any("grant#1" in b for b in blocked["blocking"])
    assert allowed["target_privilege"] == "ADMIN_EQUIVALENT"
    assert allowed["evidence"] and blocked["evidence"]
    assert v.dashboard(workspace, second)["represented_accounts"] == 2
    assert v.dashboard(workspace, second)["referenced_accounts"] >= 1
    assert assess_enterprise(workspace, [target, source]) == second
    for run in (target, source, first, second):
        workspace.require_run(run)
    assert (
        v.relationships(workspace, first, kind="ASSUME_ROLE")["items"][0]["authorization"]
        == "UNKNOWN"
    )


def test_shared_investigation_identity_and_review_are_persistent(workspace, tmp_path):
    import_package(workspace, package(tmp_path / "target", iam_rows(B), B))
    run = assess_enterprise(workspace)
    v = EnterpriseViews()
    issue = v.issues(workspace, run)["items"][0]
    assert set(issue["affected_accounts"]) == {A, B, C}
    v.review(
        workspace,
        run,
        issue["investigation_id"],
        "NEEDS_VALIDATION",
        "synthetic",
        "Verify original trust conditions",
    )
    assert v.issues(workspace, run)["items"][0]["analyst_status"] == "NEEDS_VALIDATION"
    assert assess_enterprise(workspace) == run
    assert v.issues(workspace, run)["items"][0]["analyst_status"] == "NEEDS_VALIDATION"
    assert (
        workspace.db.execute(
            "SELECT count(*) FROM decisions WHERE json_extract_string(data,'$.status')='CONFIRMED'"
        ).fetchone()[0]
        == 0
    )


def test_registry_alias_groups_and_missing_accounts_are_not_inferred(workspace, tmp_path):
    registry = tmp_path / "registry.json"
    registry.write_text(
        AccountRegistry.model_validate(
            {
                "schema_version": "account-registry-1",
                "source": "synthetic approved inventory",
                "approved_by": "synthetic",
                "approved_at": TIME,
                "entries": [
                    {
                        "account_id": B,
                        "in_scope": True,
                        "alias": "Central security",
                        "groups": ["Security"],
                    },
                    {"account_id": X, "in_scope": True},
                ],
            }
        ).model_dump_json()
    )
    install_registry(workspace.workspace, registry)
    import_package(workspace, package(tmp_path / "target", iam_rows(B), B))
    run = assess_enterprise(workspace)
    accounts = EnterpriseViews().accounts(workspace, run)
    b = next(a for a in accounts if a["account"] == B)
    x = next(a for a in accounts if a["account"] == X)
    assert b["alias"] == "Central security" and b["groups"] == ["Security"]
    assert not x["represented"] and not x["assessed"] and x["identity_basis"] == "APPROVED_REGISTRY"
    assert EnterpriseViews().dashboard(workspace, run)["missing_registry_accounts"] == 1
    assert (
        EnterpriseViews().relationships(workspace, run, group="Security", outside=True)["total"] > 0
    )


def test_graph_and_page_limits_are_explicit(workspace, tmp_path):
    import_package(workspace, package(tmp_path / "target", iam_rows(B), B))
    run = assess_enterprise(workspace)
    v = EnterpriseViews()
    graph = v.graph(workspace, run, node_limit=2, edge_limit=2)
    assert len(graph["nodes"]) <= 2 and len(graph["edges"]) <= 2 and graph["truncated"]
    with pytest.raises(InputError):
        v.graph(workspace, run, node_limit=151)
    with pytest.raises(InputError):
        v.relationships(workspace, run, limit=201)
    with pytest.raises(InputError):
        v.relationships(workspace, run, account="not-an-account")
    page = v.relationships(workspace, run, limit=1)
    assert len(page["items"]) == 1 and page["total"] > 1


PROFILE_CASES = {
    "ec2:DescribeSubnets": (
        {
            "Subnets": [
                {
                    "SubnetId": "subnet-00000001",
                    "OwnerId": B,
                    "VpcId": "vpc-00000001",
                    "CidrBlock": "10.0.1.0/24",
                    "AvailabilityZone": "us-east-1a",
                    "MapPublicIpOnLaunch": False,
                }
            ]
        },
        {},
    ),
    "ec2:DescribeNetworkAcls": (
        {
            "NetworkAcls": [
                {
                    "NetworkAclId": "acl-00000001",
                    "OwnerId": B,
                    "VpcId": "vpc-00000001",
                    "Entries": [],
                    "Associations": [],
                    "IsDefault": False,
                }
            ]
        },
        {},
    ),
    "ec2:DescribeInternetGateways": (
        {
            "InternetGateways": [
                {"InternetGatewayId": "igw-00000001", "OwnerId": B, "Attachments": []}
            ]
        },
        {},
    ),
    "s3api:GetBucketPolicy": (
        {
            "Policy": json.dumps(
                doc(
                    {
                        "Effect": "Allow",
                        "Principal": {"AWS": f"arn:aws:iam::{A}:root"},
                        "Action": "s3:GetObject",
                        "Resource": "arn:aws:s3:::synthetic-shared/*",
                    }
                )
            )
        },
        {"Bucket": "synthetic-shared", "ExpectedBucketOwner": B},
    ),
    "kms:GetKeyPolicy": (
        {
            "Policy": json.dumps(
                doc(
                    {
                        "Effect": "Allow",
                        "Principal": {"AWS": f"arn:aws:iam::{A}:root"},
                        "Action": "kms:Decrypt",
                        "Resource": "*",
                    }
                )
            )
        },
        {
            "KeyId": f"arn:aws:kms:us-east-1:{B}:key/00000000-0000-0000-0000-000000000001",
            "PolicyName": "default",
        },
    ),
    "kms:ListGrants": (
        {
            "Grants": [
                {
                    "GrantId": "synthetic-grant",
                    "KeyId": f"arn:aws:kms:us-east-1:{B}:key/00000000-0000-0000-0000-000000000001",
                    "GranteePrincipal": f"arn:aws:iam::{A}:role/allowed",
                    "RetiringPrincipal": f"arn:aws:iam::{B}:root",
                    "Operations": ["Decrypt"],
                    "Constraints": {},
                }
            ]
        },
        {"KeyId": f"arn:aws:kms:us-east-1:{B}:key/00000000-0000-0000-0000-000000000001"},
    ),
    "lambda:ListFunctions": (
        {
            "Functions": [
                {
                    "FunctionArn": f"arn:aws:lambda:us-east-1:{B}:function:synthetic",
                    "Role": f"arn:aws:iam::{B}:role/target",
                }
            ]
        },
        {},
    ),
    "ecs:DescribeTaskDefinition": (
        {
            "taskDefinition": {
                "taskDefinitionArn": f"arn:aws:ecs:us-east-1:{B}:task-definition/synthetic:1",
                "taskRoleArn": f"arn:aws:iam::{B}:role/target",
                "executionRoleArn": f"arn:aws:iam::{B}:role/target",
            }
        },
        {"taskDefinition": "synthetic:1"},
    ),
    "ec2:DescribeVpcs": (
        {
            "Vpcs": [
                {
                    "VpcId": "vpc-00000000000000001",
                    "OwnerId": B,
                    "CidrBlock": "10.0.0.0/16",
                    "State": "available",
                }
            ]
        },
        {},
    ),
    "ec2:DescribeVpcPeeringConnections": (
        {
            "VpcPeeringConnections": [
                {
                    "VpcPeeringConnectionId": "pcx-00000000000000001",
                    "RequesterVpcInfo": {
                        "VpcId": "vpc-00000000000000001",
                        "OwnerId": A,
                        "Region": "us-east-1",
                    },
                    "AccepterVpcInfo": {
                        "VpcId": "vpc-00000000000000002",
                        "OwnerId": B,
                        "Region": "us-east-1",
                    },
                    "Status": {"Code": "active"},
                }
            ]
        },
        {},
    ),
    "ec2:DescribeTransitGateways": (
        {
            "TransitGateways": [
                {"TransitGatewayId": "tgw-00000000000000001", "OwnerId": B, "State": "available"}
            ]
        },
        {},
    ),
    "ec2:DescribeTransitGatewayAttachments": (
        {
            "TransitGatewayAttachments": [
                {
                    "TransitGatewayAttachmentId": "tgw-attach-00000000000000001",
                    "TransitGatewayId": "tgw-00000000000000001",
                    "TransitGatewayOwnerId": B,
                    "ResourceId": "vpc-00000000000000001",
                    "ResourceOwnerId": A,
                    "ResourceType": "vpc",
                    "State": "available",
                    "Association": {},
                }
            ]
        },
        {},
    ),
    "ec2:DescribeTransitGatewayRouteTables": (
        {
            "TransitGatewayRouteTables": [
                {
                    "TransitGatewayRouteTableId": "tgw-rtb-00000000000000001",
                    "TransitGatewayId": "tgw-00000000000000001",
                    "State": "available",
                }
            ]
        },
        {},
    ),
    "ec2:GetTransitGatewayRouteTableAssociations": (
        {
            "Associations": [
                {
                    "TransitGatewayAttachmentId": "tgw-attach-00000000000000001",
                    "ResourceId": "vpc-00000000000000001",
                    "ResourceType": "vpc",
                    "State": "associated",
                }
            ]
        },
        {"TransitGatewayRouteTableId": "tgw-rtb-00000000000000001"},
    ),
    "ec2:GetTransitGatewayRouteTablePropagations": (
        {
            "TransitGatewayRouteTablePropagations": [
                {
                    "TransitGatewayAttachmentId": "tgw-attach-00000000000000001",
                    "ResourceId": "vpc-00000000000000001",
                    "ResourceType": "vpc",
                    "State": "enabled",
                }
            ]
        },
        {"TransitGatewayRouteTableId": "tgw-rtb-00000000000000001"},
    ),
    "ec2:SearchTransitGatewayRoutes": (
        {
            "Routes": [
                {
                    "DestinationCidrBlock": "10.0.0.0/16",
                    "Type": "static",
                    "State": "active",
                    "TransitGatewayAttachments": [
                        {"TransitGatewayAttachmentId": "tgw-attach-00000000000000001"}
                    ],
                    "PrefixListId": None,
                }
            ],
            "AdditionalRoutesAvailable": False,
        },
        {"TransitGatewayRouteTableId": "tgw-rtb-00000000000000001", "Filters": []},
    ),
    "ram:GetResourceShares": (
        {
            "resourceShares": [
                {
                    "resourceShareArn": f"arn:aws:ram:us-east-1:{B}:resource-share/synthetic",
                    "owningAccountId": B,
                    "status": "ACTIVE",
                    "allowExternalPrincipals": True,
                }
            ]
        },
        {"resourceOwner": "SELF"},
    ),
    "ram:ListResources": (
        {
            "resources": [
                {
                    "arn": f"arn:aws:ec2:us-east-1:{B}:transit-gateway/tgw-00000000000000001",
                    "resourceShareArn": f"arn:aws:ram:us-east-1:{B}:resource-share/synthetic",
                    "status": "AVAILABLE",
                    "type": "ec2:TransitGateway",
                }
            ]
        },
        {"resourceOwner": "SELF"},
    ),
    "ram:ListPrincipals": (
        {
            "principals": [
                {
                    "id": A,
                    "resourceShareArn": f"arn:aws:ram:us-east-1:{B}:resource-share/synthetic",
                    "external": True,
                }
            ]
        },
        {"resourceOwner": "SELF"},
    ),
}


@pytest.mark.parametrize("operation", list(PROFILES))
def test_additional_reference_profiles_are_synthetic_only_and_keep_authorization_qualified(
    workspace, tmp_path, operation
):
    body, params = PROFILE_CASES[operation]
    folder = package(tmp_path / "input", [(operation, body, params)], B)
    diagnostic = verify_package(folder, origin="synthetic")
    assert diagnostic["sources"][0]["parser_accepted"]
    assert diagnostic["real_export_validation"] == "NOT_ESTABLISHED"
    import_package(workspace, folder)
    run = assess_enterprise(workspace)
    for edge in EnterpriseViews().relationships(workspace, run)["items"]:
        assert edge["status"] != "DERIVED_SUPPORTED"
        assert edge["authorization"] is None
        assert edge["evidence"]


def test_excel_enterprise_sheets_reconcile_and_have_valid_names(workspace, tmp_path):
    import_package(workspace, package(tmp_path / "target", iam_rows(B), B))
    import_package(workspace, package(tmp_path / "source", iam_rows(A, True), A))
    run = assess_enterprise(workspace)
    path = export_excel(workspace, run, tmp_path / "enterprise.xlsx")
    wb = load_workbook(path)
    assert all(len(s) <= 31 for s in wb.sheetnames)
    assert (
        wb["Cross-Account Relationships"].max_row - 3
        == EnterpriseViews().relationships(workspace, run)["total"]
    )
    assert (
        wb["Enterprise Security Gaps"].max_row - 3
        == EnterpriseViews().issues(workspace, run)["total"]
    )
    assert (
        wb["Findings Register"].max_row - 3
        == workspace.db.execute("SELECT count(*) FROM findings WHERE run_id=?", [run]).fetchone()[0]
    )
    assert all(cell.data_type != "f" for sheet in wb for row in sheet for cell in row)
    wb.close()


def test_conflicting_trust_references_survive_as_unknown(workspace, tmp_path):
    import_package(workspace, package(tmp_path / "first", iam_rows(B), B))
    changed = iam_rows(B)
    changed[0][1]["RoleDetailList"][0]["AssumeRolePolicyDocument"] = trust(
        {"AWS": f"arn:aws:iam::{A}:root"}
    )
    import_package(workspace, package(tmp_path / "second", changed, B))
    run = assess_enterprise(workspace)
    refs = EnterpriseViews().relationships(workspace, run, kind="TRUST_POLICY_ALLOW_REFERENCE")[
        "items"
    ]
    assert refs and all(e["status"] == "UNKNOWN" for e in refs)
    assert any(e["source_account"] == C for e in refs)
    assert all(len(e["evidence"]) >= 2 for e in refs)


def test_mixed_parser_versions_are_rejected_without_mutating_imports(workspace, tmp_path):
    run = import_package(workspace, package(tmp_path / "first", iam_rows(B), B))
    run2 = import_package(workspace, package(tmp_path / "second", iam_rows(A, True), A))
    row = workspace.db.execute(
        "SELECT observation_id,data FROM observations WHERE run_id=? LIMIT 1", [run2]
    ).fetchone()
    data = json.loads(row[1])
    data["normalization_version"] = "r0.1.4"
    workspace.db.execute(
        "UPDATE observations SET data=? WHERE run_id=? AND observation_id=?",
        [json.dumps(data), run2, row[0]],
    )
    with pytest.raises(InputError, match="ENTERPRISE_NORMALIZER_VERSION_CONFLICT"):
        assess_enterprise(workspace, [run, run2])
    assert workspace.db.execute("SELECT count(*) FROM runs").fetchone()[0] == 2


def test_no_cross_partition_identity_collapse_in_graph(workspace, tmp_path):
    rows = iam_rows(B)
    arn = f"arn:aws-cn:iam::{B}:role/target"
    rows[0][1]["RoleDetailList"][0]["Arn"] = arn
    rows[0][1]["RoleDetailList"][0]["AssumeRolePolicyDocument"] = trust(
        {"AWS": f"arn:aws-cn:iam::{A}:root"}
    )
    import_package(workspace, package(tmp_path / "aws", iam_rows(B), B))
    import_package(workspace, package(tmp_path / "cn", rows, B))
    run = assess_enterprise(workspace)
    refs = EnterpriseViews().relationships(workspace, run, kind="TRUST_POLICY_ALLOW_REFERENCE")[
        "items"
    ]
    assert {"aws", "aws-cn"} <= {e["destination_partition"] for e in refs}
    graph = EnterpriseViews().graph(workspace, run, node_limit=100)
    assert any(n["partition"] == "aws-cn" for n in graph["nodes"])
    assert len({n["id"] for n in graph["nodes"]}) == len(graph["nodes"])


def test_tgw_attachment_never_establishes_network_reachability(workspace, tmp_path):
    op = "ec2:DescribeTransitGatewayAttachments"
    body, params = PROFILE_CASES[op]
    import_package(workspace, package(tmp_path / "network", [(op, body, params)], B))
    run = assess_enterprise(workspace)
    edge = EnterpriseViews().relationships(workspace, run, kind="TGW_ATTACHMENT_REFERENCE")[
        "items"
    ][0]
    assert edge["source_account"] == A and edge["destination_account"] == B
    assert (
        edge["status"] == "DIRECTLY_OBSERVED" and edge["authorization"] is None and edge["missing"]
    )


def test_shared_vpc_owner_is_resolved_from_inventory_not_interface_owner(workspace, tmp_path):
    from tools.fixtures.r1_demo import interface

    vpc = "vpc-00000000000000001"
    eni = interface(A, 1, [1])
    eni["VpcId"] = vpc
    import_package(
        workspace,
        package(
            tmp_path / "participant",
            [("ec2:DescribeNetworkInterfaces", {"NetworkInterfaces": [eni]}, {})],
            A,
        ),
    )
    first = assess_enterprise(workspace)
    first_ref = EnterpriseViews().relationships(workspace, first, kind="RESOURCE_IN_VPC")["items"][
        0
    ]
    assert first_ref["destination_account"] is None and first_ref["missing"]
    import_package(
        workspace,
        package(
            tmp_path / "owner",
            [
                (
                    "ec2:DescribeVpcs",
                    {
                        "Vpcs": [
                            {
                                "VpcId": vpc,
                                "OwnerId": B,
                                "CidrBlock": "10.0.0.0/16",
                                "State": "available",
                            }
                        ]
                    },
                    {},
                )
            ],
            B,
        ),
    )
    second = assess_enterprise(workspace)
    ref = EnterpriseViews().relationships(workspace, second, kind="RESOURCE_IN_VPC")["items"][0]
    assert ref["source_account"] == A and ref["destination_account"] == B
    assert len(ref["evidence"]) == 2 and ref["authorization"] is None


def test_identical_raw_bytes_with_new_collection_time_keep_observation_traceability(
    workspace, tmp_path
):
    from offline_analyzer.api.services import Views

    rows = [
        (
            "ec2:DescribeInstances",
            {
                "Reservations": [
                    {
                        "Instances": [
                            {
                                "InstanceId": "i-0123456789abcdef0",
                                "State": {"Name": "running"},
                                "MetadataOptions": {
                                    "State": "applied",
                                    "HttpTokens": "optional",
                                    "HttpEndpoint": "enabled",
                                },
                            }
                        ]
                    }
                ]
            },
            {},
        )
    ]
    first = package(tmp_path / "first", rows, A)
    second = package(tmp_path / "second", rows, A)
    manifest = second / "manifest.json"
    data = json.loads(manifest.read_text())
    data["sources"][0]["collected_at"] = "2026-10-11T12:00:00Z"
    manifest.write_text(json.dumps(data))
    one = import_package(workspace, first)
    two = import_package(workspace, second)
    run = assess_enterprise(workspace, [one, two])
    observations = list(workspace.observations(run))
    assert len(observations) == 2
    assert len({o.collected_at for o in observations}) == 2
    assert len({o.evidence.artifact_id for o in observations}) == 2
    for obs in observations:
        assert workspace.refs(run, [obs.observation_id]) == (obs.evidence,)
        detail = Views(workspace.workspace).evidence_detail(workspace, run, obs.observation_id)
        assert detail["provenance"] and all(p["integrity"] for p in detail["provenance"])
        assert detail["evaluations"] and detail["findings"]
    assert all(r.evidence for r in workspace.results(run) if r.resource_id)
