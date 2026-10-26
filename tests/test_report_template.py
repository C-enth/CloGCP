"""Assessment report template (Word + Excel) on synthetic data: sections, hop rows, STRIDE
matrix, frozen headers and filters, GCP / AWS / GCP+AWS scopes, exclusions stay excluded."""

import io
import zipfile
from xml.etree.ElementTree import fromstring

import pytest
from openpyxl import load_workbook

from offline_analyzer.contracts.exclusion import ExclusionRequest
from offline_analyzer.enterprise.snapshot import assess_enterprise
from offline_analyzer.findings.exclusions import record as exclude
from offline_analyzer.ingestion.importer import import_package
from offline_analyzer.reporting.assessment_report import export_assessment, model
from offline_analyzer.reporting.report_document import EXCEL_SHEETS, STRIDE
from offline_analyzer.rules.engine import assess
from offline_analyzer.storage.store import Store
from tools.fixtures.r12_demo import build as aws_build
from tools.gcp_fixture import build as gcp_build

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
SECTIONS = (
    "1 Executive summary",
    "2 Coverage and limitations",
    "3 Attack surface",
    "4 Attack paths",
    "5 STRIDE and direction justification",
    "6 Findings",
    "Appendix A Glossary and evidence basis",
)
TABLE_SHEETS = (
    "Attack Surface",
    "STRIDE Matrix",
    "Attack Paths",
    "Attack Path Hops",
    "Technical Findings",
    "Findings Validation",
    "Edge STRIDE Mapping",
)


@pytest.fixture(scope="module")
def estate(tmp_path_factory):
    root = tmp_path_factory.mktemp("report")
    aws_build(root / "aws")
    gcp_build(root / "gcp")
    with Store(root / "work") as store:
        members = [import_package(store, root / "aws"), import_package(store, root / "gcp")]
        for run in members:
            assess(store, run)
        run = assess_enterprise(store, members)
        linked = next(
            f
            for p in model(store, run, include="valid")["paths"]
            for f in p["finding_ids"]
            if p.get("hops")
        )
        exclude(
            store,
            run,
            linked,
            ExclusionRequest(
                action="EXCLUDE",
                reason="FALSE_POSITIVE",
                note="Synthetic reason",
                analyst="Synthetic analyst",
            ),
        )
        yield store, run, linked, root


def paragraphs(docx: bytes) -> list[tuple[str, str]]:
    document = fromstring(zipfile.ZipFile(io.BytesIO(docx)).read("word/document.xml"))
    out = []
    for p in document.iter(W + "p"):
        style = p.find(f"{W}pPr/{W}pStyle")
        text = "".join(t.text or "" for t in p.iter(W + "t"))
        out.append((style.get(W + "val") if style is not None else "", text))
    return out


@pytest.mark.parametrize("provider", ["", "GCP", "AWS"])
def test_word_and_excel_follow_the_template_for_each_cloud_scope(estate, provider):
    store, run, hidden, root = estate
    data = model(store, run, "valid", provider)
    target = root / f"report-{provider or 'all'}.zip"
    target.unlink(missing_ok=True)
    export_assessment(store, run, target, include="valid", provider=provider)
    bundle = zipfile.ZipFile(target)
    names = bundle.namelist()
    (book,) = [n for n in names if n.endswith(".xlsx")]
    (word,) = [n for n in names if n.endswith(".docx")]
    scope = provider or "GCP + AWS"

    # Word: cover, numbered sections in order, real tables, hop tables, UNKNOWN areas.
    content = paragraphs(bundle.read(word))
    texts = [t for _s, t in content]
    assert f"Cloud scope: {scope}" in texts
    headings = [t for s, t in content if s == "Heading1"]
    assert [h for h in headings if h in SECTIONS] == list(SECTIONS)
    raw = zipfile.ZipFile(io.BytesIO(bundle.read(word))).read("word/document.xml")
    tree = fromstring(raw)
    xml = raw.decode()
    assert (
        len(list(tree.iter(W + "tbl"))) >= 6 and next(tree.iter(W + "tblHeader"), None) is not None
    )
    assert "UNKNOWN and manual areas" in texts
    hops = sum(len(p.get("hops") or []) for p in data["paths"])
    if hops:
        assert "Direction justification" in texts and "Hops" in texts
    assert hidden not in xml
    try:  # optional: reopen with an independent OOXML reader when it is installed
        import docx as python_docx
    except ImportError:
        python_docx = None
    if python_docx:
        reopened = python_docx.Document(io.BytesIO(bundle.read(word)))
        assert len(reopened.tables) >= 6
        assert any(p.text == "4 Attack paths" for p in reopened.paragraphs)

    # Excel: sheet order, frozen headers and filters on every table, one row per hop.
    workbook = load_workbook(io.BytesIO(bundle.read(book)))
    from offline_analyzer.reporting.gcp_register import SHEETS as GCP_SHEETS

    assert tuple(workbook.sheetnames) == EXCEL_SHEETS + (() if provider == "AWS" else GCP_SHEETS)
    for name in TABLE_SHEETS:
        sheet = workbook[name]
        assert sheet.freeze_panes == "A2", name
        assert sheet.auto_filter.ref and sheet.auto_filter.ref.startswith("A1:"), name
    hop_sheet = workbook["Attack Path Hops"]
    header = [c.value for c in hop_sheet[1]]
    assert {"Direction justification", "Hop status", *STRIDE} <= set(header)
    rows = [
        r
        for r in hop_sheet.iter_rows(min_row=2, values_only=True)
        if r[0] != "No rows for this report scope."
    ]
    assert len(rows) == hops
    for row in rows:
        values = dict(zip(header, row, strict=True))
        assert values["Direction justification"] and "→" in values["Direction"]
        assert sum(values[s] for s in STRIDE) >= 1
        if provider:
            assert values["Account / project"].startswith(provider)
    matrix = workbook["STRIDE Matrix"]
    assert [c.value for c in matrix[1]][-6:] == list(STRIDE)
    assert workbook["Cover"]["A2"].value == f"Cloud scope: {scope}"
    for sheet in workbook.worksheets:
        if sheet.title == "Excluded Findings":
            continue  # Transparency appendix lists excluded IDs (tests/test_false_positive.py).
        for row in sheet.iter_rows(values_only=True):
            assert not any(isinstance(v, str) and hidden in v for v in row), sheet.title
    workbook.close()


def test_excluded_findings_stay_out_of_paths_and_hops(estate):
    store, run, hidden, _root = estate
    data = model(store, run)
    assert hidden not in {f["id"] for f in data["findings"]}
    assert data["excluded"] >= 1
    for path in data["paths"]:
        assert hidden not in path["finding_ids"]
        assert all(hidden not in h["finding_ids"] for h in path.get("hops") or [])
    assert data["path_coverage"] and data["provider_scope"] == "GCP + AWS"
