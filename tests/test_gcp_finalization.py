"""GCP finalization: claim consistency, presentation identity, register columns and QA.

Positive, negative and ambiguous synthetic fixtures for each change. Every change is display
or export text: stored findings, IDs, fingerprints, decisions and editions are compared
before and after all exports.
"""

import hashlib
import json

import pytest
from openpyxl import Workbook, load_workbook

from clogcp_report_finalizer.core import finalize
from clogcp_report_finalizer.safety import digest
from offline_analyzer.enterprise.snapshot import assess_enterprise
from offline_analyzer.findings.detail import finding_detail
from offline_analyzer.findings.explanations import MISUSE_QUALIFIER, explain
from offline_analyzer.findings.finalisation import freeze
from offline_analyzer.findings.gcp_recommendations import EXPECTED
from offline_analyzer.findings.generate import fingerprint, record_decision
from offline_analyzer.findings.location import display_region
from offline_analyzer.findings.recommendations import recommendation
from offline_analyzer.findings.sre import record_sre
from offline_analyzer.ingestion.importer import import_package
from offline_analyzer.reporting.assessment_report import model
from offline_analyzer.reporting.gcp_register import (
    AUDIT_SHEET,
    COLUMNS,
    SHEETS,
    _focus,
    add_sheets,
    build,
)
from offline_analyzer.reporting.sre_reports import finalized
from offline_analyzer.rules.gcp_domains import SPECS
from offline_analyzer.rules.gcp_org import ORG_RULES
from offline_analyzer.storage.store import Store
from tests.test_clogcp_domains import POLICY, SERVICE, run
from tools.fixtures.finalizer_workbook import build as finalizer_sample
from tools.gcp_fixture import build as gcp_build

GENERIC = (
    "applicability and required security posture reviewed",
    "Secure baseline",
    "Analyst to define expected state",
)


# ---------------------------------------------------------------- Cloud Run semantics


@pytest.mark.parametrize(
    ("config", "policies", "outcome"),
    [
        # Positive: an open service with a public binding, or the IAM check explicitly off.
        ({"ingress": "INGRESS_TRAFFIC_ALL"}, [POLICY], "FAIL"),
        ({"invokerIamDisabled": True, "ingress": "INGRESS_TRAFFIC_ALL"}, [], "FAIL"),
        ({"invokerIamDisabled": "true", "ingress": "INGRESS_TRAFFIC_ALL"}, [], "FAIL"),
        (
            {
                "metadata": {"annotations": {"run.googleapis.com/invoker-iam-disabled": "true"}},
                "ingress": "INGRESS_TRAFFIC_ALL",
            },
            [],
            "FAIL",
        ),
        # Negative: None / False never mean "disabled"; internal ingress limits the claim.
        ({"invokerIamDisabled": None, "ingress": "INGRESS_TRAFFIC_ALL"}, [], "MANUAL"),
        ({"invokerIamDisabled": False, "ingress": "INGRESS_TRAFFIC_ALL"}, [], "MANUAL"),
        ({"ingress": "INGRESS_TRAFFIC_ALL"}, [], "MANUAL"),
        ({"invokerIamDisabled": True, "ingress": "INGRESS_TRAFFIC_INTERNAL_ONLY"}, [], "MANUAL"),
        (
            {"invokerIamDisabled": None, "ingress": "INGRESS_TRAFFIC_INTERNAL_LOAD_BALANCER"},
            [POLICY],
            "MANUAL",
        ),
        # Ambiguous: a non-boolean flag is never interpreted.
        ({"invokerIamDisabled": "yes", "ingress": "INGRESS_TRAFFIC_ALL"}, [POLICY], "UNKNOWN"),
        ({"invokerIamDisabled": 1, "ingress": "INGRESS_TRAFFIC_ALL"}, [], "UNKNOWN"),
    ],
)
def test_invoker_iam_disabled_semantics(config, policies, outcome):
    value = run("GCP-RUN-PUBLIC", config, policies, native=SERVICE)
    expected = "MANUAL_VALIDATION_REQUIRED" if outcome == "MANUAL" else outcome
    assert value.outcome == expected, value.reasoning
    if outcome == "FAIL":
        # The reasoning records the observed values the claim rests on.
        assert "invokerIamDisabled=" in value.reasoning and "ingress=" in value.reasoning


# ------------------------------------------------------- narrative and claim consistency


def test_misuse_never_repeats_impact_for_any_domain_rule():
    for rule, spec in SPECS.items():
        text = explain({"rule_id": rule, "technical_observation": "synthetic"})
        assert text["misuse"].endswith(f"({MISUSE_QUALIFIER})"), rule
        assert spec.risk not in text["misuse"], rule
        assert text["impact"] == spec.risk


def test_org_public_serverless_claims_match_its_unresolved_evidence():
    text = explain({"rule_id": "GCP-ORG-PUBLIC-SERVERLESS"})
    risk = next(r for r in ORG_RULES if r.rule_id == "GCP-ORG-PUBLIC-SERVERLESS").risk
    assert "unresolved" in risk
    # Positive: the scenario is conditional on ingress; negative: no unconditional claims.
    assert "If ingress" in text["misuse"]
    for claim in ("Anyone on the internet can call", "run up cost", "without authentication"):
        assert claim not in text["misuse"]
    assert "Cloud Armor" not in text["recommendation"]
    assert "approved public endpoint" in text["recommendation"]


def test_every_gcp_domain_and_org_rule_has_a_checkable_expected_state():
    rules = [*SPECS, *(r.rule_id for r in ORG_RULES)]
    for rule in rules:
        posture = recommendation(rule).get("expected_posture") or ""
        assert rule in EXPECTED and posture == EXPECTED[rule], rule
        assert not any(g in posture for g in GENERIC), rule
    # Ambiguous: an unknown rule keeps the explicit analyst-authored status, never a guess.
    assert recommendation("GCP-NOT-A-RULE")["status"] == "ANALYST_AUTHORED_REMEDIATION_REQUIRED"


# --------------------------------------------------------------- presentation identity


@pytest.mark.parametrize(
    ("stored", "shown"),
    [
        (
            "https://www.googleapis.com/compute/v1/projects/synth/zones/us-central1-a",
            "us-central1-a",
        ),
        ("projects/synth/regions/europe-west1", "europe-west1"),
        ("us-central1", "us-central1"),
        ("global", "global"),
        ("US", "US"),
        ("https://example.invalid/not-a-location", "unrecorded"),
        ("projects/synth/zones/", "unrecorded"),
        ("", "unrecorded"),
        (None, "unrecorded"),
        ("zone with spaces", "unrecorded"),
    ],
)
def test_display_region_uses_only_explicit_location_segments(stored, shown):
    assert display_region(stored) == shown


# ------------------------------------------------------------------ evidence projection


def test_focus_shows_only_cited_bindings_and_never_guesses():
    policy = {
        "resource": "//run.googleapis.com/projects/synth/locations/us-central1/services/api",
        "policy": {
            "etag": "x",
            "bindings": [
                {"role": "roles/run.invoker", "members": ["allUsers", "user:a@example.com"]},
                {"role": "roles/run.admin", "members": ["group:ops@example.com"]},
                {"role": "roles/viewer", "members": ["deleted:user:old@example.com?uid=1"]},
            ],
        },
    }
    focused = _focus(policy, "roles/run.invoker -> allUsers")
    assert focused["bindings"] == [{"role": "roles/run.invoker", "members": ["allUsers"]}]
    assert focused["omitted"].startswith("3 other members")
    deleted = _focus(policy, "roles/viewer -> deleted:user:old@example.com")
    assert deleted["bindings"][0]["members"] == ["deleted:user:old@example.com?uid=1"]
    # Negative / ambiguous: uncited, unmatched or non-policy evidence is unchanged.
    assert _focus(policy, "no binding cited") == policy
    assert _focus(policy, "roles/owner -> user:nobody@example.com") == policy
    assert _focus({"ingress": "all"}, "roles/run.invoker -> allUsers") == {"ingress": "all"}


# ------------------------------------------------- register columns and stored-data equality


def database_digest(store):
    tables = [r[0] for r in store.rows("SHOW TABLES")]
    payload = [(t, list(store.rows(f'SELECT * FROM "{t}" ORDER BY ALL'))) for t in tables]
    return hashlib.sha256(json.dumps(payload, default=str, sort_keys=True).encode()).hexdigest()


@pytest.fixture(scope="module")
def assessed(tmp_path_factory):
    root = tmp_path_factory.mktemp("gcp-finalization")
    with Store(root / "workspace") as store:
        imports = [
            import_package(store, gcp_build(root / "a")),
            import_package(store, gcp_build(root / "b", "demo-second", "222222222222")),
        ]
        run_id = assess_enterprise(store, imports)
        findings = [f for f in store.candidates(run_id) if f.provider == "GCP"]
        record_decision(store, run_id, findings[0].finding_id, "CONFIRMED", "Synthetic", "note")
        record_sre(store, run_id, [findings[0].finding_id], "GAP_CONFIRMED", "Synthetic SRE")
        yield store, run_id, root, findings


def test_register_columns_follow_the_owner_rules(assessed, tmp_path):
    store, run_id, _, findings = assessed
    payload = build(store, run_id, model(store, run_id, "all", "GCP"), provider="GCP")
    rows = {r[0]: dict(zip(COLUMNS, r, strict=True)) for r in payload["register"]}
    candidates = [rows[f.finding_id] for f in findings]
    for row in candidates:
        # Columns 15 and 22 come from different coverage sources and are never identical.
        assert row[COLUMNS[14]] != row[COLUMNS[21]]
        assert row[COLUMNS[12]] and row[COLUMNS[12]] != row[COLUMNS[19]]
        assert row[COLUMNS[15]]
        assert "records)" in row[COLUMNS[28]]
        assert row[COLUMNS[32]] == len(row[COLUMNS[23]].split("; ")) + 1 or (
            row[COLUMNS[32]] == 1 and not row[COLUMNS[23]]
        )
        for column in (33, 34, 35, 37, 38):
            assert row[COLUMNS[column]] in ("", None)
        assert "Legacy projection" in row[COLUMNS[24]] or not row[COLUMNS[24]]
    audit = {r[0]: r for r in payload["audit"]}
    assert audit[findings[0].finding_id][2] == "Synthetic"
    assert audit[findings[0].finding_id][5] == "GAP_CONFIRMED"
    workbook = Workbook(write_only=True)
    add_sheets(workbook, payload)
    workbook.save(tmp_path / "register.xlsx")
    assert load_workbook(tmp_path / "register.xlsx").sheetnames == [*SHEETS, AUDIT_SHEET]


def test_exports_change_no_stored_identity_decision_or_edition(assessed, tmp_path):
    store, run_id, _, findings = assessed
    stored = sorted((f.finding_id, fingerprint(f)) for f in store.candidates(run_id))
    freeze(store, run_id, "Synthetic", True)
    edition = list(store.rows("SELECT data FROM finalized_assessments WHERE run_id=?", [run_id]))
    before = database_digest(store)
    raw = {p.name: digest(p) for p in (store.workspace / "raw").glob("*.json")}
    build(store, run_id, model(store, run_id, "all", "GCP"), provider="GCP")
    finalized(store, run_id, "GCP")
    for finding in findings[:5]:
        finding_detail(store, run_id, finding.finding_id)
    assert database_digest(store) == before
    assert {p.name: digest(p) for p in (store.workspace / "raw").glob("*.json")} == raw
    assert sorted((f.finding_id, fingerprint(f)) for f in store.candidates(run_id)) == stored
    after = list(store.rows("SELECT data FROM finalized_assessments WHERE run_id=?", [run_id]))
    assert after == edition


# ---------------------------------------------------------------- finalizer QA register


def test_finalizer_lists_manual_validation_and_prefilled_review_cells(tmp_path):
    sample = finalizer_sample(tmp_path / "input")
    before = digest(sample["source"])
    finalize(sample["source"], tmp_path / "out", guidance=sample["guidance"], origin="SYNTHETIC")
    assert digest(sample["source"]) == before
    qa = json.loads((tmp_path / "out" / "Assessor-QA-Exceptions.json").read_text())
    categories = {}
    for item in qa["exceptions"]:
        categories.setdefault(item[2], set()).add(item[1])
    source = load_workbook(sample["source"], read_only=True)
    rows = list(source["GCP_Observation_Register"].iter_rows(min_row=2, values_only=True))
    manual = {str(r[0]) for r in rows if "MANUAL_VALIDATION_REQUIRED" in str(r[25] or "")}
    prefilled = {str(r[0]) for r in rows if r[37]}
    source.close()
    assert manual and categories["MANUAL_VALIDATION_ITEM"] == manual
    assert prefilled and categories["PREFILLED_REVIEW_FIELDS"] == prefilled
