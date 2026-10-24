"""Collector packages (closec-collector-1), operation summaries, export jobs and the writer."""

import importlib.util
import json
import threading
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from offline_analyzer.api.app import create_app
from offline_analyzer.api.coordinator import Coordinator
from offline_analyzer.contracts.r0 import AWS_OPERATIONS
from offline_analyzer.ingestion import collector
from offline_analyzer.ingestion.importer import import_package, load_package
from offline_analyzer.ingestion.json_input import InputError
from offline_analyzer.ingestion.safety import sha256_file
from offline_analyzer.reporting.excel import REPORT_SHEETS, export_excel
from offline_analyzer.rules.engine import assess
from offline_analyzer.storage.store import Store
from tests.test_api import ORIGIN, TOKEN, authenticated, await_job
from tools.fixtures.r12_demo import build

spec = importlib.util.spec_from_file_location("collector_fixture", "tools/collector_fixture.py")
fixture = importlib.util.module_from_spec(spec)
spec.loader.exec_module(fixture)


@pytest.fixture
def packages(tmp_path):
    build(tmp_path / "r12")
    return fixture.convert(tmp_path / "r12", tmp_path / "collector")


def manifest(package: Path) -> dict:
    return json.loads((package / collector.COLLECTOR_MANIFEST).read_text())


def rewrite(package: Path, data: dict) -> None:
    (package / collector.COLLECTOR_MANIFEST).write_text(json.dumps(data))


def test_collector_package_imports_with_operation_summary(tmp_path, packages):
    first = packages[0]
    with Store(tmp_path / "work") as store:
        run = import_package(store, first)
        assert import_package(store, first) == run  # identical package reimports to same run
        sources = {(s["operation"], s["effective_coverage"]) for s in store.source_data(run)}
        assert ("cloudtrail:DescribeTrails", "ERROR") in sources
        assert all(not op.startswith("securityhub") for op, _ in sources)
        assess(store, run)
        outcomes = {
            str(r.outcome) for r in store.results(run) if r.rule_id == "R0-CLOUDTRAIL-VALIDATION"
        }
        assert outcomes and "PASS" not in outcomes  # failed collection is never clean
    summary = collector.stored_summary(tmp_path / "work", run)
    assert summary["totals"]["unsupported_operations"] == 1
    assert summary["totals"]["failed"] == 1
    unsupported = [o for o in summary["operations"] if o["operation"] == "securityhub:GetFindings"]
    assert unsupported[0]["import_handling"] == "UNSUPPORTED_OPERATION_RETAINED"
    retained = tmp_path / "work/raw" / f"{unsupported[0]['artifact_sha256']}.json"
    assert sha256_file(retained) == unsupported[0]["artifact_sha256"]  # not silently dropped


def test_truncated_operation_is_partial_and_unlisted_files_are_reported(tmp_path, packages):
    package = packages[1]
    data = manifest(package)
    data["operations"][0]["pagination"] = "TRUNCATED"
    rewrite(package, data)
    (package / "responses/extra.json").write_text("{}")
    loaded, summary = collector.inspect_collector(package)
    assert loaded.sources[0].coverage == "PARTIAL"
    assert summary["unlisted_files"] == ["responses/extra.json"]
    assert summary["totals"]["truncated"] == 1


@pytest.mark.parametrize(
    ("change", "code"),
    [
        (lambda d, p: (p / d["operations"][0]["path"]).write_text("{}"), "ARTIFACT_CHECKSUM"),
        (lambda d, p: (p / d["operations"][0]["path"]).unlink(), "COLLECTOR_ARTIFACT_MISSING"),
        (lambda d, p: d.update(started_at="2026-10-10T07:00:00+02:00"), "MANIFEST_INVALID"),
        (lambda d, p: d["operations"][0].update(status="FAILED"), "MANIFEST_INVALID"),
        (lambda d, p: d.update(account_id="12345"), "MANIFEST_INVALID"),
        (lambda d, p: d.update(schema="closec-collector-0"), "MANIFEST_INVALID"),
        (lambda d, p: d.update(unexpected=True), "MANIFEST_INVALID"),
        (lambda d, p: (p / "link.json").symlink_to(p / d["operations"][0]["path"]), "SYMLINK"),
    ],
)
def test_collector_integrity_failures_reject_package(tmp_path, packages, change, code):
    package = packages[0]
    data = manifest(package)
    change(data, package)
    rewrite(package, data)
    with pytest.raises(InputError, match=code), Store(tmp_path / "work") as store:
        import_package(store, package)
    with Store(tmp_path / "work") as store:
        assert not store.db.execute("SELECT count(*) FROM runs").fetchone()[0]


def test_prowler_version_and_operation_format_are_enforced():
    base = {
        "region": "us-east-1",
        "status": "SUCCEEDED",
        "path": "p.json",
        "sha256": "0" * 64,
        "started_at": "2026-10-10T07:00:00Z",
        "completed_at": "2026-10-10T07:00:00Z",
    }
    with pytest.raises(ValueError, match="Prowler"):
        collector.Operation(
            operation="prowler:Scan", format="prowler-ocsf-5.10", tool_version="4.0.0", **base
        )
    with pytest.raises(ValueError, match="format"):
        collector.Operation(operation="prowler:Scan", **base)
    with pytest.raises(ValueError, match="before it started"):
        collector.Operation(
            operation="ec2:DescribeInstances", **{**base, "started_at": "2026-10-10T08:00:00Z"}
        )


def test_contract_lists_every_supported_operation(tmp_path):
    output = tmp_path / "contract/ANALYZER_COLLECTOR_CONTRACT.json"
    assert collector.main(["--output", str(output)]) == 0
    contract = json.loads(output.read_text())
    listed = {o["operation"] for o in contract["supported_operations"]}
    assert listed == set(AWS_OPERATIONS) | {"prowler:Scan"}
    assert contract["package_schema"] == "closec-collector-1"
    assert "operations" in contract["manifest_json_schema"]["properties"]
    instances = next(
        o for o in contract["supported_operations"] if o["operation"] == "ec2:DescribeInstances"
    )
    assert (
        "MetadataOptions.HttpTokens"
        in instances["required_response_fields"]["Reservations.item.Instances.item"]
    )
    rds = next(
        o for o in contract["supported_operations"] if o["operation"] == "rds:DescribeDBInstances"
    )
    assert rds["boto3_method"] == "describe_db_instances"


def test_legacy_manifest_takes_precedence_and_still_loads(tmp_path):
    build(tmp_path / "pkg")
    (tmp_path / "pkg" / collector.COLLECTOR_MANIFEST).write_text("{}")
    assert len(load_package(tmp_path / "pkg").sources) > 1


def test_api_collector_upload_collection_view_and_evidence_requests(tmp_path, packages):
    app = create_app(tmp_path / "work", origin=ORIGIN, bootstrap=TOKEN)
    package = packages[0]
    names = sorted(
        p.relative_to(package).as_posix() for p in package.rglob("*.json") if p.is_file()
    )
    files = [("files", (n, (package / n).read_bytes(), "application/json")) for n in names]
    with TestClient(app, base_url=ORIGIN) as client:
        headers = authenticated(client)
        preview = client.post(
            "/api/uploads",
            headers=headers,
            files=files,
            data={"paths": json.dumps(names), "context": '{"origin":"synthetic"}'},
        ).json()
        assert preview["can_import"] and preview["collection"]["totals"]["unsupported_operations"]
        report = client.get(f"/api/reports/{preview['compatibility_report']}/download")
        assert report.json()["collector_package"]["totals"]["failed"] == 1
        assert "111111111111" not in report.text
        job = client.post("/api/import", headers=headers, json={"upload_id": preview["upload_id"]})
        run = await_job(client, job)["run_id"]
        view = client.get(f"/api/runs/{run}/collection").json()
        assert view["collector_packages"][0]["totals"]["failed"] == 1
        trails = next(
            o for o in view["operations"] if o["operation"] == "cloudtrail:DescribeTrails"
        )
        assert trails["ERROR"] == 1 and trails["errors"] == ["AccessDenied"]
        requests = view["evidence_requests"]
        assert any(
            r["operation"] == "cloudtrail:DescribeTrails" and r["state"] == "ERROR"
            for r in requests
        )
        csv_text = client.get(f"/api/runs/{run}/evidence-requests.csv").text
        assert csv_text.splitlines()[0] == "account_id,region,operation,state,reason,request"
        assert "AccessDenied" in csv_text


def test_evidence_request_csv_neutralizes_formulas(tmp_path):
    from offline_analyzer.api.services import Views

    views = Views(tmp_path)
    views.collection = lambda store, run: {
        "evidence_requests": [
            {
                "account_id": "=HYPERLINK()",
                "region": "+1",
                "operation": "@x",
                "state": "-1",
                "reason": None,
                "request": "ok",
            }
        ]
    }
    row = views.evidence_requests_csv(None, "run").splitlines()[1]
    assert row.startswith("'=HYPERLINK(),'+1,'@x,'-1,,ok")


def test_report_jobs_deduplicate_publish_atomically_and_report_progress(tmp_path, monkeypatch):
    coordinator = Coordinator(tmp_path)
    release = threading.Event()
    first = coordinator.submit("Report", lambda store: release.wait(10), key="report:r:excel")
    second = coordinator.submit("Report", lambda store: None, key="report:r:excel")
    assert second["id"] == first["id"] and second["deduplicated"] is True
    other = coordinator.submit("Report", lambda store: None, key="report:r:findings")
    assert other["id"] != first["id"]
    release.set()
    coordinator.close()

    from offline_analyzer.api import services

    build(tmp_path / "pkg")
    work = tmp_path / "work"
    with Store(work) as store:
        run = import_package(store, tmp_path / "pkg")
        assess(store, run)
        views = services.Views(work)

        def broken(*_args, **_kwargs):
            _args[2].write_text("partial")
            raise RuntimeError("interrupted")

        monkeypatch.setattr(services, "export_excel", broken)
        with pytest.raises(RuntimeError):
            views.export(store, run, "excel")
        assert not list((work / "reports").glob("*.meta"))
        assert not list((work / "reports").glob("*.xlsx"))
        assert not list((work / "reports").glob("*.partial"))
        monkeypatch.undo()
        stages = []
        result = views.export(store, run, "excel", progress=lambda *a: stages.append(a))
    assert stages[-1] == ("Saving workbook", REPORT_SHEETS, REPORT_SHEETS)
    assert [s[1] for s in stages] == sorted(s[1] for s in stages)
    assert result["name"] == "CloSec-assessment-excel.xlsx"
    (work / "reports" / "stale.xlsx.partial").write_text("interrupted")
    create_app(work, origin=ORIGIN, bootstrap=TOKEN)  # startup discards stale partials
    assert not list((work / "reports").glob("*.partial"))
    workbook = load_workbook(work / "reports" / f"{result['id']}.xlsx")
    assert len(workbook.sheetnames) == REPORT_SHEETS
    assert workbook.properties.title == "CloSec assessment"


def test_row_writer_preserves_types_whitespace_and_formula_safety(tmp_path):
    from offline_analyzer.reporting import excel

    build(tmp_path / "pkg")
    with Store(tmp_path / "work") as store:
        run = import_package(store, tmp_path / "pkg")
        assess(store, run)
        original = excel.assessment_metrics

        def metrics(store, run_id):
            return {
                **original(store, run_id),
                "z_formula": "=1+1",
                "z_space": " padded ",
                "z_empty": "",
                "z_bool": True,
                "z_int": 7,
                "z_control": "a\x01b",
            }

        excel.assessment_metrics = metrics
        try:
            export_excel(store, run, tmp_path / "out.xlsx")
        finally:
            excel.assessment_metrics = original
    ws = load_workbook(tmp_path / "out.xlsx")["Assessment Summary"]
    values = {r[0].value: r[1] for r in ws.iter_rows(min_row=4) if r[0].value}
    assert values["z_formula"].value == "'=1+1" and values["z_formula"].data_type == "s"
    assert values["z_space"].value == " padded "
    assert values["z_empty"].value is None
    assert values["z_bool"].value is True and values["z_int"].value == 7
    assert values["z_control"].value == "ab"
    assert values["z_int"].alignment.wrap_text and values["z_int"].alignment.vertical == "top"
    header = ws["A3"]
    assert header.font.b and header.fill.fgColor.rgb.endswith("17365D")
