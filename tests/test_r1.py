"""R1 methodology-mapped rules: truth tables, evidence gaps, traceability and reporting."""

import json
import re
from collections import Counter
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from offline_analyzer.attack_paths.prerequisites import path_records
from offline_analyzer.attack_paths.specs import BY_PATH, KINDS, PATHS, THREAT_SCOPE
from offline_analyzer.cli import main, summary
from offline_analyzer.contracts.r0 import Package, Source
from offline_analyzer.ingestion.importer import import_package
from offline_analyzer.ingestion.json_input import InputError
from offline_analyzer.ingestion.safety import sha256_file
from offline_analyzer.methodology.coverage_spec import (
    BY_CONTROL,
    COVERAGE,
    REASONS,
    REVIEWED_METHODOLOGY_SHA256,
    STATUSES,
)
from offline_analyzer.reporting.excel import export_excel
from offline_analyzer.reporting.tracker import TRACKER_COLUMNS, load_account_keys, tracker_rows
from offline_analyzer.rules.context import scope_coverage
from offline_analyzer.rules.engine import assess
from offline_analyzer.rules.iam_trust import trust_conditions
from offline_analyzer.rules.network import sensitive_ingress
from offline_analyzer.rules.r1 import account_bpa, hop_limit, rds_public
from offline_analyzer.rules.registry import ALL_RULES, BY_ID
from offline_analyzer.storage.store import Store
from tools.fixtures.r1_demo import ACCOUNT_KEYS, TIME, A, B, build_r1_package, rid


@pytest.fixture
def r1(tmp_path):
    directory = tmp_path / "input"
    build_r1_package(directory)
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, directory)
        assess(store, run)
        yield store, run, directory


def label(store, run, result) -> str:
    if result.resource_id is None:
        return "<scope>"
    obs = next(store.observations(run, resource_id=result.resource_id))
    return obs.configuration.get("RoleName") or obs.native_id


EXPECTED = {
    ("R0-EC2-IMDSV2", A, rid("i", 1)): "FAIL",
    ("R0-EC2-IMDSV2", A, rid("i", 2)): "PASS",
    ("R1-EC2-IMDS-HOP-LIMIT", A, rid("i", 1)): "FAIL",
    ("R1-EC2-IMDS-HOP-LIMIT", A, rid("i", 2)): "PASS",
    ("R1-EC2-SG-SENSITIVE-INGRESS", A, rid("sg", 1)): "FAIL",
    ("R1-EC2-SG-SENSITIVE-INGRESS", A, rid("sg", 2)): "PASS",
    ("R1-EC2-SG-SENSITIVE-INGRESS", A, rid("sg", 3)): "PASS",
    ("R1-EC2-SG-SENSITIVE-INGRESS", A, rid("sg", 4)): "MANUAL_VALIDATION_REQUIRED",
    ("R1-EC2-SG-SENSITIVE-INGRESS", A, rid("sg", 5)): "MANUAL_VALIDATION_REQUIRED",
    ("R1-EC2-SG-SENSITIVE-INGRESS", A, rid("sg", 6)): "MANUAL_VALIDATION_REQUIRED",
    ("R1-EC2-SG-SENSITIVE-INGRESS", A, rid("sg", 7)): "MANUAL_VALIDATION_REQUIRED",
    ("R1-EC2-SG-SENSITIVE-INGRESS", B, rid("sg", 11)): "UNKNOWN",
    ("R1-IAM-ROLE-TRUST-CONDITIONS", A, "synthetic-vendor-root"): "FAIL",
    ("R1-IAM-ROLE-TRUST-CONDITIONS", A, "synthetic-public-trust"): "FAIL",
    ("R1-IAM-ROLE-TRUST-CONDITIONS", A, "synthetic-org-bound"): "PASS",
    ("R1-IAM-ROLE-TRUST-CONDITIONS", A, "synthetic-external-id"): "PASS",
    ("R1-IAM-ROLE-TRUST-CONDITIONS", A, "synthetic-encoded-bound"): "PASS",
    ("R1-IAM-ROLE-TRUST-CONDITIONS", A, "synthetic-ec2-service"): "NOT_APPLICABLE",
    ("R1-IAM-ROLE-TRUST-CONDITIONS", A, "synthetic-same-account"): "NOT_APPLICABLE",
    ("R1-IAM-ROLE-TRUST-CONDITIONS", A, "synthetic-specific-principal"): (
        "MANUAL_VALIDATION_REQUIRED"
    ),
    ("R1-IAM-ROLE-TRUST-CONDITIONS", A, "synthetic-mfa-only"): "MANUAL_VALIDATION_REQUIRED",
    ("R1-IAM-ROLE-TRUST-CONDITIONS", A, "synthetic-ifexists"): "MANUAL_VALIDATION_REQUIRED",
    ("R1-IAM-ROLE-TRUST-CONDITIONS", B, "<scope>"): "COLLECTION_ERROR",
    ("R1-S3-ACCOUNT-BPA", A, A): "PASS",
    ("R1-S3-ACCOUNT-BPA", B, B): "FAIL",
}


def test_end_to_end_truth_table_and_traceability(r1):
    store, run, _ = r1
    results = list(store.results(run))
    observed = {
        (r.rule_id, r.account_id, label(store, run, r)): str(r.outcome)
        for r in results
        if r.rule_id in {k[0] for k in EXPECTED}
        and (r.resource_id or r.rule_id == "R1-IAM-ROLE-TRUST-CONDITIONS")
    }
    assert {k: v for k, v in observed.items() if k in EXPECTED} == EXPECTED
    for result in results:
        rule = BY_ID[result.rule_id]
        if rule.rule_id.startswith("SEC-"):
            # Additional observations are classified under a checklist check area.
            assert result.methodology_id == rule.methodology.control_id
            assert all(":CA." in c for c in result.methodology_criteria)
            assert result.mapping_status == "PROPOSED"
            continue
        # Original methodology ID and criteria survive evaluation unchanged.
        assert result.methodology_id == rule.methodology.control_id
        assert result.methodology_criteria == rule.methodology.criteria
        assert result.methodology_sha256 == REVIEWED_METHODOLOGY_SHA256
        assert result.mapping_status == "PROPOSED"
        if result.outcome == "PASS":
            assert result.evidence and not result.missing_dependencies
    candidates = list(store.candidates(run))
    # 8 methodology candidates + 2 latent internet-ingress observations (SEC, 0.6.1).
    assert len(candidates) == sum(r.outcome == "FAIL" for r in results) == 10
    for finding in candidates:
        link = BY_ID[finding.rule_id].methodology
        if finding.rule_id.startswith("SEC-"):
            # Additional observations carry a proposed check-area checklist mapping.
            assert finding.methodology_id == link.control_id and not link.full_coverage
            assert finding.evidence and finding.review_status == "CANDIDATE"
            continue
        assert finding.methodology_id == link.control_id
        assert finding.attack_paths == link.attack_paths
        assert finding.methodology_inherent_severity == "Critical"
        assert finding.suggested_severity in {"LOW", "MEDIUM", "HIGH"}  # independent field
        assert finding.evidence and finding.review_status == "CANDIDATE"
    exposure = next(f for f in candidates if f.rule_id == "R1-EC2-SG-SENSITIVE-INGRESS")
    # Relationship evidence: the security group and the public interface are both cited.
    assert len({ref.artifact_id for ref in exposure.evidence}) == 2


def test_latent_exposure_passes_only_with_complete_interface_evidence(r1):
    store, run, _ = r1
    latent = next(
        r
        for r in store.results(run)
        if r.rule_id == "R1-EC2-SG-SENSITIVE-INGRESS" and r.native_id == rid("sg", 3)
    )
    assert latent.outcome == "PASS"
    assert any(item.startswith("Latent exposure") for item in latent.limitations)
    missing = next(
        r
        for r in store.results(run)
        if r.rule_id == "R1-EC2-SG-SENSITIVE-INGRESS" and r.account_id == B
    )
    assert missing.missing_dependencies == (
        f"complete ec2:DescribeNetworkInterfaces for {B}/eu-west-1",
    )
    assert not any(
        c.account_id == B and c.rule_id == missing.rule_id for c in store.candidates(run)
    )


def test_global_scope_and_family_gating(r1):
    store, run, _ = r1
    results = list(store.results(run))
    iam_scope = [
        r
        for r in results
        if r.rule_id == "R1-IAM-ROLE-TRUST-CONDITIONS" and r.evaluation_scope == "COLLECTION_SCOPE"
    ]
    assert [(r.account_id, r.region, str(r.outcome)) for r in iam_scope] == [
        (B, "global", "COLLECTION_ERROR")
    ]
    # "global" is a resource label, never a Regional collection scope for Regional rules.
    assert not any(r.region == "global" for r in results if BY_ID[r.rule_id].scope == "REGIONAL")
    assert summary(store, run)["unique_regions"] == 2


def test_r0_package_does_not_report_unsupplied_r1_families(tmp_path):
    from tools.fixtures.demo import build_package

    directory = tmp_path / "input"
    build_package(directory)
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, directory)
        assess(store, run)
        rules = {r.rule_id for r in store.results(run)}
        assert not rules & {
            "R1-EC2-SG-SENSITIVE-INGRESS",
            "R1-IAM-ROLE-TRUST-CONDITIONS",
            "R1-S3-ACCOUNT-BPA",
        }
        output = tmp_path / "r0.xlsx"
        export_excel(store, run, output)
        wb = load_workbook(output)
        status = {
            row[0]: row[8]
            for row in wb["Methodology Coverage"].iter_rows(min_row=4, values_only=True)
        }
        assert status["R1-IAM-ROLE-TRUST-CONDITIONS"].startswith("NOT_SUPPLIED")
        assert status["R0-EC2-IMDSV2"] == "SUPPLIED"
        wb.close()


def obs(account, config, region="us-east-1", native="x"):
    return SimpleNamespace(
        account_id=account, region=region, configuration=config, native_id=native
    )


def statement(principal, condition=None, effect="Allow", action="sts:AssumeRole", **extra):
    item = {"Effect": effect, "Principal": principal, "Action": action, **extra}
    if condition is not None:
        item["Condition"] = condition
    return item


VENDOR_ROOT = {"AWS": "arn:aws:iam::333333333333:root"}


@pytest.mark.parametrize(
    "statements,expected",
    [
        ([statement(VENDOR_ROOT)], "FAIL"),
        ([statement(VENDOR_ROOT, {})], "FAIL"),  # empty Condition is no condition
        ([statement({"AWS": ["333333333333", f"arn:aws:iam::{A}:root"]})], "FAIL"),
        ([statement(VENDOR_ROOT, action="sts:*")], "FAIL"),
        ([statement(VENDOR_ROOT, action=["sts:AssumeRole", "sts:TagSession"])], "FAIL"),
        (
            [
                statement(VENDOR_ROOT, {"StringEquals": {"sts:ExternalId": "x"}}),
                statement({"AWS": "arn:aws:iam::444444444444:root"}),
            ],
            "FAIL",
        ),
        ([statement(VENDOR_ROOT, {"StringEquals": {"STS:EXTERNALID": "x"}})], "PASS"),
        ([statement(VENDOR_ROOT, {"StringLike": {"sts:ExternalId": "exact"}})], "PASS"),
        ([statement("*", {"StringEquals": {"aws:PrincipalOrgID": "o-synthetic"}})], "PASS"),
        (
            [statement(VENDOR_ROOT, {"StringLike": {"sts:ExternalId": "*"}})],
            "MANUAL_VALIDATION_REQUIRED",
        ),
        (
            [statement(VENDOR_ROOT, {"ForAnyValue:StringEquals": {"sts:ExternalId": "x"}})],
            "MANUAL_VALIDATION_REQUIRED",
        ),
        (
            [statement(VENDOR_ROOT, {"Null": {"sts:ExternalId": "false"}})],
            "MANUAL_VALIDATION_REQUIRED",
        ),
        (
            [statement(VENDOR_ROOT, {"StringEquals": {"aws:PrincipalOrgPaths": "o-x/r-y/"}})],
            "MANUAL_VALIDATION_REQUIRED",
        ),
        (
            [statement({"AWS": "AROAEXAMPLEUNIQUEID01"})],
            "MANUAL_VALIDATION_REQUIRED",
        ),
        (
            [statement(VENDOR_ROOT), statement(VENDOR_ROOT, effect="Deny")],
            "MANUAL_VALIDATION_REQUIRED",
        ),
        (
            [{"Effect": "Allow", "NotPrincipal": VENDOR_ROOT, "Action": "sts:AssumeRole"}],
            "MANUAL_VALIDATION_REQUIRED",
        ),
        (
            [statement({"Federated": "arn:aws:iam::111111111111:oidc-provider/x"},
                       action="sts:AssumeRoleWithWebIdentity")],
            "NOT_APPLICABLE",
        ),
        ([statement({"Service": "lambda.amazonaws.com"})], "NOT_APPLICABLE"),
        ([statement(VENDOR_ROOT, effect="allow")], "UNKNOWN"),
        ([statement(VENDOR_ROOT, action=[])], "UNKNOWN"),
        ([statement({"Unknown": "x"})], "UNKNOWN"),
        ([statement(VENDOR_ROOT, {"StringEquals": "not-a-map"})], "UNKNOWN"),
    ],
)  # fmt: skip
def test_trust_condition_truth_table(statements, expected):
    document = {"Version": "2012-10-17", "Statement": statements}
    assert trust_conditions(obs(A, {"AssumeRolePolicyDocument": document})).outcome == expected


@pytest.mark.parametrize("document", [None, {}, {"Statement": []}, {"Statement": ["x"]}])
def test_unparseable_trust_document_is_unknown(document):
    verdict = trust_conditions(obs(A, {"AssumeRolePolicyDocument": document}))
    assert verdict.outcome == "UNKNOWN" and verdict.missing


class FakeContext:
    def __init__(self, interfaces=(), coverage="COMPLETE", route="INTERNET_ROUTE", groups=()):
        self.interfaces = list(interfaces)
        self.state = coverage
        self.route = route
        self.groups = list(groups)
        self.registry = None

    def coverage(self, operation, account, region):
        return self.state

    def interfaces_for_group(self, region, group_id):
        return self.interfaces

    def route_status(self, account, region, vpc, subnet, ipv6=False):
        return self.route, None

    def security_groups(self, region, group_id):
        return [g for g in self.groups if g.configuration["GroupId"] == group_id]


def eni(public=True, status="in-use", managed=False, ipv6=0, n=1):
    return SimpleNamespace(
        observation_id=f"observation-{n}",
        native_id=rid("eni", n),
        account_id=A,
        region="us-east-1",
        configuration={
            "Status": status,
            "PublicIp": "203.0.113.5" if public else None,
            "AddressPublicIps": [],
            "RequesterManaged": managed,
            "InterfaceType": "interface",
            "Ipv6AddressCount": ipv6,
        },
    )


def rule(protocol="tcp", low=22, high=22, cidrs=("0.0.0.0/0",), v6=(), prefix=()):
    return {
        "IpProtocol": protocol,
        "FromPort": low,
        "ToPort": high,
        "IpRanges": [{"CidrIp": c} for c in cidrs],
        "Ipv6Ranges": [{"CidrIpv6": c} for c in v6],
        "PrefixListIds": [{"PrefixListId": p} for p in prefix],
    }


@pytest.mark.parametrize(
    "permissions,context,expected",
    [
        ([rule()], FakeContext([eni()]), "FAIL"),
        ([rule("-1", -1, -1)], FakeContext([eni()]), "FAIL"),
        ([rule("tcp", 20, 30)], FakeContext([eni()]), "FAIL"),
        ([rule("6", 3389, 3389)], FakeContext([eni()]), "FAIL"),
        ([rule()], FakeContext([eni()], "PARTIAL"), "FAIL"),  # found evidence suffices
        ([rule()], FakeContext([eni(public=False)]), "PASS"),  # latent, complete evidence
        ([rule()], FakeContext([eni(status="available")]), "PASS"),
        ([rule()], FakeContext([], "PARTIAL"), "UNKNOWN"),
        ([rule()], FakeContext([], "NOT_COLLECTED"), "UNKNOWN"),
        ([rule()], FakeContext([eni(status=None)]), "UNKNOWN"),
        ([rule()], FakeContext([eni(managed=True)]), "MANUAL_VALIDATION_REQUIRED"),
        ([rule("icmp", -1, -1)], FakeContext(), "PASS"),
        ([rule(low=443, high=443), rule(low=80, high=80)], FakeContext(), "PASS"),
        ([rule(low=22, high=22, cidrs=("10.1.0.0/16",))], FakeContext(), "PASS"),
        ([rule(low=80, high=443)], FakeContext(), "MANUAL_VALIDATION_REQUIRED"),
        ([rule("udp", 3389, 3389)], FakeContext(), "MANUAL_VALIDATION_REQUIRED"),
        ([rule(cidrs=("198.51.100.0/24",))], FakeContext(), "MANUAL_VALIDATION_REQUIRED"),
        ([rule(cidrs=(), prefix=("pl-1",))], FakeContext(), "MANUAL_VALIDATION_REQUIRED"),
        ([rule(cidrs=(), v6=("::/0",))], FakeContext([eni(ipv6=1)]), "MANUAL_VALIDATION_REQUIRED"),
        ([rule(cidrs=(), v6=("::/0",))], FakeContext([]), "PASS"),
        ([rule(cidrs=("not-a-cidr",))], FakeContext(), "UNKNOWN"),
        ([rule("tcp", None, None)], FakeContext(), "UNKNOWN"),
        ([rule("tcp", 30, 20)], FakeContext(), "UNKNOWN"),
    ],
)
def test_security_group_exposure_truth_table(permissions, context, expected):
    config = {"GroupId": rid("sg", 1), "IpPermissions": permissions}
    assert sensitive_ingress(obs(A, config), context).outcome == expected


@pytest.mark.parametrize(
    "config,expected",
    [
        ({"configured": False}, "FAIL"),
        ({"configured": True, **dict.fromkeys(("BlockPublicAcls", "IgnorePublicAcls",
          "BlockPublicPolicy", "RestrictPublicBuckets"), True)}, "PASS"),
        ({"configured": True, "BlockPublicAcls": True, "IgnorePublicAcls": True,
          "BlockPublicPolicy": False, "RestrictPublicBuckets": True}, "FAIL"),
        ({"configured": True, "BlockPublicAcls": "true", "IgnorePublicAcls": True,
          "BlockPublicPolicy": True, "RestrictPublicBuckets": True}, "UNKNOWN"),
        ({}, "UNKNOWN"),
    ],
)  # fmt: skip
def test_account_bpa_truth_table(config, expected):
    assert account_bpa(obs(A, config)).outcome == expected


@pytest.mark.parametrize(
    "engine,public,expected",
    [
        ("postgres", True, "FAIL"),
        ("aurora-mysql", False, "PASS"),
        ("docdb", True, "NOT_APPLICABLE"),
        ("custom-oracle-ee", True, "MANUAL_VALIDATION_REQUIRED"),
        ("mysql", "true", "UNKNOWN"),
        (None, True, "UNKNOWN"),
    ],
)
def test_rds_public_truth_table(engine, public, expected):
    config = {"Engine": engine, "PubliclyAccessible": public}
    assert rds_public(obs(A, config), FakeContext()).outcome == expected


@pytest.mark.parametrize(
    "limit,endpoint,expected",
    [(1, "enabled", "PASS"), (2, "enabled", "FAIL"), (None, "enabled", "UNKNOWN"),
     (True, "enabled", "UNKNOWN"), (0, "enabled", "UNKNOWN"), (3, "disabled", "NOT_APPLICABLE")],
)  # fmt: skip
def test_hop_limit_truth_table(limit, endpoint, expected):
    options = {"State": "applied", "HttpEndpoint": endpoint, "HttpTokens": "required"}
    if limit is not None:
        options["HttpPutResponseHopLimit"] = limit
    config = {"State": {"Name": "running"}, "MetadataOptions": options}
    assert hop_limit(obs(A, config)).outcome == expected


def write_source(directory, name, body, operation, account=A, region="us-east-1", **extra):
    path = directory / name
    path.write_text(json.dumps(body))
    return Source(
        path=name,
        sha256=sha256_file(path),
        format="aws-json-v1",
        operation=operation,
        account_id=account,
        region=region,
        collected_at=extra.pop("collected_at", TIME),
        coverage=extra.pop("coverage", "COMPLETE"),
        request_parameters=extra.pop("request_parameters", {}),
        **extra,
    )


def group_body(port):
    return {
        "SecurityGroups": [
            {
                "GroupId": rid("sg", 1),
                "OwnerId": A,
                "IpPermissions": [rule(low=port, high=port)],
            }
        ]
    }


def test_contradictory_snapshots_become_unknown_without_candidate(tmp_path):
    first = write_source(tmp_path, "sg1.json", group_body(22), "ec2:DescribeSecurityGroups")
    second = write_source(
        tmp_path,
        "sg2.json",
        group_body(443),
        "ec2:DescribeSecurityGroups",
        collected_at="2026-10-09T09:00:00+00:00",
    )
    public = {
        "NetworkInterfaceId": rid("eni", 1),
        "OwnerId": A,
        "Status": "in-use",
        "InterfaceType": "interface",
        "RequesterManaged": False,
        "Groups": [{"GroupId": rid("sg", 1)}],
        "Association": {"PublicIp": "203.0.113.7"},
    }
    interfaces = write_source(
        tmp_path, "eni.json", {"NetworkInterfaces": [public]}, "ec2:DescribeNetworkInterfaces"
    )
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, tmp_path, Package(sources=(first, second, interfaces)))
        assess(store, run)
        # Snapshot 1 (port 22 open on a public interface) FAILs; snapshot 2 PASSes.
        result = next(
            r
            for r in store.results(run)
            if r.resource_id and r.rule_id == "R1-EC2-SG-SENSITIVE-INGRESS"
        )
        assert result.outcome == "UNKNOWN"
        assert len(result.contradictory_observations) == 2
        assert result.confidence == "UNASSESSED"
        assert not list(store.candidates(run))


def test_unreadable_interface_groups_block_latent_pass(tmp_path):
    groups = write_source(tmp_path, "sg.json", group_body(22), "ec2:DescribeSecurityGroups")
    interfaces = write_source(
        tmp_path,
        "eni.json",
        {"NetworkInterfaces": [{"NetworkInterfaceId": rid("eni", 1), "OwnerId": A}]},
        "ec2:DescribeNetworkInterfaces",
    )
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, tmp_path, Package(sources=(groups, interfaces)))
        assess(store, run)
        assert (
            scope_coverage(
                list(store.source_data(run)), "ec2:DescribeNetworkInterfaces", A, "us-east-1"
            )
            == "PARTIAL"
        )
        result = next(r for r in store.results(run) if r.rule_id == "R1-EC2-SG-SENSITIVE-INGRESS")
        assert result.outcome == "UNKNOWN"


def test_list_roles_encoding_and_duplicate_keys(tmp_path):
    from urllib.parse import quote

    good = quote(
        json.dumps(
            {
                "Statement": [
                    {
                        "Effect": "Allow",
                        "Principal": {"AWS": "333333333333"},
                        "Action": "sts:AssumeRole",
                    }
                ]
            }
        )
    )
    duplicate = '{"Statement": [], "Statement": []}'
    body = {
        "Roles": [
            {"RoleName": "good", "Arn": f"arn:aws:iam::{A}:role/good",
             "AssumeRolePolicyDocument": good},
            {"RoleName": "dup", "Arn": f"arn:aws:iam::{A}:role/dup",
             "AssumeRolePolicyDocument": duplicate},
        ],
        "IsTruncated": False,
    }  # fmt: skip
    source = write_source(tmp_path, "roles.json", body, "iam:ListRoles", region=None)
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, tmp_path, Package(sources=(source,)))
        assess(store, run)
        by_name = {
            label(store, run, r): r
            for r in store.results(run)
            if r.rule_id == "R1-IAM-ROLE-TRUST-CONDITIONS" and r.resource_id
        }
        assert by_name["good"].outcome == "FAIL"
        assert by_name["dup"].outcome == "UNKNOWN"
        assert by_name["good"].region == "global"
        # The malformed document downgrades declared completeness.
        assert next(store.source_data(run))["effective_coverage"] == "PARTIAL"
        # A region-less global collection is assigned, not reported as an unscoped source.
        assert not any(r.evaluation_id.startswith("unassigned-source") for r in store.results(run))


def test_credential_material_in_iam_export_is_rejected(tmp_path):
    body = {
        "RoleDetailList": [{"Arn": f"arn:aws:iam::{A}:role/x", "Description": "AKIA" + "A" * 16}]
    }
    source = write_source(tmp_path, "gaad.json", body, "iam:GetAccountAuthorizationDetails")
    with Store(tmp_path / "workspace") as store:
        with pytest.raises(InputError, match="CREDENTIAL_MATERIAL_REJECTED"):
            import_package(store, tmp_path, Package(sources=(source,)))


@pytest.mark.parametrize(
    "body,errors,expected",
    [
        ({"Error": {"Code": "AccessDenied"}, "ResponseMetadata": {"HTTPStatusCode": 403}}, (),
         "COLLECTION_ERROR"),
        ({"Error": {"Code": "NoSuchPublicAccessBlockConfiguration"}}, ("Throttled",),
         "COLLECTION_ERROR"),
        ({"Error": {"Code": "NoSuchPublicAccessBlockConfiguration"},
          "ResponseMetadata": {"HTTPStatusCode": 500}}, (), "COLLECTION_ERROR"),
    ],
)  # fmt: skip
def test_bpa_absence_requires_clean_defined_response(tmp_path, body, errors, expected):
    source = write_source(
        tmp_path,
        "bpa.json",
        body,
        "s3control:GetPublicAccessBlock",
        coverage="PARTIAL" if errors else "COMPLETE",
        errors=errors,
    )
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, tmp_path, Package(sources=(source,)))
        assess(store, run)
        outcomes = {str(r.outcome) for r in store.results(run) if r.rule_id == "R1-S3-ACCOUNT-BPA"}
        assert outcomes == {expected}
        assert not list(store.candidates(run))


def test_specification_integrity():
    ids = [entry.control_id for entry in COVERAGE]
    assert len(ids) == len(set(ids)) == 204
    assert {f"AWS-{n:02}" for n in range(1, 29)} <= set(ids)
    assert all(re.fullmatch(r"\d+\.\d+\.\d+|AWS-\d\d", i) for i in ids)
    for entry in COVERAGE:
        assert entry.status in STATUSES
        assert set(entry.reasons) <= set(REASONS)
        assert set(entry.rules) <= set(BY_ID)
        assert set(entry.related) <= set(ids)
    for rule_ in ALL_RULES:
        from offline_analyzer.rules.security_observations import ADDITIONAL_RULES

        link = rule_.methodology
        if rule_ in ADDITIONAL_RULES:
            # Check-area classification: no source P/F clause is claimed.
            assert link and link.control_id in BY_CONTROL and not link.full_coverage
            assert all(c.startswith(link.control_id + ":CA.") for c in link.criteria)
            assert rule_.rule_id in BY_CONTROL[link.control_id].rules
            continue
        assert link and link.control_id in BY_CONTROL
        assert rule_.rule_id in BY_CONTROL[link.control_id].rules
        assert not set(link.criteria) & set(link.not_evaluated)
        assert all(c.startswith(link.control_id + ":") for c in link.criteria + link.not_evaluated)
        assert not link.full_coverage  # no control is claimed fully automated
        assert set(link.attack_paths) <= set(BY_PATH)
    # Controls a rule maps to are never claimed verified on synthetic evidence alone.
    assert not any(e.status == "IMPLEMENTED_AND_VERIFIED" for e in COVERAGE)


def test_attack_path_spec_integrity():
    assert [p.path_id for p in PATHS] == [f"AP-{n:02}" for n in range(1, 31)]
    keys = set()
    for path in PATHS:
        assert set(path.threats) <= set(THREAT_SCOPE)
        for prerequisite in path.prerequisites:
            assert prerequisite.key.startswith(path.path_id + ".P")
            assert prerequisite.kind in KINDS
            assert set(prerequisite.rules) <= set(BY_ID)
            keys.add(prerequisite.key)
        if path.scope == "OUTSIDE_CURRENT_AWS_SCOPE":
            assert not path.prerequisites
    assert sorted(THREAT_SCOPE) == [f"TS-{n:02}" for n in range(1, 21)]
    for _, prerequisite_keys in THREAT_SCOPE.values():
        assert set(prerequisite_keys) <= keys


def test_attack_path_records_never_claim_support(r1):
    store, run, _ = r1
    records = path_records(store, run)
    assert {r.state for r in records} <= {"UNKNOWN", "BLOCKED"}
    exposed = next(r for r in records if r.path_id == "AP-04")
    assert exposed.anchor == rid("i", 1)
    assert exposed.prerequisites["AP-04.P1"] == "VERIFIED"
    assert exposed.prerequisites["AP-04.P2"] == "VERIFIED"
    assert exposed.prerequisites["AP-04.P3"] == "NOT_EVALUATED"
    assert "AP-04.P3 (EXPLOITATION)" in exposed.missing
    storage = {r.account_id: r for r in records if r.path_id == "AP-05"}
    assert storage[A].state == "BLOCKED" and storage[A].blockers
    assert storage[B].state == "UNKNOWN" and storage[B].prerequisites["AP-05.P1"] == "VERIFIED"
    assert len(records) == 3


def test_tracker_staging_reconciles_and_never_writes_analyst_columns(r1, tmp_path):
    store, run, directory = r1
    keys = load_account_keys(directory / "account-keys.json")
    assert keys == ACCOUNT_KEYS
    rows = tracker_rows(store, run, keys)
    status = {(row[15], row[2]): row[3] for row in rows}
    assert {k for k, v in status.items() if v == "Fail"} == {
        (A, "AWS-04"),
        (A, "AWS-15"),
        (A, "AWS-19"),
        (A, "AWS-22"),
        (B, "AWS-11"),
        # Latent internet ingress (not on a public ENI) meets 1.5.2, not AWS-15.
        (A, "1.5.2"),
        (B, "1.5.2"),
    }
    # Partial criteria coverage can never become a control-level Pass.
    assert "Pass" not in status.values()
    assert status[(A, "AWS-11")] == "Not Assessed"
    analyst = [
        TRACKER_COLUMNS.index(c)
        for c in TRACKER_COLUMNS
        if not c.startswith(
            (
                "Result ID",
                "Account ID",
                "Control ID",
                "Scanner_",
                "Scan_Date",
                "Evidence Ref",
                "Resource ARN",
            )
        )
    ]
    assert all(row[i] == "" for row in rows for i in analyst)
    # Evidence and resource columns are filled only for failing controls.
    evidence = TRACKER_COLUMNS.index("Evidence Ref")
    resource = TRACKER_COLUMNS.index("Resource ARN / ID")
    assert all(bool(row[evidence]) == (row[3] == "Fail") for row in rows)
    assert all(bool(row[resource]) == (row[3] == "Fail") for row in rows)
    assert all((row[1] == "ACCT-001") == (row[15] == A) for row in rows)
    assert {row[16] for row in rows if row[15] == B} == {"ACCOUNT_KEY_MAPPING_REQUIRED"}
    output = tmp_path / "report.xlsx"
    export_excel(store, run, output, keys)
    wb = load_workbook(output)
    sheet_rows = list(wb["Tracker Staging"].iter_rows(min_row=4, values_only=True))
    assert len(sheet_rows) == len(rows)
    assert Counter(r[3] for r in sheet_rows) == Counter(r[3] for r in rows)
    findings = list(wb["Findings Register"].iter_rows(min_row=4, values_only=True))
    assert len(findings) == 10
    header = [c.value for c in wb["Findings Register"][3]]
    control = header.index("Methodology control")
    # SEC observations are classified under the checklist (latent ingress -> 1.5.2).
    assert {row[control] for row in findings} == {
        "AWS-04",
        "AWS-11",
        "AWS-15",
        "AWS-19",
        "AWS-22",
        "1.5.2",
    }
    assert all(cell.data_type != "f" for sheet in wb for row in sheet for cell in row)
    wb.close()


@pytest.mark.parametrize(
    "content,code",
    [
        ('{"111111111111": "ACCT-1"}', "INVALID_ACCOUNT_KEY_FILE"),
        ('{"1111": "ACCT-001"}', "INVALID_ACCOUNT_KEY_FILE"),
        ('{"111111111111": "ACCT-001", "222222222222": "ACCT-001"}',
         "DUPLICATE_TRACKER_ACCOUNT_KEY"),
        ('{"111111111111": "ACCT-001", "111111111111": "ACCT-002"}', "DUPLICATE_JSON_KEY"),
        ('["ACCT-001"]', "INVALID_ACCOUNT_KEY_FILE"),
    ],
)  # fmt: skip
def test_account_key_file_validation(tmp_path, content, code):
    path = tmp_path / "keys.json"
    path.write_text(content)
    with pytest.raises(InputError, match=code):
        load_account_keys(path)


def test_generated_fixture_preserves_historical_bytes_and_cli(tmp_path, capsys):
    from fixture_fingerprints import package_fingerprint

    generated = tmp_path / "generated"
    build_r1_package(generated)
    assert len(list(generated.iterdir())) == 10
    assert package_fingerprint(generated) == (
        "61308eb6e5ddd5b14f84506afc9c05f606acdde888f9a31c03e60ca49e44e118"
    )
    workspace = tmp_path / "workspace"
    assert main(["--workspace", str(workspace), "import", str(generated)]) == 0
    run = capsys.readouterr().out.strip()
    assert main(["--workspace", str(workspace), "assess", "--run", run]) == 0
    assert json.loads(capsys.readouterr().out)["candidate_findings"] == 10
    report = tmp_path / "cli.xlsx"
    keys = generated / "account-keys.json"
    args = ["--workspace", str(workspace), "export", "--run", run, "--output", str(report)]
    assert main(args + ["--account-keys", str(keys)]) == 0
    bad = tmp_path / "bad.json"
    bad.write_text('{"x": "y"}')
    assert main(args[:-1] + [str(tmp_path / "other.xlsx"), "--account-keys", str(bad)]) == 1
    assert "INVALID_ACCOUNT_KEY_FILE" in capsys.readouterr().err
