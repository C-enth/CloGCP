"""Standalone finalizer safety, exact schema, review preservation and focused evidence."""

import json
import subprocess
import sys
import zipfile

import pytest
from openpyxl import load_workbook

from clogcp_report_finalizer.core import finalize, verify_output
from clogcp_report_finalizer.evidence import exported_evidence, focus
from clogcp_report_finalizer.narratives import verification
from clogcp_report_finalizer.safety import FinalizerError, digest, parse_json, workbook
from clogcp_report_finalizer.schema import CLIENT_REGISTER, COLUMNS, REGISTER
from tools.fixtures.finalizer_workbook import build


@pytest.fixture
def sample(tmp_path):
    return build(tmp_path / "input")


def test_exact_40_columns_all_source_rows_protected_values_and_navigation(sample, tmp_path):
    source = sample["source"]
    before = digest(source)
    result = finalize(
        source,
        tmp_path / "client",
        guidance=sample["guidance"],
        attack_paths=[sample["attack_paths"]],
        history=[sample["history"]],
        origin="SYNTHETIC",
    )
    target = tmp_path / "client" / "CloGCP-client-report.xlsx"
    assert digest(source) == before
    original, output = load_workbook(source), load_workbook(target)
    try:
        assert (
            original[REGISTER].max_row
            == output[REGISTER].max_row
            == output[CLIENT_REGISTER].max_row
        )
        assert tuple(c.value for c in output[CLIENT_REGISTER][1]) == COLUMNS
        for a, b, c in zip(
            original[REGISTER], output[REGISTER], output[CLIENT_REGISTER], strict=True
        ):
            assert [x.value for x in a] == [x.value for x in b]
            for n, (x, y) in enumerate(zip(a, c, strict=True)):
                if n != 10:
                    assert x.value == y.value
        assert output[CLIENT_REGISTER].freeze_panes == "C2"
        assert output[CLIENT_REGISTER].auto_filter.ref
        assert output[CLIENT_REGISTER].tables
        assert output[CLIENT_REGISTER].print_title_rows == "$1:$1"
        assert output[CLIENT_REGISTER].print_title_cols == "$A:$B"
        text = output[CLIENT_REGISTER].cell(2, 11).value
        assert "unrelated-" not in text
        assert "serviceAccountUser" in text and "serviceAccountTokenCreator" in text
        assert "denyRules" in text
        assert all(len(str(c.value or "")) <= 32767 for s in output for row in s for c in row)
        assert all(c.data_type != "f" for s in output for row in s for c in row)
        assert output["Source_Reconciliation"].max_row == original[REGISTER].max_row
        assert result["client_observations"] == 14
        assert result["qa_issues"] > 0
        assert result["equivalent_source_duplicates"] == 1
        assert (tmp_path / "client" / "Assessor-QA-Exceptions.json").is_file()
        exceptions = json.loads((tmp_path / "client" / "Assessor-QA-Exceptions.json").read_text())[
            "exceptions"
        ]
        assert any(
            r[2] == "DUPLICATE_REPORTABLE_OBSERVATIONS" and r[3] == "EXACT_EQUIVALENCE_GROUPED"
            for r in exceptions
        )
        # Missing IAM remains visible and cannot prevent preparing the reporting copy.
        assert "Evidence detail unavailable" in " ".join(
            str(c.value) for row in output["Evidence_Limitations"] for c in row
        )
        verify_output(target)
    finally:
        original.close()
        output.close()


def test_qualifiers_modes_and_missing_context_are_not_lost(sample):
    book = workbook(sample["source"])
    try:
        evidence = exported_evidence([book])
        guides = json.loads(sample["guidance"].read_text())["records"]
        focused = [focus(record, entries, guides[record]) for record, entries in evidence.items()]
        text = json.dumps([f.entries for f in focused])
        for field in [
            "denialCondition",
            "priority",
            "targetTags",
            "INGRESS_TRAFFIC_INTERNAL_ONLY",
            "publicAccessPrevention",
            "enforced",
            "inherited",
            "autopilot",
            "workloadIdentityConfig",
            "auditConfigs",
            "central-sink",
        ]:
            assert field in text
        assert len([f for f in focused if not f.entries]) >= 1
        iam = next(iter(evidence))
        unavailable = focus(iam, evidence[iam], None)
        assert not unavailable.entries and unavailable.limitations
        whole = focus(iam, evidence[iam], {"selectors": {evidence[iam][0].mapping: [""]}})
        assert not whole.entries and "WHOLE_POLICY" in whole.limitations[0]
    finally:
        book.close()


def test_finalizer_records_material_correction_without_changing_security_fields(sample, tmp_path):
    data = json.loads(sample["guidance"].read_text())
    record, guide = next(iter(data["records"].items()))
    mapping = next(iter(guide["selectors"]))
    guide["assertions"] = [
        {"mapping": mapping, "pointer": "/bindings/400", "observed_value": "fabricated assertion"}
    ]
    sample["guidance"].write_text(json.dumps(data))
    finalize(sample["source"], tmp_path / "client", guidance=sample["guidance"], origin="SYNTHETIC")
    qa = json.loads((tmp_path / "client" / "Assessor-QA-Exceptions.json").read_text())
    assert any(
        r[1] == record and r[2] == "OBSERVATION_EVIDENCE_CONTRADICTION" for r in qa["exceptions"]
    )
    output = load_workbook(tmp_path / "client" / "CloGCP-client-report.xlsx")
    try:
        assert output[CLIENT_REGISTER].cell(2, 5).value == output[REGISTER].cell(2, 5).value
        assert output[CLIENT_REGISTER].cell(2, 40).value == output[REGISTER].cell(2, 40).value
    finally:
        output.close()


def test_untrusted_workbooks_formulas_and_external_links(sample, tmp_path):
    source = sample["source"]
    book = load_workbook(source)
    book[REGISTER].cell(2, 9).value = '=HYPERLINK("https://example.invalid", "test")'
    book[REGISTER].cell(3, 9).value = "\t=1+1"
    book.save(source)
    book.close()
    finalize(source, tmp_path / "safe", origin="SYNTHETIC")
    safe = load_workbook(tmp_path / "safe" / "CloGCP-client-report.xlsx")
    try:
        assert safe[REGISTER].cell(2, 9).data_type == "s"
        assert safe[REGISTER].cell(3, 9).value == "\t=1+1"
    finally:
        safe.close()
    book = load_workbook(source)
    book[REGISTER].cell(2, 1).hyperlink = "https://example.invalid/customer"
    book.save(source)
    book.close()
    with pytest.raises(FinalizerError, match="EXTERNAL_WORKBOOK_LINK"):
        finalize(source, tmp_path / "refused", origin="SYNTHETIC")
    assert not (tmp_path / "refused").exists()
    macro = tmp_path / "macro.xlsx"
    with zipfile.ZipFile(source) as incoming, zipfile.ZipFile(macro, "w") as outgoing:
        for name in incoming.namelist():
            outgoing.writestr(name, incoming.read(name))
        outgoing.writestr("xl/vbaProject.bin", b"synthetic")
    with pytest.raises(FinalizerError):
        workbook(macro)


@pytest.mark.parametrize(
    "text",
    [
        '{"a":1,"a":2}',
        '{"access_token":"synthetic-sensitive-value"}',
        '{"keyString":"synthetic-sensitive-value"}',
        '{"privateKeyData":"synthetic-sensitive-value"}',
        "[" * 45 + "0" + "]" * 45,
    ],
)
def test_bounded_json_rejects_duplicates_credentials_and_depth(text):
    with pytest.raises(FinalizerError):
        parse_json(text)


def test_final_release_requires_explicit_acceptance_but_review_copy_is_available(sample, tmp_path):
    with pytest.raises(FinalizerError, match="FINAL_RELEASE_REQUIRES_ENGINE_ACCEPTANCE"):
        finalize(sample["source"], tmp_path / "final", release=True, origin="SYNTHETIC")
    assert not (tmp_path / "final").exists()
    assert "required" in verification({"control_id": "unsupported-control"})
    assert "required" in verification(
        {
            "control_id": "GCP-STORAGE-PUBLIC",
            "parameters": {"project": "demo-alpha", "bucket": "$(unsafe)"},
        }
    )


def test_finalizer_runtime_is_independent_of_backend_and_uses_only_selected_files(sample, tmp_path):
    script = """
import builtins, socket, sys
original = builtins.__import__
def guarded(name, *args, **kwargs):
    if name.startswith(("offline_analyzer", "duckdb", "fastapi")):
        raise AssertionError("Forbidden analyzer dependency: " + name)
    return original(name, *args, **kwargs)
builtins.__import__ = guarded
def deny_network(*args, **kwargs):
    raise AssertionError("Network access")
socket.create_connection = deny_network
from pathlib import Path
from clogcp_report_finalizer.core import finalize
result = finalize(Path(sys.argv[1]), Path(sys.argv[2]), origin="SYNTHETIC")
assert result["client_observations"] > 0
"""
    result = subprocess.run(
        [sys.executable, "-c", script, str(sample["source"]), str(tmp_path / "independent")],
        capture_output=True,
        text=True,
    )
    assert result.returncode == 0, result.stderr


def test_evidence_pointers_are_exact_rfc6901_paths():
    from clogcp_report_finalizer.safety import pointer

    document = {"a/b": [{"x~y": "synthetic"}]}
    assert pointer(document, "/a~1b/0/x~0y") == "synthetic"
    for path in ("/a~1b/-1", "/a~1b/00", "/a~2b/0"):
        with pytest.raises(FinalizerError):
            pointer(document, path)
