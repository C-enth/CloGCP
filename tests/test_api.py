"""Synthetic functional local API and access-control regression tests."""

import json
import time

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from offline_analyzer.api.app import create_app
from offline_analyzer.api.services import redact
from offline_analyzer.launcher import WorkspaceLock
from tools.fixtures.r12_demo import build

ORIGIN = "http://127.0.0.1:8765"
TOKEN = "synthetic-launch-secret"


def authenticated(client):
    response = client.post(
        "/api/session", json={}, headers={"Origin": ORIGIN, "X-Bootstrap-Token": TOKEN}
    )
    assert response.status_code == 200
    return {"Origin": ORIGIN, "X-CSRF-Token": response.json()["csrf"]}


def test_browser_context_identity_is_stable_private_and_workspace_scoped(tmp_path):
    identities = []
    for name in ["first", "first", "second"]:
        workspace = tmp_path / name
        with TestClient(
            create_app(workspace, origin=ORIGIN, bootstrap=TOKEN), base_url=ORIGIN
        ) as client:
            authenticated(client)
            response = client.get("/api/environment")
            identity = response.json()["workspace_id"]
            assert len(identity) == 64 and all(c in "0123456789abcdef" for c in identity)
            assert str(workspace) not in response.text
            identities.append(identity)
    assert identities[0] == identities[1] != identities[2]


def await_job(client, response):
    assert response.status_code == 200, response.text
    job = response.json()
    for _ in range(200):
        if job["status"] not in {"QUEUED", "RUNNING"}:
            break
        time.sleep(0.03)
        job = client.get("/api/jobs/" + job["id"]).json()
    assert job["status"] == "SUCCEEDED", job
    return job["result"]


def upload_package(client, headers, directory):
    names = sorted(p.name for p in directory.iterdir())
    files = [
        ("files", (name, (directory / name).read_bytes(), "application/json")) for name in names
    ]
    return client.post(
        "/api/uploads",
        headers=headers,
        files=files,
        data={"paths": json.dumps(names), "context": '{"origin":"synthetic"}'},
    )


@pytest.fixture(scope="module")
def sample(tmp_path_factory):
    root = tmp_path_factory.mktemp("api")
    build(root / "package")
    app = create_app(root / "work", origin=ORIGIN, bootstrap=TOKEN)
    with TestClient(app, base_url=ORIGIN) as client:
        headers = authenticated(client)
        preview = upload_package(client, headers, root / "package")
        assert preview.status_code == 200, preview.text
        assert preview.json()["can_import"]
        assert preview.json()["diagnostics"]["real_export_validation"] == "NOT_ESTABLISHED"
        upload_id = preview.json()["upload_id"]
        run = await_job(
            client, client.post("/api/import", headers=headers, json={"upload_id": upload_id})
        )["run_id"]
        await_job(client, client.post(f"/api/runs/{run}/assess", headers=headers, json={}))
        yield root, client, headers, run, upload_id


def test_dashboard_and_account_filtering(sample):
    _, client, _, run, _ = sample
    dashboard = client.get(f"/api/runs/{run}/dashboard").json()
    assert dashboard["unique_accounts"] == 5
    assert dashboard["candidate_findings"] == 8
    assert dashboard["control_evaluation_records"] == 97
    assert sum(dashboard["severity"].values()) == 8
    assert dashboard["iam_investigations"] == 52
    accounts = client.get(f"/api/runs/{run}/records/accounts?q=111111111111").json()
    assert accounts["total"] == 1
    findings = client.get(f"/api/runs/{run}/records/findings?account=111111111111").json()
    assert findings["items"] and all(f["account_id"] == "111111111111" for f in findings["items"])
    assert client.get(f"/api/runs/{run}/records/evidence?limit=101").status_code == 400


def test_finding_facets_drive_filters_and_account_view(sample):
    _, client, _, run, _ = sample
    facets = client.get(f"/api/runs/{run}/finding-facets").json()
    assert sum(f["count"] for f in facets["severity"]) == 8
    assert sum(a["total"] for a in facets["accounts"]) == 8
    assert all(sum(a["severity"].values()) == a["total"] for a in facets["accounts"])
    # Every offered value selects the findings it counts.
    for key, param in (("domain", "domain"), ("control", "control"), ("region", "region")):
        for facet in facets[key]:
            page = client.get(
                f"/api/runs/{run}/records/findings", params={param: facet["value"], "limit": 100}
            ).json()
            assert page["total"] == facet["count"], (key, facet)
    for account in facets["accounts"]:
        param = "project" if account["provider"] == "GCP" else "account"
        page = client.get(
            f"/api/runs/{run}/records/findings", params={param: account["scope"], "limit": 100}
        ).json()
        assert page["total"] == account["total"]
    assert client.get(f"/api/runs/{run}/finding-facets?provider=X").status_code == 400
    finding = client.get(f"/api/runs/{run}/records/findings").json()["items"][0]
    explanation = finding["explanation"]
    assert explanation["why_flagged"] == finding["technical_observation"]
    assert "not demonstrated" in explanation["misuse"]
    assert explanation["impact"] and explanation["recommendation"]


def test_dashboard_uses_bulk_review_decisions(sample, monkeypatch):
    root, _, _, run, _ = sample
    workspace = root / "work"
    from offline_analyzer.api.services import Views
    from offline_analyzer.storage.store import Store

    def fail(*args):
        raise AssertionError("Dashboard must not query each finding's decision")

    with Store(workspace) as store:
        monkeypatch.setattr(store, "latest_decision", fail)
        assert Views(workspace).dashboard(store, run)["candidate_findings"] == 8


def test_bulk_finding_page_preserves_evidence_times_and_relationship_scope(sample):
    root, _, _, run, _ = sample
    from datetime import datetime

    from offline_analyzer.api.services import Views
    from offline_analyzer.storage.store import Store

    with Store(root / "work") as store:
        page = Views(store.workspace).records(store, run, "findings", "", "", "", 0, 100)
        for finding in page["items"]:
            expected = [
                rid
                for (rid,) in store.rows(
                    "SELECT relationship_id FROM enterprise_edges WHERE run_id=? AND "
                    "EXISTS (SELECT 1 FROM json_each(data,'$.related_findings') f WHERE "
                    "json_extract_string(f.value,'$')=?) ORDER BY relationship_id LIMIT 100",
                    [run, finding["finding_id"]],
                )
            ]
            assert finding["relationships"] == expected
            for stamp in finding["source_timestamps"]:
                obs = next(store.observations(run, observation_id=stamp["observation_id"]))
                assert datetime.fromisoformat(stamp["collected_at"]) == obs.collected_at
                assert stamp["operation"] == obs.operation


def test_persisted_review_survives_reimport_and_restart(sample):
    root, client, headers, run, upload_id = sample
    finding = client.get(f"/api/runs/{run}/records/findings").json()["items"][0]
    fid = finding["finding_id"]
    response = client.post(
        f"/api/runs/{run}/review/{fid}",
        headers=headers,
        json={
            "status": "NEEDS_VALIDATION",
            "analyst": "synthetic-reviewer",
            "note": "Original evidence needed",
        },
    )
    assert response.status_code == 200
    assert (
        await_job(
            client, client.post("/api/import", headers=headers, json={"upload_id": upload_id})
        )["run_id"]
        == run
    )
    await_job(client, client.post(f"/api/runs/{run}/assess", headers=headers, json={}))
    found = next(
        f
        for f in client.get(f"/api/runs/{run}/records/findings").json()["items"]
        if f["finding_id"] == fid
    )
    assert found["analyst_status"] == "NEEDS_VALIDATION"
    app = create_app(root / "work", origin=ORIGIN, bootstrap=TOKEN)
    with TestClient(app, base_url=ORIGIN) as restarted:
        authenticated(restarted)
        record = next(
            f
            for f in restarted.get(f"/api/runs/{run}/records/findings").json()["items"]
            if f["finding_id"] == fid
        )
        assert record["decision"]["note"] == "Original evidence needed"


def test_exception_is_validation_not_confirmation(sample):
    _, client, headers, run, _ = sample
    fid = client.get(f"/api/runs/{run}/records/findings").json()["items"][1]["finding_id"]
    assert (
        client.post(
            f"/api/runs/{run}/review/{fid}",
            headers=headers,
            json={
                "status": "CONFIRMED",
                "exception": True,
                "analyst": "synthetic-reviewer",
                "note": "Time-limited risk exception",
            },
        ).status_code
        == 200
    )
    finding = next(
        f
        for f in client.get(f"/api/runs/{run}/records/findings").json()["items"]
        if f["finding_id"] == fid
    )
    assert finding["analyst_status"] == "NEEDS_VALIDATION"
    assert finding["decision"]["note"].startswith("EXCEPTION RECORDED")


def test_evidence_integrity_redaction_traceability(sample):
    _, client, _, run, _ = sample
    items = client.get(f"/api/runs/{run}/records/evidence?limit=100").json()["items"]
    assert items and "configuration" not in items[0]
    details = [
        client.get(f"/api/runs/{run}/evidence/{i['observation_id']}").json() for i in items[:5]
    ]
    assert all(all(r["integrity"] for r in d["provenance"]) for d in details)
    assert "raw source" in details[0]["inspection"]
    assert redact(
        {
            "AssumeRolePolicyDocument": {"Statement": "private"},
            "Password": "secret",
            "Tags": [{"Value": "private"}],
        }
    ) == {"AssumeRolePolicyDocument": "[redacted]", "Password": "[redacted]", "Tags": "[redacted]"}


def test_report_generation_reconciliation_and_register_exports(sample, tmp_path):
    _, client, headers, run, _ = sample
    for kind in ["excel", "findings", "accounts", "coverage", "evidence"]:
        report = await_job(
            client, client.post(f"/api/runs/{run}/reports", headers=headers, json={"kind": kind})
        )
        download = client.get(f"/api/reports/{report['id']}/download")
        assert download.status_code == 200
        if kind == "excel":
            path = tmp_path / "report.xlsx"
            path.write_bytes(download.content)
            workbook = load_workbook(path)
            assert workbook["Findings Register"].max_row - 3 == 8
            assert workbook["Control Assessment"].max_row - 3 == 97
            assert workbook["IAM Investigations"].max_row - 3 == 52
            assert all(c.data_type != "f" for s in workbook for row in s for c in row)
            workbook.close()
        else:
            assert isinstance(download.json(), list) and download.json()
    assert len(client.get("/api/reports").json()) >= 5


def test_invalid_upload_does_not_create_run(sample):
    root, client, headers, _, _ = sample
    before = client.get("/api/runs").json()
    staged_before = set((root / "work" / "uploads").iterdir())
    for name, content in [
        ("../escape.json", b"{}"),
        ("broken.json", b'{"broken":'),
        ("secret.json", b'{"Password":"synthetic"}'),
    ]:
        response = client.post(
            "/api/uploads",
            headers=headers,
            files=[("files", (name, content, "application/json"))],
            data={"paths": json.dumps([name]), "context": "{}"},
        )
        assert response.status_code == 400
        assert "synthetic" not in response.text and "broken.json" not in response.text
    assert client.get("/api/runs").json() == before
    assert set((root / "work" / "uploads").iterdir()) == staged_before
    assert not list((root / "work" / ".upload-staging").iterdir())


def test_single_file_context_and_detected_format(sample):
    _, client, headers, _, _ = sample
    response = client.post(
        "/api/uploads",
        headers=headers,
        files=[("files", ("empty.json", b'{"Reservations":[]}', "application/json"))],
        data={
            "paths": '["empty.json"]',
            "context": json.dumps(
                {
                    "collected_at": "2026-10-10T07:00:00Z",
                    "account_id": "111111111111",
                    "region": "us-east-1",
                    "request_parameters": {},
                }
            ),
        },
    )
    assert response.status_code == 200, response.text
    assert response.json()["can_import"]
    assert response.json()["diagnostics"]["sources"][0]["operation"] == "ec2:DescribeInstances"
    assert response.json()["diagnostics"]["sources"][0]["effective_coverage"] == "PARTIAL"


def test_access_restrictions(tmp_path):
    app = create_app(tmp_path / "work", origin=ORIGIN, bootstrap=TOKEN)
    with TestClient(app, base_url=ORIGIN) as client:
        assert client.get("/api/runs").status_code == 401
        assert client.get("/api/health", headers={"Host": "attacker.example"}).status_code == 403
        assert (
            client.get("/api/health", headers={"Origin": "https://attacker.example"}).status_code
            == 403
        )
        assert (
            client.post(
                "/api/session", json={}, headers={"Origin": ORIGIN, "X-Bootstrap-Token": "wrong"}
            ).status_code
            == 401
        )
        headers = authenticated(client)
        assert (
            client.post(
                "/api/import", headers={"Origin": ORIGIN}, json={"upload_id": "0" * 32}
            ).status_code
            == 403
        )
        assert (
            client.post(
                "/api/import",
                headers={"X-CSRF-Token": headers["X-CSRF-Token"]},
                json={"upload_id": "0" * 32},
            ).status_code
            == 403
        )
        assert client.get("/api/runs", headers={"Sec-Fetch-Site": "cross-site"}).status_code == 403
        assert client.get("/api/runs").headers["Cache-Control"] == "no-store"
        assert (
            client.post("/api/shutdown", headers=headers, json={"excess": "x" * 70000}).status_code
            == 413
        )
        assert client.get("/api/reports/not-a-file/download").status_code == 400
        assert client.get("/raw/evidence.json").status_code in {200, 503}  # SPA, never raw file


def test_workspace_single_owner(tmp_path):
    first = WorkspaceLock(tmp_path)
    try:
        with pytest.raises(RuntimeError):
            WorkspaceLock(tmp_path)
    finally:
        first.close()
    WorkspaceLock(tmp_path).close()


def test_failed_job_and_queued_cancellation(sample):
    _, client, headers, _, _ = sample
    failed = client.post("/api/runs/missing/assess", headers=headers, json={}).json()
    for _ in range(100):
        failed = client.get("/api/jobs/" + failed["id"]).json()
        if failed["status"] == "FAILED":
            break
        time.sleep(0.02)
    assert failed["status"] == "FAILED" and failed["error"] == "RUN_NOT_FOUND"
    coordinator = client.app.state.coordinator
    from threading import Event

    started, release = Event(), Event()

    def blocking(store):
        started.set()
        release.wait(5)
        return {}

    first = coordinator.submit("Synthetic blocker", blocking)
    assert started.wait(3)
    queued = coordinator.submit("Synthetic cancellable", lambda store: {})
    try:
        cancelled = client.post("/api/jobs/" + queued["id"] + "/cancel", headers=headers, json={})
        assert cancelled.json()["status"] == "CANCELLED"
        running = client.post("/api/jobs/" + first["id"] + "/cancel", headers=headers, json={})
        assert running.json()["status"] == "RUNNING"
        assert not running.json()["cancel_supported"]
    finally:
        release.set()


def test_scoped_views_and_resource_inventory(sample):
    _, client, _, run, _ = sample
    accounts = client.get(f"/api/runs/{run}/records/accounts").json()["items"]
    account = accounts[0]["account"]
    resources = client.get(f"/api/runs/{run}/records/resources?account={account}&limit=100").json()
    assert resources["total"] == accounts[0]["resources"]
    assert len({r["resource_id"] for r in resources["items"]}) == len(resources["items"])
    iam = client.get(f"/api/runs/{run}/records/iam_findings").json()["items"]
    assert all(f["rule_id"].startswith("R1-IAM-") for f in iam)
    assert client.get(f"/api/runs/{run}/records/network_findings").json()["total"] == 0
    assert all(f["source_timestamps"] for f in iam)


def test_framework_telemetry_is_disabled_even_with_environment_opt_in(tmp_path, monkeypatch):
    monkeypatch.setenv("FASTAPI_OTEL_AUTO_CONFIGURE", "true")
    app = create_app(tmp_path / "work", origin=ORIGIN, bootstrap=TOKEN)
    if hasattr(app, "_telemetry"):
        assert all(
            app._telemetry[k] is False
            for k in ("tracing", "metrics", "logs", "operation_spans", "auto_configure")
        )
        assert app._native_telemetry.enabled() is False
    app.state.coordinator.close()


def test_evidence_projection_matches_source_links_without_full_model_scan(sample, monkeypatch):
    """Bounded SQL projection preserves supporting and provenance links exactly."""
    from offline_analyzer.api.services import Views
    from offline_analyzer.storage.store import Store

    root, _, _, run, _ = sample
    with Store(root / "work") as store:
        oid = next(store.observations(run)).observation_id
        refs = store.refs(run, [oid])
        results = [
            r.model_dump(mode="json")
            for r in store.results(run)
            if oid in r.supporting_observations or any(x in r.evidence for x in refs)
        ][:100]
        findings = [
            f.finding_id for f in store.candidates(run) if any(x in f.evidence for x in refs)
        ][:100]

        def fail(*_):
            raise AssertionError("Unbounded model scan")

        monkeypatch.setattr(store, "results", fail)
        monkeypatch.setattr(store, "candidates", fail)
        projected = Views(root / "work").evidence_detail(store, run, oid)
        assert projected["evaluations"] == results
        assert projected["findings"] == findings
