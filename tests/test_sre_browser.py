"""Opt-in SRE validation workflow in Chromium against the production bundle (synthetic only)."""

import io
import os
import re
import socket
import threading
import time
import zipfile
from pathlib import Path

import pytest

from tests.browser_diagnostics import word_text

pytestmark = pytest.mark.skipif(
    os.environ.get("RUN_GUI_BROWSER") != "1", reason="opt-in browser runtime required"
)


def test_sre_workflow_to_finalized_findings(tmp_path):
    import uvicorn
    from openpyxl import load_workbook
    from playwright.sync_api import expect, sync_playwright

    from offline_analyzer.api.app import create_app
    from tools.fixtures.r12_demo import build

    build(tmp_path / "package")
    listener = socket.socket()
    listener.bind(("127.0.0.1", 0))
    port = listener.getsockname()[1]
    listener.listen(128)
    origin = f"http://127.0.0.1:{port}"
    app = create_app(tmp_path / "work", origin=origin, bootstrap="browser-test")
    server = uvicorn.Server(
        uvicorn.Config(app, host="127.0.0.1", port=port, log_level="critical", access_log=False)
    )
    threading.Thread(target=lambda: server.run(sockets=[listener]), daemon=True).start()
    for _ in range(100):
        if server.started:
            break
        time.sleep(0.05)
    assert server.started
    try:
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(headless=True)
            context = browser.new_context(
                accept_downloads=True, viewport={"width": 1440, "height": 1000}
            )
            external, errors = [], []

            def route(request):
                if not request.request.url.startswith(origin):
                    external.append(request.request.url)
                    request.abort()
                else:
                    request.continue_()

            context.route("**/*", route)
            page = context.new_page()
            page.on("pageerror", lambda exc: errors.append(str(exc)))
            page.goto(origin + "/#token=browser-test")
            page.get_by_role("button", name="Evidence import", exact=True).click()
            page.get_by_label("Select package folder").set_input_files(str(tmp_path / "package"))
            page.get_by_role("button", name="Validate evidence", exact=True).click()
            page.get_by_role("heading", name="Validation preview").wait_for()
            page.get_by_role("button", name="Import validated evidence").click()
            page.get_by_role("button", name="Run assessment", exact=True).click()
            page.get_by_label("Local operation details").get_by_role("status").filter(
                has_text="SUCCEEDED"
            ).wait_for(timeout=60000)

            page.get_by_role("button", name="Findings workbench", exact=True).click()
            page.get_by_role("button", name=re.compile("Show findings by account")).click()
            with page.expect_download() as bundle:
                page.get_by_role("link", name="Download all SRE workbooks (ZIP)").click()
            archive = zipfile.ZipFile(Path(bundle.value.path()))
            sheet = load_workbook(io.BytesIO(archive.read(archive.namelist()[0])))[
                "Findings for SRE"
            ]
            assert sheet.cell(1, 14).value == "SRE verdict"
            page.get_by_role("button", name=re.compile(r"Mark \d+ sent to SRE")).first.click()
            expect(page.locator(".badge.stage-with_sre").first).to_be_visible()

            page.locator("table").last.get_by_role("button", name="View").first.click()
            # Details: quick summary first, then the collected evidence; no name prompt.
            expect(page.get_by_label("Quick summary")).to_contain_text("What is wrong")
            expect(page.locator(".evidence-snippet").first).to_be_visible()
            page.get_by_role("tab", name="Client Reporting", exact=True).click()
            expect(page.get_by_label("Analyst identity", exact=True)).not_to_have_value("")
            page.get_by_text("SRE response (optional)").click()
            # A BAU conclusion cannot be recorded without the SRE's justification.
            page.get_by_label("SRE verdict").select_option("NO_GAP_BAU")
            expect(page.get_by_role("button", name="Record SRE verdict")).to_be_disabled()
            page.get_by_label("SRE verdict").select_option("GAP_CONFIRMED")
            page.get_by_label("SRE contact").fill("synthetic-sre")
            page.get_by_role("button", name="Record SRE verdict").click()
            page.locator("table").last.get_by_role("button", name="View").first.click()
            page.get_by_role("tab", name="Client Reporting", exact=True).click()
            page.get_by_text("SRE response (optional)").click()
            expect(page.locator(".sre-section")).to_contain_text("analyst to confirm")
            page.get_by_role("button", name="Confirm as vulnerability").click()
            # SRE is optional: a plain Confirm also finalizes a finding.
            page.locator("table").last.get_by_role("button", name="View").nth(1).click()
            page.get_by_role("tab", name="Client Reporting", exact=True).click()
            page.get_by_role("button", name="Confirm", exact=True).click()

            page.get_by_role("button", name="Finalized findings", exact=True).click()
            page.get_by_role("heading", name="Confirmed vulnerabilities", exact=True).wait_for()
            expect(page.locator(".finalized-detail")).to_have_count(2)
            expect(page.locator(".finalized-detail").first).to_contain_text("How it can be misused")
            with page.expect_download() as excel:
                page.get_by_role("link", name="Download Excel").click()
            confirmed = load_workbook(io.BytesIO(Path(excel.value.path()).read_bytes()))[
                "Technical Findings"
            ]
            assert confirmed.max_row == 3
            with page.expect_download() as word:
                page.get_by_role("link", name="Download Word report").click()
            assert "Confirmed Vulnerabilities Report" in word_text(word.value.path())

            # Bulk "not valid": the marked finding leaves the default report.
            page.get_by_role("button", name="Findings workbench", exact=True).click()
            page.get_by_label("Stage").select_option("WITH_SRE")
            table = page.locator("table").last
            expect(table.locator("tbody tr").first).to_be_visible()
            rejected_title = table.locator("tbody tr").first.locator("td").nth(1).inner_text()
            table.locator("tbody tr").first.get_by_role("checkbox").check()
            bar = page.get_by_label("Bulk marking")
            expect(bar.get_by_role("button", name="Mark not valid")).to_be_disabled()
            bar.get_by_label("Bulk note").fill("Known and accepted by the owner")
            bar.get_by_role("button", name="Mark not valid").click()
            page.get_by_label("Stage").select_option("REJECTED")
            expect(page.locator("table").last.locator("tbody tr")).to_have_count(1)

            # Per-account Excel + Word report follows the marking (default: exclude not valid).
            page.get_by_role("button", name=re.compile("Show findings by account")).click()
            with page.expect_download(timeout=60000) as report:
                page.get_by_role("button", name="Excel + Word report").first.click()
            bundle = zipfile.ZipFile(Path(report.value.path()))
            sheet = load_workbook(
                io.BytesIO(bundle.read(next(n for n in bundle.namelist() if n.endswith(".xlsx"))))
            )["Technical Findings"]
            ids = {row[0] for row in sheet.iter_rows(min_row=2, values_only=True)}
            run_id = page.get_by_label("Active evidence run").input_value()
            account = context.request.get(
                origin + "/api/runs/" + run_id + "/finding-facets"
            ).json()["accounts"][0]
            assert rejected_title and account["stage"]["REJECTED"] == 1
            assert len(ids) == account["total"] - 1
            page.get_by_role("button", name="Remove / replace data").first.click()
            expect(page.get_by_label("Selected assessment scope")).to_be_visible()

            # Attack paths are drawn as a chain with the rule behind each step.
            page.get_by_role("button", name="Dashboard", exact=True).click()
            page.get_by_role("tab", name="Attack Paths", exact=True).click()
            diagram = page.locator(".path-diagram").first
            expect(diagram).to_be_visible(timeout=30000)
            assert diagram.locator(".path-node").count() >= 2
            expect(page.locator(".path-rules").first).to_contain_text("Rule:")
            assert errors == [] and external == []
            browser.close()
    finally:
        server.should_exit = True
