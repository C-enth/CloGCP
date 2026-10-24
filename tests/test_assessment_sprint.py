"""Assessment sprint regression tests. All inputs are generated synthetic fixtures."""

import io
import json
import zipfile
from datetime import UTC, datetime
from xml.etree.ElementTree import fromstring

import pytest
from openpyxl import load_workbook

from offline_analyzer.api.facets import choices, facets
from offline_analyzer.api.services import Views
from offline_analyzer.attack_paths.feasibility import attack_paths, gcp_feasibility
from offline_analyzer.attack_paths.review import record as review_path
from offline_analyzer.contracts.r0 import Source
from offline_analyzer.contracts.report_finding import ReportEdit
from offline_analyzer.enterprise.snapshot import assess_enterprise
from offline_analyzer.enterprise.validation import validate_snapshot
from offline_analyzer.findings.finalisation import filtered, freeze, frozen
from offline_analyzer.findings.generate import record_decision
from offline_analyzer.findings.quality import finding_quality, run_quality
from offline_analyzer.findings.report_ready import report_detail, save_revision
from offline_analyzer.ingestion.importer import import_package
from offline_analyzer.ingestion.json_input import InputError
from offline_analyzer.normalization.gcp import validate_identity
from offline_analyzer.reporting.assessment_report import export_assessment, model
from offline_analyzer.reporting.cloud import dashboard, workspace_scope_summary
from offline_analyzer.reporting.sre_reports import finalized, finalized_document, finalized_files
from offline_analyzer.rules.engine import assess
from offline_analyzer.storage.store import Store
from tools.fixtures.r12_demo import build as aws_build
from tools.gcp_fixture import build as gcp_build


def source(**kwargs):
    return Source(
        path="inventory/project.json",
        sha256="0" * 64,
        format="gcp-json-v1",
        operation="gcp:projects.get",
        project_id="demo-alpha",
        project_number="123456",
        collected_at=datetime(2026, 10, 1, tzinfo=UTC),
        **kwargs,
    )


def test_identity_compares_id_and_number_without_display_name_joins():
    declared = source()
    for number in (123456, "123456"):
        validate_identity(
            {
                "projectId": "demo-alpha",
                "projectNumber": number,
                "name": "projects/a-display-label",
            },
            declared,
        )
        validate_identity({"name": "projects/123456/locations/global"}, declared)
    for record in (
        {"projectId": "demo-beta"},
        {"projectNumber": 999999},
        {"name": "projects/999999"},
        {"selfLink": "projects/demo-beta/resource"},
    ):
        with pytest.raises(InputError) as caught:
            validate_identity(record, declared)
        assert str(caught.value) == "GCP_PROJECT_IDENTITY_MISMATCH"
        assert caught.value.identity_context["file"] == "inventory/project.json"
        assert caught.value.identity_context["declared_project"] == "demo-alpha"
        assert "demo-alpha" not in str(caught.value)


def test_identity_details_are_session_only_and_never_saved_or_exported(tmp_path):
    from offline_analyzer.api.coordinator import Coordinator
    from offline_analyzer.api.diagnostics import job_diagnostic
    from offline_analyzer.api.runtime import RuntimeIdentity
    from offline_analyzer.normalization.gcp_identity import mismatch

    coordinator = Coordinator(tmp_path / "synthetic")
    try:

        def fail(_store):
            raise mismatch("synthetic-dataset.json", "demo-alpha", "projectId", "demo-beta")

        job = coordinator.submit("Synthetic validation", fail)
        coordinator.futures[job["id"]].result(timeout=10)
        status = coordinator.status(job["id"])
        assert status["status"] == "FAILED"
        assert status["identity_context"]["file"] == "synthetic-dataset.json"
        for path in coordinator.journal.directory.iterdir():
            assert "synthetic-dataset.json" not in path.read_text()
            assert "demo-alpha" not in path.read_text()
        diagnostic = job_diagnostic(RuntimeIdentity(), status)
        assert "synthetic-dataset.json" not in json.dumps(diagnostic)
        assert "demo-alpha" not in json.dumps(diagnostic)
    finally:
        coordinator.close()


@pytest.fixture()
def estate(tmp_path):
    aws_build(tmp_path / "aws")
    gcp_build(tmp_path / "gcp")
    with Store(tmp_path / "work") as store:
        aws = import_package(store, tmp_path / "aws")
        gcp = import_package(store, tmp_path / "gcp")
        run = assess_enterprise(store)
        yield store, run, aws, gcp


def test_enterprise_reconciles_exact_ids_scopes_and_omitted_imports(estate):
    store, run, aws, _gcp = estate
    union = validate_snapshot(store, run)
    assert union["status"] == "PASS"
    assert all(c["expected"] == c["actual"] and c["missing"] == 0 for c in union["checks"])
    assert all(
        union["scopes"][p]["expected"] == union["scopes"][p]["actual"] for p in ("AWS", "GCP")
    )
    static_before = workspace_scope_summary(store)
    for row in static_before["domains"]:
        for cloud in ("AWS", "GCP"):
            assert set(row[cloud]) <= set(static_before["providers"][cloud])
    subset = assess_enterprise(store, [aws])
    assert validate_snapshot(store, subset)["status"] == "WARNING"
    assert workspace_scope_summary(store) == static_before
    with pytest.raises(InputError, match="FINALISATION_REQUIRES_COMPLETE"):
        freeze(store, subset, "Synthetic analyst", True)
    store.db.execute(
        "DELETE FROM observations WHERE run_id=? AND observation_id="
        "(SELECT min(observation_id) FROM observations WHERE run_id=?)",
        [run, run],
    )
    result = validate_snapshot(store, run)
    assert result["status"] == "ERROR"
    assert next(c for c in result["checks"] if c["stage"] == "observations")["missing"] == 1


def test_cloud_reports_drafts_native_office_and_freeze_are_consistent(estate, tmp_path):
    store, run, aws, _gcp = estate
    candidates = list(store.candidates(run))
    chosen = [next(f for f in candidates if f.provider == p) for p in ("AWS", "GCP")]
    for f in chosen:
        record_decision(
            store, run, f.finding_id, "CONFIRMED", "Synthetic analyst", "Synthetic review"
        )
    detail = report_detail(store, run, chosen[1].finding_id)
    detail["sections"]["title"] = "Synthetic edited report title"
    save_revision(
        store,
        run,
        chosen[1].finding_id,
        ReportEdit(
            expected_revision=0,
            analyst="Synthetic analyst",
            sections=detail["sections"],
            member_ids=[chosen[1].finding_id],
            include_in_report=True,
        ),
    )
    paths = attack_paths(store, run)["items"]
    assert paths and all("stride" in p and "analyst_review" in p for p in paths)
    linked = next((p for p in paths if chosen[1].finding_id in p.get("finding_ids", [])), paths[0])
    review_path(store, run, linked["path_id"], "CONFIRMED", "Synthetic analyst")
    refreshed = attack_paths(store, run)["items"]
    assert (
        next(p for p in refreshed if p["path_id"] == linked["path_id"])["analyst_review"]
        == "CONFIRMED"
    )
    for provider in ("AWS", "GCP"):
        data = model(store, run, "confirmed", provider)
        assert data["findings"] and all(f["provider"] == provider for f in data["findings"])
        assert all(p["provider"] == provider for p in data["paths"])
        assert all(s["provider"] == provider for s in data["scopes"])
        if provider == "GCP":
            assert data["findings"][0]["title"] == "Synthetic edited report title"
        target = tmp_path / (provider + ".zip")
        export_assessment(store, run, target, include="confirmed", provider=provider)
        with zipfile.ZipFile(target) as archive:
            assert not any(n.endswith(".rtf") for n in archive.namelist())
            word = next(n for n in archive.namelist() if n.endswith(".docx"))
            pdf = next(n for n in archive.namelist() if n.endswith(".pdf"))
            assert archive.read(pdf).startswith(b"%PDF-1.4")
            with zipfile.ZipFile(io.BytesIO(archive.read(word))) as docx:
                xml = fromstring(docx.read("word/document.xml"))
                assert "3 Attack paths" in " ".join(xml.itertext())
                assert not any(
                    'TargetMode="External"' in docx.read(n).decode()
                    for n in docx.namelist()
                    if n.endswith(".rels")
                )
            workbook_name = next(n for n in archive.namelist() if n.endswith(".xlsx"))
            book = load_workbook(io.BytesIO(archive.read(workbook_name)))
            assert "Findings Validation" in book.sheetnames
            assert not any(c.data_type == "f" for ws in book for row in ws for c in row)
            book.close()
    with pytest.raises(InputError, match="FINALISATION_CONFIRMATION_REQUIRED"):
        freeze(store, run, "Synthetic analyst", False)
    record = freeze(store, run, "Synthetic analyst", True)
    assert record == freeze(store, run, "Synthetic analyst", True)
    assert frozen(store, run) is not None
    frozen_before = finalized(store, run)
    assert {f["provider"] for f in frozen_before["findings"]} == {"AWS", "GCP"}
    for action in (
        lambda: record_decision(
            store, run, chosen[0].finding_id, "REJECTED", "Synthetic analyst", "test"
        ),
        lambda: assess(store, run),
        lambda: review_path(store, run, linked["path_id"], "REJECTED", "Synthetic analyst"),
    ):
        with pytest.raises(InputError, match="ASSESSMENT_FINALIZED"):
            action()
    # A later decision elsewhere cannot rewrite the preserved frozen text or counts.
    assess(store, aws)
    record_decision(
        store, aws, chosen[0].finding_id, "REJECTED", "Synthetic analyst", "Later review"
    )
    assert finalized(store, run) == frozen_before
    for provider in ("GCP", "AWS", ""):
        data = finalized_document(store, run, provider)
        assert data["frozen"]
        assert all(not provider or f["provider"] == provider for f in data["findings"])
        assert finalized_files(data, "docx").startswith(b"PK")
        assert finalized_files(data, "pdf").startswith(b"%PDF-1.4")


def test_findings_quality_coverage_choices_and_strict_feasibility(estate):
    store, run, _, _ = estate
    quality = run_quality(store, run)
    assert len(quality["items"]) == len(list(store.candidates(run)))
    assert quality["ready"] + quality["needs_validation"] == len(quality["items"])
    data = dashboard(store, run)
    assert len(data["stride"]) == 6
    assert all(s["state"] in {"PARTIAL", "NOT_ASSESSED"} for s in data["stride"])
    assert "cannot support a PASS" in data["coverage_definitions"]["ERROR"]
    choices = facets(Views(store.workspace), store, run, "coverage", "GCP")
    assert choices["choices"]["project"]
    assert choices["choices"]["service"]
    assert "configuration" not in json.dumps(choices["choices"])
    path = {"transitions": [{"relationship": "Observed service-account privilege binding"}]}
    assert gcp_feasibility(path, "GCP-KEY-PRIVILEGED", "roles/owner")[0] == "INSUFFICIENT_EVIDENCE"
    assert gcp_feasibility(path, "GCP-IAM-PUBLIC", None)[0] == "INSUFFICIENT_EVIDENCE"


def test_quality_never_auto_readies_contradictions_or_low_confidence(estate):
    store, run, _, _ = estate
    finding = next(iter(store.candidates(run)))
    artifacts = {ref.artifact_id for ref in finding.evidence}
    result = next(
        r
        for r in store.results(run)
        if (r.rule_id, r.resource_id) == (finding.rule_id, finding.resource_id)
    )
    for confidence in ("LOW", "UNASSESSED"):
        changed = finding.model_copy(update={"evidence_confidence": confidence})
        quality = finding_quality(changed, result, artifacts)
        assert quality["state"] == "NEEDS_VALIDATION"
        assert "evidence_confidence_requires_validation" in quality["unresolved"]
    changed = finding.model_copy(update={"contradictory_observations": ["synthetic-conflict"]})
    assert "synthetic-conflict" in finding_quality(changed, result, artifacts)["unresolved"]
    assert "evidence_references" in finding_quality(finding, result, set())["missing_fields"]


def test_choices_include_relationship_identifiers_alongside_titles():
    selected = choices(
        [
            {
                "title": "Synthetic relationship",
                "relationship_id": "synthetic-edge",
                "relationship_type": "ASSUME_ROLE",
            }
        ]
    )["choices"]["q"]
    assert set(selected) == {"Synthetic relationship", "synthetic-edge", "ASSUME_ROLE"}


def test_frozen_path_scope_is_identity_not_relationship_scope():
    finding = {"id": "synthetic", "provider": "AWS", "scope": "111111111111", "severity": "HIGH"}
    path = {
        "provider": "AWS",
        "scope": "cross-account",
        "account_id": "111111111111",
        "finding_ids": ["synthetic"],
    }
    record = {
        "created_at": "2026-10-01T00:00:00Z",
        "analyst": "Synthetic reviewer",
        "document": {"findings": [finding], "paths": [path]},
    }
    assert filtered(record, "document", "AWS", "111111111111")["paths"] == [path]
    assert not filtered(record, "document", "AWS", "222222222222")["paths"]


def test_aws_only_estate_shares_reconciliation_quality_review_and_freeze(tmp_path):
    aws_build(tmp_path / "aws")
    with Store(tmp_path / "work") as store:
        member = import_package(store, tmp_path / "aws")
        run = assess_enterprise(store)
        union = validate_snapshot(store, run)
        assert union["status"] == "PASS"
        assert union["scopes"]["AWS"]["expected"] == union["scopes"]["AWS"]["actual"]
        assert all(c["expected"] == c["actual"] and c["missing"] == 0 for c in union["checks"])
        candidates = list(store.candidates(run))
        assert candidates and {f.provider for f in candidates} == {"AWS"}
        quality = run_quality(store, run)
        assert quality["ready"] + quality["needs_validation"] == len(candidates)
        finding = candidates[0]
        result = next(
            r
            for r in store.results(run)
            if (r.rule_id, r.resource_id) == (finding.rule_id, finding.resource_id)
        )
        artifacts = {ref.artifact_id for ref in finding.evidence}
        low = finding.model_copy(update={"evidence_confidence": "LOW"})
        assert finding_quality(low, result, artifacts)["state"] == "NEEDS_VALIDATION"
        record_decision(store, run, finding.finding_id, "CONFIRMED", "Synthetic analyst", "AWS")
        paths = attack_paths(store, run)["items"]
        if paths:
            review_path(store, run, paths[0]["path_id"], "CONFIRMED", "Synthetic analyst")
        with pytest.raises(InputError, match="FINALISATION_CONFIRMATION_REQUIRED"):
            freeze(store, run, "Synthetic analyst", False)
        freeze(store, run, "Synthetic analyst", True)
        before = finalized(store, run)
        assert {f["provider"] for f in before["findings"]} == {"AWS"}
        with pytest.raises(InputError, match="ASSESSMENT_FINALIZED"):
            record_decision(store, run, finding.finding_id, "REJECTED", "Synthetic analyst", "x")
        assess(store, member)
        assert finalized(store, run) == before
        data = finalized_document(store, run, "AWS")
        assert data["frozen"] and finalized_files(data, "pdf").startswith(b"%PDF-1.4")
