"""Synthetic integration/security/regression tests for the actual vertical slice."""

import json
import socket
from collections import Counter

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from offline_analyzer.cli import main, summary
from offline_analyzer.contracts.r0 import Package, Source
from offline_analyzer.findings.generate import record_decision, validation_status
from offline_analyzer.ingestion.importer import import_package, single_file_package
from offline_analyzer.ingestion.json_input import InputError, inspect_json
from offline_analyzer.ingestion.safety import sha256_file
from offline_analyzer.reporting.excel import export_excel, safe_text
from offline_analyzer.rules.engine import assess
from offline_analyzer.storage.store import Store
from tools.fixtures.demo import TIME, A, B, build_package


@pytest.fixture
def pipeline(tmp_path):
    directory = tmp_path / "input"
    build_package(directory)
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, directory)
        assess(store, run)
        yield store, run, directory


def test_end_to_end(pipeline, tmp_path):
    store, run, directory = pipeline
    all_results = list(store.results(run))
    # R0 regression: the three original rules keep their exact outcomes.
    results = [r for r in all_results if r.rule_id.startswith("R0-")]
    assert Counter(r.outcome for r in results) == {
        "PASS": 2,
        "FAIL": 3,
        "UNKNOWN": 3,
        "COLLECTION_ERROR": 1,
        "MANUAL_VALIDATION_REQUIRED": 1,
    }
    # R1 rules sharing these operations only report missing fields here, never PASS.
    assert Counter(r.outcome for r in all_results if r.rule_id.startswith("R1-")) == {
        "UNKNOWN": 5,
        "COLLECTION_ERROR": 1,
    }
    assert {r.account_id for r in results} == {A, B}
    assert len(list(store.candidates(run))) == 3
    expected = {
        "R0-EC2-IMDSV2": "AWS-19",
        "R0-RDS-ENCRYPTION": "AWS-22",
        "R0-CLOUDTRAIL-VALIDATION": "AWS-07",
    }
    assert all(r.methodology_id == expected[r.rule_id] for r in results)
    assert all(r.mapping_status == "PROPOSED" and r.methodology_criteria for r in results)
    assert not any(r.outcome == "PASS" and r.resource_id is None for r in results)
    rds_fail = next(r for r in results if r.rule_id == "R0-RDS-ENCRYPTION" and r.outcome == "FAIL")
    assert rds_fail.contradictory_observations
    assert rds_fail.confidence == "MEDIUM"
    assert summary(store, run)["evidence_files"] == 7
    assert len(summary(store, run)["represented_scopes"]) == 3
    path = tmp_path / "report.xlsx"
    export_excel(store, run, path)
    wb = load_workbook(path, data_only=False)
    assert set(wb.sheetnames) == {
        "Findings Register",
        "Confirmed Findings",
        "Account Coverage",
        "Control Assessment",
        "Evidence Register",
        "Review Queue",
        "Assessment Summary",
        "Methodology Coverage",
        "Tracker Staging",
        "Attack Path Prerequisites",
        "Security Gaps",
        "Potential Exposures",
        "Baseline Deviations",
        "Incomplete Investigations",
        "Assessment Coverage",
        "IAM Investigations",
        "IAM Relationships",
        "Effective Permissions",
        "IAM Evidence Gaps",
        "Enterprise Account Summary",
        "Cross-Account Relationships",
        "Cross-Account IAM Investigation",
        "Shared Resource Access",
        "Shared Network Relationships",
        "Enterprise Security Gaps",
        "Cross-Account Evidence Gaps",
        "Affected Accounts and Resources",
    }
    report_results = list(wb["Control Assessment"].iter_rows(min_row=4, values_only=True))
    assert Counter(row[7] for row in report_results) == Counter(r.outcome for r in all_results)
    assert {row[0] for row in report_results} == {r.evaluation_id for r in all_results}
    assert wb["Findings Register"].max_row - 3 == 3
    assert wb["Confirmed Findings"].max_row == 3
    assert any(
        "NOT_COLLECTED" in str(row) or "UNKNOWN" in str(row)
        for row in wb["Account Coverage"].values
    )
    assert all(cell.data_type != "f" for sheet in wb for row in sheet for cell in row)
    wb.close()
    with pytest.raises(InputError, match="REPORT_ALREADY_EXISTS"):
        export_excel(store, run, path)


def test_deduplication_provenance_and_raw_retention(pipeline):
    store, run, directory = pipeline
    assert (
        store.db.execute("SELECT count(*) FROM observations WHERE run_id=?", [run]).fetchone()[0]
        == 8
    )
    assert (
        store.db.execute("SELECT count(*) FROM evidence_refs WHERE run_id=?", [run]).fetchone()[0]
        == 10
    )
    for obs in store.observations(run):
        for ref in store.refs(run, [obs.observation_id]):
            digest = store.db.execute(
                "SELECT sha256 FROM artifacts WHERE run_id=? AND artifact_id=?",
                [run, ref.artifact_id],
            ).fetchone()[0]
            raw = store.workspace / "raw" / f"{digest}.json"
            assert sha256_file(raw) == digest
            value = json.loads(raw.read_text())
            for part in ref.locator.lstrip("/").split("/"):
                value = value[int(part)] if isinstance(value, list) else value[part]
            assert isinstance(value, dict)
    scanner = next(o for o in store.observations(run) if o.source_tool == "prowler")
    original = json.loads(
        (store.workspace / "raw" / f"{scanner.integrity_sha256}.json").read_text()
    )
    assert original[0]["unrecognized_extension"]["retained"] == "synthetic raw detail"
    assert scanner.source_severity == "High"
    assert scanner.resource_type == "AWS::RDS::DBInstance"


def test_repeat_import_reassess_and_decisions_survive(pipeline, tmp_path):
    store, run, directory = pipeline
    before = [r.model_dump_json() for r in store.results(run)]
    candidate = next(store.candidates(run))
    record_decision(
        store, run, candidate.finding_id, "CONFIRMED", "synthetic-analyst", "=untrusted note"
    )
    assert validation_status(store, candidate)[0] == "CONFIRMED"
    assert import_package(store, directory) == run
    assess(store, run)
    assert [r.model_dump_json() for r in store.results(run)] == before
    assert validation_status(store, candidate)[0] == "CONFIRMED"
    assert store.db.execute("SELECT count(*) FROM decisions").fetchone()[0] == 1
    export_excel(store, run, tmp_path / "confirmed.xlsx")
    wb = load_workbook(tmp_path / "confirmed.xlsx")
    confirmed = wb["Confirmed Findings"]
    assert confirmed.max_row == 4
    headers = [c for row in confirmed.iter_rows(max_row=3) for c in row]
    note = next(c for c in headers if c.value == "Analyst note")
    # The untrusted "=..." note is stored as literal text, never as a formula.
    assert confirmed.cell(4, note.column).value.startswith("'")
    wb.close()
    # A changed source snapshot creates a new run but retains the issue fingerprint.
    package = Package.model_validate_json((directory / "manifest.json").read_bytes())
    changed = package.model_copy(
        update={
            "sources": tuple(
                s.model_copy(update={"collected_at": s.collected_at.replace(hour=9)})
                for s in package.sources
            )
        }
    )
    run2 = import_package(store, directory, changed)
    assert run2 != run
    assess(store, run2)
    updated = next(f for f in store.candidates(run2) if f.finding_id == candidate.finding_id)
    assert validation_status(store, updated)[0] == "NEEDS_REVALIDATION"
    assert store.latest_decision(candidate.finding_id).status == "CONFIRMED"


def test_reopen_storage(tmp_path):
    directory = tmp_path / "input"
    build_package(directory)
    workspace = tmp_path / "workspace"
    with Store(workspace) as store:
        run = import_package(store, directory)
        assess(store, run)
        finding = next(store.candidates(run))
        record_decision(store, run, finding.finding_id, "REJECTED", "synthetic", "reason")
    with Store(workspace) as store:
        assert summary(store, run)["candidate_findings"] == 3
        assert import_package(store, directory) == run
        assert validation_status(store, finding)[0] == "REJECTED"


@pytest.mark.parametrize(
    "payload,code",
    [
        ('{"Reservations":[],"Reservations":[]}', "DUPLICATE_JSON_KEY"),
        ('{"aws_secret_access_key":"synthetic-never-persisted"}', "CREDENTIAL_MATERIAL_REJECTED"),
        ('{"Reservations":[', "MALFORMED_JSON"),
        ("[" * 45 + "0" + "]" * 45, "INPUT_DEPTH_LIMIT"),
        ('{"x":NaN}', "MALFORMED_JSON"),
        ("{} {}", "MALFORMED_JSON"),
    ],
)
def test_malicious_json(tmp_path, payload, code):
    path = tmp_path / "malicious.json"
    path.write_text(payload)
    with pytest.raises(InputError, match=code):
        inspect_json(path)


def one_file(tmp_path, data, operation="rds:DescribeDBInstances", account=A, region="us-east-1"):
    path = tmp_path / "source.json"
    path.write_text(json.dumps(data))
    package = single_file_package(
        path,
        format="aws-json-v1",
        operation=operation,
        account_id=account,
        region=region,
        collected_at=TIME,
    )
    return path, package


@pytest.mark.parametrize("data", [{}, {"DBInstances": {}}, {"DBInstances": "invalid"}, []])
def test_wrong_response_shape_rolls_back(tmp_path, data):
    path, package = one_file(tmp_path, data)
    with Store(tmp_path / "workspace") as store:
        with pytest.raises(InputError):
            import_package(store, path.parent, package)
        for table in ["runs", "sources", "artifacts", "observations", "evidence_refs"]:
            assert store.db.execute(f"SELECT count(*) FROM {table}").fetchone()[0] == 0
        assert not (store.workspace / "raw").exists()


def test_partial_import_rolls_back_and_recovers(tmp_path):
    directory = tmp_path / "input"
    package = build_package(directory)
    scanner_path = directory / "prowler.json"
    scanner = json.loads(scanner_path.read_text())
    scanner[0]["metadata"]["product"]["version"] = "99.0.0"
    scanner_path.write_text(json.dumps(scanner))
    source = package.sources[-1].model_copy(update={"sha256": sha256_file(scanner_path)})
    changed = package.model_copy(update={"sources": package.sources[:-1] + (source,)})
    with Store(tmp_path / "workspace") as store:
        with pytest.raises(InputError, match="UNSUPPORTED_PROWLER_VERSION"):
            import_package(store, directory, changed)
        assert store.db.execute("SELECT count(*) FROM runs").fetchone()[0] == 0
        assert store.db.execute("SELECT count(*) FROM observations").fetchone()[0] == 0
        scanner[0]["metadata"]["product"]["version"] = "5.10.0"
        scanner_path.write_text(json.dumps(scanner))
        source = source.model_copy(update={"sha256": sha256_file(scanner_path)})
        changed = package.model_copy(update={"sources": package.sources[:-1] + (source,)})
        run = import_package(store, directory, changed)
        assess(store, run)
        assert summary(store, run)["candidate_findings"] == 3


def test_checksum_failure_and_raw_tamper(pipeline, tmp_path):
    store, run, directory = pipeline
    package = Package.model_validate_json((directory / "manifest.json").read_bytes())
    path = directory / package.sources[0].path
    path.write_text("{}")
    with pytest.raises(InputError, match="ARTIFACT_CHECKSUM_MISMATCH"):
        import_package(store, directory)
    assert len(list(store.results(run))) == 16  # 10 R0 + 6 R1 missing-field/scope rows
    digest = store.db.execute("SELECT sha256 FROM artifacts LIMIT 1").fetchone()[0]
    (store.workspace / "raw" / f"{digest}.json").write_text("{}")
    with pytest.raises(InputError, match="STORED_RAW_INTEGRITY_FAILURE"):
        assess(store, run)
    with pytest.raises(InputError, match="STORED_RAW_INTEGRITY_FAILURE"):
        export_excel(store, run, tmp_path / "tampered.xlsx")


@pytest.mark.parametrize(
    "path", ["../outside.json", "C:/outside.json", "\\\\server\\file", "/outside"]
)
def test_manifest_paths(path):
    with pytest.raises(ValidationError):
        Source(
            path=path,
            sha256="0" * 64,
            format="aws-json-v1",
            operation="ec2:DescribeInstances",
            collected_at=TIME,
        )


def test_package_symlink_escape(tmp_path):
    path, package = one_file(tmp_path, {"DBInstances": []})
    external = tmp_path / "external.json"
    external.write_text(path.read_text())
    directory = tmp_path / "inside"
    directory.mkdir()
    (directory / path.name).symlink_to(external)
    with Store(tmp_path / "workspace") as store:
        with pytest.raises(ValueError, match="escapes"):
            import_package(store, directory, package)


@pytest.mark.parametrize(
    "value,expected",
    [(True, "PASS"), (False, "FAIL"), ("false", "UNKNOWN"), (None, "UNKNOWN"), (0, "UNKNOWN")],
)
def test_rds_strict_boolean(tmp_path, value, expected):
    path, package = one_file(
        tmp_path,
        {
            "DBInstances": [
                {
                    "DBInstanceArn": f"arn:aws:rds:us-east-1:{A}:db:db",
                    "Engine": "postgres",
                    "StorageEncrypted": value,
                }
            ]
        },
    )
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, path.parent, package)
        assess(store, run)
        result = next(
            r for r in store.results(run) if r.resource_id and r.rule_id.startswith("R0-")
        )
        assert result.outcome == expected


@pytest.mark.parametrize(
    "engine,outcome",
    [
        ("aurora-postgresql", "UNKNOWN"),
        ("docdb", "NOT_APPLICABLE"),
        (None, "UNKNOWN"),
        ("new-engine", "MANUAL_VALIDATION_REQUIRED"),
    ],
)
def test_rds_applicability(tmp_path, engine, outcome):
    path, package = one_file(
        tmp_path,
        {
            "DBInstances": [
                {
                    "DBInstanceArn": f"arn:aws:rds:us-east-1:{A}:db:db",
                    "Engine": engine,
                    "StorageEncrypted": True,
                }
            ]
        },
    )
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, path.parent, package)
        assess(store, run)
        assert (
            next(
                r for r in store.results(run) if r.resource_id and r.rule_id.startswith("R0-")
            ).outcome
            == outcome
        )


@pytest.mark.parametrize(
    "options,state,outcome",
    [
        (
            {"State": "applied", "HttpEndpoint": "disabled", "HttpTokens": "optional"},
            "running",
            "NOT_APPLICABLE",
        ),
        (
            {"State": "pending", "HttpEndpoint": "enabled", "HttpTokens": "required"},
            "running",
            "UNKNOWN",
        ),
        ({"State": "applied", "HttpEndpoint": "enabled"}, "running", "UNKNOWN"),
        ({}, "terminated", "NOT_APPLICABLE"),
    ],
)
def test_ec2_prerequisites(tmp_path, options, state, outcome):
    path, package = one_file(
        tmp_path,
        {
            "Reservations": [
                {
                    "Instances": [
                        {
                            "InstanceId": "i-12345678",
                            "State": {"Name": state},
                            "MetadataOptions": options,
                        }
                    ]
                }
            ]
        },
        "ec2:DescribeInstances",
    )
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, path.parent, package)
        assess(store, run)
        assert (
            next(
                r for r in store.results(run) if r.resource_id and r.rule_id.startswith("R0-")
            ).outcome
            == outcome
        )


def test_native_conflict_not_arbitrary_latest(tmp_path):
    path, package = one_file(
        tmp_path,
        {
            "DBInstances": [
                {
                    "DBInstanceArn": f"arn:aws:rds:us-east-1:{A}:db:db",
                    "Engine": "postgres",
                    "StorageEncrypted": v,
                }
                for v in [True, False]
            ]
        },
    )
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, path.parent, package)
        assess(store, run)
        assert (
            next(
                r for r in store.results(run) if r.resource_id and r.rule_id.startswith("R0-")
            ).outcome
            == "UNKNOWN"
        )
        assert not list(store.candidates(run))


@pytest.mark.parametrize(
    "data,operation",
    [
        ({"DBInstances": [None, {"StorageEncrypted": True}]}, "rds:DescribeDBInstances"),
        ({"Reservations": [None]}, "ec2:DescribeInstances"),
    ],
)
def test_missing_identity_retained_not_pass(tmp_path, data, operation):
    path, package = one_file(tmp_path, data, operation, account=None, region=None)
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, path.parent, package)
        assess(store, run)
        observations = list(store.observations(run))
        assert observations and all(o.resource_id is None for o in observations)
        assert all(r.outcome == "UNKNOWN" for r in store.results(run))
        assert not list(store.candidates(run))


def test_no_network_runtime(tmp_path, monkeypatch):
    def denied(*args, **kwargs):
        raise AssertionError("Network connection attempted")

    monkeypatch.setattr(socket.socket, "connect", denied)
    directory = tmp_path / "input"
    build_package(directory)
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, directory)
        assess(store, run)
        export_excel(store, run, tmp_path / "offline.xlsx")


@pytest.mark.parametrize("value", ["=1+1", "+SUM(A1)", "-cmd", "@cmd", "\x00 \t=cmd", "\n=cmd"])
def test_report_formula_safety(value):
    assert safe_text(value).startswith("'")


def test_cli_workflow_and_safe_error(tmp_path, capsys):
    directory = tmp_path / "input"
    build_package(directory)
    workspace = tmp_path / "workspace"
    args = ["--workspace", str(workspace)]
    assert main(args + ["import", str(directory)]) == 0
    run = capsys.readouterr().out.strip()
    assert run.startswith("run-")
    assert main(args + ["assess", "--run", run]) == 0
    assert json.loads(capsys.readouterr().out)["candidate_findings"] == 3
    assert main(args + ["export", "--run", run, "--output", str(tmp_path / "cli.xlsx")]) == 0
    capsys.readouterr()
    assert main(args + ["assess", "--run", "PRIVATE_VALUE_DO_NOT_LOG"]) == 1
    output = capsys.readouterr()
    assert "PRIVATE_VALUE_DO_NOT_LOG" not in output.err
    assert "RUN_NOT_FOUND" in output.err


def test_budget_enforcement(tmp_path, monkeypatch):
    import offline_analyzer.ingestion.json_input as module

    path = tmp_path / "source.json"
    path.write_text('[{"a": "synthetic"}]')
    monkeypatch.setattr(module, "MAX_BYTES", 2)
    with pytest.raises(InputError, match="INPUT_BYTE_LIMIT"):
        inspect_json(path)
    monkeypatch.setattr(module, "MAX_BYTES", 1024)
    monkeypatch.setattr(module, "MAX_STRING", 2)
    with pytest.raises(InputError, match="INPUT_STRING_LIMIT"):
        inspect_json(path)
    monkeypatch.setattr(module, "MAX_STRING", 128)
    monkeypatch.setattr(module, "MAX_EVENTS", 2)
    with pytest.raises(InputError, match="INPUT_EVENT_LIMIT"):
        inspect_json(path)


def test_workspace_raw_symlink_not_followed(tmp_path):
    directory = tmp_path / "input"
    build_package(directory)
    outside = tmp_path / "outside"
    outside.mkdir()
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    (workspace / "raw").symlink_to(outside, target_is_directory=True)
    with Store(workspace) as store:
        with pytest.raises(ValueError, match="escapes"):
            import_package(store, directory)
        assert store.db.execute("SELECT count(*) FROM runs").fetchone()[0] == 0
    assert not list(outside.iterdir())


def test_workspace_database_symlink_not_followed(tmp_path):
    outside = tmp_path / "outside.duckdb"
    outside.write_bytes(b"not a database")
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    (workspace / "assessment.duckdb").symlink_to(outside)
    with pytest.raises(ValueError, match="escapes"):
        Store(workspace)
    assert outside.read_bytes() == b"not a database"


def test_explicit_aws_error_response(tmp_path):
    path, package = one_file(tmp_path, {"Error": {"Code": "AccessDenied", "Message": "synthetic"}})
    source = package.sources[0].model_copy(
        update={"coverage": "ERROR", "errors": ("AccessDenied",)}
    )
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, path.parent, package.model_copy(update={"sources": (source,)}))
        assess(store, run)
        assert not list(store.observations(run))
        assert any(
            r.rule_id == "R0-RDS-ENCRYPTION" and r.outcome == "COLLECTION_ERROR"
            for r in store.results(run)
        )


def test_duplicate_descriptor_rejected(tmp_path):
    _, package = one_file(tmp_path, {"DBInstances": []})
    with pytest.raises(ValidationError, match="duplicate source"):
        Package(sources=(package.sources[0], package.sources[0]))


def test_uncollected_empty_listing_never_pass(tmp_path):
    path, package = one_file(tmp_path, {"DBInstances": []})
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, path.parent, package)
        assess(store, run)
        assert all(r.outcome == "UNKNOWN" for r in store.results(run))
        assert not list(store.candidates(run))


def test_missing_ec2_lifecycle_requires_validation(tmp_path):
    path, package = one_file(
        tmp_path,
        {
            "Reservations": [
                {
                    "Instances": [
                        {
                            "InstanceId": "i-12345678",
                            "MetadataOptions": {
                                "State": "applied",
                                "HttpTokens": "required",
                                "HttpEndpoint": "enabled",
                            },
                        }
                    ]
                }
            ]
        },
        "ec2:DescribeInstances",
    )
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, path.parent, package)
        assess(store, run)
        result = next(
            r for r in store.results(run) if r.resource_id and r.rule_id.startswith("R0-")
        )
        assert result.outcome == "UNKNOWN" and "State.Name" in result.missing_dependencies


def test_cli_import_file(tmp_path, capsys):
    path, _ = one_file(
        tmp_path,
        {
            "DBInstances": [
                {
                    "DBInstanceArn": f"arn:aws:rds:us-east-1:{A}:db:db",
                    "Engine": "postgres",
                    "StorageEncrypted": False,
                }
            ]
        },
    )
    args = ["--workspace", str(tmp_path / "workspace")]
    assert (
        main(
            args
            + [
                "import-file",
                str(path),
                "--format",
                "aws-json-v1",
                "--operation",
                "rds:DescribeDBInstances",
                "--account",
                A,
                "--region",
                "us-east-1",
                "--collected-at",
                TIME,
            ]
        )
        == 0
    )
    run = capsys.readouterr().out.strip()
    assert main(args + ["assess", "--run", run]) == 0
    assert json.loads(capsys.readouterr().out)["outcomes"]["FAIL"] == 1


def test_assessment_rollback_preserves_previous_results(pipeline, monkeypatch):
    import offline_analyzer.rules.engine as module

    store, run, _ = pipeline
    before = [r.model_dump_json() for r in store.results(run)]
    original = module.persist_result
    calls = 0

    def failure(store, run, result):
        nonlocal calls
        original(store, run, result)
        calls += 1
        if calls == 2:
            raise RuntimeError("synthetic interruption")

    monkeypatch.setattr(module, "persist_result", failure)
    with pytest.raises(RuntimeError, match="synthetic interruption"):
        assess(store, run)
    assert [r.model_dump_json() for r in store.results(run)] == before
    assert len(list(store.candidates(run))) == 3


def test_full_scope_empty_is_not_applicable_not_pass(tmp_path):
    path, package = one_file(tmp_path, {"DBInstances": []})
    source = package.sources[0].model_copy(
        update={"coverage": "COMPLETE", "request_parameters": {}}
    )
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, path.parent, package.model_copy(update={"sources": (source,)}))
        assess(store, run)
        result = next(r for r in store.results(run) if r.rule_id == "R0-RDS-ENCRYPTION")
        assert result.outcome == "NOT_APPLICABLE"
        assert result.evidence


def test_pagination_overrides_complete(tmp_path):
    path, package = one_file(tmp_path, {"DBInstances": [], "Marker": "synthetic-more"})
    source = package.sources[0].model_copy(
        update={"coverage": "COMPLETE", "request_parameters": {}}
    )
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, path.parent, package.model_copy(update={"sources": (source,)}))
        assess(store, run)
        result = next(r for r in store.results(run) if r.rule_id == "R0-RDS-ENCRYPTION")
        assert result.outcome == "UNKNOWN"
        assert next(store.source_data(run))["effective_coverage"] == "PARTIAL"


def test_cross_account_trail_ownership_not_collection_scope(tmp_path):
    path, package = one_file(
        tmp_path,
        {
            "trailList": [
                {
                    "TrailARN": f"arn:aws:cloudtrail:eu-west-1:{B}:trail/t",
                    "HomeRegion": "eu-west-1",
                    "LogFileValidationEnabled": False,
                }
            ]
        },
        "cloudtrail:DescribeTrails",
        account=A,
        region="us-east-1",
    )
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, path.parent, package)
        assess(store, run)
        obs = next(store.observations(run))
        assert obs.account_id == B and obs.region == "eu-west-1"
        assert "collection account differs from resource owner" in obs.unresolved
        finding = next(store.candidates(run))
        assert finding.account_id == B
        assert next(store.source_data(run))["effective_coverage"] == "PARTIAL"


def test_report_can_display_malicious_resource_label(tmp_path):
    path, package = one_file(
        tmp_path,
        {
            "DBInstances": [
                {
                    "DBInstanceArn": f"arn:aws:rds:us-east-1:{A}:db:=label",
                    "Engine": "postgres",
                    "StorageEncrypted": False,
                }
            ]
        },
    )
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, path.parent, package)
        assess(store, run)
        export_excel(store, run, tmp_path / "report.xlsx")
    wb = load_workbook(tmp_path / "report.xlsx", data_only=False)
    assert all(cell.data_type != "f" for sheet in wb for row in sheet for cell in row)
    wb.close()


def test_second_cli_process_cannot_write_open_workspace(pipeline):
    import subprocess
    import sys

    store, run, _ = pipeline
    result = subprocess.run(
        [
            sys.executable,
            "-m",
            "offline_analyzer",
            "--workspace",
            str(store.workspace),
            "summary",
            "--run",
            run,
        ],
        capture_output=True,
        text=True,
        timeout=20,
    )
    assert result.returncode == 1
    assert "WORKSPACE_UNAVAILABLE_OR_BUSY" in result.stderr


@pytest.mark.parametrize("change", ["account", "region", "resource", "timestamp", "suppressed"])
def test_incomplete_scanner_is_retained_unresolved(tmp_path, change):
    directory = tmp_path / "input"
    package = build_package(directory)
    path = directory / "prowler.json"
    scanner = json.loads(path.read_text())[0]
    if change == "account":
        scanner["cloud"]["account"].pop("uid")
        scanner["resources"][0]["uid"] = "synthetic-without-arn"
    elif change == "region":
        scanner["cloud"].pop("region")
        scanner["resources"][0].pop("region")
        scanner["resources"][0]["uid"] = "synthetic-without-arn"
    elif change == "resource":
        scanner["resources"][0].pop("uid")
    elif change == "timestamp":
        scanner.pop("time_dt")
    else:
        scanner["status_id"] = 3
    path.write_text(json.dumps([scanner]))
    source = package.sources[-1].model_copy(update={"sha256": sha256_file(path)})
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, directory, Package(sources=(source,)))
        assess(store, run)
        obs = next(store.observations(run))
        assert obs.source_status == "PASS"
        assert not list(store.candidates(run))
        assert all(
            r.outcome in {"UNKNOWN", "MANUAL_VALIDATION_REQUIRED"} for r in store.results(run)
        )


def test_rule_severity_is_independent_from_scanner_severity(pipeline):
    store, run, _ = pipeline
    finding = next(f for f in store.candidates(run) if f.rule_id == "R0-RDS-ENCRYPTION")
    assert finding.suggested_severity == "MEDIUM"
    assert finding.evidence_confidence == "MEDIUM"
    assert finding.review_status == "CANDIDATE"
    assert (
        next(o for o in store.observations(run) if o.source_tool == "prowler").source_severity
        == "High"
    )


def test_contradictory_pass_result_appears_in_review_queue(tmp_path):
    directory = tmp_path / "input"
    package = build_package(directory)
    scanner = json.loads((directory / "prowler.json").read_text())[0]
    scanner["status_code"] = "FAIL"
    scanner["resources"][0]["uid"] = f"arn:aws:rds:us-east-1:{A}:db:synthetic-db"
    scanner["resources"][0]["region"] = "us-east-1"
    scanner["cloud"]["account"]["uid"] = A
    scanner["cloud"]["region"] = "us-east-1"
    path = directory / "prowler.json"
    path.write_text(json.dumps([scanner]))
    source = package.sources[-1].model_copy(update={"sha256": sha256_file(path)})
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, directory, Package(sources=package.sources[:-1] + (source,)))
        assess(store, run)
        passed = next(
            r
            for r in store.results(run)
            if r.rule_id == "R0-RDS-ENCRYPTION" and r.outcome == "PASS"
        )
        assert passed.contradictory_observations
        export_excel(store, run, tmp_path / "report.xlsx")
    wb = load_workbook(tmp_path / "report.xlsx")
    assert any(
        row[1] == passed.evaluation_id
        for row in wb["Review Queue"].iter_rows(min_row=4, values_only=True)
    )
    wb.close()


def test_non_ascii_account_identifier_rejected():
    with pytest.raises(ValidationError):
        Source(
            format="aws-json-v1",
            operation="rds:DescribeDBInstances",
            coverage="ERROR",
            errors=("synthetic",),
            account_id="١" * 12,
            collected_at=TIME,
        )


def test_manifest_is_bounded_before_decode(tmp_path):
    directory = tmp_path / "input"
    directory.mkdir()
    (directory / "manifest.json").write_bytes(b" " * (1024 * 1024 + 1))
    with Store(tmp_path / "workspace") as store:
        with pytest.raises(InputError, match="MANIFEST_BYTE_LIMIT"):
            import_package(store, directory)
        assert not store.db.execute("SELECT 1 FROM runs").fetchone()


def test_unscoped_collection_error_is_not_silent(tmp_path):
    source = Source(
        format="aws-json-v1",
        operation="ec2:DescribeInstances",
        collected_at=TIME,
        coverage="ERROR",
        errors=("synthetic-failure",),
    )
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, tmp_path, Package(sources=(source,)))
        assess(store, run)
        results = list(store.results(run))
        # One unscoped error row per rule that consumes ec2:DescribeInstances.
        assert {r.rule_id for r in results} == {"R0-EC2-IMDSV2", "R1-EC2-IMDS-HOP-LIMIT"}
        assert all(r.account_id is None and r.region is None for r in results)
        assert all(r.outcome == "COLLECTION_ERROR" for r in results)
        assert summary(store, run)["unresolved_collection_gaps"] == 2
        export_excel(store, run, tmp_path / "unscoped.xlsx")


def test_unsupported_only_scanner_can_export_evidence_report(tmp_path):
    directory = tmp_path / "input"
    package = build_package(directory)
    path = directory / "prowler.json"
    row = json.loads(path.read_text())[-1]
    row["cloud"]["account"].pop("uid")
    row["cloud"].pop("region")
    row["resources"][0].pop("region")
    path.write_text(json.dumps([row]))
    source = package.sources[-1].model_copy(update={"sha256": sha256_file(path)})
    with Store(tmp_path / "workspace") as store:
        run = import_package(store, directory, Package(sources=(source,)))
        with pytest.raises(InputError, match="RUN_NOT_ASSESSED"):
            export_excel(store, run, tmp_path / "before.xlsx")
        assess(store, run)
        assert summary(store, run)["assessed"]
        assert summary(store, run)["outcomes"] == {}
        assert summary(store, run)["unsupported_scanner_observations"] == 1
        export_excel(store, run, tmp_path / "evidence-only.xlsx")
    wb = load_workbook(tmp_path / "evidence-only.xlsx")
    assert wb["Control Assessment"].max_row == 3
    assert wb["Evidence Register"].max_row == 5
    assert "UNSUPPORTED_SCANNER_CHECK" in str(list(wb["Evidence Register"].values))
    assert wb["Evidence Register"]["B2"].value == run
    assert wb.properties.identifier == run
    assert wb["Evidence Register"].auto_filter.ref.endswith("5")
    wb.close()
